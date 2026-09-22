"""Ручной benchmark на 15 слайдах; запускает настоящую генерацию через LLM API."""

from __future__ import annotations

import argparse
import json
import os
import signal
import subprocess
import sys
import time
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from score import score_response

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
CASE_ID = "fast-product-review-15"
TOTAL_BUDGET_SECONDS = 300


def load_case() -> dict[str, Any]:
    cases = json.loads(Path(__file__).with_name("generation_cases.json").read_text(encoding="utf-8"))
    case = next(case for case in cases if case["id"] == CASE_ID)
    sections = case["sections"]
    script = "\n\n".join(section["title"] + "\n" + section["text"] for section in sections)
    if len(sections) != 15 or case["settings"]["max_slides"] != 15 or script != case["script"]:
        raise ValueError("Сценарий benchmark должен содержать согласованные 15 разделов.")
    return case


def write_inputs(directory: Path, case: dict[str, Any]) -> None:
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    mapping = {}
    for index, section in enumerate(case["sections"], start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F5F3EC")
        slide.shapes.title.text = (
            section["title"] + " — основные факты и выводы для презентации команды продукта"
        )
        slide.placeholders[1].text = (
            section["text"] + "\n" + "Подробности и практические действия команды продукта. " * 7
        )
        for shape, size in ((slide.shapes.title, 30), (slide.placeholders[1], 22)):
            for paragraph in shape.text_frame.paragraphs:
                paragraph.font.size = Pt(size)
                paragraph.font.name = "Arial"
        mapping[str(index)] = {"0": section["title"]}
    presentation.save(directory / "template.pptx")
    (directory / "script.txt").write_text(case["script"], encoding="utf-8")
    (directory / "mapping.json").write_text(
        json.dumps(mapping, ensure_ascii=False, indent=2), encoding="utf-8",
    )


def _kill_process_group(process: subprocess.Popen) -> None:
    try:
        if os.name == "posix":
            os.killpg(process.pid, signal.SIGKILL)
        else:
            process.kill()
    except ProcessLookupError:
        pass
    try:
        process.wait(timeout=5)
    except subprocess.TimeoutExpired:
        pass


def run_benchmark(directory: Path) -> dict[str, Any]:
    """Каждый сервис работает в своём процессе; ошибки не выводят содержимое логов."""
    directory = directory.expanduser().resolve()
    directory.mkdir(parents=True, exist_ok=True)
    if any(directory.iterdir()):
        raise ValueError("Выберите новый или пустой каталог: benchmark не перезаписывает результаты.")
    started = time.monotonic()
    summary: dict[str, Any] = {"case_id": CASE_ID, "mode": "fast", "stages": {}, "success": False}
    try:
        case = load_case()
        summary.update(slide_target=15, script_characters=len(case["script"]))
        write_inputs(directory, case)
        environment = os.environ.copy()
        environment["LOG_LEVEL"] = "INFO"
        environment["LOG_FILE"] = str(directory / "content.log")
        settings = case["settings"]
        steps = (
            ("parser", "parsing-service", [
                "--input-pptx", str(directory / "template.pptx"),
                "--output-json", str(directory / "template.json"),
            ]),
            ("content", "content-service", [
                "--cli", "--template-json", str(directory / "template.json"),
                "--script", str(directory / "script.txt"),
                "--output-json", str(directory / "content.json"),
                "--user-mapping", str(directory / "mapping.json"),
                "--max-slides", "15", "--generation-mode", "fast",
                "--language", settings["language"], "--tone", settings["tone"],
                "--complexity", settings["complexity"],
            ]),
            ("builder", "builder-service", [
                "--template-pptx", str(directory / "template.pptx"),
                "--template-json", str(directory / "template.json"),
                "--content-json", str(directory / "content.json"),
                "--output-pptx", str(directory / "result.pptx"),
            ]),
        )
        with (directory / "pipeline.log").open("w", encoding="utf-8") as log:
            for stage, service, arguments in steps:
                remaining = TOTAL_BUDGET_SECONDS - (time.monotonic() - started)
                if remaining <= 0:
                    summary["error"] = "time_budget_exceeded"
                    break
                stage_started = time.monotonic()
                process = subprocess.Popen(
                    [sys.executable, "-m", "app.main", *arguments],
                    cwd=REPOSITORY_ROOT / "services" / service,
                    env=environment,
                    stdin=subprocess.DEVNULL,
                    stdout=log,
                    stderr=subprocess.STDOUT,
                    start_new_session=os.name == "posix",
                )
                stage_result: dict[str, Any] = {}
                try:
                    stage_result["exit_code"] = process.wait(timeout=remaining)
                except subprocess.TimeoutExpired:
                    _kill_process_group(process)
                    stage_result["timed_out"] = True
                    summary["error"] = "time_budget_exceeded"
                except BaseException:
                    _kill_process_group(process)
                    raise
                finally:
                    stage_result["seconds"] = round(time.monotonic() - stage_started, 2)
                    summary["stages"][stage] = stage_result
                print(json.dumps({stage: stage_result}, ensure_ascii=False), flush=True)
                if stage_result.get("timed_out"):
                    break
                if stage_result["exit_code"] != 0:
                    summary["error"] = f"{stage}_failed"
                    break
            else:
                result = Presentation(directory / "result.pptx")
                generated = json.loads((directory / "content.json").read_text(encoding="utf-8"))
                quality = score_response(case, generated)
                summary.update(slides=len(result.slides), quality=quality)
                summary["success"] = (
                    len(result.slides) == 15 and quality["slide_count"] == 15
                    and quality["score"] == 100
                    and generated.get("validation_report", {}).get("ok") is True
                )
                if not summary["success"]:
                    summary["error"] = "output_validation_failed"
    except KeyboardInterrupt:
        summary["error"] = "interrupted"
    except Exception:
        summary["error"] = "benchmark_failed"
    finally:
        elapsed = time.monotonic() - started
        summary["total_seconds"] = round(elapsed, 2)
        summary["within_budget"] = elapsed <= TOTAL_BUDGET_SECONDS
        if not summary["within_budget"]:
            summary.setdefault("error", "time_budget_exceeded")
        summary["success"] = bool(summary["success"] and summary["within_budget"])
        (directory / "summary.json").write_text(
            json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8",
        )
    return summary


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True, help="Новый или пустой каталог")
    args = parser.parse_args(argv)
    try:
        summary = run_benchmark(args.output_dir)
    except (OSError, ValueError) as error:
        print(f"Не удалось подготовить каталог benchmark: {error}", file=sys.stderr)
        return 1
    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)
    return 0 if summary["success"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
