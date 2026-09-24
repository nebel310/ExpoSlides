import asyncio
import io
import json
from typing import Any

from aiokafka import AIOKafkaConsumer, AIOKafkaProducer
from pptx import Presentation
from pptx.util import Inches

HELLOWORLD = "helloworld"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def make_template_pptx() -> bytes:
    """Генерирует минимальный PPTX с двумя текстовыми слайдами"""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    title_layout = prs.slide_layouts[0]
    slide_one = prs.slides.add_slide(title_layout)
    slide_one.shapes.title.text = "Demo title placeholder"
    if len(slide_one.placeholders) > 1:
        slide_one.placeholders[1].text = "Demo subtitle placeholder"

    content_layout = prs.slide_layouts[1]
    slide_two = prs.slides.add_slide(content_layout)
    slide_two.shapes.title.text = "Demo heading placeholder"
    if len(slide_two.placeholders) > 1:
        slide_two.placeholders[1].text = "First bullet placeholder\nSecond bullet placeholder"

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def make_script_txt() -> bytes:
    """Возвращает текст-заглушку для скрипта"""
    return "Интеграционный тест gateway. Источник не используется.".encode("utf-8")


def build_hello_world(target_length: int) -> str:
    """Возвращает строку из helloworld, обрезанную до target_length символов"""
    if target_length <= 0:
        return HELLOWORLD
    repeats = target_length // len(HELLOWORLD) + 1
    return (HELLOWORLD * repeats)[:target_length]


def build_content_json(structure: dict[str, Any]) -> dict[str, Any]:
    """Строит content.json с helloworld вместо текста каждого placeholder"""
    slides = structure.get("slides") or []
    content: dict[str, dict[str, Any]] = {}

    for slide in slides:
        slide_index = slide.get("index")
        if slide_index is None:
            continue
        placeholders: dict[str, str] = {}
        for element in slide.get("elements") or []:
            if element.get("type") != "text":
                continue
            placeholder_idx = element.get("placeholder_idx")
            placeholder_name = element.get("placeholder_name")
            key = str(placeholder_idx) if placeholder_idx is not None else placeholder_name
            if key is None:
                continue
            text_element = element.get("text") or {}
            original_length = _text_length(text_element)
            if original_length <= 0:
                original_length = 20
            placeholders[key] = build_hello_world(original_length)
        if placeholders:
            content[str(slide_index)] = {"placeholders": placeholders, "notes": None}

    return {"content": content, "validation_report": {"ok": True, "issues": []}, "error": None}


def _text_length(text_element: dict[str, Any]) -> int:
    """Приблизительная длина исходного текста из structure.json"""
    paragraphs = text_element.get("paragraphs") or []
    total = 0
    for paragraph in paragraphs:
        for run in paragraph.get("runs") or []:
            total += len(run.get("text") or "")
        total += 1
    return total


async def upload_bytes(
    client,
    filename: str,
    content: bytes,
    content_type: str,
    cookies: dict[str, str],
) -> dict[str, Any]:
    """Загружает байты в gateway и возвращает JSON-ответ"""
    response = await client.post(
        "/api/files/upload",
        files={"file": (filename, content, content_type)},
        cookies=cookies,
    )
    response.raise_for_status()
    return response.json()


async def create_task(
    client,
    template_file_id: str,
    script_file_id: str,
    cookies: dict[str, str],
) -> str:
    """Создаёт задачу и возвращает task_id"""
    response = await client.post(
        "/api/tasks",
        json={
            "template_file_id": template_file_id,
            "script_file_id": script_file_id,
        },
        cookies=cookies,
    )
    response.raise_for_status()
    return response.json()["task_id"]


async def wait_for_event(
    consumer: AIOKafkaConsumer,
    task_id: str,
    topics: set[str],
    timeout: float = 90.0,
    ignore_content_stage_failures: bool = True,
) -> dict[str, Any]:
    """Ждёт Kafka-событие для task_id из указанных топиков

    По умолчанию игнорирует task.failed со stage='content' —
    это фейл content-service, который тест намеренно не вызывает.
    """
    deadline = asyncio.get_event_loop().time() + timeout
    while True:
        remaining = deadline - asyncio.get_event_loop().time()
        if remaining <= 0:
            raise TimeoutError(f"Не дождались события {topics} для {task_id}")
        message = await asyncio.wait_for(consumer.getone(), timeout=remaining)
        if message.topic not in topics:
            continue
        payload = json.loads(message.value.decode("utf-8"))
        if payload.get("task_id") != task_id:
            continue
        if (
            ignore_content_stage_failures
            and message.topic == "task.failed"
            and (payload.get("payload") or {}).get("stage") == "content"
        ):
            continue
        return {"topic": message.topic, "payload": payload}


async def publish_content_ready(
    producer: AIOKafkaProducer,
    task_id: str,
    payload: dict[str, str],
) -> None:
    """Публикует task.content_ready с заданным payload"""
    message = {
        "task_id": task_id,
        "attempt": 1,
        "payload": payload,
        "error": None,
    }
    await producer.send_and_wait(
        "task.content_ready",
        json.dumps(message, ensure_ascii=False).encode("utf-8"),
    )


async def poll_task_status(client, task_id: str, cookies: dict[str, str], timeout: float = 60.0) -> dict[str, Any]:
    """Опрашивает статус задачи до терминального состояния"""
    deadline = asyncio.get_event_loop().time() + timeout
    while True:
        response = await client.get(f"/api/tasks/{task_id}", cookies=cookies)
        response.raise_for_status()
        task = response.json()
        if task["status"] in {"done", "failed"}:
            return task
        if asyncio.get_event_loop().time() >= deadline:
            return task
        await asyncio.sleep(1.0)