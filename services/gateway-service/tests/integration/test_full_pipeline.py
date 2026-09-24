import asyncio
import io
from typing import Any

import pytest
from pptx import Presentation

from tests.integration.helpers import (
    HELLOWORLD,
    PPTX_MIME,
    build_content_json,
    create_task,
    make_script_txt,
    make_template_pptx,
    poll_task_status,
    publish_content_ready,
    upload_bytes,
    wait_for_event,
)

pytestmark = pytest.mark.integration


async def _bootstrap(client) -> dict[str, str]:
    """Создаёт сессию и возвращает cookies"""
    response = await client.post("/api/session/bootstrap")
    response.raise_for_status()
    sid = response.json()["sid"]
    return {"exposlides_sid": sid}


async def _download_bytes(client, file_id: str, cookies: dict[str, str]) -> bytes:
    """Скачивает байты файла через gateway"""
    response = await client.get(f"/api/files/{file_id}", cookies=cookies)
    response.raise_for_status()
    return response.content


async def test_full_pipeline_from_http_to_pptx(
    gateway_client,
    kafka_producer,
    kafka_consumer,
):
    """Полный прогон от HTTP до готового PPTX без вызова LLM"""
    cookies = await _bootstrap(gateway_client)

    template_bytes = make_template_pptx()
    script_bytes = make_script_txt()

    template_upload = await upload_bytes(
        gateway_client, "template.pptx", template_bytes, PPTX_MIME, cookies,
    )
    script_upload = await upload_bytes(
        gateway_client, "script.txt", script_bytes, "text/plain", cookies,
    )

    task_id = await create_task(
        gateway_client,
        template_upload["file_id"],
        script_upload["file_id"],
        cookies,
    )

    parsed_event = await wait_for_event(kafka_consumer, task_id, {"task.parsed"}, timeout=90.0)
    structure_file_id = parsed_event["payload"]["payload"]["structure_file_id"]

    structure_bytes = await _download_bytes(gateway_client, structure_file_id, cookies)
    structure = _load_json(structure_bytes)
    content_json = build_content_json(structure)
    assert content_json["content"], "content.json не должен быть пустым"

    content_bytes = _dump_json(content_json)
    content_upload = await upload_bytes(
        gateway_client, "content.json", content_bytes, "application/json", cookies,
    )

    await publish_content_ready(
        kafka_producer,
        task_id,
        {
            "structure_file_id": structure_file_id,
            "content_file_id": content_upload["file_id"],
            "template_file_id": template_upload["file_id"],
            "script_file_id": script_upload["file_id"],
        },
    )

    built_event = await wait_for_event(kafka_consumer, task_id, {"task.built", "task.failed"}, timeout=120.0)
    assert built_event["topic"] == "task.built", built_event

    task = await poll_task_status(gateway_client, task_id, cookies, timeout=30.0)
    assert task["status"] == "done", task
    assert task["result_file_id"], "result_file_id не заполнен"

    result_bytes = await _download_bytes(gateway_client, task["result_file_id"], cookies)
    _assert_pptx_contains_helloworld(result_bytes)


def _load_json(raw: bytes) -> dict[str, Any]:
    """Загружает JSON из байтов"""
    import json
    return json.loads(raw.decode("utf-8"))


def _dump_json(value: dict[str, Any]) -> bytes:
    """Сериализует dict в UTF-8 JSON"""
    import json
    return json.dumps(value, ensure_ascii=False).encode("utf-8")


def _assert_pptx_contains_helloworld(raw: bytes) -> None:
    """Проверяет, что pptx открывается и содержит helloworld"""
    prs = Presentation(io.BytesIO(raw))
    assert len(prs.slides) >= 2, f"Ожидалось минимум 2 слайда, получено {len(prs.slides)}"

    text_fragments: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    text_fragments.append(run.text)

    assert text_fragments, "В результате нет текстовых run"
    matched = [text for text in text_fragments if HELLOWORLD in text]
    assert matched, f"Ни один run не содержит {HELLOWORLD}: {text_fragments[:5]}"