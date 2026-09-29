from pathlib import Path

import pytest
from app.export.html import HtmlExportError, convert_pptx_to_html
from pptx import Presentation
from pptx.util import Inches

pytestmark = pytest.mark.unit


def _make_minimal_pptx(path: Path, title: str = "Hello", subtitle: str = "World") -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle
    prs.save(path)


async def test_convert_html_success(tmp_path):
    """HTML создаётся и содержит текст со слайда"""
    source = tmp_path / "in.pptx"
    _make_minimal_pptx(source, title="Заголовок теста", subtitle="Подзаголовок")

    output = tmp_path / "out.html"
    result = await convert_pptx_to_html(source, output)

    assert result == output
    html = result.read_text(encoding="utf-8")
    assert html.startswith("<!DOCTYPE html>")
    assert 'class="slide"' in html
    assert "Заголовок теста" in html
    assert "Подзаголовок" in html
    assert "<style>" in html


async def test_convert_html_source_missing(tmp_path):
    """Ошибка, если исходный PPTX отсутствует"""
    with pytest.raises(HtmlExportError, match="не найден"):
        await convert_pptx_to_html(tmp_path / "missing.pptx", tmp_path / "out.html")


async def test_convert_html_escapes_special_chars(tmp_path):
    """HTML-спецсимволы экранируются"""
    source = tmp_path / "in.pptx"
    _make_minimal_pptx(source, title="<script>alert(1)</script>", subtitle="a & b")

    output = tmp_path / "out.html"
    await convert_pptx_to_html(source, output)
    html = output.read_text(encoding="utf-8")

    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "a &amp; b" in html


async def test_convert_html_with_image(tmp_path):
    """Картинка встраивается как base64"""
    png_bytes = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
        "0000000a49444154789c6300010000000500010d0a2db40000000049454e44ae426082"
    )
    image_path = tmp_path / "pixel.png"
    image_path.write_bytes(png_bytes)

    source = tmp_path / "in.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(1), width=Inches(1))
    prs.save(source)

    output = tmp_path / "out.html"
    await convert_pptx_to_html(source, output)
    html = output.read_text(encoding="utf-8")

    assert "data:image/png;base64," in html
    assert "<img" in html


async def test_convert_html_with_table(tmp_path):
    """Таблица рендерится как <table>"""
    source = tmp_path / "in.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rows, cols = 2, 2
    table = slide.shapes.add_table(
        rows, cols, Inches(1), Inches(1), Inches(4), Inches(2),
    ).table
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"
    prs.save(source)

    output = tmp_path / "out.html"
    await convert_pptx_to_html(source, output)
    html = output.read_text(encoding="utf-8")

    assert "<table>" in html
    assert "<td>A1</td>" in html
    assert "<td>B2</td>" in html


async def test_convert_html_multiple_slides(tmp_path):
    """Каждый слайд — отдельная секция"""
    source = tmp_path / "in.pptx"
    prs = Presentation()
    for index in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = f"Slide {index + 1}"
    prs.save(source)

    output = tmp_path / "out.html"
    await convert_pptx_to_html(source, output)
    html = output.read_text(encoding="utf-8")

    assert html.count('class="slide"') == 3
    assert "Slide 1" in html
    assert "Slide 3" in html
    assert 'id="slide-1"' in html
    assert 'id="slide-3"' in html


async def test_convert_html_empty_slide_marked(tmp_path):
    """Пустой слайд помечается заглушкой"""
    source = tmp_path / "in.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(source)

    output = tmp_path / "out.html"
    await convert_pptx_to_html(source, output)
    html = output.read_text(encoding="utf-8")

    assert "Пустой слайд" in html


async def test_convert_html_preserves_alignment(tmp_path):
    """Центрирование из PPTX переносится в CSS"""
    from pptx.enum.text import PP_ALIGN
    source = tmp_path / "in.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    para = box.text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = "Центр"
    prs.save(source)

    output = tmp_path / "out.html"
    await convert_pptx_to_html(source, output)
    html = output.read_text(encoding="utf-8")

    assert "text-align:center" in html
    assert "Центр" in html