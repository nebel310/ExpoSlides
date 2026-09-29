from pathlib import Path

import pytest
from app.builder import PPTXBuilder
from app.errors import BuilderInputError
from app.models.content import GeneratedContent
from app.models.presentation import Presentation as PresentationModel
from pptx import Presentation
from tests.helpers import build_content_dict, build_template_dict, make_pptx


async def _build(tmp_path: Path, slides=None, text: str = "helloworld"):
    pptx_path = tmp_path / "template.pptx"
    make_pptx(pptx_path, slides=slides)
    template_dict = build_template_dict(pptx_path)
    content_dict = build_content_dict(template_dict, text=text)
    template = PresentationModel.model_validate(template_dict)
    content = GeneratedContent.model_validate(content_dict)
    output = tmp_path / "result.pptx"
    await PPTXBuilder.build(pptx_path, template, content, output)
    return output


async def test_build_creates_output(tmp_path):
    """Сборка создаёт итоговый файл"""
    output = await _build(tmp_path)
    assert output.is_file()
    assert output.stat().st_size > 0


async def test_build_preserves_slide_count(tmp_path):
    """Число слайдов равно числу в контенте"""
    output = await _build(tmp_path, slides=[
        {"title": "A", "body": "B"},
        {"title": "C", "body": "D"},
    ])
    prs = Presentation(str(output))
    assert len(prs.slides) == 2


async def test_build_replaces_text(tmp_path):
    """Текст заменяется на helloworld"""
    output = await _build(tmp_path, text="helloworld")
    prs = Presentation(str(output))
    texts = [
        run.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for para in shape.text_frame.paragraphs
        for run in para.runs
    ]
    assert any("helloworld" in text for text in texts)
    assert all("Sample" not in text for text in texts)


async def test_build_custom_order(tmp_path):
    """Слайды идут в порядке из content"""
    pptx_path = tmp_path / "template.pptx"
    make_pptx(pptx_path, slides=[
        {"title": "Slide one", "body": "Body one"},
        {"title": "Slide two", "body": "Body two"},
    ])
    template_dict = build_template_dict(pptx_path)
    template = PresentationModel.model_validate(template_dict)
    content = GeneratedContent.model_validate({
        "content": {
            "2": {"placeholders": {"0": "second-first", "1": "second-body"}},
            "1": {"placeholders": {"0": "first-second", "1": "first-body"}},
        }
    })
    output = tmp_path / "out.pptx"
    await PPTXBuilder.build(pptx_path, template, content, output)

    prs = Presentation(str(output))
    texts = [slide.shapes.title.text for slide in prs.slides]
    assert texts == ["second-first", "first-second"]


async def test_build_rejects_same_input_output(tmp_path):
    """Нельзя сохранить поверх входного pptx"""
    pptx_path = tmp_path / "template.pptx"
    make_pptx(pptx_path)
    template_dict = build_template_dict(pptx_path)
    content_dict = build_content_dict(template_dict)
    template = PresentationModel.model_validate(template_dict)
    content = GeneratedContent.model_validate(content_dict)
    with pytest.raises(BuilderInputError, match="отличаться"):
        await PPTXBuilder.build(pptx_path, template, content, pptx_path)


async def test_build_rejects_missing_pptx(tmp_path):
    """Отсутствующий pptx отклоняется"""
    pptx_path = tmp_path / "template.pptx"
    make_pptx(pptx_path)
    template_dict = build_template_dict(pptx_path)
    content_dict = build_content_dict(template_dict)
    template = PresentationModel.model_validate(template_dict)
    content = GeneratedContent.model_validate(content_dict)
    with pytest.raises(BuilderInputError, match="открыть"):
        await PPTXBuilder.build(tmp_path / "missing.pptx", template, content, tmp_path / "out.pptx")


async def test_build_rejects_content_error(tmp_path):
    """error в контенте останавливает сборку"""
    pptx_path = tmp_path / "template.pptx"
    make_pptx(pptx_path)
    template_dict = build_template_dict(pptx_path)
    template = PresentationModel.model_validate(template_dict)
    content = GeneratedContent(error="boom")
    with pytest.raises(BuilderInputError, match="ошибку"):
        await PPTXBuilder.build(pptx_path, template, content, tmp_path / "out.pptx")


async def test_build_output_dir_created(tmp_path):
    """Каталог результата создаётся автоматически"""
    pptx_path = tmp_path / "template.pptx"
    make_pptx(pptx_path)
    template_dict = build_template_dict(pptx_path)
    content_dict = build_content_dict(template_dict)
    template = PresentationModel.model_validate(template_dict)
    content = GeneratedContent.model_validate(content_dict)
    output = tmp_path / "nested" / "dir" / "result.pptx"
    await PPTXBuilder.build(pptx_path, template, content, output)
    assert output.is_file()