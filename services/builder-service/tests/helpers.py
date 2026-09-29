from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation


def make_pptx(path: Path, slides: list[dict[str, str]] | None = None) -> None:
    """Создаёт PPTX с layout Title and Content и заданными текстами"""
    prs = Presentation()
    layout = prs.slide_layouts[1]
    if slides is None:
        slides = [{"title": "Sample title", "body": "Sample body"}]
    for spec in slides:
        slide = prs.slides.add_slide(layout)
        for shape in slide.placeholders:
            if shape.placeholder_format.idx == 0:
                shape.text = spec.get("title") or "Sample title"
            elif shape.placeholder_format.idx == 1:
                shape.text = spec.get("body") or "Sample body"
    prs.save(str(path))


def build_template_dict(pptx_path: Path) -> dict[str, Any]:
    """Строит v1-шаблон Presentation JSON из реального PPTX"""
    prs = Presentation(str(pptx_path))
    slides: list[dict[str, Any]] = []
    for idx, slide in enumerate(prs.slides, start=1):
        elements: list[dict[str, Any]] = []
        for shape in slide.shapes:
            if not shape.is_placeholder or not shape.has_text_frame:
                continue
            elements.append({
                "id": f"slide-{idx}-shape-{shape.shape_id}",
                "type": "text",
                "bbox": {
                    "left": shape.left or 0,
                    "top": shape.top or 0,
                    "width": shape.width or 0,
                    "height": shape.height or 0,
                },
                "placeholder_type": shape.placeholder_format.type.name,
                "placeholder_idx": shape.placeholder_format.idx,
                "placeholder_name": shape.name,
            })
        slides.append({
            "index": idx,
            "layout_name": slide.slide_layout.name,
            "elements": elements,
        })
    return {
        "slide_width": prs.slide_width,
        "slide_height": prs.slide_height,
        "slides": slides,
    }


def build_content_dict(
    template_dict: dict[str, Any],
    text: str = "helloworld",
) -> dict[str, Any]:
    """Строит GeneratedContent JSON с одинаковым текстом во всех placeholder'ах"""
    content: dict[str, dict[str, Any]] = {}
    for slide in template_dict["slides"]:
        placeholders: dict[str, str] = {}
        for element in slide.get("elements", []):
            if element.get("type") != "text":
                continue
            idx = element.get("placeholder_idx")
            name = element.get("placeholder_name")
            key = str(idx) if idx is not None else name
            if key is None:
                continue
            placeholders[key] = text
        content[str(slide["index"])] = {"placeholders": placeholders}
    return {"content": content}


def make_template_and_content(
    tmp_path: Path,
    slides: list[dict[str, str]] | None = None,
    text: str = "helloworld",
) -> tuple[Path, dict[str, Any], dict[str, Any]]:
    """Создаёт pptx и парные template/content словари"""
    pptx_path = tmp_path / "template.pptx"
    make_pptx(pptx_path, slides=slides)
    template = build_template_dict(pptx_path)
    content = build_content_dict(template, text=text)
    return pptx_path, template, content