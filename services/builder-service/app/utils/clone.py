import copy
from pptx import Presentation as PPTXPresentation
from pptx.util import Inches




def clone_slide(source_prs, target_prs, source_slide):
    """Клонирует слайд из source_prs в target_prs"""
    slide_layout = source_slide.slide_layout
    new_slide = target_prs.slides.add_slide(slide_layout)

    # копируем все shape-элементы
    for shape in source_slide.shapes:
        el = copy.deepcopy(shape._element)
        new_slide.shapes._spTree.append(el)

    return new_slide