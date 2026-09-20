import copy

from pptx.shapes.base import BaseShape
from pptx.slide import Slide


def get_text_placeholder_key(shape: BaseShape) -> str | None:
    """Возвращает контрактный ключ текстового placeholder."""
    if not shape.is_placeholder or not shape.has_text_frame:
        return None

    placeholder_idx = shape.placeholder_format.idx
    return str(placeholder_idx) if placeholder_idx is not None else shape.name


def replace_placeholder_text(
    slide: Slide,
    placeholder_key: str,
    new_text: str,
) -> bool:
    """Заменяет текст placeholder и сохраняет свойства его первого run."""
    for shape in slide.shapes:
        key = get_text_placeholder_key(shape)
        if key != placeholder_key:
            continue

        text_frame = shape.text_frame
        first_paragraph = text_frame.paragraphs[0]
        run_properties = None
        if first_paragraph.runs:
            first_run_properties = first_paragraph.runs[0]._r.rPr
            if first_run_properties is not None:
                run_properties = copy.deepcopy(first_run_properties)

        paragraph_properties = None
        if first_paragraph._p.pPr is not None:
            paragraph_properties = copy.deepcopy(first_paragraph._p.pPr)

        # TextFrame.clear() сохраняет свойства первого абзаца. Каждая строка
        # становится отдельным абзацем, чтобы list-placeholder применял свой
        # унаследованный маркер к каждому пункту.
        text_frame.clear()
        lines = new_text.splitlines() or [new_text]
        for index, line in enumerate(lines):
            paragraph = (
                text_frame.paragraphs[0]
                if index == 0
                else text_frame.add_paragraph()
            )
            if index > 0 and paragraph_properties is not None:
                current_paragraph_properties = paragraph._p.pPr
                if current_paragraph_properties is not None:
                    paragraph._p.remove(current_paragraph_properties)
                paragraph._p.insert(0, copy.deepcopy(paragraph_properties))

            run = paragraph.add_run()
            run.text = line
            if run_properties is not None:
                current_properties = run._r.rPr
                if current_properties is not None:
                    run._r.remove(current_properties)
                run._r.insert(0, copy.deepcopy(run_properties))
        return True

    return False
