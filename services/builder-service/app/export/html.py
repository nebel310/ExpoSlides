from __future__ import annotations

import asyncio
import base64
import html
import logging
from pathlib import Path

from pptx import Presentation as PPTXPresentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

logger = logging.getLogger(__name__)

ALIGN_CSS = {
    PP_ALIGN.LEFT: "left",
    PP_ALIGN.CENTER: "center",
    PP_ALIGN.RIGHT: "right",
    PP_ALIGN.JUSTIFY: "justify",
}

BASE_CSS = """
* { box-sizing: border-box; }
body { margin: 0; background: #1e1e1e; font-family: system-ui, -apple-system, sans-serif; }
.slide {
    position: relative;
    width: 1280px;
    height: 720px;
    margin: 24px auto;
    background: #ffffff;
    box-shadow: 0 4px 24px rgba(0, 0, 0, 0.4);
    overflow: hidden;
}
.slide .text { overflow: hidden; }
.slide .text p { margin: 0 0 4pt 0; }
.slide img { object-fit: contain; }
.slide table { border-collapse: collapse; width: 100%; height: 100%; }
.slide table td { border: 1px solid #c0c0c0; padding: 4pt; font-size: 12pt; }
.slide .empty { color: #999; font-style: italic; }
"""


class HtmlExportError(RuntimeError):
    """Не удалось сконвертировать PPTX в HTML"""


async def convert_pptx_to_html(source_pptx: Path, output_path: Path) -> Path:
    """Конвертирует PPTX в один HTML-файл с встроенными картинками"""
    return await asyncio.to_thread(_convert_sync, source_pptx, output_path)


def _convert_sync(source_pptx: Path, output_path: Path) -> Path:
    if not source_pptx.is_file():
        raise HtmlExportError(f"PPTX не найден: {source_pptx}")

    try:
        presentation = PPTXPresentation(str(source_pptx))
    except Exception as error:
        raise HtmlExportError(f"Не удалось открыть PPTX: {error}") from error

    slide_width = presentation.slide_width
    slide_height = presentation.slide_height
    if not slide_width or not slide_height:
        raise HtmlExportError("В PPTX не заданы размеры слайда")

    slides_html = []
    for index, slide in enumerate(presentation.slides, start=1):
        rendered = _render_slide(slide, slide_width, slide_height)
        slides_html.append(
            f'<section class="slide" id="slide-{index}">{rendered}</section>'
        )

    document = (
        "<!DOCTYPE html>\n"
        '<html lang="ru">\n'
        "<head>\n"
        '<meta charset="utf-8">\n'
        "<title>Presentation</title>\n"
        f"<style>{BASE_CSS}</style>\n"
        "</head>\n"
        "<body>\n"
        + "\n".join(slides_html)
        + "\n</body>\n</html>\n"
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(document, encoding="utf-8")
    logger.info("HTML создан: %s (%d байт)", output_path, output_path.stat().st_size)
    return output_path


def _render_slide(slide, slide_width: int, slide_height: int) -> str:
    """Рендерит один слайд в набор абсолютно позиционированных блоков"""
    parts: list[str] = []
    for shape in slide.shapes:
        try:
            rendered = _render_shape(shape, slide_width, slide_height)
        except Exception:
            logger.debug("Не удалось отрендерить фигуру %s", shape.shape_id, exc_info=True)
            rendered = None
        if rendered:
            parts.append(rendered)
    if not parts:
        parts.append('<div class="text empty" style="left:40%;top:45%;">Пустой слайд</div>')
    return "".join(parts)


def _render_shape(shape, slide_width: int, slide_height: int) -> str | None:
    position = _position_css(shape, slide_width, slide_height)
    if position is None:
        return None

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return _render_picture(shape, position)
    if getattr(shape, "has_table", False):
        return _render_table(shape.table, position)
    if getattr(shape, "has_text_frame", False):
        return _render_text(shape.text_frame, position)
    return None


def _render_picture(shape, position: str) -> str | None:
    try:
        image = shape.image
    except Exception:
        return None
    mime = image.content_type or "image/png"
    encoded = base64.b64encode(image.blob).decode("ascii")
    return f'<img style="{position}" src="data:{mime};base64,{encoded}" alt="">'


def _render_table(table, position: str) -> str:
    rows_html = []
    for row in table.rows:
        cells = [f"<td>{html.escape(cell.text)}</td>" for cell in row.cells]
        rows_html.append(f"<tr>{''.join(cells)}</tr>")
    return f'<div style="{position}"><table>{"".join(rows_html)}</table></div>'


def _render_text(text_frame, position: str) -> str | None:
    paragraphs_html = []
    for paragraph in text_frame.paragraphs:
        runs_html = [_run_to_html(run) for run in paragraph.runs]
        line = "".join(runs_html)
        if not line.strip():
            continue
        alignment = ALIGN_CSS.get(paragraph.alignment, "")
        style = f"text-align:{alignment};" if alignment else ""
        paragraphs_html.append(f'<p style="{style}">{line}</p>')
    if not paragraphs_html:
        return None
    return f'<div class="text" style="{position}">{"".join(paragraphs_html)}</div>'


def _run_to_html(run) -> str:
    text = html.escape(run.text)
    if not text:
        return ""
    styles = _run_style(run.font)
    if not styles:
        return text
    return f'<span style="{styles}">{text}</span>'


def _run_style(font) -> str:
    styles: list[str] = []
    if font.size is not None:
        styles.append(f"font-size:{font.size.pt}pt")
    if font.bold:
        styles.append("font-weight:bold")
    if font.italic:
        styles.append("font-style:italic")
    if font.name:
        styles.append(f"font-family:'{font.name}',sans-serif")
    try:
        if font.color is not None and font.color.type is not None and font.color.rgb:
            styles.append(f"color:#{font.color.rgb}")
    except Exception:
        pass
    return ";".join(styles)


def _position_css(shape, slide_width: int, slide_height: int) -> str | None:
    left = shape.left
    top = shape.top
    width = shape.width
    height = shape.height
    if None in (left, top, width, height):
        return None
    if width <= 0 or height <= 0:
        return None

    left_pct = left / slide_width * 100
    top_pct = top / slide_height * 100
    width_pct = width / slide_width * 100
    height_pct = height / slide_height * 100
    return (
        f"position:absolute;"
        f"left:{left_pct:.3f}%;"
        f"top:{top_pct:.3f}%;"
        f"width:{width_pct:.3f}%;"
        f"height:{height_pct:.3f}%;"
    )