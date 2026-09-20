import asyncio
import os
import tempfile
from pathlib import Path
from typing import Any

from app.errors import BuilderError, BuilderInputError, BuilderOutputError
from app.models.content import GeneratedContent
from app.models.presentation import Presentation, Slide
from app.utils.text import get_text_placeholder_key, replace_placeholder_text
from pptx import Presentation as PPTXPresentation
from pptx.oxml.ns import qn

P14_NAMESPACE = "http://schemas.microsoft.com/office/powerpoint/2010/main"
P14_SLIDE_ID_TAG = f"{{{P14_NAMESPACE}}}sldId"
LIST_PLACEHOLDER_TYPES = {
    "BODY",
    "OBJECT",
    "TEXT",
    "VERTICAL_BODY",
    "VERTICAL_OBJECT",
}


class PPTXBuilder:
    """Сборка итоговой презентации из шаблона и контента."""

    @classmethod
    async def build(
        cls,
        template_pptx_path: str | Path,
        template_data: Presentation,
        content_data: GeneratedContent,
        output_path: str | Path,
    ) -> Path:
        """Создаёт и атомарно публикует итоговый PPTX."""
        return await asyncio.to_thread(
            cls._build_sync,
            Path(template_pptx_path),
            template_data,
            content_data,
            Path(output_path),
        )

    @classmethod
    def _build_sync(
        cls,
        template_pptx_path: Path,
        template_data: Presentation,
        content_data: GeneratedContent,
        output_path: Path,
    ) -> Path:
        if template_pptx_path.expanduser().resolve() == output_path.expanduser().resolve():
            raise BuilderInputError(
                "Путь результата должен отличаться от пути PPTX-шаблона"
            )

        _validate_content_response(content_data)

        try:
            presentation = PPTXPresentation(str(template_pptx_path))
        except Exception as error:
            raise BuilderInputError(
                f"Не удалось открыть PPTX-шаблон {template_pptx_path}: {error}"
            ) from error

        requested_indices = list(content_data.content)
        slides_by_index, slide_ids_by_index = _index_slides(presentation)
        _validate_template_data(presentation, template_data, slides_by_index)
        _validate_requested_indices(requested_indices, slides_by_index)
        _validate_placeholder_mappings(
            requested_indices,
            slides_by_index,
            content_data,
        )

        _keep_and_order_slides(presentation, requested_indices, slide_ids_by_index)
        _replace_slide_content(requested_indices, slides_by_index, content_data)
        _save_atomically(presentation, output_path, len(requested_indices))
        return output_path


def _validate_content_response(content_data: GeneratedContent) -> None:
    if content_data.error:
        raise BuilderInputError(
            f"Content-service вернул ошибку, сборка отменена: {content_data.error}"
        )
    if not content_data.content:
        raise BuilderInputError("Content-service не вернул ни одного слайда")
    if (
        content_data.validation_report is not None
        and content_data.validation_report.get("ok") is False
    ):
        raise BuilderInputError("Контент не прошёл валидацию content-service")


def _index_slides(presentation: Any) -> tuple[dict[int, Any], dict[int, Any]]:
    slides = list(presentation.slides)
    slide_ids = list(presentation.slides._sldIdLst)
    return (
        {index: slide for index, slide in enumerate(slides, start=1)},
        {index: slide_id for index, slide_id in enumerate(slide_ids, start=1)},
    )


def _validate_template_data(
    presentation: Any,
    template_data: Presentation,
    slides_by_index: dict[int, Any],
) -> None:
    if (
        template_data.slide_width != presentation.slide_width
        or template_data.slide_height != presentation.slide_height
    ):
        raise BuilderInputError("Размер слайда в template.json не совпадает с PPTX-шаблоном")

    template_indices = [slide.index for slide in template_data.slides]
    if len(template_indices) != len(set(template_indices)):
        raise BuilderInputError("template.json содержит повторяющиеся индексы слайдов")

    actual_indices = set(slides_by_index)
    if set(template_indices) != actual_indices:
        raise BuilderInputError(
            "Состав слайдов в template.json не совпадает с PPTX-шаблоном"
        )

    template_slides = {slide.index: slide for slide in template_data.slides}
    for index, actual_slide in slides_by_index.items():
        template_slide = template_slides[index]
        actual_layout_name = actual_slide.slide_layout.name
        if template_slide.layout_name != actual_layout_name:
            raise BuilderInputError(
                f"Слайд {index}: layout в template.json не совпадает с PPTX-шаблоном"
            )

        expected_keys = _template_placeholder_keys(template_slide)
        actual_keys = _slide_placeholder_keys(actual_slide)
        if expected_keys != actual_keys:
            raise BuilderInputError(
                f"Слайд {index}: placeholders в template.json не совпадают "
                "с PPTX-шаблоном"
            )


def _template_placeholder_keys(slide: Slide) -> set[str]:
    keys = []
    for element in slide.elements:
        if element.type != "text" or not element.placeholder_type:
            continue
        key = (
            str(element.placeholder_idx)
            if element.placeholder_idx is not None
            else element.placeholder_name
        )
        if key is not None:
            keys.append(key)
    if len(keys) != len(set(keys)):
        raise BuilderInputError(
            f"Слайд {slide.index}: template.json содержит повторяющиеся placeholders"
        )
    return set(keys)


def _slide_placeholder_keys(slide: Any) -> set[str]:
    return set(_slide_placeholders_by_key(slide))


def _slide_placeholders_by_key(slide: Any) -> dict[str, Any]:
    placeholders = [
        (key, shape)
        for shape in slide.shapes
        if (key := get_text_placeholder_key(shape)) is not None
    ]
    if len(placeholders) != len({key for key, _ in placeholders}):
        raise BuilderInputError("PPTX-шаблон содержит повторяющиеся ключи placeholders")
    return dict(placeholders)


def _validate_requested_indices(
    requested_indices: list[int],
    slides_by_index: dict[int, Any],
) -> None:
    # dict не может содержать один ключ дважды, но проверка фиксирует запрет на
    # повторное использование слайда, если формат content в будущем изменится.
    if len(requested_indices) != len(set(requested_indices)):
        raise BuilderInputError(
            "Повторное использование одного слайда шаблона не поддерживается"
        )

    unknown_indices = sorted(set(requested_indices) - set(slides_by_index))
    if unknown_indices:
        values = ", ".join(map(str, unknown_indices))
        raise BuilderInputError(f"В PPTX-шаблоне нет слайдов с индексами: {values}")


def _validate_placeholder_mappings(
    requested_indices: list[int],
    slides_by_index: dict[int, Any],
    content_data: GeneratedContent,
) -> None:
    for index in requested_indices:
        placeholders = content_data.content[index].placeholders
        if not placeholders:
            raise BuilderInputError(
                f"Слайд {index}: content-service вернул пустой набор placeholders"
            )

        empty_keys = sorted(
            key for key, value in placeholders.items() if not value.strip()
        )
        if empty_keys:
            values = ", ".join(empty_keys)
            raise BuilderInputError(
                f"Слайд {index}: placeholders без текста: {values}"
            )

        available_placeholders = _slide_placeholders_by_key(slides_by_index[index])
        keys_with_blank_lines = sorted(
            key
            for key, value in placeholders.items()
            if key in available_placeholders
            and available_placeholders[key].placeholder_format.type.name
            in LIST_PLACEHOLDER_TYPES
            and _contains_blank_lines(value)
        )
        if keys_with_blank_lines:
            values = ", ".join(keys_with_blank_lines)
            raise BuilderInputError(
                f"Слайд {index}: placeholders содержат пустые строки: {values}"
            )

        available_keys = set(available_placeholders)
        supplied_keys = set(placeholders)
        missing_keys = sorted(available_keys - supplied_keys)
        if missing_keys:
            values = ", ".join(missing_keys)
            raise BuilderInputError(
                f"Слайд {index}: отсутствуют обязательные placeholders: {values}"
            )

        unknown_keys = sorted(supplied_keys - available_keys)
        if unknown_keys:
            values = ", ".join(unknown_keys)
            raise BuilderInputError(
                f"Слайд {index}: неизвестные ключи placeholders: {values}"
            )


def _contains_blank_lines(text: str) -> bool:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    return "\n" in normalized and any(not line.strip() for line in normalized.split("\n"))


def _keep_and_order_slides(
    presentation: Any,
    requested_indices: list[int],
    slide_ids_by_index: dict[int, Any],
) -> None:
    """Изолирует приватную XML-операцию удаления и перестановки слайдов."""
    slide_id_list = presentation.slides._sldIdLst
    requested_set = set(requested_indices)
    removed_slide_ids = [
        slide_id
        for index, slide_id in slide_ids_by_index.items()
        if index not in requested_set
    ]
    _validate_removable_slide_references(presentation, removed_slide_ids)

    for index, slide_id in slide_ids_by_index.items():
        if index in requested_set:
            continue
        relationship_id = slide_id.rId
        presentation.part.drop_rel(relationship_id)
        slide_id_list.remove(slide_id)
        if relationship_id in presentation.part.rels:
            raise BuilderInputError(
                "Не удалось безопасно удалить слайд: PPTX содержит "
                "неподдерживаемые дополнительные ссылки"
            )

    for index in requested_indices:
        slide_id = slide_ids_by_index[index]
        slide_id_list.remove(slide_id)
        slide_id_list.append(slide_id)


def _validate_removable_slide_references(
    presentation: Any,
    removed_slide_ids: list[Any],
) -> None:
    """Запрещает удаление слайдов, на которые ссылаются custom shows/sections."""
    if not removed_slide_ids:
        return

    presentation_xml = presentation.part._element
    relationship_attribute = qn("r:id")
    relationship_references = [
        element.get(relationship_attribute)
        for element in presentation_xml.iter()
        if element.get(relationship_attribute) is not None
    ]
    extra_relationship_refs = {
        slide_id.rId
        for slide_id in removed_slide_ids
        if relationship_references.count(slide_id.rId) > 1
    }

    removed_numeric_ids = {str(slide_id.id) for slide_id in removed_slide_ids}
    section_references = {
        element.get("id")
        for element in presentation_xml.iter(P14_SLIDE_ID_TAG)
        if element.get("id") is not None
    }
    referenced_sections = removed_numeric_ids & section_references

    if extra_relationship_refs or referenced_sections:
        raise BuilderInputError(
            "Нельзя исключить выбранные слайды: PPTX содержит custom shows, "
            "sections или другие presentation-level ссылки на них"
        )


def _replace_slide_content(
    requested_indices: list[int],
    slides_by_index: dict[int, Any],
    content_data: GeneratedContent,
) -> None:
    for index in requested_indices:
        slide = slides_by_index[index]
        for placeholder_key, text in content_data.content[index].placeholders.items():
            if not replace_placeholder_text(slide, placeholder_key, text):
                raise BuilderInputError(
                    f"Слайд {index}: placeholder {placeholder_key!r} не найден"
                )


def _save_atomically(
    presentation: Any,
    output_path: Path,
    expected_slide_count: int,
) -> None:
    try:
        output_path.parent.mkdir(parents=True, exist_ok=True)
    except OSError as error:
        raise BuilderOutputError(
            f"Не удалось создать каталог результата {output_path.parent}: {error}"
        ) from error

    try:
        file_descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{output_path.name}.",
            suffix=".tmp",
            dir=output_path.parent,
        )
        os.close(file_descriptor)
    except OSError as error:
        raise BuilderOutputError(
            f"Не удалось создать временный файл рядом с {output_path}: {error}"
        ) from error

    temporary_path = Path(temporary_name)
    try:
        presentation.save(str(temporary_path))
        _validate_saved_presentation(temporary_path, expected_slide_count)
        os.replace(temporary_path, output_path)
    except BuilderError:
        raise
    except Exception as error:
        raise BuilderOutputError(
            f"Не удалось сохранить и проверить итоговый PPTX: {error}"
        ) from error
    finally:
        temporary_path.unlink(missing_ok=True)


def _validate_saved_presentation(path: Path, expected_slide_count: int) -> None:
    try:
        reopened = PPTXPresentation(str(path))
    except Exception as error:
        raise BuilderOutputError(
            f"Созданный PPTX не открывается повторно: {error}"
        ) from error

    actual_slide_count = len(reopened.slides)
    if actual_slide_count != expected_slide_count:
        raise BuilderOutputError(
            "После повторного открытия число слайдов не совпало: "
            f"ожидалось {expected_slide_count}, получено {actual_slide_count}"
        )
