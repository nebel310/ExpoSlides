import base64
import io
import json
import struct
import uuid

import grpc
import pytest
from pptx import Presentation

from file_service_pb2 import (
    DeleteFileRequest,
    DownloadFileRequest,
    GetFileInfoRequest,
    UploadFileRequest,
)




# ---------------------------------------------------------------------------
# Генераторы валидных файлов: исходные типы
# ---------------------------------------------------------------------------

def create_valid_pptx() -> bytes:
    """Создать валидный pptx в памяти и вернуть bytes"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Test"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def create_valid_pptx_with_text(text: str) -> bytes:
    """Создать валидный pptx с заданным текстом заголовка"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = text
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def create_valid_pdf() -> bytes:
    """Создать минимальный валидный pdf в памяти и вернуть bytes"""
    return b"%PDF-1.4\n1 0 obj\n<< /Type /Catalog >>\nendobj\ntrailer\n<< /Root 1 0 R >>\n%%EOF"


def create_valid_json() -> bytes:
    """Создать валидный json в памяти и вернуть bytes"""
    data = {"key": "value", "list": [1, 2, 3]}
    return json.dumps(data).encode("utf-8")


def create_valid_txt() -> bytes:
    """Создать валидный txt в памяти и вернуть bytes"""
    return "Hello, world!".encode("utf-8")


# ---------------------------------------------------------------------------
# Генераторы валидных файлов: изображения и векторные форматы
# ---------------------------------------------------------------------------

def create_valid_png() -> bytes:
    """Минимальный валидный PNG 1x1 (прозрачный)"""
    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
        "+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
    )


def create_valid_jpeg() -> bytes:
    """Минимальный валидный JPEG 1x1"""
    return base64.b64decode(
        "/9j/4AAQSkZJRgABAQEAYABgAAD/2wBDAAgGBgcGBQgHBwcJCQgKDBQNDAsLDBkSEw8U"
        "HRofHh0aHBwgJC4nICIsIxwcKDcpLDAxNDQ0Hyc5PTgyPC4zNDL/wAALCAABAAEBAREA"
        "/8QAFAABAAAAAAAAAAAAAAAAAAAACf/EABQQAQAAAAAAAAAAAAAAAAAAAAD/2gAIAQEA"
        "AD8AKp//2Q=="
    )


def create_valid_gif87a() -> bytes:
    """Минимальный валидный GIF87a 1x1"""
    # GIF87a header + logical screen descriptor + image descriptor + trailer
    return (
        b"GIF87a"
        + b"\x01\x00\x01\x00"          # ширина/высота 1x1
        + b"\x80\x00\x00"              # packed, bg color, aspect
        + b"\x00\x00\x00\xff\xff\xff"  # global color table (2 цвета)
        + b"\x2c\x00\x00\x00\x00\x01\x00\x01\x00\x00"
        + b"\x02\x02\x44\x01\x00"
        + b"\x3b"
    )


def create_valid_gif89a() -> bytes:
    """Минимальный валидный GIF89a 1x1"""
    return base64.b64decode("R0lGODlhAQABAIAAAAAAAP///yH5BAEAAAAALAAAAAABAAEAAAIBRAA7")


def create_valid_bmp() -> bytes:
    """Минимальный валидный BMP 1x1 24-bit"""
    file_header = b"BM" + struct.pack("<IHHI", 58, 0, 0, 54)
    dib_header = struct.pack(
        "<IiiHHIIiiII", 40, 1, 1, 1, 24, 0, 4, 2835, 2835, 0, 0
    )
    pixel = b"\x00\x00\x00\x00"  # одна строка, выровненная до 4 байт
    return file_header + dib_header + pixel


def create_valid_tiff_le() -> bytes:
    """Минимальный валидный TIFF (little-endian, II*\\x00)"""
    return (
        b"II*\x00"
        + struct.pack("<I", 8)  # offset до IFD
        + struct.pack("<H", 0)  # 0 записей в IFD
        + struct.pack("<I", 0)  # next IFD offset = 0
    )


def create_valid_tiff_be() -> bytes:
    """Минимальный валидный TIFF (big-endian, MM\\x00*)"""
    return (
        b"MM\x00*"
        + struct.pack(">I", 8)
        + struct.pack(">H", 0)
        + struct.pack(">I", 0)
    )


def create_valid_emf() -> bytes:
    """Минимальный валидный EMF: signature + ' EMF' на смещении 40"""
    # 4 байта magic + 36 байт заголовка = 40 -> затем ' EMF'
    return b"\x01\x00\x00\x00" + b"\x00" * 36 + b" EMF" + b"\x00" * 16


def create_valid_wmf_placeable() -> bytes:
    """Минимальный валидный WMF (placeable, \\xD7\\xCD\\xC6\\x9A)"""
    return b"\xD7\xCD\xC6\x9A" + b"\x00" * 18


def create_valid_wmf_standard() -> bytes:
    """Минимальный валидный WMF (standard, \\x01\\x00\\x09\\x00)"""
    return b"\x01\x00\x09\x00" + b"\x00" * 18


def create_valid_svg_xml() -> bytes:
    """Валидный SVG, начинающийся с <?xml"""
    return b'<?xml version="1.0"?><svg xmlns="http://www.w3.org/2000/svg"/>'


def create_valid_svg_tag() -> bytes:
    """Валидный SVG, начинающийся сразу с <svg"""
    return b'<svg xmlns="http://www.w3.org/2000/svg"/>'


def create_valid_webp() -> bytes:
    """Минимальный валидный WEBP (RIFF....WEBP)"""
    payload = b"VP8L" + struct.pack("<I", 0) + b"\x00" * 8
    riff_size = 4 + len(payload)
    return b"RIFF" + struct.pack("<I", riff_size) + b"WEBP" + payload


# ---------------------------------------------------------------------------
# Фикстуры
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def valid_files():
    """Базовые валидные файлы исходных типов"""
    return {
        "test.pptx": create_valid_pptx(),
        "test.pdf": create_valid_pdf(),
        "test.json": create_valid_json(),
        "test.txt": create_valid_txt(),
    }


@pytest.fixture(scope="module")
def valid_image_files():
    """Валидные файлы новых типов (по одному представителю)"""
    return {
        "test.png": create_valid_png(),
        "test.jpg": create_valid_jpeg(),
        "test.jpeg": create_valid_jpeg(),
        "test.gif": create_valid_gif89a(),
        "test.bmp": create_valid_bmp(),
        "test.tiff": create_valid_tiff_le(),
        "test.tif": create_valid_tiff_le(),
        "test.emf": create_valid_emf(),
        "test.wmf": create_valid_wmf_placeable(),
        "test.svg": create_valid_svg_xml(),
        "test.webp": create_valid_webp(),
    }


@pytest.fixture(scope="module")
def invalid_files():
    """Невалидные файлы исходных типов"""
    return [
        ("fake.pptx", b"not a pptx"),
        ("fake.pdf", b"not a pdf"),
        ("fake.json", b"not a json"),
        ("fake.txt", b"\x00\x01\x02\x03"),
        ("no_extension", b"random bytes"),
        ("evil.exe", b"MZ\x90\x00"),
    ]


@pytest.fixture(scope="module")
def invalid_image_files():
    """Невалидные файлы новых типов (правильное расширение, но битый контент)"""
    return [
        ("fake.png", b"not a png"),
        ("fake.jpg", b"not a jpeg"),
        ("fake.jpeg", b"not a jpeg"),
        ("fake.gif", b"not a gif"),
        ("fake.bmp", b"not a bmp"),
        ("fake.tiff", b"not a tiff"),
        ("fake.tif", b"not a tiff"),
        ("fake.emf", b"not an emf"),
        ("fake.wmf", b"not a wmf"),
        ("fake.svg", b"\x00\x01\x02\x03"),
        ("fake.webp", b"not a webp"),
    ]


# ---------------------------------------------------------------------------
# Базовые тесты (исходные типы)
# ---------------------------------------------------------------------------

def test_upload_valid_files(file_service_stub, valid_files):
    """Проверить успешную загрузку файлов допустимых исходных типов"""
    for filename, content in valid_files.items():
        response = file_service_stub.UploadFile(
            UploadFileRequest(
                filename=filename,
                content=content,
                content_type="application/octet-stream",
            )
        )
        assert response.file_id
        assert response.original_name == filename
        assert response.size == len(content)
        assert response.file_type in ("pptx", "pdf", "json", "txt")
        assert response.version == 1

        file_service_stub.DeleteFile(DeleteFileRequest(file_id=response.file_id))


def test_upload_invalid_files(file_service_stub, invalid_files):
    """Проверить что невалидные файлы исходных типов отклоняются"""
    for filename, content in invalid_files:
        with pytest.raises(grpc.RpcError) as exc_info:
            file_service_stub.UploadFile(
                UploadFileRequest(
                    filename=filename,
                    content=content,
                    content_type="application/octet-stream",
                )
            )
        assert exc_info.value.code() == grpc.StatusCode.INVALID_ARGUMENT


# ---------------------------------------------------------------------------
# Новые типы: базовые валидные / невалидные
# ---------------------------------------------------------------------------

def test_upload_valid_image_files(file_service_stub, valid_image_files):
    """Проверить успешную загрузку всех новых типов"""
    expected_types = {
        "png", "jpg", "gif", "bmp", "tiff", "emf", "wmf", "svg", "webp",
    }
    for filename, content in valid_image_files.items():
        response = file_service_stub.UploadFile(
            UploadFileRequest(
                filename=filename,
                content=content,
                content_type="application/octet-stream",
            )
        )
        assert response.file_id
        assert response.original_name == filename
        assert response.size == len(content)
        assert response.file_type in expected_types
        assert response.version == 1

        file_service_stub.DeleteFile(DeleteFileRequest(file_id=response.file_id))


def test_upload_invalid_image_files(file_service_stub, invalid_image_files):
    """Невалидный контент для новых типов должен быть отклонён"""
    for filename, content in invalid_image_files:
        with pytest.raises(grpc.RpcError) as exc_info:
            file_service_stub.UploadFile(
                UploadFileRequest(
                    filename=filename,
                    content=content,
                    content_type="application/octet-stream",
                )
            )
        assert exc_info.value.code() == grpc.StatusCode.INVALID_ARGUMENT


# ---------------------------------------------------------------------------
# Алиасы расширений
# ---------------------------------------------------------------------------

def test_jpg_and_jpeg_both_accepted(file_service_stub):
    """Один и тот же JPEG должен приниматься с .jpg и .jpeg"""
    content = create_valid_jpeg()
    for filename in ("alias.jpg", "alias.jpeg"):
        response = file_service_stub.UploadFile(
            UploadFileRequest(
                filename=filename,
                content=content,
                content_type="image/jpeg",
            )
        )
        assert response.file_type == "jpg"
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=response.file_id))


def test_tif_and_tiff_both_accepted(file_service_stub):
    """Один и тот же TIFF должен приниматься с .tif и .tiff"""
    content = create_valid_tiff_le()
    for filename in ("alias.tif", "alias.tiff"):
        response = file_service_stub.UploadFile(
            UploadFileRequest(
                filename=filename,
                content=content,
                content_type="image/tiff",
            )
        )
        assert response.file_type == "tiff"
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=response.file_id))


# ---------------------------------------------------------------------------
# Краевые случаи
# ---------------------------------------------------------------------------

def test_uppercase_extension_accepted(file_service_stub):
    """Расширение в верхнем регистре должно нормализоваться"""
    cases = [
        ("UPPER.PNG", create_valid_png(), "png"),
        ("UPPER.JPEG", create_valid_jpeg(), "jpg"),
        ("UPPER.WEBP", create_valid_webp(), "webp"),
    ]
    for filename, content, expected_type in cases:
        response = file_service_stub.UploadFile(
            UploadFileRequest(
                filename=filename,
                content=content,
                content_type="application/octet-stream",
            )
        )
        assert response.file_type == expected_type
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=response.file_id))


def test_extension_content_mismatch_rejected(file_service_stub):
    """Контент одного типа с расширением другого должен отклоняться"""
    mismatches = [
        ("mismatch.jpg", create_valid_png()),    # png с расширением .jpg
        ("mismatch.png", create_valid_jpeg()),   # jpeg с расширением .png
        ("mismatch.gif", create_valid_bmp()),    # bmp с расширением .gif
        ("mismatch.bmp", create_valid_gif89a()), # gif с расширением .bmp
        ("mismatch.wmf", create_valid_emf()),    # emf с расширением .wmf
        ("mismatch.emf", create_valid_wmf_placeable()),  # wmf с расширением .emf
        ("mismatch.svg", create_valid_webp()),   # webp с расширением .svg
        ("mismatch.webp", create_valid_svg_xml()),  # svg с расширением .webp
    ]
    for filename, content in mismatches:
        with pytest.raises(grpc.RpcError) as exc_info:
            file_service_stub.UploadFile(
                UploadFileRequest(
                    filename=filename,
                    content=content,
                    content_type="application/octet-stream",
                )
            )
        assert exc_info.value.code() == grpc.StatusCode.INVALID_ARGUMENT


def test_gif_87a_and_89a_both_accepted(file_service_stub):
    """GIF87a и GIF89a — оба варианта допустимы"""
    for filename, content in (
        ("gif87.gif", create_valid_gif87a()),
        ("gif89.gif", create_valid_gif89a()),
    ):
        response = file_service_stub.UploadFile(
            UploadFileRequest(
                filename=filename,
                content=content,
                content_type="image/gif",
            )
        )
        assert response.file_type == "gif"
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=response.file_id))


def test_tiff_little_and_big_endian_both_accepted(file_service_stub):
    """TIFF с II*\\x00 и MM\\x00* — оба допустимы"""
    for filename, content in (
        ("le.tiff", create_valid_tiff_le()),
        ("be.tiff", create_valid_tiff_be()),
    ):
        response = file_service_stub.UploadFile(
            UploadFileRequest(
                filename=filename,
                content=content,
                content_type="image/tiff",
            )
        )
        assert response.file_type == "tiff"
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=response.file_id))


def test_wmf_placeable_and_standard_both_accepted(file_service_stub):
    """WMF placeable и standard — оба допустимы"""
    for filename, content in (
        ("placeable.wmf", create_valid_wmf_placeable()),
        ("standard.wmf", create_valid_wmf_standard()),
    ):
        response = file_service_stub.UploadFile(
            UploadFileRequest(
                filename=filename,
                content=content,
                content_type="image/x-wmf",
            )
        )
        assert response.file_type == "wmf"
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=response.file_id))


def test_svg_xml_declaration_and_svg_tag_both_accepted(file_service_stub):
    """SVG как с <?xml, так и сразу с <svg — оба допустимы"""
    for filename, content in (
        ("xml.svg", create_valid_svg_xml()),
        ("tag.svg", create_valid_svg_tag()),
    ):
        response = file_service_stub.UploadFile(
            UploadFileRequest(
                filename=filename,
                content=content,
                content_type="image/svg+xml",
            )
        )
        assert response.file_type == "svg"
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=response.file_id))


def test_emf_too_short_rejected(file_service_stub):
    """EMF с корректным magic, но короче 44 байт, должен отклоняться"""
    short_emf = b"\x01\x00\x00\x00" + b"\x00" * 10  # всего 14 байт
    with pytest.raises(grpc.RpcError) as exc_info:
        file_service_stub.UploadFile(
            UploadFileRequest(
                filename="short.emf",
                content=short_emf,
                content_type="image/x-emf",
            )
        )
    assert exc_info.value.code() == grpc.StatusCode.INVALID_ARGUMENT


def test_emf_missing_signature_at_offset_40_rejected(file_service_stub):
    """EMF с корректным magic, но без ' EMF' на смещении 40, отклоняется"""
    bad_emf = b"\x01\x00\x00\x00" + b"\x00" * 36 + b"XXXX" + b"\x00" * 16
    with pytest.raises(grpc.RpcError) as exc_info:
        file_service_stub.UploadFile(
            UploadFileRequest(
                filename="bad_sig.emf",
                content=bad_emf,
                content_type="image/x-emf",
            )
        )
    assert exc_info.value.code() == grpc.StatusCode.INVALID_ARGUMENT


def test_webp_missing_format_marker_rejected(file_service_stub):
    """RIFF без 'WEBP' на смещении 8 — не WEBP"""
    bad_webp = b"RIFF" + struct.pack("<I", 12) + b"XXXX" + b"\x00" * 8
    with pytest.raises(grpc.RpcError) as exc_info:
        file_service_stub.UploadFile(
            UploadFileRequest(
                filename="bad.webp",
                content=bad_webp,
                content_type="image/webp",
            )
        )
    assert exc_info.value.code() == grpc.StatusCode.INVALID_ARGUMENT


def test_empty_content_rejected(file_service_stub):
    """Пустой контент не должен приниматься ни для одного расширения"""
    for filename in ("empty.png", "empty.jpg", "empty.svg", "empty.webp"):
        with pytest.raises(grpc.RpcError) as exc_info:
            file_service_stub.UploadFile(
                UploadFileRequest(
                    filename=filename,
                    content=b"",
                    content_type="application/octet-stream",
                )
            )
        assert exc_info.value.code() == grpc.StatusCode.INVALID_ARGUMENT


def test_webp_too_short_rejected(file_service_stub):
    """WEBP короче 12 байт должен отклоняться (даже если начинается с RIFF)"""
    short_webp = b"RIFF" + b"\x00" * 4
    with pytest.raises(grpc.RpcError) as exc_info:
        file_service_stub.UploadFile(
            UploadFileRequest(
                filename="short.webp",
                content=short_webp,
                content_type="image/webp",
            )
        )
    assert exc_info.value.code() == grpc.StatusCode.INVALID_ARGUMENT


# ---------------------------------------------------------------------------
# Интеграционные тесты: скачивание/версии для нового типа
# ---------------------------------------------------------------------------

def test_download_after_upload_png(file_service_stub):
    """Скачивание PNG возвращает байт-в-байт исходное содержимое"""
    content = create_valid_png()
    response = file_service_stub.UploadFile(
        UploadFileRequest(
            filename="roundtrip.png",
            content=content,
            content_type="image/png",
        )
    )
    file_id = response.file_id
    try:
        download = file_service_stub.DownloadFile(DownloadFileRequest(file_id=file_id))
        assert download.filename == "roundtrip.png"
        assert download.content == content
        assert download.version == 1
    finally:
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=file_id))


def test_upload_new_version_wmf(file_service_stub):
    """Версионирование работает и для новых типов (на примере WMF)"""
    content_v1 = create_valid_wmf_placeable()
    content_v2 = create_valid_wmf_standard()

    response_v1 = file_service_stub.UploadFile(
        UploadFileRequest(
            filename="versioned.wmf",
            content=content_v1,
            content_type="image/x-wmf",
        )
    )
    file_id = response_v1.file_id
    try:
        response_v2 = file_service_stub.UploadFile(
            UploadFileRequest(
                filename="versioned.wmf",
                content=content_v2,
                content_type="image/x-wmf",
                file_id=file_id,
            )
        )
        assert response_v2.file_id == file_id
        assert response_v2.version == 2

        info = file_service_stub.GetFileInfo(GetFileInfoRequest(file_id=file_id, version=0))
        assert info.version == 2
        assert info.object_key.endswith(".wmf")
    finally:
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=file_id))


# ---------------------------------------------------------------------------
# Остальные исходные тесты (без изменений)
# ---------------------------------------------------------------------------

def test_upload_and_get_file_info(file_service_stub):
    """Проверить сохранение метаданных в БД после загрузки"""
    content = create_valid_pptx()
    task_id = str(uuid.uuid4())
    response = file_service_stub.UploadFile(
        UploadFileRequest(
            filename="info_test.pptx",
            content=content,
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            task_id=task_id,
        )
    )
    file_id = response.file_id
    try:
        info = file_service_stub.GetFileInfo(GetFileInfoRequest(file_id=file_id))
        assert info.file_id == file_id
        assert info.original_name == "info_test.pptx"
        assert info.size == len(content)
        assert info.task_id == task_id
        assert info.version == 1
        assert f"/v1.pptx" in info.object_key
    finally:
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=file_id))


def test_upload_without_task_id(file_service_stub):
    """Проверить что task_id может быть пустым"""
    content = create_valid_txt()
    response = file_service_stub.UploadFile(
        UploadFileRequest(
            filename="no_task.txt",
            content=content,
            content_type="text/plain",
        )
    )
    file_id = response.file_id
    try:
        info = file_service_stub.GetFileInfo(GetFileInfoRequest(file_id=file_id))
        assert info.task_id == ""
    finally:
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=file_id))


def test_delete_removes_db_record(file_service_stub):
    """Проверить что после удаления файла запись в БД удаляется"""
    content = create_valid_json()
    response = file_service_stub.UploadFile(
        UploadFileRequest(
            filename="delete_test.json",
            content=content,
            content_type="application/json",
        )
    )
    file_id = response.file_id

    file_service_stub.DeleteFile(DeleteFileRequest(file_id=file_id))

    with pytest.raises(grpc.RpcError) as exc_info:
        file_service_stub.GetFileInfo(GetFileInfoRequest(file_id=file_id))
    assert exc_info.value.code() == grpc.StatusCode.NOT_FOUND


def test_download_after_upload(file_service_stub):
    """Проверить скачивание файла после загрузки"""
    content = create_valid_pptx()
    response = file_service_stub.UploadFile(
        UploadFileRequest(
            filename="download_test.pptx",
            content=content,
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
    )
    file_id = response.file_id
    try:
        download = file_service_stub.DownloadFile(DownloadFileRequest(file_id=file_id))
        assert download.filename == "download_test.pptx"
        assert download.content == content
        assert download.version == 1
    finally:
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=file_id))


def test_upload_new_version(file_service_stub):
    """Проверить что загрузка с существующим file_id создаёт новую версию"""
    content_v1 = create_valid_pptx_with_text("Version 1")
    content_v2 = create_valid_pptx_with_text("Version 2")

    response_v1 = file_service_stub.UploadFile(
        UploadFileRequest(
            filename="versioned.pptx",
            content=content_v1,
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
    )
    file_id = response_v1.file_id
    assert response_v1.version == 1

    try:
        response_v2 = file_service_stub.UploadFile(
            UploadFileRequest(
                filename="versioned.pptx",
                content=content_v2,
                content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                file_id=file_id,
            )
        )
        assert response_v2.file_id == file_id
        assert response_v2.version == 2

        info_latest = file_service_stub.GetFileInfo(GetFileInfoRequest(file_id=file_id, version=0))
        assert info_latest.version == 2

        info_v1 = file_service_stub.GetFileInfo(GetFileInfoRequest(file_id=file_id, version=1))
        assert info_v1.version == 1

        info_v2 = file_service_stub.GetFileInfo(GetFileInfoRequest(file_id=file_id, version=2))
        assert info_v2.version == 2
        assert info_v1.object_key != info_v2.object_key
    finally:
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=file_id))


def test_download_specific_version(file_service_stub):
    """Проверить скачивание конкретной версии файла"""
    content_v1 = create_valid_pptx_with_text("Content V1")
    content_v2 = create_valid_pptx_with_text("Content V2")

    response_v1 = file_service_stub.UploadFile(
        UploadFileRequest(
            filename="dl_versioned.pptx",
            content=content_v1,
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
    )
    file_id = response_v1.file_id

    try:
        file_service_stub.UploadFile(
            UploadFileRequest(
                filename="dl_versioned.pptx",
                content=content_v2,
                content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                file_id=file_id,
            )
        )

        download_v1 = file_service_stub.DownloadFile(DownloadFileRequest(file_id=file_id, version=1))
        assert download_v1.version == 1
        assert download_v1.content == content_v1

        download_v2 = file_service_stub.DownloadFile(DownloadFileRequest(file_id=file_id, version=2))
        assert download_v2.version == 2
        assert download_v2.content == content_v2

        download_latest = file_service_stub.DownloadFile(DownloadFileRequest(file_id=file_id, version=0))
        assert download_latest.version == 2
        assert download_latest.content == content_v2
    finally:
        file_service_stub.DeleteFile(DeleteFileRequest(file_id=file_id))


def test_upload_new_version_for_nonexistent_file(file_service_stub):
    """Проверить что загрузка с несуществующим file_id отклоняется"""
    content = create_valid_pptx()
    fake_file_id = str(uuid.uuid4())
    with pytest.raises(grpc.RpcError) as exc_info:
        file_service_stub.UploadFile(
            UploadFileRequest(
                filename="fake.pptx",
                content=content,
                content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                file_id=fake_file_id,
            )
        )
    assert exc_info.value.code() == grpc.StatusCode.NOT_FOUND


def test_delete_removes_all_versions(file_service_stub):
    """Проверить что удаление убирает все версии файла"""
    content_v1 = create_valid_pptx_with_text("V1")
    content_v2 = create_valid_pptx_with_text("V2")

    response_v1 = file_service_stub.UploadFile(
        UploadFileRequest(
            filename="multi_version.pptx",
            content=content_v1,
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
    )
    file_id = response_v1.file_id

    file_service_stub.UploadFile(
        UploadFileRequest(
            filename="multi_version.pptx",
            content=content_v2,
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            file_id=file_id,
        )
    )

    file_service_stub.DeleteFile(DeleteFileRequest(file_id=file_id))

    with pytest.raises(grpc.RpcError) as exc_info:
        file_service_stub.GetFileInfo(GetFileInfoRequest(file_id=file_id, version=0))
    assert exc_info.value.code() == grpc.StatusCode.NOT_FOUND

    with pytest.raises(grpc.RpcError) as exc_info:
        file_service_stub.GetFileInfo(GetFileInfoRequest(file_id=file_id, version=1))
    assert exc_info.value.code() == grpc.StatusCode.NOT_FOUND

    with pytest.raises(grpc.RpcError) as exc_info:
        file_service_stub.GetFileInfo(GetFileInfoRequest(file_id=file_id, version=2))
    assert exc_info.value.code() == grpc.StatusCode.NOT_FOUND