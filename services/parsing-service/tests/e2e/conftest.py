from __future__ import annotations

import os
from collections.abc import AsyncIterator
from io import BytesIO
from uuid import uuid4

import grpc
import pytest
import pytest_asyncio
from aiokafka import AIOKafkaConsumer, AIOKafkaProducer
from pptx import Presentation as PPTXPresentation

from app.grpc.file_service_client import FileServiceClient

FILE_SERVICE_HOST = os.environ.get("E2E_FILE_SERVICE_HOST", "127.0.0.1")
FILE_SERVICE_PORT = int(os.environ.get("E2E_FILE_SERVICE_PORT", "50051"))
PARSER_SERVICE_HOST = os.environ.get("E2E_PARSER_SERVICE_HOST", "127.0.0.1")
PARSER_SERVICE_PORT = int(os.environ.get("E2E_PARSER_SERVICE_PORT", "50052"))
KAFKA_BOOTSTRAP_SERVERS = os.environ.get("E2E_KAFKA_BOOTSTRAP", "localhost:9093")

TOPIC_PARSED = os.environ.get("E2E_TOPIC_TASK_PARSED", "task.parsed")
TOPIC_FAILED = os.environ.get("E2E_TOPIC_TASK_FAILED", "task.failed")


def pytest_configure(config: pytest.Config) -> None:
    """Регистрирует маркер e2e"""
    config.addinivalue_line(
        "markers",
        "e2e: тесты, требующие запущенные kafka, file-service и parser-service",
    )


def _grpc_reachable(host: str, port: int) -> bool:
    """Проверяет доступность gRPC-порта коротким подключением"""
    try:
        channel = grpc.insecure_channel(f"{host}:{port}")
        grpc.channel_ready_future(channel).result(timeout=5)
        channel.close()
        return True
    except Exception:
        return False


@pytest.fixture(scope="session")
def e2e_environment() -> None:
    """Скипает e2e-тесты, если нужные сервисы не подняты"""
    missing: list[str] = []
    if not _grpc_reachable(FILE_SERVICE_HOST, FILE_SERVICE_PORT):
        missing.append(f"file-service ({FILE_SERVICE_HOST}:{FILE_SERVICE_PORT})")
    if not _grpc_reachable(PARSER_SERVICE_HOST, PARSER_SERVICE_PORT):
        missing.append(f"parser-service ({PARSER_SERVICE_HOST}:{PARSER_SERVICE_PORT})")
    if missing:
        pytest.skip(
            "E2E требует запущенных сервисов: "
            + ", ".join(missing)
            + "; подними `docker compose up -d`",
        )


@pytest_asyncio.fixture
async def file_client(e2e_environment: None) -> AsyncIterator[FileServiceClient]:
    """Подключённый FileServiceClient с авто-закрытием"""
    client = FileServiceClient(host=FILE_SERVICE_HOST, port=FILE_SERVICE_PORT)
    await client.connect()
    try:
        yield client
    finally:
        await client.close()


@pytest_asyncio.fixture
async def kafka_producer() -> AsyncIterator[AIOKafkaProducer]:
    """Продюсер для публикации в task.created"""
    producer = AIOKafkaProducer(bootstrap_servers=KAFKA_BOOTSTRAP_SERVERS)
    try:
        await producer.start()
    except Exception as error:
        pytest.skip(f"Kafka недоступна на {KAFKA_BOOTSTRAP_SERVERS}: {error}")
    try:
        yield producer
    finally:
        await producer.stop()


@pytest_asyncio.fixture
async def results_consumer() -> AsyncIterator[AIOKafkaConsumer]:
    """Консьюмер task.parsed + task.failed с уникальной группой"""
    consumer = AIOKafkaConsumer(
        TOPIC_PARSED,
        TOPIC_FAILED,
        bootstrap_servers=KAFKA_BOOTSTRAP_SERVERS,
        group_id=f"parser-service-e2e-{uuid4().hex}",
        auto_offset_reset="latest",
        enable_auto_commit=False,
    )
    try:
        await consumer.start()
    except Exception as error:
        pytest.skip(f"Kafka недоступна на {KAFKA_BOOTSTRAP_SERVERS}: {error}")
    try:
        yield consumer
    finally:
        await consumer.stop()


@pytest.fixture
def simple_pptx_bytes() -> bytes:
    """Минимальный валидный pptx в байтах"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "E2E Test"
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "Parser Service"
    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()