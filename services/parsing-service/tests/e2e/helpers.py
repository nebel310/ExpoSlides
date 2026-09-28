from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation as PPTXPresentation
from pptx.util import Emu

PNG_1X1 = (
    b"\x89PNG\r\n\x1a\n"
    b"\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
    b"\x00\x00\x00\rIDATx\x9cc\xf8\xcf\xc0\x00\x00\x00\x03\x00\x01"
    b"\x00\x00\x00\x00IEND\xaeB`\x82"
)


def make_simple_pptx(title: str = "E2E Title", subtitle: str = "E2E Subtitle") -> bytes:
    """Байты pptx с титульным слайдом"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def make_pptx_with_bullets(bullets: int = 3) -> bytes:
    """Байты pptx со слайдом и списком"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Bullets"
    body = slide.placeholders[1].text_frame
    body.text = "Bullet 1"
    for i in range(2, bullets + 1):
        p = body.add_paragraph()
        p.text = f"Bullet {i}"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def make_pptx_with_table(rows: int = 2, cols: int = 2) -> bytes:
    """Байты pptx со слайдом и таблицей"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_table(
        rows, cols, Emu(0), Emu(0), Emu(4572000), Emu(2286000)
    )
    table = shape.table
    for r in range(rows):
        for c in range(cols):
            table.cell(r, c).text = f"r{r}c{c}"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def make_pptx_with_image(tmp_dir: Path, slides_with_image: int = 1) -> bytes:
    """Байты pptx с картинкой на N слайдах"""
    img = tmp_dir / "tiny.png"
    img.write_bytes(PNG_1X1)

    prs = PPTXPresentation()
    for _ in range(slides_with_image):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(
            str(img), Emu(0), Emu(0), width=Emu(914400), height=Emu(914400)
        )
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()