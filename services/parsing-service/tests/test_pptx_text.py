from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation as PPTXPresentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from app.models.presentation import Alignment, PlaceholderKind
from app.parsers.pptx import text as text_module


# ---------- Базовые валидные случаи ----------


def _make_slide_with_text_box(prs, text: str):
    """Хелпер: создаёт слайд и textbox с заданным текстом"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(9144000), Emu(914400))
    box.text_frame.text = text
    return slide, box


def test_parse_text_frame_simple(tmp_path: Path) -> None:
    """Один параграф, один run — текст и стиль по умолчанию"""
    prs = PPTXPresentation()
    _, box = _make_slide_with_text_box(prs, "Hello")

    result = text_module.parse_text_frame(box.text_frame, None, None)

    assert len(result.paragraphs) == 1
    assert len(result.paragraphs[0].runs) == 1
    assert result.paragraphs[0].runs[0].text == "Hello"
    assert result.paragraphs[0].level == 0
    assert result.paragraphs[0].bullet is False


def test_parse_text_frame_multiple_paragraphs(tmp_path: Path) -> None:
    """Несколько параграфов — все попадают в результат"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(9144000), Emu(914400))
    tf = box.text_frame
    tf.text = "one"
    tf.add_paragraph().text = "two"
    tf.add_paragraph().text = "three"

    result = text_module.parse_text_frame(tf, None, None)

    assert len(result.paragraphs) == 3
    assert [p.runs[0].text for p in result.paragraphs] == ["one", "two", "three"]


def test_parse_text_frame_paragraph_level(tmp_path: Path) -> None:
    """Уровень вложенности параграфа сохраняется"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(9144000), Emu(914400))
    tf = box.text_frame
    tf.text = "root"
    p2 = tf.add_paragraph()
    p2.text = "child"
    p2.level = 2

    result = text_module.parse_text_frame(tf, None, None)

    assert result.paragraphs[0].level == 0
    assert result.paragraphs[1].level == 2


# ---------- Стили ----------


def test_extract_style_bold_italic_underline(tmp_path: Path) -> None:
    """Жирность, курсив и подчёркивание переносятся в стиль"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(9144000), Emu(914400))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "styled"
    run.font.bold = True
    run.font.italic = True
    run.font.underline = True

    result = text_module.parse_text_frame(box.text_frame, None, None)

    style = result.paragraphs[0].runs[0].style
    assert style.bold is True
    assert style.italic is True
    assert style.underline is True


def test_extract_style_size_and_font(tmp_path: Path) -> None:
    """Кегль и имя шрифта берутся из run"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(9144000), Emu(914400))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "sized"
    run.font.size = Pt(28)
    run.font.name = "Arial"

    result = text_module.parse_text_frame(box.text_frame, None, None)

    style = result.paragraphs[0].runs[0].style
    assert style.size_pt == pytest.approx(28.0)
    assert style.font_name == "Arial"


def test_extract_style_rgb_color(tmp_path: Path) -> None:
    """Цвет в формате RGB сохраняется как HEX без токена"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(9144000), Emu(914400))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "red"
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    result = text_module.parse_text_frame(box.text_frame, None, None)

    style = result.paragraphs[0].runs[0].style
    assert style.color_hex == "FF0000"
    assert style.color_token is None


def test_extract_style_alignment_and_spacing(tmp_path: Path) -> None:
    """Выравнивание и межстрочный интервал переносятся"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(9144000), Emu(914400))
    para = box.text_frame.paragraphs[0]
    run = para.add_run()
    run.text = "x"
    para.alignment = PP_ALIGN.CENTER
    para.line_spacing = 1.5

    result = text_module.parse_text_frame(box.text_frame, None, None)

    style = result.paragraphs[0].runs[0].style
    assert style.alignment == Alignment.CENTER
    assert style.line_spacing == pytest.approx(1.5)


# ---------- Шрифт и токены темы ----------


def test_resolve_font_with_explicit_name() -> None:
    """Если у run задан шрифт — возвращается он, без токена"""

    class FakeFont:
        name = "Arial"

    name, token = text_module.resolve_font(FakeFont(), None, None)
    assert name == "Arial"
    assert token is None


def test_resolve_font_theme_major_for_title() -> None:
    """Для заголовка без явного шрифта берётся major из темы"""

    class FakeFont:
        name = None

    class FakeTheme:
        fonts = {"major": "Inter", "minor": "Roboto"}

    name, token = text_module.resolve_font(
        FakeFont(), PlaceholderKind.TITLE, FakeTheme()
    )
    assert name == "Inter"
    assert token == "+mj-lt"


def test_resolve_font_theme_minor_for_body() -> None:
    """Для тела без явного шрифта берётся minor из темы"""

    class FakeFont:
        name = None

    class FakeTheme:
        fonts = {"major": "Inter", "minor": "Roboto"}

    name, token = text_module.resolve_font(FakeFont(), None, FakeTheme())
    assert name == "Roboto"
    assert token == "+mn-lt"


def test_resolve_font_major_placeholder_prefix() -> None:
    """Строка '+mj-lt' разворачивается в major из темы"""

    class FakeFont:
        name = "+mj-lt"

    class FakeTheme:
        fonts = {"major": "Inter", "minor": "Roboto"}

    name, token = text_module.resolve_font(FakeFont(), None, FakeTheme())
    assert name == "Inter"
    assert token == "+mj-lt"


# ---------- Краевые случаи ----------


def test_parse_text_frame_empty(tmp_path: Path) -> None:
    """Пустой text_frame возвращает None"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(9144000), Emu(914400))

    result = text_module.parse_text_frame(box.text_frame, None, None)

    assert result is None


def test_parse_text_frame_run_without_text(tmp_path: Path) -> None:
    """Run без текста не падает"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(9144000), Emu(914400))
    para = box.text_frame.paragraphs[0]
    run = para.add_run()
    run.text = ""

    result = text_module.parse_text_frame(box.text_frame, None, None)

    assert result.paragraphs[0].runs[0].text == ""


def test_extract_hyperlink_without_link(tmp_path: Path) -> None:
    """Если гиперссылки нет — возвращается None"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(9144000), Emu(914400))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "no link"

    assert text_module.extract_hyperlink(run) is None


# ---------- Невалидные данные ----------


def test_resolve_font_without_theme_returns_none() -> None:
    """Без темы и без имени шрифта возвращается (None, None)"""

    class FakeFont:
        name = None

    name, token = text_module.resolve_font(FakeFont(), None, None)
    assert name is None
    assert token is None


def test_theme_color_to_token_unknown_returns_lowercased() -> None:
    """Неизвестный цвет темы возвращается как lowercase-имя"""

    class FakeThemeColor:
        name = "SOMETHING_NEW"

    from app.parsers.pptx.helpers import theme_color_to_token

    assert theme_color_to_token(FakeThemeColor()) == "something_new"