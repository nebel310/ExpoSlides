from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pytest
from pptx import Presentation as PPTXPresentation

from app.main import _parse_args, main, run_cli


def _save_simple(tmp_path: Path) -> Path:
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Title"
    path = tmp_path / "in.pptx"
    prs.save(str(path))
    return path


# ---------- _parse_args ----------


def test_parse_args_defaults() -> None:
    args = _parse_args([])
    assert args.input_pptx is None
    assert args.output_json is None


def test_parse_args_both() -> None:
    args = _parse_args(["--input-pptx", "a.pptx", "--output-json", "a.json"])
    assert args.input_pptx == Path("a.pptx")
    assert args.output_json == Path("a.json")


def test_parse_args_only_input() -> None:
    args = _parse_args(["--input-pptx", "a.pptx"])
    assert args.input_pptx == Path("a.pptx")
    assert args.output_json is None


# ---------- run_cli ----------


@pytest.mark.asyncio
async def test_run_cli_creates_json(tmp_path: Path) -> None:
    input_path = _save_simple(tmp_path)
    output_path = tmp_path / "out.json"

    await run_cli(input_path, output_path)

    assert output_path.is_file()
    data = json.loads(output_path.read_text(encoding="utf-8"))
    assert data["file_type"] == "pptx"
    assert "schema_version" in data


@pytest.mark.asyncio
async def test_run_cli_missing_input(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        await run_cli(tmp_path / "missing.pptx", tmp_path / "out.json")


@pytest.mark.asyncio
async def test_run_cli_wrong_input_extension(tmp_path: Path) -> None:
    bad = tmp_path / "in.txt"
    bad.write_text("x")
    with pytest.raises(ValueError):
        await run_cli(bad, tmp_path / "out.json")


@pytest.mark.asyncio
async def test_run_cli_wrong_output_extension(tmp_path: Path) -> None:
    input_path = _save_simple(tmp_path)
    with pytest.raises(ValueError):
        await run_cli(input_path, tmp_path / "out.txt")


@pytest.mark.asyncio
async def test_run_cli_creates_parent_dirs(tmp_path: Path) -> None:
    input_path = _save_simple(tmp_path)
    output_path = tmp_path / "nested" / "deep" / "out.json"

    await run_cli(input_path, output_path)

    assert output_path.is_file()


# ---------- main ----------


def test_main_cli_success(tmp_path: Path) -> None:
    input_path = _save_simple(tmp_path)
    output_path = tmp_path / "out.json"

    code = main(["--input-pptx", str(input_path), "--output-json", str(output_path)])

    assert code == 0
    assert output_path.is_file()


def test_main_cli_missing_input_returns_1(tmp_path: Path) -> None:
    code = main(
        [
            "--input-pptx",
            str(tmp_path / "missing.pptx"),
            "--output-json",
            str(tmp_path / "out.json"),
        ]
    )
    assert code == 1


def test_main_only_input_returns_2(tmp_path: Path) -> None:
    code = main(["--input-pptx", str(tmp_path / "a.pptx")])
    assert code == 2


def test_main_only_output_returns_2(tmp_path: Path) -> None:
    code = main(["--output-json", str(tmp_path / "a.json")])
    assert code == 2


def test_main_service_success() -> None:
    with patch("app.main.run_service", new=AsyncMock()) as run_service:
        code = main([])
    assert code == 0
    run_service.assert_awaited_once()


def test_main_service_exception_returns_1() -> None:
    with patch("app.main.run_service", new=AsyncMock(side_effect=RuntimeError("x"))):
        code = main([])
    assert code == 1