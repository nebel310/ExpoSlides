from __future__ import annotations

import pytest
from pptx import Presentation as PPTXPresentation
from pptx.util import Emu

from app.parsers.pptx import tables as tables_module


def _make_table(prs, rows: int, cols: int):
    """Хелпер: создаёт слайд с таблицей rows×cols"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_table(
        rows,
        cols,
        Emu(0),
        Emu(0),
        Emu(4572000),
        Emu(2286000),
    )
    return shape.table


# ---------- Базовые валидные случаи ----------


def test_parse_table_dimensions() -> None:
    """Размерность таблицы соответствует заданной"""
    prs = PPTXPresentation()
    table = _make_table(prs, 3, 4)

    result = tables_module.parse_table(table)

    assert result.rows == 3
    assert result.cols == 4
    assert len(result.cells) == 3
    assert all(len(row) == 4 for row in result.cells)


def test_parse_table_cells_text() -> None:
    """Текст ячеек сохраняется"""
    prs = PPTXPresentation()
    table = _make_table(prs, 2, 2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "C"
    table.cell(1, 1).text = "D"

    result = tables_module.parse_table(table)

    assert result.cells[0][0].text == "A"
    assert result.cells[1][1].text == "D"


def test_parse_table_widths_and_heights() -> None:
    """Ширины столбцов и высоты строк переносятся"""
    prs = PPTXPresentation()
    table = _make_table(prs, 2, 3)

    result = tables_module.parse_table(table)

    assert len(result.column_widths) == 3
    assert len(result.row_heights) == 2


# ---------- Краевые случаи ----------


def test_parse_table_1x1() -> None:
    """Минимальная таблица 1×1"""
    prs = PPTXPresentation()
    table = _make_table(prs, 1, 1)
    table.cell(0, 0).text = "only"

    result = tables_module.parse_table(table)

    assert result.rows == 1
    assert result.cols == 1
    assert result.cells[0][0].text == "only"


def test_parse_table_empty_cells() -> None:
    """Пустые ячейки возвращают пустой текст"""
    prs = PPTXPresentation()
    table = _make_table(prs, 2, 2)

    result = tables_module.parse_table(table)

    for row in result.cells:
        for cell in row:
            assert cell.text == ""
            assert cell.row_span == 1
            assert cell.col_span == 1


def test_parse_cell_span_default() -> None:
    """Обычная ячейка без объединений"""
    prs = PPTXPresentation()
    table = _make_table(prs, 1, 1)
    cell = table.cell(0, 0)

    row_span, col_span, is_origin, is_spanned = tables_module.parse_cell_span(cell)

    assert row_span == 1
    assert col_span == 1
    assert is_origin is False
    assert is_spanned is False


# ---------- Флаги оформления ----------


def test_is_first_row_header_returns_bool() -> None:
    """Флаг header-строки — всегда bool"""
    prs = PPTXPresentation()
    table = _make_table(prs, 2, 2)

    result = tables_module.is_first_row_header(table)

    assert isinstance(result, bool)


def test_is_banded_rows_returns_bool() -> None:
    """Флаг banded rows — всегда bool"""
    prs = PPTXPresentation()
    table = _make_table(prs, 2, 2)

    result = tables_module.is_banded_rows(table)

    assert isinstance(result, bool)


def test_is_first_row_header_broken_table_returns_false() -> None:
    """Сломанная таблица даёт False"""

    class BrokenTable:
        pass

    assert tables_module.is_first_row_header(BrokenTable()) is False


def test_is_banded_rows_broken_table_returns_false() -> None:
    """Сломанная таблица даёт False"""

    class BrokenTable:
        pass

    assert tables_module.is_banded_rows(BrokenTable()) is False


# ---------- Невалидные данные ----------


def test_parse_cell_span_broken_returns_default() -> None:
    """Если у ячейки нет _tc, возвращаются дефолтные значения"""

    class BrokenCell:
        pass

    row_span, col_span, is_origin, is_spanned = tables_module.parse_cell_span(
        BrokenCell()
    )
    assert (row_span, col_span, is_origin, is_spanned) == (1, 1, False, False)


def test_is_first_row_header_broken_table() -> None:
    """Сломанная таблица не роняет парсер"""

    class BrokenTable:
        pass

    assert tables_module.is_first_row_header(BrokenTable()) is False
    assert tables_module.is_banded_rows(BrokenTable()) is False