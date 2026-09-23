from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation as PPTXPresentation
from pptx.util import Emu, Pt

from app.models.presentation import Fill, FillType
from app.models.presentation import (
    BBox,
    ElementType,
    LayoutInfo,
    LayoutPattern,
    PlaceholderInfo,
    PlaceholderKind,
    Slide,
    SlideElement,
    SlotSignature,
    TextStyle,
    TypographyEntry,
    TypographyScale,
)
from app.parsers.pptx import tokens as tokens_module


# ---------- extract_theme ----------


def test_extract_theme_from_real_presentation() -> None:
    """Тема реальной презентации содержит цвета и шрифты"""
    prs = PPTXPresentation()

    theme = tokens_module.extract_theme(prs)

    assert theme is not None
    assert isinstance(theme.colors, dict)
    assert isinstance(theme.fonts, dict)
    # дефолтная тема office содержит accent1 и хотя бы один шрифт
    assert "accent1" in theme.colors or "dk1" in theme.colors


def test_extract_theme_colors_have_hex() -> None:
    """Все значения цветов — 6-символьные HEX-строки"""
    prs = PPTXPresentation()

    theme = tokens_module.extract_theme(prs)

    assert theme is not None
    for token, hex_val in theme.colors.items():
        assert isinstance(hex_val, str)
        assert len(hex_val) == 6


def test_extract_theme_broken_presentation_returns_none() -> None:
    """Сломанный объект презентации возвращает None"""

    class Broken:
        pass

    assert tokens_module.extract_theme(Broken()) is None


def test_extract_theme_fonts_have_major_minor() -> None:
    """Тема содержит major и minor шрифты"""
    prs = PPTXPresentation()

    theme = tokens_module.extract_theme(prs)

    assert theme is not None
    assert "major" in theme.fonts or "minor" in theme.fonts


# ---------- extract_masters ----------


def test_extract_masters_lists_layout_indices() -> None:
    """Мастер ссылается на индексы макетов"""
    prs = PPTXPresentation()
    layout_indices = {
        str(l.part.partname): idx for idx, l in enumerate(prs.slide_layouts, start=1)
    }

    masters = tokens_module.extract_masters(prs, layout_indices)

    assert len(masters) >= 1
    assert len(masters[0].layout_indices) >= 1


def test_extract_masters_empty_layout_mapping() -> None:
    """Без маппинга макетов списки пустые"""
    prs = PPTXPresentation()

    masters = tokens_module.extract_masters(prs, {})

    assert len(masters) >= 1
    assert masters[0].layout_indices == []


# ---------- compute_patterns ----------


def _layout(name: str, index: int, kinds: list[PlaceholderKind]) -> LayoutInfo:
    """Хелпер: собирает LayoutInfo с заданными placeholder-типами"""
    placeholders = [
        PlaceholderInfo(kind=k, name=f"ph-{i}", idx=i, bbox=BBox(left=0, top=0, width=100, height=100))
        for i, k in enumerate(kinds)
    ]
    return LayoutInfo(name=name, index=index, placeholders=placeholders)


def test_compute_patterns_groups_same_signature() -> None:
    """Два макета с одинаковыми kind'ами попадают в один паттерн"""
    layouts = [
        _layout("A", 1, [PlaceholderKind.TITLE, PlaceholderKind.BODY]),
        _layout("B", 2, [PlaceholderKind.TITLE, PlaceholderKind.BODY]),
    ]
    sigs = {1: [PlaceholderKind.TITLE, PlaceholderKind.BODY], 2: [PlaceholderKind.TITLE, PlaceholderKind.BODY]}

    patterns = tokens_module.compute_patterns(layouts, sigs)

    assert len(patterns) == 1
    assert sorted(patterns[0].layout_indices) == [1, 2]


def test_compute_patterns_different_signatures() -> None:
    """Разные сигнатуры — разные паттерны"""
    layouts = [
        _layout("A", 1, [PlaceholderKind.TITLE]),
        _layout("B", 2, [PlaceholderKind.TITLE, PlaceholderKind.PICTURE]),
    ]
    sigs = {1: [PlaceholderKind.TITLE], 2: [PlaceholderKind.TITLE, PlaceholderKind.PICTURE]}

    patterns = tokens_module.compute_patterns(layouts, sigs)

    assert len(patterns) == 2


def test_compute_patterns_slots_from_first_layout() -> None:
    """Слоты паттерна — из первого макета группы"""
    layouts = [_layout("A", 1, [PlaceholderKind.TITLE, PlaceholderKind.BODY])]
    sigs = {1: [PlaceholderKind.TITLE, PlaceholderKind.BODY]}

    patterns = tokens_module.compute_patterns(layouts, sigs)

    assert len(patterns[0].slots) == 2


def test_compute_patterns_empty() -> None:
    """Без макетов паттернов нет"""
    assert tokens_module.compute_patterns([], {}) == []


def test_map_layouts_to_patterns() -> None:
    """Все layout_indices маппятся на pattern.id"""
    layouts = [_layout("A", 1, []), _layout("B", 2, [])]
    patterns = [
        LayoutPattern(id="p1", name="a", layout_indices=[1]),
        LayoutPattern(id="p2", name="b", layout_indices=[2]),
    ]

    mapping = tokens_module.map_layouts_to_patterns(layouts, patterns)

    assert mapping == {1: "p1", 2: "p2"}


# ---------- element_signature ----------


def _simple_element(
    element_id: str = "e1",
    width: int = 100,
    height: int = 100,
    element_type: ElementType = ElementType.SHAPE,
) -> SlideElement:
    """Хелпер: собирает минимальный SlideElement"""
    return SlideElement(
        id=element_id,
        type=element_type,
        bbox=BBox(left=0, top=0, width=width, height=height),
    )


def test_element_signature_shape_returns_string() -> None:
    """Для SHAPE возвращается непустая сигнатура"""
    elem = _simple_element()

    sig = tokens_module.element_signature(elem)

    assert sig is not None
    assert "shape" in sig


def test_element_signature_text_returns_none() -> None:
    """Для TEXT сигнатура не строится"""
    elem = _simple_element(element_type=ElementType.TEXT)

    assert tokens_module.element_signature(elem) is None


def test_element_signature_chart_returns_none() -> None:
    """Для CHART сигнатура не строится"""
    elem = _simple_element(element_type=ElementType.CHART)

    assert tokens_module.element_signature(elem) is None


def test_element_signature_same_for_identical() -> None:
    """Идентичные элементы дают одинаковую сигнатуру"""
    a = _simple_element("a")
    b = _simple_element("b")

    assert tokens_module.element_signature(a) == tokens_module.element_signature(b)


def test_element_signature_differs_by_size() -> None:
    """Разный размер — разная сигнатура"""
    a = _simple_element("a", width=100)
    b = _simple_element("b", width=200)

    assert tokens_module.element_signature(a) != tokens_module.element_signature(b)


# ---------- compute_components ----------


def _slide_with_elements(index: int, elements: list[SlideElement]) -> Slide:
    """Хелпер: собирает минимальный слайд"""
    return Slide(index=index, elements=elements)


def test_compute_components_finds_duplicates() -> None:
    """Одинаковые элементы на 2+ слайдах становятся компонентом"""
    slides = [
        _slide_with_elements(1, [_simple_element("a")]),
        _slide_with_elements(2, [_simple_element("b")]),
    ]

    components = tokens_module.compute_components(slides)

    assert len(components) == 1
    assert len(components[0].occurrences) == 2


def test_compute_components_skips_unique() -> None:
    """Уникальные элементы не становятся компонентами"""
    slides = [
        _slide_with_elements(1, [_simple_element("a", width=100)]),
        _slide_with_elements(2, [_simple_element("b", width=200)]),
    ]

    components = tokens_module.compute_components(slides)

    assert components == []


def test_compute_components_empty_slides() -> None:
    """Пустые слайды — пустой список компонентов"""
    assert tokens_module.compute_components([]) == []
    assert tokens_module.compute_components([_slide_with_elements(1, [])]) == []


def test_compute_components_ids_increment() -> None:
    """Идентификаторы компонентов идут по порядку component-1, component-2"""
    slides = [
        _slide_with_elements(1, [_simple_element("a", width=100), _simple_element("c", width=500)]),
        _slide_with_elements(2, [_simple_element("b", width=100), _simple_element("d", width=500)]),
    ]

    components = tokens_module.compute_components(slides)

    ids = sorted(c.id for c in components)
    assert ids == ["component-1", "component-2"]


# ---------- compute_typography ----------


def _text_element(element_id: str, size_pt: float) -> SlideElement:
    """Хелпер: элемент с одним run заданного размера"""
    from app.models.presentation import Paragraph, Run, TextElement

    return SlideElement(
        id=element_id,
        type=ElementType.TEXT,
        bbox=BBox(left=0, top=0, width=100, height=100),
        text=TextElement(
            paragraphs=[
                Paragraph(runs=[Run(text="x", style=TextStyle(size_pt=size_pt))])
            ]
        ),
    )


def test_compute_typography_single_size() -> None:
    """Один размер — одна запись с ролью body"""
    slides = [_slide_with_elements(1, [_text_element("a", 18.0)])]

    result = tokens_module.compute_typography(slides)

    assert len(result.entries) == 1
    assert result.entries[0].size_pt == 18.0
    assert result.entries[0].role == "body"
    assert result.min_pt == 18.0
    assert result.max_pt == 18.0


def test_compute_typography_multiple_sizes() -> None:
    """Несколько размеров — роли назначаются по рангу"""
    slides = [
        _slide_with_elements(
            1,
            [
                _text_element("a", 44.0),
                _text_element("b", 28.0),
                _text_element("c", 18.0),
                _text_element("d", 18.0),
                _text_element("e", 18.0),
            ],
        )
    ]

    result = tokens_module.compute_typography(slides)

    sizes = [e.size_pt for e in result.entries]
    assert sizes == [44.0, 28.0, 18.0]
    # 18.0 самый частый — body
    by_size = {e.size_pt: e.role for e in result.entries}
    assert by_size[18.0] == "body"


def test_compute_typography_empty() -> None:
    """Без текстовых элементов шкала пустая"""
    result = tokens_module.compute_typography([])

    assert result.entries == []
    assert result.min_pt is None
    assert result.max_pt is None


def test_compute_typography_visits_group_children() -> None:
    """Размеры внутри групп тоже учитываются"""
    from app.models.presentation import GroupElement

    inner = _text_element("child", 20.0)
    group = SlideElement(
        id="g",
        type=ElementType.GROUP,
        bbox=BBox(left=0, top=0, width=100, height=100),
        group=GroupElement(children=[inner]),
    )
    slides = [_slide_with_elements(1, [group])]

    result = tokens_module.compute_typography(slides)

    assert len(result.entries) == 1
    assert result.entries[0].size_pt == 20.0


# ---------- compute_grid ----------


def _layout_with_placeholders(
    index: int,
    bboxes: list[tuple[int, int, int, int]],
) -> LayoutInfo:
    """Хелпер: LayoutInfo с placeholder'ами заданной геометрии"""
    placeholders = [
        PlaceholderInfo(
            kind=PlaceholderKind.BODY,
            name=f"ph-{i}",
            idx=i,
            bbox=BBox(left=l, top=t, width=w, height=h),
        )
        for i, (l, t, w, h) in enumerate(bboxes)
    ]
    return LayoutInfo(name=f"L{index}", index=index, placeholders=placeholders)


def test_compute_grid_margins() -> None:
    """Поля вычисляются от крайних placeholder'ов"""
    layouts = [_layout_with_placeholders(1, [(914400, 457200, 4572000, 3000000)])]

    grid = tokens_module.compute_grid(layouts, slide_width=9144000, slide_height=6858000)

    assert grid.margin_left == 914400
    assert grid.margin_top == 457200


def test_compute_grid_empty_layouts() -> None:
    """Без placeholder'ов поля не вычисляются"""
    grid = tokens_module.compute_grid([], slide_width=9144000, slide_height=6858000)

    assert grid.margin_left is None
    assert grid.column_positions == []


def test_compute_grid_column_positions_snapped() -> None:
    """Позиции колонок округляются до сетки GRID_STEP_EMU"""
    step = tokens_module.GRID_STEP_EMU
    layouts = [
        _layout_with_placeholders(
            1,
            [
                (step, 0, 100, 100),
                (step * 5 + 100, 0, 100, 100),
            ],
        )
    ]

    grid = tokens_module.compute_grid(layouts, slide_width=9144000, slide_height=6858000)

    # обе позиции должны попасть в разные колонки
    assert len(grid.column_positions) >= 2


def test_compute_grid_snap_function() -> None:
    """_snap округляет к ближайшему шагу"""
    step = tokens_module.GRID_STEP_EMU
    assert tokens_module._snap(step) == step
    assert tokens_module._snap(step * 3) == step * 3
    assert tokens_module._snap(0) == 0


def test_build_layout_size_map_returns_dict() -> None:
    prs = PPTXPresentation()
    layout_indices = {
        str(l.part.partname): idx for idx, l in enumerate(prs.slide_layouts, start=1)
    }
    result = tokens_module.build_layout_size_map(prs, layout_indices)
    assert isinstance(result, dict)


def test_collect_all_fonts_empty() -> None:
    assert tokens_module.collect_all_fonts([]) == []


def test_collect_all_fonts_ignores_empty() -> None:
    from app.models.presentation import Slide, SlideElement, BBox, ElementType, TextElement, Paragraph, Run, TextStyle

    elem = SlideElement(
        id="a",
        type=ElementType.TEXT,
        bbox=BBox(left=0, top=0, width=100, height=100),
        text=TextElement(
            paragraphs=[Paragraph(runs=[Run(text="x", style=TextStyle())])]
        ),
    )
    assert tokens_module.collect_all_fonts([Slide(index=1, elements=[elem])]) == []


def test_element_signature_shape_without_fill_color() -> None:
    """Фигура с заливкой без цвета не должна ронять построение сигнатуры"""

    elem = SlideElement(
        id="e1",
        type=ElementType.SHAPE,
        bbox=BBox(left=0, top=0, width=100, height=100),
        fill=Fill(type=FillType.NONE),
    )

    sig = tokens_module.element_signature(elem)

    assert sig is not None
    assert "shape" in sig