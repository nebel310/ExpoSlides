from __future__ import annotations

from pathlib import Path

import pytest
from app.models.presentation import (
    BackgroundKind,
    ElementType,
    FillType,
)
from app.parsers.pptx import shapes as shapes_module
from pptx import Presentation as PPTXPresentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt


def _new_prs() -> PPTXPresentation:
    """Пустая презентация со стандартным размером слайда"""
    return PPTXPresentation()


def _blank_slide(prs):
    """Слайд с пустым макетом"""
    return prs.slides.add_slide(prs.slide_layouts[6])


# ---------- parse_shape: базовые типы ----------


def test_parse_shape_text_box() -> None:
    """Textbox распознаётся как элемент TEXT"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(914400), Emu(914400))
    box.text_frame.text = "hello"

    result = shapes_module.parse_shape(box, None, "el-1", {})

    assert result is not None
    assert result.type == ElementType.TEXT
    assert result.id == "el-1"
    assert result.bbox.width == 914400
    assert result.text is not None


def test_parse_shape_auto_shape() -> None:
    """Автофигура без текста распознаётся как SHAPE"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Emu(0),
        Emu(0),
        Emu(914400),
        Emu(914400),
    )

    result = shapes_module.parse_shape(shape, None, "el-1", {})

    assert result is not None
    assert result.type == ElementType.SHAPE
    assert result.geometry is not None
    assert result.geometry.shape_type == "ROUNDED_RECTANGLE"


def test_parse_shape_picture(tmp_path: Path) -> None:
    """Картинка распознаётся как IMAGE и регистрирует ассет"""
    png = (
        b"\x89PNG\r\n\x1a\n"
        b"\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
        b"\x00\x00\x00\rIDATx\x9cc\xf8\xcf\xc0\x00\x00\x00\x03\x00\x01"
        b"\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    img = tmp_path / "t.png"
    img.write_bytes(png)

    prs = _new_prs()
    slide = _blank_slide(prs)
    shape = slide.shapes.add_picture(
        str(img), Emu(0), Emu(0), width=Emu(914400), height=Emu(914400)
    )

    assets: dict = {}
    result = shapes_module.parse_shape(shape, None, "el-1", assets)

    assert result is not None
    assert result.type == ElementType.IMAGE
    assert result.image is not None
    assert result.image.asset_id is not None
    assert result.image.asset_id in assets


def test_parse_shape_table() -> None:
    """Таблица распознаётся как TABLE"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    shape = slide.shapes.add_table(
        2, 2, Emu(0), Emu(0), Emu(914400), Emu(914400)
    )

    result = shapes_module.parse_shape(shape, None, "el-1", {})

    assert result is not None
    assert result.type == ElementType.TABLE
    assert result.table is not None
    assert result.table.rows == 2


def test_parse_shape_placeholder_kind_title() -> None:
    """Placeholder заголовка получает kind=TITLE"""
    prs = _new_prs()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title

    result = shapes_module.parse_shape(title, None, "el-1", {})

    assert result is not None
    assert result.placeholder_kind is not None
    assert result.placeholder_kind.value == "title"


# ---------- z-order и rotation ----------


def test_parse_shape_z_order_increases() -> None:
    """Каждая следующая фигура имеет больший z-order"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    s1 = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(100), Emu(100))
    s2 = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(100), Emu(100))

    r1 = shapes_module.parse_shape(s1, None, "a", {})
    r2 = shapes_module.parse_shape(s2, None, "b", {})

    assert r1.z_order is not None and r2.z_order is not None
    assert r2.z_order > r1.z_order


def test_parse_shape_rotation() -> None:
    """Rotation переносится в модель"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    shape = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(100), Emu(100))
    shape.rotation = 45.0

    result = shapes_module.parse_shape(shape, None, "el-1", {})

    assert result.rotation == pytest.approx(45.0)


def test_parse_shape_rotation_default_none() -> None:
    """Без rotation поле остаётся None"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    shape = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(100), Emu(100))

    result = shapes_module.parse_shape(shape, None, "el-1", {})

    assert result.rotation is None


# ---------- Заливки ----------


def test_parse_fill_solid_rgb() -> None:
    """Solid заливка сохраняет HEX"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), Emu(914400), Emu(914400)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)

    result = shapes_module.parse_fill(shape, None, {})

    assert result is not None
    assert result.type == FillType.SOLID
    assert result.color_hex == "123456"


def test_parse_fill_no_fill() -> None:
    """Без заливки возвращается Fill с типом NONE"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), Emu(914400), Emu(914400)
    )
    shape.fill.background()

    result = shapes_module.parse_fill(shape, None, {})

    assert result is not None
    assert result.type in (FillType.NONE, FillType.BACKGROUND)


def test_parse_fill_broken_returns_none() -> None:
    """Фигура без .fill не роняет парсер"""

    class Broken:
        pass

    assert shapes_module.parse_fill(Broken(), None, {}) is None


# ---------- Линии ----------


def test_parse_line_color_and_width() -> None:
    """Обводка сохраняет цвет и толщину"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), Emu(914400), Emu(914400)
    )
    shape.line.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    shape.line.width = Pt(2.0)

    result = shapes_module.parse_line(shape, None)

    assert result is not None
    assert result.color_hex == "FF0000"
    assert result.width_pt == pytest.approx(2.0)


def test_parse_line_empty_returns_none() -> None:
    """Без цвета и ширины линия не возвращается"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), Emu(914400), Emu(914400)
    )

    result = shapes_module.parse_line(shape, None)

    assert result is None


def test_parse_line_broken_returns_none() -> None:
    """Сломанная фигура не роняет парсер"""

    class Broken:
        pass

    assert shapes_module.parse_line(Broken(), None) is None


# ---------- Геометрия ----------


def test_parse_geometry_auto_shape() -> None:
    """AutoShape сохраняет тип и adjustments"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Emu(0),
        Emu(0),
        Emu(914400),
        Emu(914400),
    )

    result = shapes_module.parse_geometry(shape)

    assert result is not None
    assert result.shape_type == "ROUNDED_RECTANGLE"
    assert isinstance(result.adjustments, list)


def test_parse_geometry_textbox_returns_none() -> None:
    """Textbox не имеет auto_shape_type и adjustments"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(914400), Emu(914400))

    result = shapes_module.parse_geometry(box)

    assert result is None


# ---------- Группы ----------


def test_parse_group_with_two_children() -> None:
    """Группа из двух фигур даёт двух потомков с префиксом id"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    try:
        group = slide.shapes.add_group_shape()
    except AttributeError:
        pytest.skip("add_group_shape недоступен")

    group.shapes.add_textbox(Emu(0), Emu(0), Emu(100), Emu(100))
    group.shapes.add_textbox(Emu(100), Emu(0), Emu(100), Emu(100))

    result = shapes_module.parse_group(group, None, "grp", {})

    assert len(result.children) == 2
    assert result.children[0].id == "grp-child-0"
    assert result.children[1].id == "grp-child-1"


def test_parse_shape_group_type() -> None:
    """Фигура-группа распознаётся как GROUP"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    try:
        group = slide.shapes.add_group_shape()
    except AttributeError:
        pytest.skip("add_group_shape недоступен")

    result = shapes_module.parse_shape(group, None, "grp", {})

    assert result is not None
    assert result.type == ElementType.GROUP
    assert result.group is not None


# ---------- Фон ----------


def test_parse_background_none() -> None:
    """None-фон не падает"""
    assert shapes_module.parse_background(None, None, {}) is None


def test_parse_background_solid_neutral() -> None:
    """Solid фон классифицируется как NEUTRAL"""
    from app.models.presentation import Fill

    fill = Fill(type=FillType.SOLID, color_hex="FFFFFF")
    kind = shapes_module.classify_background(fill)

    assert kind == BackgroundKind.NEUTRAL


def test_parse_background_picture_unknown() -> None:
    """Picture фон классифицируется как UNKNOWN (Analyzer уточнит)"""
    from app.models.presentation import Fill

    fill = Fill(type=FillType.PICTURE)
    kind = shapes_module.classify_background(fill)

    assert kind == BackgroundKind.UNKNOWN


def test_parse_background_gradient_neutral() -> None:
    """Gradient фон — NEUTRAL"""
    from app.models.presentation import Fill

    fill = Fill(type=FillType.GRADIENT)
    kind = shapes_module.classify_background(fill)

    assert kind == BackgroundKind.NEUTRAL


# ---------- Коннекторы ----------


def test_parse_connector_returns_none_for_textbox() -> None:
    """Textbox не является коннектором"""
    prs = _new_prs()
    slide = _blank_slide(prs)
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(100), Emu(100))

    assert shapes_module.parse_connector(box) is None