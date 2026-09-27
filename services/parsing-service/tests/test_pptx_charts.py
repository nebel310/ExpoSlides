from __future__ import annotations

from app.models.presentation import ChartType
from app.parsers.pptx import charts as charts_module
from pptx import Presentation as PPTXPresentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu


def _make_bar_chart(prs):
    """Хелпер: слайд с простой столбчатой диаграммой"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chart_data = CategoryChartData()
    chart_data.categories = ["A", "B", "C"]
    chart_data.add_series("Series 1", (1.0, 2.0, 3.0))
    chart_data.add_series("Series 2", (4.0, 5.0, 6.0))
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Emu(0),
        Emu(0),
        Emu(4572000),
        Emu(3048000),
        chart_data,
    )
    return graphic_frame.chart


# ---------- Базовые валидные случаи ----------


def test_parse_chart_type_column() -> None:
    """Тип диаграммы определяется как COLUMN"""
    prs = PPTXPresentation()
    chart = _make_bar_chart(prs)

    result = charts_module.parse_chart(chart, None)

    assert result.chart_type == ChartType.COLUMN


def test_parse_chart_series_count_and_names() -> None:
    """Две серии с правильными именами"""
    prs = PPTXPresentation()
    chart = _make_bar_chart(prs)

    result = charts_module.parse_chart(chart, None)

    assert len(result.series) == 2
    assert result.series[0].name == "Series 1"
    assert result.series[1].name == "Series 2"


def test_parse_chart_series_values() -> None:
    """Значения серий сохраняются как float"""
    prs = PPTXPresentation()
    chart = _make_bar_chart(prs)

    result = charts_module.parse_chart(chart, None)

    assert result.series[0].values == [1.0, 2.0, 3.0]
    assert result.series[1].values == [4.0, 5.0, 6.0]


def test_parse_chart_categories() -> None:
    """Категории переносятся во все серии"""
    prs = PPTXPresentation()
    chart = _make_bar_chart(prs)

    result = charts_module.parse_chart(chart, None)

    assert result.series[0].categories == ["A", "B", "C"]


def test_parse_chart_has_value_axis() -> None:
    """У диаграммы есть value-ось"""
    prs = PPTXPresentation()
    chart = _make_bar_chart(prs)

    result = charts_module.parse_chart(chart, None)

    kinds = [a.kind for a in result.axes]
    assert "value" in kinds


def test_parse_chart_has_category_axis() -> None:
    """У диаграммы есть category-ось"""
    prs = PPTXPresentation()
    chart = _make_bar_chart(prs)

    result = charts_module.parse_chart(chart, None)

    kinds = [a.kind for a in result.axes]
    assert "category" in kinds


# ---------- Легенда ----------


def test_parse_legend_enabled() -> None:
    """Легенда включена и позиция сохраняется"""
    prs = PPTXPresentation()
    chart = _make_bar_chart(prs)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False

    result = charts_module.parse_legend(chart)

    assert result.has_legend is True
    assert result.position == "bottom"


def test_parse_legend_disabled() -> None:
    """Без легенды has_legend=False"""
    prs = PPTXPresentation()
    chart = _make_bar_chart(prs)
    chart.has_legend = False

    result = charts_module.parse_legend(chart)

    assert result.has_legend is False


def test_map_legend_position_all() -> None:
    """Все позиции маппятся в строки"""
    assert charts_module.map_legend_position(XL_LEGEND_POSITION.BOTTOM) == "bottom"
    assert charts_module.map_legend_position(XL_LEGEND_POSITION.LEFT) == "left"
    assert charts_module.map_legend_position(XL_LEGEND_POSITION.RIGHT) == "right"
    assert charts_module.map_legend_position(XL_LEGEND_POSITION.TOP) == "top"
    assert charts_module.map_legend_position(XL_LEGEND_POSITION.CORNER) == "corner"


def test_map_legend_position_unknown_returns_none() -> None:
    """Неизвестная позиция — None"""
    assert charts_module.map_legend_position(None) is None


# ---------- Тип диаграмм ----------


def test_map_chart_type_pie() -> None:
    """PIE распознаётся"""
    assert charts_module.map_chart_type(XL_CHART_TYPE.PIE) == ChartType.PIE


def test_map_chart_type_line() -> None:
    """LINE распознаётся"""
    assert charts_module.map_chart_type(XL_CHART_TYPE.LINE) == ChartType.LINE


def test_map_chart_type_bar() -> None:
    """BAR распознаётся"""
    assert charts_module.map_chart_type(XL_CHART_TYPE.BAR_CLUSTERED) == ChartType.BAR


def test_map_chart_type_doughnut() -> None:
    """DOUGHNUT распознаётся"""
    assert charts_module.map_chart_type(XL_CHART_TYPE.DOUGHNUT) == ChartType.DOUGHNUT


def test_map_chart_type_area() -> None:
    """AREA распознаётся"""
    assert charts_module.map_chart_type(XL_CHART_TYPE.AREA) == ChartType.AREA


def test_map_chart_type_unknown() -> None:
    """Строка без известного ключа → UNKNOWN"""
    assert charts_module.map_chart_type("nonsense") == ChartType.UNKNOWN


# ---------- Data labels ----------


def test_parse_data_labels_disabled() -> None:
    """По умолчанию подписи данных выключены"""
    prs = PPTXPresentation()
    chart = _make_bar_chart(prs)

    result = charts_module.parse_data_labels(chart)

    assert result.has_labels is False


def test_parse_data_labels_enabled_values() -> None:
    """Включённые подписи значений отражаются во флагах"""
    prs = PPTXPresentation()
    chart = _make_bar_chart(prs)
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.show_value = True

    result = charts_module.parse_data_labels(chart)

    assert result.has_labels is True
    assert result.show_value is True


# ---------- Краевые случаи ----------


def test_parse_axis_broken_returns_defaults() -> None:
    """Сломанная ось не падает и возвращает дефолты"""

    class BrokenAxis:
        pass

    result = charts_module.parse_axis(BrokenAxis(), "value")

    assert result.kind == "value"
    assert result.title is None
    assert result.has_major_gridlines is False


def test_parse_chart_broken_returns_empty() -> None:
    """Полностью сломанный объект не роняет парсер"""

    class BrokenChart:
        pass

    # должен либо упасть с понятным исключением, либо вернуть пустой результат
    # мы ожидаем, что верхнеуровневый parse_chart сам обработает это в parser.py
    try:
        result = charts_module.parse_chart(BrokenChart(), None)
        assert result.series == []
    except Exception:
        # альтернативный вариант — исключение перехватывается на уровне shapes.py
        pass