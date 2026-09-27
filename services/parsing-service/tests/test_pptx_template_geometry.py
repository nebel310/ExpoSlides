"""Реальные PPTX: диаграммы без осей, несколько masters и вложенные группы."""

from copy import deepcopy
from pathlib import Path

import pytest
from app.models.presentation import BBox, LayoutInfo, PlaceholderInfo, PlaceholderKind
from app.parsers.pptx import PPTXParser
from app.parsers.pptx.tokens import compute_patterns
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.parts.slide import SlideLayoutPart, SlideMasterPart
from pptx.util import Inches


@pytest.mark.asyncio
@pytest.mark.parametrize("chart_type", [XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT])
async def test_native_chart_without_axes_roundtrip(tmp_path: Path, chart_type) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Share", [25, 75])
    slide.shapes.add_chart(chart_type, Inches(1), Inches(1), Inches(5), Inches(3), data)
    path = tmp_path / "without-axes.pptx"
    prs.save(path)
    result = await PPTXParser.parse(path)
    chart = result.presentation.model_dump(mode="json")["slides"][0]["elements"][0]["chart"]
    assert chart["axes"] == []
    assert chart["series"][0]["categories"] == ["A", "B"]
    assert chart["series"][0]["values"] == [25, 75]
    assert len(Presentation(path).slides) == 1


def _second_master(prs):
    """Создаёт настоящий второй master, свой layout и отличающуюся тему."""
    source = prs.slide_masters[0].part
    package = prs.part.package
    master = SlideMasterPart.load(
        PackURI("/ppt/slideMasters/slideMaster2.xml"), source.content_type, package, source.blob
    )
    source_layout = prs.slide_layouts[0].part
    layout = SlideLayoutPart.load(
        PackURI("/ppt/slideLayouts/slideLayout12.xml"),
        source_layout.content_type,
        package,
        source_layout.blob,
    )
    theme_source = source.part_related_by(RT.THEME)
    theme_xml = etree.fromstring(theme_source.blob)
    theme_xml.find(".//" + qn("a:accent1") + "/" + qn("a:srgbClr")).set("val", "CC3311")
    theme_xml.find(".//" + qn("a:majorFont") + "/" + qn("a:latin")).set("typeface", "Courier New")
    theme = Part(
        PackURI("/ppt/theme/theme2.xml"),
        theme_source.content_type,
        package,
        etree.tostring(theme_xml),
    )
    master.relate_to(theme, RT.THEME)
    rid = master.relate_to(layout, RT.SLIDE_LAYOUT)
    layout.relate_to(master, RT.SLIDE_MASTER)
    layout_ids = master._element.find(qn("p:sldLayoutIdLst"))
    for child in list(layout_ids):
        layout_ids.remove(child)
    child = OxmlElement("p:sldLayoutId")
    child.set("id", "2147483659")
    child.set(qn("r:id"), rid)
    layout_ids.append(child)
    master_ids = prs._element.get_or_add_sldMasterIdLst()
    child = OxmlElement("p:sldMasterId")
    child.set("id", "2147483649")
    child.set(qn("r:id"), prs.part.relate_to(master, RT.SLIDE_MASTER))
    master_ids.append(child)
    return layout.slide_layout


@pytest.mark.asyncio
async def test_all_masters_and_actual_slide_theme(tmp_path: Path) -> None:
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = "First"
    second = prs.slides.add_slide(_second_master(prs))
    second.shapes.title.text = "Second"
    font = second.shapes.title.text_frame.paragraphs[0].runs[0].font
    font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
    font.name = "+mj-lt"
    path = tmp_path / "masters.pptx"
    prs.save(path)
    parsed = (await PPTXParser.parse(path)).presentation.model_dump(mode="json")
    assert len(parsed["masters"]) == 2
    assert len(parsed["layouts"]) == 12
    assert parsed["masters"][1]["layout_indices"] == [12]
    assert parsed["masters"][1]["theme"]["colors"]["accent1"] == "CC3311"
    slide = parsed["slides"][1]
    assert slide["master_index"] == 2
    assert slide["layout_index"] == 12
    assert slide["theme"]["fonts"]["major"] == "Courier New"
    style = slide["elements"][0]["text"]["paragraphs"][0]["runs"][0]["style"]
    assert style["font_name"] == "Courier New"
    assert style["color_hex"] == "CC3311"
    assert parsed["layouts"][11]["master_index"] == 2
    assert parsed["layouts"][11]["pattern_id"] == slide["pattern_id"]


@pytest.mark.asyncio
async def test_nested_group_scaled_and_rotated_world_geometry(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    text = inner.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1))
    text.text = "Editable nested content"
    outer.shapes._recalculate_extents()
    outer.left = Inches(3)
    outer.top = Inches(2)
    outer.width = Inches(4)
    outer.height = Inches(2)
    outer.rotation = 90
    path = tmp_path / "group.pptx"
    prs.save(path)
    first = (await PPTXParser.parse(path)).presentation.model_dump(mode="json")
    second = (await PPTXParser.parse(path)).presentation.model_dump(mode="json")
    assert first == second
    leaf = first["slides"][0]["elements"][0]["group"]["children"][0]["group"]["children"][0]
    assert leaf["shape_path"] == [outer.shape_id, inner.shape_id, text.shape_id]
    assert leaf["shape_id"] == text.shape_id
    assert leaf["shape_name"] == text.name
    assert leaf["bbox"] == dict(left=Inches(1), top=Inches(2), width=Inches(2), height=Inches(1))
    assert leaf["slide_bbox"] == dict(
        left=Inches(4), top=Inches(1), width=Inches(2), height=Inches(4)
    )
    assert leaf["text"]["paragraphs"][0]["runs"][0]["text"] == "Editable nested content"


@pytest.mark.asyncio
async def test_layout_without_placeholders_retains_text_slots(tmp_path: Path) -> None:
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for shape in list(layout.shapes):
        layout.shapes._spTree.remove(shape._element)
    slide = prs.slides.add_slide(layout)
    text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    text.text = "Ordinary title"
    layout.shapes._spTree.insert_element_before(deepcopy(text._element), "p:extLst")
    path = tmp_path / "textbox-layout.pptx"
    prs.save(path)
    parsed = (await PPTXParser.parse(path)).presentation.model_dump(mode="json")
    actual_layout = next(item for item in parsed["layouts"] if item["index"] == 7)
    assert actual_layout["placeholders"] == []
    assert (
        actual_layout["elements"][0]["text"]["paragraphs"][0]["runs"][0]["text"] == "Ordinary title"
    )
    pattern = next(item for item in parsed["patterns"] if 7 in item["layout_indices"])
    assert pattern["name"] != "blank"
    assert pattern["slots"][0]["element_id"] == actual_layout["elements"][0]["id"]


def test_same_placeholder_roles_with_different_geometry_are_distinct() -> None:
    layouts = [
        LayoutInfo(
            name="A",
            index=1,
            placeholders=[
                PlaceholderInfo(
                    kind=PlaceholderKind.BODY,
                    idx=1,
                    bbox=BBox(left=0, top=0, width=1_000_000, height=2_000_000),
                )
            ],
        ),
        LayoutInfo(
            name="B",
            index=2,
            placeholders=[
                PlaceholderInfo(
                    kind=PlaceholderKind.BODY,
                    idx=1,
                    bbox=BBox(left=3_000_000, top=0, width=1_000_000, height=2_000_000),
                )
            ],
        ),
    ]
    patterns = compute_patterns(layouts, {1: [PlaceholderKind.BODY], 2: [PlaceholderKind.BODY]})
    assert len(patterns) == 2
    assert patterns[0].slots[0].bbox != patterns[1].slots[0].bbox


@pytest.mark.asyncio
async def test_empty_and_filled_placeholder_inherit_color_bold_and_font(tmp_path):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    prs = Presentation()
    layout = prs.slide_layouts[0]
    font = layout.placeholders[0].text_frame.paragraphs[0].font
    font.name = "Arial"
    font.size = Pt(38)
    font.color.rgb = RGBColor(255, 255, 255)
    font.bold = True
    empty = prs.slides.add_slide(layout)
    empty.shapes.title.text = ""
    filled = prs.slides.add_slide(layout)
    filled.shapes.title.text = "Inherited white title"
    path = tmp_path / "inherited.pptx"
    prs.save(path)
    parsed = (await PPTXParser.parse(path)).presentation.model_dump(mode="json")
    default = parsed["slides"][0]["elements"][0]["default_text_style"]
    actual = parsed["slides"][1]["elements"][0]["text"]["paragraphs"][0]["runs"][0]["style"]
    for style in [default, actual]:
        assert style["color_hex"] == "FFFFFF"
        assert style["font_name"] == "Arial"
        assert style["size_pt"] == 38
        assert style["bold"] is True
    assert parsed["slides"][0]["elements"][0]["has_text_frame"] is True
    assert parsed["masters"][0]["elements"]
