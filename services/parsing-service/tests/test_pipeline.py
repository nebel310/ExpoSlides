from __future__ import annotations

from pathlib import Path

import pytest
from app.kafka.schemas import TaskCreatedPayload
from app.services.parser_pipeline import (
    STRUCTURE_CONTENT_TYPE,
    STRUCTURE_FILENAME,
    ParserPipeline,
)
from pptx import Presentation as PPTXPresentation


def _make_pptx_bytes() -> bytes:
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "T"
    import io

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_pipeline_process_full_flow(fake_file_client) -> None:
    pptx_bytes = _make_pptx_bytes()
    fake_file_client.put("template-id", pptx_bytes)

    pipeline = ParserPipeline(fake_file_client)
    payload = TaskCreatedPayload(template_file_id="template-id", script_file_id="script-id")

    result = await pipeline.process("task-1", payload)

    assert result.template_file_id == "template-id"
    assert result.script_file_id == "script-id"
    assert result.structure_file_id


@pytest.mark.asyncio
async def test_pipeline_uploads_structure_json(fake_file_client) -> None:
    fake_file_client.put("template-id", _make_pptx_bytes())
    pipeline = ParserPipeline(fake_file_client)
    payload = TaskCreatedPayload(template_file_id="template-id", script_file_id="script-id")

    await pipeline.process("task-1", payload)

    uploads = fake_file_client.upload_calls
    structure_uploads = [u for u in uploads if u["filename"] == STRUCTURE_FILENAME]
    assert len(structure_uploads) == 1
    assert structure_uploads[0]["content_type"] == STRUCTURE_CONTENT_TYPE
    assert structure_uploads[0]["task_id"] == "task-1"


@pytest.mark.asyncio
async def test_pipeline_downloads_template(fake_file_client) -> None:
    fake_file_client.put("template-id", _make_pptx_bytes())
    pipeline = ParserPipeline(fake_file_client)
    payload = TaskCreatedPayload(template_file_id="template-id", script_file_id="script-id")

    await pipeline.process("task-1", payload)

    downloads = fake_file_client.download_calls
    assert any(d["file_id"] == "template-id" for d in downloads)


@pytest.mark.asyncio
async def test_pipeline_missing_template_raises(fake_file_client) -> None:
    pipeline = ParserPipeline(fake_file_client)
    payload = TaskCreatedPayload(template_file_id="missing", script_file_id="s")

    with pytest.raises(FileNotFoundError):
        await pipeline.process("task-1", payload)


@pytest.mark.asyncio
async def test_pipeline_invalid_pptx_bytes_raises(fake_file_client) -> None:
    fake_file_client.put("template-id", b"not a pptx")
    pipeline = ParserPipeline(fake_file_client)
    payload = TaskCreatedPayload(template_file_id="template-id", script_file_id="s")

    with pytest.raises(Exception):
        await pipeline.process("task-1", payload)


@pytest.mark.asyncio
async def test_pipeline_uploads_assets(fake_file_client, tmp_path: Path) -> None:
    import io

    from pptx.util import Emu

    png = (
        b"\x89PNG\r\n\x1a\n"
        b"\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
        b"\x00\x00\x00\rIDATx\x9cc\xf8\xcf\xc0\x00\x00\x00\x03\x00\x01"
        b"\x00\x00\x00\x00IEND\xaeB`\x82"
    )

    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    img_path = tmp_path / "tiny-test.png"
    img_path.write_bytes(png)
    slide.shapes.add_picture(
        str(img_path), Emu(0), Emu(0), width=Emu(914400), height=Emu(914400)
    )

    buf = io.BytesIO()
    prs.save(buf)
    fake_file_client.put("template-id", buf.getvalue())

    pipeline = ParserPipeline(fake_file_client)
    payload = TaskCreatedPayload(template_file_id="template-id", script_file_id="s")

    await pipeline.process("task-1", payload)

    uploads = fake_file_client.upload_calls
    asset_uploads = [u for u in uploads if u["filename"] != STRUCTURE_FILENAME]
    assert len(asset_uploads) >= 1