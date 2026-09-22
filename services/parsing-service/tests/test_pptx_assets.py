"""Тесты работы с бинарными ассетами: картинки, OLE, picture-fill."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation as PPTXPresentation
from pptx.util import Emu

from app.parsers.pptx import assets as assets_module
from app.parsers.pptx.assets import AssetBlob


PNG_1X1 = (
    b"\x89PNG\r\n\x1a\n"
    b"\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
    b"\x00\x00\x00\rIDATx\x9cc\xf8\xcf\xc0\x00\x00\x00\x03\x00\x01"
    b"\x00\x00\x00\x00IEND\xaeB`\x82"
)


# ---------- register_asset ----------


def test_register_asset_new() -> None:
    """Новый ассет добавляется в словарь и возвращает свой id"""
    assets: dict[str, AssetBlob] = {}

    asset_id = assets_module.register_asset(
        data=b"hello",
        content_type="text/plain",
        original_name="a.txt",
        assets=assets,
    )

    assert asset_id in assets
    assert assets[asset_id].data == b"hello"
    assert assets[asset_id].content_type == "text/plain"
    assert assets[asset_id].original_name == "a.txt"


def test_register_asset_deduplicates() -> None:
    """Одинаковые байты не дублируются в словаре"""
    assets: dict[str, AssetBlob] = {}

    id1 = assets_module.register_asset(b"same", "text/plain", "a.txt", assets)
    id2 = assets_module.register_asset(b"same", "text/plain", "b.txt", assets)

    assert id1 == id2
    assert len(assets) == 1


def test_register_asset_different_bytes() -> None:
    """Разные байты — разные id"""
    assets: dict[str, AssetBlob] = {}

    id1 = assets_module.register_asset(b"one", "text/plain", None, assets)
    id2 = assets_module.register_asset(b"two", "text/plain", None, assets)

    assert id1 != id2
    assert len(assets) == 2


def test_register_asset_default_content_type() -> None:
    """Если content_type = None — ставится application/octet-stream"""
    assets: dict[str, AssetBlob] = {}

    asset_id = assets_module.register_asset(b"data", None, None, assets)

    assert assets[asset_id].content_type == "application/octet-stream"


# ---------- parse_image ----------


def test_parse_image_extracts_blob(tmp_path: Path) -> None:
    """Картинка сохраняется в assets, метаданные заполнены"""
    img_path = tmp_path / "tiny.png"
    img_path.write_bytes(PNG_1X1)

    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_picture(
        str(img_path),
        Emu(0),
        Emu(0),
        width=Emu(914400),
        height=Emu(914400),
    )

    assets: dict[str, AssetBlob] = {}
    image = assets_module.parse_image(shape, assets)

    assert image.asset_id is not None
    assert image.asset_id in assets
    assert assets[image.asset_id].data == PNG_1X1
    assert image.content_type == "image/png"
    assert image.original_width == 1
    assert image.original_height == 1


def test_parse_image_broken_shape_returns_empty() -> None:
    """Если у shape нет .image, возвращается пустой ImageElement"""

    class BrokenShape:
        alt_text = None

    result = assets_module.parse_image(BrokenShape(), {})

    assert result.asset_id is None
    assert result.content_type is None


# ---------- parse_ole ----------


def test_parse_ole_broken_returns_empty() -> None:
    """Сломанный OLE-объект не роняет парсер"""

    class BrokenShape:
        class _element:
            @staticmethod
            def get(key):
                return None

        class part:
            rels = {}

    result = assets_module.parse_ole(BrokenShape(), {})

    assert result.prog_id is None
    assert result.asset_id is None


# ---------- parse_picture_fill_asset ----------


def test_parse_picture_fill_asset_broken_returns_none() -> None:
    """Если blip отсутствует — возвращается None"""

    class BrokenShape:
        class _element:
            @staticmethod
            def find(path):
                return None

    assert assets_module.parse_picture_fill_asset(BrokenShape(), {}) is None


# ---------- AssetBlob ----------


def test_asset_blob_fields() -> None:
    """AssetBlob хранит все переданные поля"""
    blob = AssetBlob(
        asset_id="id",
        content_type="image/png",
        data=b"\x89PNG",
        original_name="a.png",
    )

    assert blob.asset_id == "id"
    assert blob.content_type == "image/png"
    assert blob.data == b"\x89PNG"
    assert blob.original_name == "a.png"