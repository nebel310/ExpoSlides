from __future__ import annotations

from io import BytesIO

from pptx import Presentation as PPTXPresentation
from pptx.util import Inches, Pt


def build_simple_pptx() -> bytes:
    """Генерит pptx с титульным слайдом и слайдом с текстом"""
    prs = PPTXPresentation()
    slide_width = prs.slide_width
    slide_height = prs.slide_height

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = "Integration Test"
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = "Parser Service"

    content_slide = prs.slides.add_slide(prs.slide_layouts[1])
    content_slide.shapes.title.text = "Key points"
    body = content_slide.placeholders[1].text_frame
    body.text = "First point"
    paragraph = body.add_paragraph()
    paragraph.text = "Second point"
    paragraph.level = 1

    assert prs.slide_width == slide_width
    assert prs.slide_height == slide_height

    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def build_pptx_with_table() -> bytes:
    """Генерит pptx со слайдом, содержащим таблицу"""
    prs = PPTXPresentation()
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    rows, cols = 3, 3
    left, top = Inches(1), Inches(1)
    width, height = Inches(6), Inches(3)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    for r in range(rows):
        for c in range(cols):
            table.cell(r, c).text = f"cell-{r}-{c}"

    textbox = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(1))
    tf = textbox.text_frame
    tf.text = "Note about the table"
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(18)

    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()