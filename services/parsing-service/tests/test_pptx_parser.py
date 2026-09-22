from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation as PPTXPresentation
from pptx.util import Emu

from app.models.presentation import SCHEMA_VERSION
from app.parsers.pptx import PPTXParser, ParseResult
from app.parsers.pptx.assets import AssetBlob


PNG_1X1 = (
    b"\x89PNG\r\n\x1a\n"
    b"\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
    b"\x00\x00\x00\rIDATx\x9cc\xf8\xcf\xc0\x00\x00\x00\x03\x00\x01"
    b"\x00\x00\x00\x00IEND\xaeB`\x82"
)


def _save_simple(tmp_path: Path) -> Path:
    """Хелпер: сохраняет pptx с одним титульным слайдом"""
    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Title"
    path = tmp_path / "s.pptx"
    prs.save(str(path))
    return path


# ---------- Базовый парсинг ----------


@pytest.mark.asyncio
async def test_parse_returns_parse_result(tmp_path: Path) -> None:
    """parse возвращает ParseResult с presentation и assets"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    assert isinstance(result, ParseResult)
    assert result.presentation is not None
    assert isinstance(result.assets, dict)


@pytest.mark.asyncio
async def test_parse_schema_version(tmp_path: Path) -> None:
    """В presentation проставлен schema_version"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    assert result.presentation.schema_version == SCHEMA_VERSION


@pytest.mark.asyncio
async def test_parse_slide_dimensions(tmp_path: Path) -> None:
    """Размер слайда переносится"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    assert result.presentation.slide_width > 0
    assert result.presentation.slide_height > 0


@pytest.mark.asyncio
async def test_parse_one_slide(tmp_path: Path) -> None:
    """Презентация содержит ровно один слайд"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    assert len(result.presentation.slides) == 1
    assert result.presentation.slides[0].index == 1


@pytest.mark.asyncio
async def test_parse_slide_has_layout_index(tmp_path: Path) -> None:
    """Слайд ссылается на индекс макета"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    assert result.presentation.slides[0].layout_index is not None


@pytest.mark.asyncio
async def test_parse_slide_has_pattern_id(tmp_path: Path) -> None:
    """Слайд получает pattern_id из паттерна макета"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    assert result.presentation.slides[0].pattern_id is not None


@pytest.mark.asyncio
async def test_parse_slide_content_hash_present(tmp_path: Path) -> None:
    """content_hash всегда вычислен"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    assert result.presentation.slides[0].content_hash is not None
    assert len(result.presentation.slides[0].content_hash) == 16


@pytest.mark.asyncio
async def test_parse_layouts_present(tmp_path: Path) -> None:
    """Макеты присутствуют в модели"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    assert len(result.presentation.layouts) >= 1
    assert result.presentation.layouts[0].name


@pytest.mark.asyncio
async def test_parse_patterns_present(tmp_path: Path) -> None:
    """Паттерны построены"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    assert len(result.presentation.patterns) >= 1


@pytest.mark.asyncio
async def test_parse_masters_present(tmp_path: Path) -> None:
    """Мастера присутствуют"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    assert len(result.presentation.masters) >= 1


# ---------- Токены ----------


@pytest.mark.asyncio
async def test_parse_tokens_theme_present(tmp_path: Path) -> None:
    """В токенах есть тема с цветами"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    tokens = result.presentation.tokens
    assert tokens.theme is not None
    assert isinstance(tokens.theme.colors, dict)


@pytest.mark.asyncio
async def test_parse_tokens_typography_present(tmp_path: Path) -> None:
    """В токенах есть типографическая шкала"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    typography = result.presentation.tokens.typography
    assert typography is not None
    assert typography.min_pt is None or typography.min_pt > 0


@pytest.mark.asyncio
async def test_parse_tokens_grid_present(tmp_path: Path) -> None:
    """В токенах есть сетка"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    assert result.presentation.tokens.grid is not None


# ---------- Ассеты ----------


@pytest.mark.asyncio
async def test_parse_assets_from_picture(tmp_path: Path) -> None:
    """Картинка на слайде становится ассетом"""
    img = tmp_path / "t.png"
    img.write_bytes(PNG_1X1)

    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(
        str(img), Emu(0), Emu(0), width=Emu(914400), height=Emu(914400)
    )
    path = tmp_path / "pic.pptx"
    prs.save(str(path))

    result = await PPTXParser.parse(path)

    assert len(result.assets) >= 1
    blob = next(iter(result.assets.values()))
    assert isinstance(blob, AssetBlob)
    assert blob.data == PNG_1X1


@pytest.mark.asyncio
async def test_parse_assets_refs_in_presentation(tmp_path: Path) -> None:
    """asset_refs в presentation ссылаются на ассеты"""
    img = tmp_path / "t.png"
    img.write_bytes(PNG_1X1)

    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(
        str(img), Emu(0), Emu(0), width=Emu(914400), height=Emu(914400)
    )
    path = tmp_path / "pic.pptx"
    prs.save(str(path))

    result = await PPTXParser.parse(path)

    assert len(result.presentation.assets) >= 1
    asset_ref = result.presentation.assets[0]
    assert asset_ref.asset_id in result.assets


@pytest.mark.asyncio
async def test_parse_assets_deduplication(tmp_path: Path) -> None:
    """Две одинаковые картинки — один ассет"""
    img = tmp_path / "t.png"
    img.write_bytes(PNG_1X1)

    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for _ in range(2):
        slide.shapes.add_picture(
            str(img), Emu(0), Emu(0), width=Emu(914400), height=Emu(914400)
        )
    path = tmp_path / "dup.pptx"
    prs.save(str(path))

    result = await PPTXParser.parse(path)

    assert len(result.assets) == 1


# ---------- Краевые случаи ----------


@pytest.mark.asyncio
async def test_parse_empty_presentation(tmp_path: Path) -> None:
    """Презентация без слайдов парсится без ошибок"""
    prs = PPTXPresentation()
    path = tmp_path / "empty.pptx"
    prs.save(str(path))

    result = await PPTXParser.parse(path)

    assert result.presentation.slides == []
    assert result.assets == {}


@pytest.mark.asyncio
async def test_parse_no_assets_no_pictures(tmp_path: Path) -> None:
    """Без картинок ассетов нет"""
    path = _save_simple(tmp_path)

    result = await PPTXParser.parse(path)

    assert result.assets == {}
    assert result.presentation.assets == []


@pytest.mark.asyncio
async def test_parse_missing_file_raises(tmp_path: Path) -> None:
    """Отсутствующий файл → исключение"""
    missing = tmp_path / "nope.pptx"

    with pytest.raises(Exception):
        await PPTXParser.parse(missing)


@pytest.mark.asyncio
async def test_parse_not_pptx_file_raises(tmp_path: Path) -> None:
    """Невалидный pptx → исключение"""
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a real pptx")

    with pytest.raises(Exception):
        await PPTXParser.parse(bad)


# ---------- Детерминированность ----------


@pytest.mark.asyncio
async def test_parse_two_runs_same_source_path(tmp_path: Path) -> None:
    """Два прогона одного файла дают одинаковый слайд-хеш"""
    path = _save_simple(tmp_path)

    first = await PPTXParser.parse(path)
    second = await PPTXParser.parse(path)

    assert first.presentation.slides[0].content_hash == second.presentation.slides[0].content_hash


@pytest.mark.asyncio
async def test_parse_two_runs_same_asset_ids(tmp_path: Path) -> None:
    """asset_id детерминирован по содержимому"""
    img = tmp_path / "t.png"
    img.write_bytes(PNG_1X1)

    prs = PPTXPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(
        str(img), Emu(0), Emu(0), width=Emu(914400), height=Emu(914400)
    )
    path = tmp_path / "pic.pptx"
    prs.save(str(path))

    first = await PPTXParser.parse(path)
    second = await PPTXParser.parse(path)

    assert list(first.assets.keys()) == list(second.assets.keys())