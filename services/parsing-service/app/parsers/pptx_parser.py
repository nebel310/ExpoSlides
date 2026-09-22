import asyncio
import logging
from pathlib import Path

from app.models.legacy_presentation import (
    BBox,
    ElementType,
    ImageElement,
    LayoutInfo,
    LayoutType,
    Paragraph,
    PlaceholderInfo,
    Presentation,
    Run,
    Slide,
    SlideBackground,
    SlideElement,
    TableElement,
    TextElement,
    TextStyle,
    ThemeInfo,
)
from app.parsers.base import BaseParser
from app.parsers.text_style import inherited_font
from lxml import etree
from pptx import Presentation as PPTXPresentation
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.shapes.base import BaseShape
from pptx.slide import Slide as PPTXSlide
from pptx.slide import SlideLayout as PPTXSlideLayout
from pptx.text.text import TextFrame

logger = logging.getLogger(__name__)




class PPTXParser(BaseParser):
    """Парсер файлов PowerPoint (.pptx)"""

    @classmethod
    async def parse(cls, file_path: str | Path) -> Presentation:
        """Асинхронно парсит pptx файл в модель Presentation"""
        path = Path(file_path)
        return await asyncio.to_thread(cls._parse_sync, path)

    @classmethod
    def _parse_sync(cls, path: Path) -> Presentation:
        """Синхронно извлекает структуру презентации"""
        prs = PPTXPresentation(str(path))

        theme = cls._extract_theme(prs)

        layouts: list[LayoutInfo] = []
        layout_indices_by_partname: dict[str, int] = {}
        for idx, layout in enumerate(prs.slide_layouts, start=1):
            layout_indices_by_partname[str(layout.part.partname)] = idx
            layouts.append(cls._parse_layout(layout, idx, theme))

        slides: list[Slide] = []
        for idx, slide in enumerate(prs.slides, start=1):
            slides.append(
                cls._parse_slide(
                    slide,
                    idx,
                    theme,
                    layout_indices_by_partname,
                )
            )

        return Presentation(
            source_path=str(path),
            file_type="pptx",
            slide_width=prs.slide_width,
            slide_height=prs.slide_height,
            slides=slides,
            layouts=layouts,
            theme=theme,
        )

    @classmethod
    def _parse_slide(
        cls,
        slide,
        index: int,
        theme: ThemeInfo | None,
        layout_indices_by_partname: dict[str, int],
    ) -> Slide:
        """Извлекает данные одного слайда"""
        elements: list[SlideElement] = []
        for shape in slide.shapes:
            element_id = f"slide-{index}-shape-{shape.shape_id}"
            elem = cls._parse_shape(shape, theme, element_id, slide)
            if elem:
                elements.append(elem)

        layout_index = None
        if slide.slide_layout:
            layout_partname = str(slide.slide_layout.part.partname)
            layout_index = layout_indices_by_partname.get(layout_partname)

        background = cls._parse_background(slide.background)

        notes = None
        if slide.has_notes_slide:
            notes = (
                slide.notes_slide.notes_text_frame.text
                if slide.notes_slide.notes_text_frame
                else None
            )

        return Slide(
            index=index,
            layout_type=cls._classify_layout(slide),
            layout_name=slide.slide_layout.name if slide.slide_layout else None,
            layout_index=layout_index,
            placeholder_type=cls._get_slide_placeholder_type(slide),
            elements=elements,
            background=background,
            notes=notes,
        )

    @classmethod
    def _parse_layout(cls, layout, index: int, theme: ThemeInfo | None) -> LayoutInfo:
        """Извлекает информацию о макете"""
        elements: list[SlideElement] = []
        placeholders: list[PlaceholderInfo] = []

        for shape in layout.shapes:
            element_id = f"layout-{index}-shape-{shape.shape_id}"
            elem = cls._parse_shape(shape, theme, element_id, layout)
            if elem:
                elements.append(elem)
                if shape.is_placeholder:
                    placeholders.append(
                        PlaceholderInfo(
                            placeholder_type=shape.placeholder_format.type.name,
                            name=shape.name,
                            idx=shape.placeholder_format.idx,
                            bbox=BBox(
                                left=shape.left,
                                top=shape.top,
                                width=shape.width,
                                height=shape.height,
                            ),
                            element=elem,
                        )
                    )

        background = cls._parse_background(layout.background)

        return LayoutInfo(
            name=layout.name,
            index=index,
            elements=elements,
            placeholders=placeholders,
            background=background,
        )

    @classmethod
    def _parse_shape(
        cls,
        shape,
        theme: ThemeInfo | None,
        element_id: str,
        owner: PPTXSlide | PPTXSlideLayout,
    ) -> SlideElement | None:
        """Извлекает данные из одной фигуры слайда"""
        bbox = BBox(
            left=shape.left,
            top=shape.top,
            width=shape.width,
            height=shape.height,
        )

        # Порядок как позиция элемента среди соседей в родительском контейнере
        z_order = None
        try:
            parent = shape._element.getparent()
            if parent is not None:
                z_order = list(parent).index(shape._element)
        except Exception:
            logger.debug("Не удалось определить z-order элемента %s", element_id, exc_info=True)

        placeholder_type = None
        placeholder_idx = None
        placeholder_name = None
        if shape.is_placeholder:
            placeholder_type = shape.placeholder_format.type.name
            placeholder_idx = shape.placeholder_format.idx
            placeholder_name = shape.name

        if shape.has_table:
            table = cls._parse_table(shape.table)
            return SlideElement(
                id=element_id,
                type=ElementType.TABLE,
                bbox=bbox,
                z_order=z_order,
                placeholder_type=placeholder_type,
                placeholder_idx=placeholder_idx,
                placeholder_name=placeholder_name,
                table=table,
            )

        if shape.has_text_frame:
            text = cls._parse_text_frame(
                shape.text_frame, placeholder_type, theme, shape, owner
            )
            return SlideElement(
                id=element_id,
                type=ElementType.TEXT,
                bbox=bbox,
                z_order=z_order,
                placeholder_type=placeholder_type,
                placeholder_idx=placeholder_idx,
                placeholder_name=placeholder_name,
                text=text,
            )

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image = ImageElement(
                image_path=getattr(shape.image, "filename", None),
                content_type=shape.image.content_type if hasattr(shape, "image") else None,
                alt_text=getattr(shape, "alt_text", None),
            )
            return SlideElement(
                id=element_id,
                type=ElementType.IMAGE,
                bbox=bbox,
                z_order=z_order,
                placeholder_type=placeholder_type,
                placeholder_idx=placeholder_idx,
                placeholder_name=placeholder_name,
                image=image,
            )

        return SlideElement(
            id=element_id,
            type=ElementType.OTHER,
            bbox=bbox,
            z_order=z_order,
            placeholder_type=placeholder_type,
            placeholder_idx=placeholder_idx,
            placeholder_name=placeholder_name,
        )

    @classmethod
    def _parse_text_frame(
        cls,
        text_frame: TextFrame,
        placeholder_type: str | None,
        theme: ThemeInfo | None,
        shape: BaseShape,
        owner: PPTXSlide | PPTXSlideLayout,
    ) -> TextElement:
        """Извлекает текст и стили из текстовой рамки"""
        paragraphs: list[Paragraph] = []
        full_text_parts: list[str] = []

        for para in text_frame.paragraphs:
            runs: list[Run] = []
            para_text = ""
            inherited_size, inherited_name = inherited_font(shape, para, owner)

            for run in para.runs:
                runs.append(
                    Run(
                        text=run.text,
                        style=cls._extract_style(
                            run.font, para, placeholder_type, theme, inherited_size, inherited_name
                        ),
                    )
                )
                para_text += run.text

            # Определяем наличие маркера
            bullet = False
            pPr = para._p.find(qn('a:pPr'))
            if pPr is not None:
                if (
                    pPr.find(qn('a:buChar')) is not None
                    or pPr.find(qn('a:buAutoNum')) is not None
                ):
                    bullet = True

            paragraph = Paragraph(
                text=para_text,
                level=para.level if para.level is not None else 0,
                bullet=bullet,
                runs=runs,
            )

            paragraphs.append(paragraph)
            full_text_parts.append(para_text)

        return TextElement(
            paragraphs=paragraphs,
            full_text="\n".join(full_text_parts),
        )

    @classmethod
    def _extract_style(
        cls,
        font,
        paragraph,
        placeholder_type: str | None,
        theme: ThemeInfo | None,
        inherited_size: float | None = None,
        inherited_name: str | None = None,
    ) -> TextStyle:
        """Извлекает стиль текста с учётом темы"""
        # Шрифт
        font_name = font.name or inherited_name
        if font_name is None:
            if placeholder_type in ("TITLE", "CENTER_TITLE", "SUBTITLE", "SECTION_HEADER"):
                font_name = theme.fonts.get("major") if theme else None
            else:
                font_name = theme.fonts.get("minor") if theme else None
        elif font_name.startswith("+mj"):
            font_name = theme.fonts.get("major", font_name) if theme else font_name
        elif font_name.startswith("+mn"):
            font_name = theme.fonts.get("minor", font_name) if theme else font_name

        # Цвет
        color_hex = None
        try:
            if font.color and font.color.rgb:
                color_hex = str(font.color.rgb)
            elif font.color and font.color.type == MSO_COLOR_TYPE.SCHEME:
                theme_color = font.color.theme_color
                if theme_color is not None:
                    # Пробуем несколько вариантов ключа
                    key = theme_color.name.lower().replace("_", "")
                    color_hex = theme.colors.get(key) if theme else None
                    if color_hex is None and theme:
                        # Альтернативные ключи: dk1 -> dark1, lt1 -> light1 и т.д.
                        alt_map = {
                            "dark1": "dk1",
                            "light1": "lt1",
                            "dark2": "dk2",
                            "light2": "lt2",
                            "hyperlink": "hlink",
                            "followedhyperlink": "folHlink",
                        }
                        alt_key = alt_map.get(key)
                        if alt_key:
                            color_hex = theme.colors.get(alt_key)
        except Exception:
            logger.debug("Не удалось извлечь цвет текста", exc_info=True)

        # Жирность, курсив, подчёркивание
        bold = bool(font.bold)
        italic = bool(font.italic)
        underline = bool(font.underline)

        # Размер
        size_pt = font.size.pt if font.size else inherited_size

        # Выравнивание из paragraph
        alignment = None
        if paragraph.alignment is not None:
            alignment_map = {
                PP_ALIGN.LEFT: "left",
                PP_ALIGN.CENTER: "center",
                PP_ALIGN.RIGHT: "right",
                PP_ALIGN.JUSTIFY: "justify",
            }
            alignment = alignment_map.get(paragraph.alignment, None)

        line_spacing = paragraph.line_spacing if paragraph.line_spacing else None

        return TextStyle(
            font_name=font_name,
            size_pt=size_pt,
            bold=bold,
            italic=italic,
            underline=underline,
            color_hex=color_hex,
            alignment=alignment,
            line_spacing=line_spacing,
        )

    @classmethod
    def _parse_table(cls, table) -> TableElement:
        """Извлекает данные таблицы"""
        rows = len(table.rows)
        cols = len(table.columns)

        cells: list[list[str]] = []
        for row in table.rows:
            row_cells: list[str] = []
            for cell in row.cells:
                row_cells.append(cell.text)
            cells.append(row_cells)

        return TableElement(rows=rows, cols=cols, cells=cells)

    @classmethod
    def _parse_background(cls, background) -> SlideBackground | None:
        """Извлекает фон слайда или макета"""
        if background is None:
            return None

        fill = background.fill
        fill_type = str(fill.type) if fill.type is not None else None
        color_hex = None
        image_path = None

        try:
            if fill.type == 1:  # MSO_FILL_TYPE.SOLID
                if fill.fore_color and fill.fore_color.rgb:
                    color_hex = str(fill.fore_color.rgb)
            elif fill.type == 6:  # MSO_FILL_TYPE.PICTURE
                # Получение картинки фона не реализовано
                pass
        except Exception:
            logger.debug("Не удалось извлечь фон", exc_info=True)

        if color_hex or image_path or fill_type:
            return SlideBackground(
                fill_type=fill_type,
                color_hex=color_hex,
                image_path=image_path,
            )
        return None

    @classmethod
    def _extract_theme(cls, prs) -> ThemeInfo | None:
        """Извлекает тему презентации (цвета и шрифты)"""
        try:
            # Пробуем получить тему через презентацию (надёжнее)
            theme_part = None
            try:
                theme_part = prs.part.part_related_by(RT.THEME)
            except KeyError:
                pass

            if theme_part is None:
                for slide_master in prs.slide_masters:
                    try:
                        theme_part = slide_master.part.part_related_by(RT.THEME)
                        break
                    except KeyError:
                        continue
            if theme_part is None:
                return None

            theme_element = etree.fromstring(theme_part.blob)
            nsmap = {
                "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
            }

            colors: dict[str, str] = {}
            clrScheme = theme_element.find(".//a:clrScheme", namespaces=nsmap)
            if clrScheme is not None:
                for color_entry in clrScheme:
                    localname = etree.QName(color_entry).localname
                    srgb = color_entry.find("a:srgbClr", namespaces=nsmap)
                    if srgb is not None and srgb.get("val"):
                        hex_val = srgb.get("val")
                    else:
                        sys_clr = color_entry.find("a:sysClr", namespaces=nsmap)
                        if sys_clr is not None:
                            hex_val = sys_clr.get("lastClr", sys_clr.get("val"))
                        else:
                            hex_val = None
                    if hex_val:
                        colors[localname] = hex_val
                        # Добавляем нормализованное имя
                        norm_name = localname.lower().replace("_", "")
                        colors[norm_name] = hex_val

            fonts: dict[str, str] = {}
            fontScheme = theme_element.find(".//a:fontScheme", namespaces=nsmap)
            if fontScheme is not None:
                major_font = fontScheme.find("a:majorFont/a:latin", namespaces=nsmap)
                minor_font = fontScheme.find("a:minorFont/a:latin", namespaces=nsmap)
                if major_font is not None and major_font.get("typeface"):
                    fonts["major"] = major_font.get("typeface")
                if minor_font is not None and minor_font.get("typeface"):
                    fonts["minor"] = minor_font.get("typeface")

            if colors or fonts:
                return ThemeInfo(colors=colors, fonts=fonts)
        except Exception:
            logger.debug("Не удалось извлечь тему презентации", exc_info=True)
        return None

    @classmethod
    def _classify_layout(cls, slide) -> LayoutType:
        """Классифицирует макет слайда по имени layout или placeholder"""
        layout_name = slide.slide_layout.name.lower() if slide.slide_layout else ""

        if "title slide" in layout_name or "титульный" in layout_name:
            return LayoutType.TITLE
        if "section header" in layout_name or "заголовок раздела" in layout_name:
            return LayoutType.SECTION_HEADER
        if "two content" in layout_name or "два содержимого" in layout_name:
            return LayoutType.TWO_CONTENT
        if "picture" in layout_name or "изображение" in layout_name:
            return LayoutType.IMAGE_TEXT
        if "table" in layout_name or "таблица" in layout_name:
            return LayoutType.TABLE
        if "thank you" in layout_name or "благодарность" in layout_name:
            return LayoutType.THANK_YOU
        if "title and content" in layout_name or "заголовок и содержимое" in layout_name:
            return LayoutType.BULLETS
        if "title only" in layout_name or "только заголовок" in layout_name:
            return LayoutType.TITLE

        placeholders = [sh for sh in slide.shapes if sh.is_placeholder]
        types = [ph.placeholder_format.type.name for ph in placeholders]
        if "TITLE" in types and ("BODY" in types or "CONTENT" in types):
            return LayoutType.BULLETS
        if "TITLE" in types and len(types) == 1:
            return LayoutType.TITLE
        if "PICTURE" in types:
            return LayoutType.IMAGE_TEXT

        return LayoutType.UNKNOWN

    @classmethod
    def _get_slide_placeholder_type(cls, slide) -> str | None:
        """Возвращает тип основного placeholder слайда"""
        for shape in slide.shapes:
            if shape.is_placeholder:
                return shape.placeholder_format.type.name
        return None
