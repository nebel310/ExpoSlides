"""Разрешение наследуемых шрифтов без изменения XML презентации."""

from collections.abc import Iterator

from lxml.etree import _Element
from pptx.oxml.ns import qn
from pptx.shapes.base import BaseShape
from pptx.slide import Slide, SlideLayout
from pptx.text.text import _Paragraph

_BODY_TYPES = {
    "BODY", "SUBTITLE", "OBJECT", "CHART", "TABLE", "BITMAP",
    "ORG_CHART", "MEDIA_CLIP", "PICTURE",
}
_TITLE_TYPES = {"TITLE", "CENTER_TITLE"}


def _paragraph_defaults(paragraph: _Paragraph) -> Iterator[_Element]:
    properties = paragraph._p.find(qn("a:pPr"))
    if properties is not None:
        defaults = properties.find(qn("a:defRPr"))
        if defaults is not None:
            yield defaults


def _list_defaults(style: _Element | None, level: int) -> Iterator[_Element]:
    if style is None:
        return
    for name in (f"a:lvl{level + 1}pPr", "a:defPPr"):
        properties = style.find(qn(name))
        if properties is not None:
            defaults = properties.find(qn("a:defRPr"))
            if defaults is not None:
                yield defaults


def _shape_defaults(shape: BaseShape, level: int) -> Iterator[_Element]:
    if not shape.has_text_frame:
        return
    # Унаследованный абзац выбирается по уровню, а не по его позиции в образце.
    for paragraph in shape.text_frame.paragraphs:
        if paragraph.level == level:
            yield from _paragraph_defaults(paragraph)
            break
    text_body = shape._element.find(qn("p:txBody"))
    if text_body is not None:
        yield from _list_defaults(text_body.find(qn("a:lstStyle")), level)


def _font_sources(
    shape: BaseShape, paragraph: _Paragraph, owner: Slide | SlideLayout
) -> Iterator[_Element]:
    level = paragraph.level
    yield from _paragraph_defaults(paragraph)
    text_body = shape._element.find(qn("p:txBody"))
    if text_body is not None:
        yield from _list_defaults(text_body.find(qn("a:lstStyle")), level)

    master = owner.slide_layout.slide_master if isinstance(owner, Slide) else owner.slide_master
    placeholder_type = shape.placeholder_format.type.name if shape.is_placeholder else None
    if shape.is_placeholder:
        if isinstance(owner, Slide):
            # На слайде связь с макетом задаёт idx, тип может отличаться.
            layout_shape = owner.slide_layout.placeholders.get(idx=shape.placeholder_format.idx)
            if layout_shape is not None:
                yield from _shape_defaults(layout_shape, level)
                placeholder_type = layout_shape.placeholder_format.type.name

        # У макета связь с образцом задаётся типом, а не idx.
        master_type = placeholder_type
        if placeholder_type in _TITLE_TYPES:
            master_type = "TITLE"
        elif placeholder_type in _BODY_TYPES:
            master_type = "BODY"
        for master_shape in master.placeholders:
            if master_shape.placeholder_format.type.name == master_type:
                yield from _shape_defaults(master_shape, level)
                break

    style_name = "p:otherStyle"
    if placeholder_type in _TITLE_TYPES:
        style_name = "p:titleStyle"
    elif placeholder_type in _BODY_TYPES:
        style_name = "p:bodyStyle"
    master_styles = master._element.find(qn("p:txStyles"))
    if master_styles is not None:
        yield from _list_defaults(master_styles.find(qn(style_name)), level)

    presentation = owner.part.package.presentation_part.presentation
    yield from _list_defaults(presentation._element.find(qn("p:defaultTextStyle")), level)


def inherited_font(
    shape: BaseShape, paragraph: _Paragraph, owner: Slide | SlideLayout
) -> tuple[float | None, str | None]:
    """Возвращает размер и гарнитуру; явные свойства run применяет вызывающий код."""
    size_pt = None
    font_name = None
    for properties in _font_sources(shape, paragraph, owner):
        if size_pt is None and properties.get("sz") is not None:
            size_pt = int(properties.get("sz")) / 100
        if font_name is None:
            latin = properties.find(qn("a:latin"))
            if latin is not None:
                font_name = latin.get("typeface") or None
        if size_pt is not None and font_name is not None:
            break
    return size_pt, font_name
