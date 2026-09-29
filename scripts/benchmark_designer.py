"""Офлайн benchmark реального PPTX → PDF/HTML pipeline на синтетических шаблонах.

Запуск из корня: uv run python -m scripts.benchmark_designer --output-dir /new/path
Это проверка реализации, а не качества LLM или предоставленных организатором шаблонов.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import platform
import time
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

from exposlides.design_content import source_excerpts, validate_story
from exposlides.design_models import AuditReport, Dataset, DeckPlan, DesignRequest, VisualRequest
from exposlides.design_native_text import TEXT_KINDS
from exposlides.design_pipeline import DesignPipeline, save_model

CASES = (
    ("ordinary_textboxes", 12),
    ("dark", 12),
    ("grouped_protected", 12),
    ("synthetic_holdout", 10),
)


def implementation_hashes() -> dict[str, str]:
    root = Path(__file__).resolve().parents[1]
    paths = [
        Path(__file__).resolve(), root / "uv.lock",
        *sorted((root / "exposlides").glob("*.py")),
        *sorted((root / "services/parsing-service/app").rglob("*.py")),
        root / "services/content-service/app/utils/fact_grounding.py",
    ]
    return {str(path.relative_to(root)): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in paths}


SECTIONS = (
    "Результаты пилота измерены. Выручка базового периода составила 12 млн рублей. "
    "Выручка пилотного периода составила 18 млн рублей.",
    "План внедрения согласован. Анализ потребностей. Настройка сервиса. Проверка результата.",
    "Ответственность команды определена. Команда отвечает за настройку. "
    "Заказчик принимает результат.",
    "Поддержка помогает пользователям. Обращения регистрируются в едином канале. "
    "Ответственный сопровождает запрос до решения.",
    "Обучение встроено в запуск. Практические занятия объясняют рабочие сценарии. "
    "Материалы доступны сотрудникам после обучения.",
    "Качество проверяется по источникам. Числа сверяются с исходными данными. "
    "Редактор подтверждает формулировки перед публикацией.",
    "Доступ регулируется ролями. Участник видит согласованные материалы. "
    "Администратор контролирует изменение прав.",
    "Обратная связь улучшает продукт. Команда собирает замечания участников. "
    "Приоритет исправления зависит от влияния на работу.",
    "Надёжность требует наблюдения. Ответственные контролируют выполнение этапов. "
    "Сбой фиксируется с понятным описанием причины.",
    "Следующий шаг — запустить пилот. Команда выбирает участников. "
    "Заказчик подтверждает критерии приёмки.",
    "Риски обсуждаются заранее. Ограничение: качество зависит от полноты материалов. "
    "Команда запрашивает недостающие данные до сборки.",
    "Результат готов к передаче. Редактируемые материалы получает заказчик. "
    "Ответственный сохраняет согласованную версию.",
)


def _textbox(shapes, text, bounds, font, size, color, name):
    box = shapes.add_textbox(*(Inches(value) for value in bounds))
    box.name = name
    box.text = text
    box.text_frame.word_wrap = True
    for paragraph in box.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor.from_string(color)
    return box


def make_template(path: Path, case: str) -> None:
    """Проверяем обычные textbox, тёмную тему, группы/защиту и иной формат кадра."""
    heldout = case == "synthetic_holdout"
    dark = case == "dark"
    width, height = (10, 7.5) if heldout else (13.333, 7.5)
    font = "Georgia" if heldout else "Arial"
    background = "16243A" if dark else ("FFF9ED" if heldout else "FFFFFF")
    foreground = "FFFFFF" if dark else "202124"
    accent = "45C7EA" if dark else ("BB5427" if heldout else "0077FF")
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(width), Inches(height)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(background)
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12 if heldout else width),
        Inches(height if heldout else 0.08),
    )
    band.name = "brand-decoration"
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor.from_string(accent)
    band.line.fill.background()
    _textbox(slide.shapes, "Образец заголовка", (0.6, 0.42, width-1.8, 0.82),
             font, 30 if heldout else 32, foreground, "Editorial title")
    body_shapes = (slide.shapes.add_group_shape().shapes
                   if case == "grouped_protected" else slide.shapes)
    _textbox(body_shapes, "Основное содержание шаблона", (0.6, 1.6, width-1.2, 4.95),
             font, 20, foreground, "Editable body")
    footer = _textbox(body_shapes, "ExpoSlides · synthetic benchmark",
                      (0.6, 7.0, width-1.8, 0.25), font, 10, foreground, "copyright footer")
    footer.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.com/benchmark"
    _textbox(slide.shapes, "01", (width-0.9, 7.0, 0.4, 0.25),
             font, 10, foreground, "Slide number")
    if case == "grouped_protected":
        picture = io.BytesIO()
        Image.new("RGB", (60, 30), "#"+accent).save(picture, format="PNG")
        picture.seek(0)
        logo = slide.shapes.add_picture(picture, Inches(width-1.2), Inches(0.45), width=Inches(0.6))
        logo.name = "protected logo"
    presentation.save(path)


def make_request(slide_count: int) -> DesignRequest:
    return DesignRequest(
        script="\n\n".join(SECTIONS[:slide_count]), slide_count=slide_count,
        mode="extractive", contextual_audit=False,
        purpose="Проверить перенос содержания, редактируемость и разные композиции",
        datasets=[Dataset(
            id="revenue", name="Выручка", columns=["Период", "Выручка"],
            rows=[["Базовый", 12], ["Пилот", 18]], unit="млн рублей",
            source="Первый раздел синтетического brief; данные вымышлены для проверки",
        )],
    )


def _walk(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from _walk(shape.shapes)


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def verify_variant(directory: Path, story, request: DesignRequest) -> dict:
    plan = DeckPlan.model_validate_json((directory / "plan.json").read_text(encoding="utf-8"))
    audit = AuditReport.model_validate_json((directory / "audit.json").read_text(encoding="utf-8"))
    exports = json.loads((directory / "exports.json").read_text(encoding="utf-8"))
    pptx = directory / "presentation.pptx"
    presentation = Presentation(pptx)
    _require(len(presentation.slides) == request.slide_count, "Неверное число слайдов PPTX")
    _require(plan.datasets == request.datasets, "Изменились исходные наборы данных")
    counts = Counter()
    for source, instance, slide in zip(story.slides, plan.slides, presentation.slides, strict=True):
        _require(source.id == instance.story_slide_id, "Изменился порядок истории")
        shapes = list(_walk(slide.shapes))
        named = {shape.name: shape for shape in shapes}
        by_id = {shape.shape_id: shape for shape in shapes}
        actual_paragraphs = []
        for block in instance.blocks:
            shape = (by_id.get(block.source_shape_id)
                     if block.kind in TEXT_KINDS
                     and block.source_shape_id is not None
                     else named.get(f"exposlides:{block.id}"))
            _require(shape is not None, f"Потерян нативный объект {block.id}")
            if block.kind in {"title", "text"}:
                _require(shape.has_text_frame, f"Текст {block.id} не редактируется")
                texts = [p.text for p in shape.text_frame.paragraphs]
                if block.kind == "title":
                    _require(shape.text == source.title, "Изменился заголовок")
                else:
                    actual_paragraphs.extend(texts)
                counts["text"] += 1
            elif block.kind == "table":
                dataset = next(d for d in request.datasets if d.id == block.dataset_id)
                _require(shape.has_table, "Таблица не является нативной")
                expected = [dataset.columns, *[[str(value) for value in row] for row in dataset.rows]]
                actual = [[cell.text for cell in row.cells] for row in shape.table.rows]
                _require(actual == expected, "Данные нативной таблицы изменились")
                counts["table"] += 1
            elif block.kind == "chart":
                dataset = next(d for d in request.datasets if d.id == block.dataset_id)
                _require(shape.has_chart, "График не является нативным")
                _require(list(shape.chart.series[0].values) == [row[1] for row in dataset.rows],
                         "Числа графика изменились")
                _require(shape.chart.part.part_related_by(RT.PACKAGE).blob.startswith(b"PK"),
                         "Отсутствует редактируемая книга графика")
                counts["chart"] += 1
            elif block.kind == "process":
                _require(hasattr(shape, "shapes"), "Схема не является группой фигур")
                labels = [item.text for item in shape.shapes if item.has_text_frame]
                _require(labels == block.items, "Изменились подписи схемы")
                counts["process"] += 1
        _require(actual_paragraphs == source.paragraphs, "Изменились абзацы истории")
    with ZipFile(pptx) as archive:
        _require(archive.testzip() is None, "Повреждён ZIP-контейнер PPTX")
        _require(len(archive.namelist()) == len(set(archive.namelist())), "Дубли частей OOXML")
    _require(counts["table"] > 0 and counts["process"] > 0, "Нет таблицы или схемы")
    _require(plan.variant_id == "evidence" or counts["chart"] > 0, "Нет нативного графика")
    for filename in ("presentation.pdf", "presentation.html"):
        _require((directory / filename).is_file() and (directory / filename).stat().st_size > 0,
                 f"Не создан {filename}")
    _require((directory / "presentation.pdf").read_bytes().startswith(b"%PDF"), "Некорректный PDF")
    images = [directory / name for name in exports["images"]]
    _require(len(images) == request.slide_count, "Неверное число рендеров")
    hashes = [hashlib.sha256(image.read_bytes()).hexdigest() for image in images]
    _require(audit.contextual_status == "not_run", "В офлайн benchmark запущен модельный аудит")
    return {
        "id": plan.variant_id, "slide_count": len(presentation.slides),
        "native_objects": dict(counts), "png_sha256": hashes,
        "audit": {"errors": sum(i.severity == "error" for i in audit.issues),
                  "warnings": sum(i.severity == "warning" for i in audit.issues),
                  "rules": dict(Counter(i.rule for i in audit.issues)),
                  "contextual_status": audit.contextual_status},
        "html_visual": exports["html_visual"], "artifact_directory": str(directory),
    }


def distinct_deck_count(variants: list[dict]) -> int:
    """ТЗ требует разные колоды; общие обложки и текстовые слайды допустимы."""
    return len({tuple(variant["png_sha256"]) for variant in variants})


def run_case(output: Path, case: str, slide_count: int) -> dict:
    directory = output / case
    directory.mkdir()
    template = directory / "template.pptx"
    make_template(template, case)
    template_hash = hashlib.sha256(template.read_bytes()).hexdigest()
    request = make_request(slide_count)
    (directory / "script.txt").write_text(request.script, encoding="utf-8")
    pipeline = DesignPipeline(directory / "job", timeout=300)
    started = time.monotonic()
    try:
        profile, story = pipeline.plan(template, request)
        story.slides[0].visual = VisualRequest(kind="bar", dataset_id="revenue")
        story.slides[1].visual = VisualRequest(
            kind="process", labels=["Анализ потребностей", "Настройка сервиса", "Проверка результата"],
        )
        story.slides[2].visual = VisualRequest(kind="table", dataset_id="revenue")
        _require(not validate_story(story, request, source_excerpts(request.script)),
                 "Синтетическая история не прошла проверку")
        variants = pipeline.build(template, request, profile, story)
        pipeline_seconds = time.monotonic() - started
    finally:
        pipeline.close()
    _require(hashlib.sha256(template.read_bytes()).hexdigest() == template_hash,
             "Исходный шаблон изменился")
    _require(len(variants) == 3, "Не создано три варианта")
    checked = [verify_variant(directory / "job/variants" / variant["id"] / "1", story, request)
               for variant in variants]
    distinct = [len({variant["png_sha256"][index] for variant in checked}) == 3
                for index in range(slide_count)]
    _require(pipeline_seconds <= 300, "Превышен лимит пяти минут")
    distinct_decks = distinct_deck_count(checked)
    result = {
        "case": case, "status": "passed" if distinct_decks == 3 else "failed",
        "distinct_deck_count": distinct_decks, "slide_count": slide_count,
        "pipeline_seconds": round(pipeline_seconds, 3),
        "template_sha256": template_hash, "distinct_variants_per_slide": distinct,
        "variants": checked,
        "claim_scope": "synthetic extractive pipeline; no live model or official template evaluation",
    }
    if distinct_decks != 3:
        result["error"] = "Не получены три визуально различные колоды"
    save_model(directory / "benchmark-result.json", result)
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args(argv)
    output = args.output_dir.expanduser().resolve()
    if output.exists() and (not output.is_dir() or any(output.iterdir())):
        parser.error("Каталог результата должен быть новым или пустым")
    output.mkdir(parents=True, exist_ok=True)
    results = []
    source_hashes = implementation_hashes()
    started = time.monotonic()
    for case, count in CASES:
        print(f"Проверка {case}: {count} слайдов × 3 варианта", flush=True)
        try:
            result = run_case(output, case, count)
        except Exception as error:
            result = {"case": case, "status": "failed", "error": str(error)}
        results.append(result)
        print(json.dumps({key: result[key] for key in ("case", "status", "pipeline_seconds", "error")
                          if key in result}, ensure_ascii=False), flush=True)
    report = {
        "schema_version": "1.0", "python": platform.python_version(),
        "platform": platform.platform(), "mode": "extractive", "api_calls_requested": False,
        "implementation_sha256": source_hashes,
        "implementation_changed_during_run": source_hashes != implementation_hashes(),
        "seconds": round(time.monotonic()-started, 3), "cases": results,
        "limitations": [
            "Все шаблоны и данные синтетические; это не шаблоны организатора.",
            "Нет оценки живой Qwen, VK endpoint, VLM или генератора изображений.",
            "Разные PNG подтверждают различие рендеров, но не дизайнерское качество.",
            "Audit warnings/errors сохраняются для разбора; passed означает структурные проверки.",
        ],
    }
    save_model(output / "benchmark-report.json", report)
    return int(any(result["status"] != "passed" for result in results))


if __name__ == "__main__":
    raise SystemExit(main())
