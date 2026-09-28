from __future__ import annotations

import asyncio
import importlib
import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Pt

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
PARSING_SERVICE_ROOT = REPOSITORY_ROOT / "services" / "parsing-service"
CONTENT_SERVICE_ROOT = REPOSITORY_ROOT / "services" / "content-service"
BUILDER_SERVICE_ROOT = REPOSITORY_ROOT / "services" / "builder-service"
SOURCE = "Команда готовит запуск продукта и описывает стратегию."


def write_template(path: Path) -> None:
    presentation = Presentation()
    for title in ("Первый исходный заголовок презентации", "Второй исходный заголовок"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = (
            "Исходный текст достаточно длинный для безопасной оценки нового содержимого"
        )
        run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
        run.font.name = "Arial"
        run.font.size = Pt(28)
        run.font.bold = True
    presentation.save(path)


class FastFakeLLM:
    """Два детерминированных ответа для настоящего графа без сетевых запросов."""

    def __init__(self, numbered_plan: bool = False) -> None:
        self.calls: list[str] = []
        self.numbered_plan = numbered_plan

    async def generate_json(self, prompt, model, strict=True):
        self.calls.append("outline")
        return model.model_validate({
            "analysis": {
                "topic": "Запуск продукта",
                "audience": "",
                "objective": "",
                "blocks": [{
                    "index": 1,
                    "heading": "Подготовка запуска",
                    "summary": SOURCE,
                    "key_points": ["Команда готовит запуск продукта"],
                    "facts": [],
                }],
                "key_messages": ["Команда готовит запуск продукта"],
                "facts": [],
            },
            "plan": {"slides": [{
                "template_slide_index": index,
                "layout_type": "bullets",
                "title": title,
                "content": (
                    "1. Команда готовит запуск продукта. 2. Команда описывает стратегию."
                    if self.numbered_plan else "Команда готовит запуск продукта"
                ),
                "purpose": "Описать стратегию",
                "key_message": "Команда готовит запуск продукта",
                "source_block_indices": [1],
            } for index, title in ((2, "Стратегия запуска"), (1, "Команда запуска"))]},
        })

    async def generate_json_object(self, prompt, schema, strict=True, *, model=None):
        self.calls.append("content")
        assert schema["required"] == [f"field_{index:04d}" for index in range(4)]
        return dict(zip(schema["required"], (
            "Стратегия запуска",
            "Команда готовит запуск продукта",
            "Команда запуска",
            "Стратегия продукта описана командой",
        ), strict=True))


@pytest.fixture
def parsed_inputs(tmp_path, service_importer):
    template = tmp_path / "template.pptx"
    parsed = tmp_path / "template.json"
    script = tmp_path / "script.txt"
    write_template(template)
    script.write_text(SOURCE, encoding="utf-8")
    parser = service_importer(PARSING_SERVICE_ROOT, "app.main")
    asyncio.run(parser.run(template, parsed))
    return template, parsed, script


@pytest.mark.parametrize("numbered_plan", [False, True])
def test_fast_pipeline_preserves_order_editability_and_fonts_with_two_requests(
    parsed_inputs, tmp_path, monkeypatch, service_importer, numbered_plan,
) -> None:
    template, parsed, script = parsed_inputs
    content = tmp_path / "generated_content.json"
    output = tmp_path / "result.pptx"
    content_main = service_importer(CONTENT_SERVICE_ROOT, "app.main")
    fast = importlib.import_module("app.graph.fast")
    request = importlib.import_module("app.models.request")
    llm = FastFakeLLM(numbered_plan=numbered_plan)
    monkeypatch.setattr(fast, "fast_llm_client", llm)
    asyncio.run(content_main.run(
        parsed, script, content,
        request.GenerationSettings(max_slides=None, generation_mode="fast"),
    ))
    assert llm.calls == ["outline", "content"]

    builder = service_importer(BUILDER_SERVICE_ROOT, "app.main")
    asyncio.run(builder.run(template, parsed, content, output))
    generated = json.loads(content.read_text(encoding="utf-8"))
    assert list(generated["content"]) == ["2", "1"]
    assert generated["validation_report"] == {"ok": True, "issues": []}
    result = Presentation(output)
    assert len(result.slides) == 2
    assert [slide.shapes.title.text for slide in result.slides] == [
        "Стратегия запуска", "Команда запуска",
    ]
    assert [slide.placeholders[1].text for slide in result.slides] == [
        "Команда готовит запуск продукта", "Стратегия продукта описана командой",
    ]
    for slide in result.slides:
        run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
        assert run.font.name == "Arial"
        assert run.font.size == Pt(28)
        assert run.font.bold is True


@pytest.mark.parametrize("blocked_phase", ["outline", "content"])
def test_fast_deadline_cancels_awaited_request_without_writing_partial_content(
    parsed_inputs, tmp_path, monkeypatch, service_importer, blocked_phase,
) -> None:
    _, parsed, script = parsed_inputs
    content_main = service_importer(CONTENT_SERVICE_ROOT, "app.main")
    fast = importlib.import_module("app.graph.fast")
    graph = importlib.import_module("app.graph.builder")
    request = importlib.import_module("app.models.request")
    errors = importlib.import_module("app.errors")

    class SlowLLM(FastFakeLLM):
        started = False
        cancelled = False

        async def wait_until_cancelled(self):
            self.started = True
            try:
                await asyncio.Future()
            except asyncio.CancelledError:
                self.cancelled = True
                raise

        async def generate_json(self, *args, **kwargs):
            if blocked_phase == "outline":
                await self.wait_until_cancelled()
            return await super().generate_json(*args, **kwargs)

        async def generate_json_object(self, *args, **kwargs):
            if blocked_phase == "content":
                await self.wait_until_cancelled()
            return await super().generate_json_object(*args, **kwargs)

    llm = SlowLLM()
    monkeypatch.setattr(fast, "fast_llm_client", llm)
    monkeypatch.setattr(graph.settings, "fast_generation_timeout", 0.02)
    content = tmp_path / "generated_content.json"
    with pytest.raises(errors.ContentValidationError, match="не успела"):
        asyncio.run(content_main.run(
            parsed, script, content,
            request.GenerationSettings(max_slides=2, generation_mode="fast"),
        ))
    assert llm.started
    assert llm.cancelled
    assert llm.calls == ([] if blocked_phase == "outline" else ["outline"])
    assert not content.exists()
    assert not list(tmp_path.glob(".generated_content.json.*.tmp"))


def test_content_cli_forwards_fast_mode_to_generation_settings(
    tmp_path, monkeypatch, service_importer,
) -> None:
    content_main = service_importer(CONTENT_SERVICE_ROOT, "app.main")
    captured = {}

    async def fake_run(template, script, output, generation_settings, user_mapping):
        captured["settings"] = generation_settings
        captured["output"] = output
        return output

    monkeypatch.setattr(content_main, "run", fake_run)
    monkeypatch.setattr(content_main, "setup_logging", lambda: None)
    output = tmp_path / "generated_content.json"
    assert content_main.main([
        "--cli", "--generation-mode", "fast", "--max-slides", "2", "--output-json", str(output),
    ]) == 0
    assert captured["settings"].generation_mode == "fast"
    assert captured["settings"].max_slides == 2
    assert captured["output"] == output
    assert content_main._parse_args(["--cli"]).generation_mode == "standard"
