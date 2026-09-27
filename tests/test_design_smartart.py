"""Регрессии OOXML dataModel/drawing: подписи и независимые части SmartArt.

Синтетический нативный OOXML-образец проверяет контракт и связи; визуальное
качество реальных SmartArt-макетов требует отдельного PowerPoint-прогона.
"""

from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches

from exposlides.design_builder import build_deck
from exposlides.design_models import Box, DeckPlan, PlacedBlock, SlideInstance, TextStyle
from exposlides.design_pptx_parts import NativeBuildError
from exposlides.design_smartart import DGM, DSP

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"


def _smartart_template(path: Path) -> int:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    package = prs.part.package
    data = etree.Element(f"{{{DGM}}}dataModel", nsmap={"dgm": DGM, "a": A})
    points = etree.SubElement(data, f"{{{DGM}}}ptLst")
    drawing = etree.Element(f"{{{DSP}}}drawing", nsmap={"dsp": DSP, "a": A})
    shapes = etree.SubElement(drawing, f"{{{DSP}}}spTree")
    for index in range(2):
        node = etree.SubElement(points, f"{{{DGM}}}pt", modelId=f"node-{index}", type="node")
        body = etree.SubElement(node, f"{{{DGM}}}t")
        paragraph = etree.SubElement(body, f"{{{A}}}p")
        run = etree.SubElement(paragraph, f"{{{A}}}r")
        etree.SubElement(run, f"{{{A}}}t").text = f"Old {index}"
        shape = etree.SubElement(shapes, f"{{{DSP}}}sp", modelId=f"node-{index}")
        body = etree.SubElement(shape, f"{{{DSP}}}txBody")
        paragraph = etree.SubElement(body, f"{{{A}}}p")
        run = etree.SubElement(paragraph, f"{{{A}}}r")
        etree.SubElement(run, f"{{{A}}}t").text = f"Old {index}"
    data_part = Part(
        PackURI("/ppt/diagrams/data1.xml"),
        "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
        package,
        etree.tostring(data),
    )
    drawing_part = Part(
        PackURI("/ppt/diagrams/drawing1.xml"),
        "application/vnd.ms-office.drawingml.diagramDrawing+xml",
        package,
        etree.tostring(drawing),
    )
    data_part.relate_to(
        drawing_part, "http://schemas.microsoft.com/office/2007/relationships/diagramDrawing"
    )
    rid = slide.part.relate_to(data_part, REL + "diagramData")
    element = parse_xml(f'''<p:graphicFrame xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}" xmlns:dgm="{DGM}">
      <p:nvGraphicFramePr><p:cNvPr id="2" name="Native SmartArt"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>
      <p:xfrm><a:off x="914400" y="914400"/><a:ext cx="4572000" cy="2743200"/></p:xfrm>
      <a:graphic><a:graphicData uri="{DGM}"><dgm:relIds r:dm="{rid}"/></a:graphicData></a:graphic>
    </p:graphicFrame>''')
    slide.shapes._spTree.insert_element_before(element, "p:extLst")
    prs.save(path)
    return 2


def _plan(items, source_id=2):
    return DeckPlan(
        variant_id="cards",
        name="Native",
        description="Native template reuse",
        width=Inches(10),
        height=Inches(7.5),
        slides=[
            SlideInstance(
                id=f"slide-{index}",
                story_slide_id=f"story-{index}",
                source_slide_index=1,
                blocks=[
                    PlacedBlock(
                        id="smartart",
                        kind="smartart",
                        source_shape_id=source_id,
                        box=Box(left=Inches(1), top=Inches(1), width=Inches(5), height=Inches(3)),
                        style=TextStyle(),
                        items=[f"{item} {index}" for item in items],
                    )
                ],
            )
            for index in range(2)
        ],
    )


def test_smartart_labels_update_both_native_data_and_drawing_independently(tmp_path):
    template, result = tmp_path / "source.pptx", tmp_path / "result.pptx"
    _smartart_template(template)
    original = template.read_bytes()
    build_deck(template, _plan(["Discover", "Deliver"]), result)
    reopened = Presentation(result)
    data_parts = []
    for index, slide in enumerate(reopened.slides):
        shape = next(s for s in slide.shapes if s.name == "exposlides:smartart")
        rel_ids = shape._element.find(f".//{{{DGM}}}relIds")
        data_part = slide.part.related_part(rel_ids.get(qn("r:dm")))
        data_parts.append(data_part)
        texts = [node.text for node in etree.fromstring(data_part.blob).iter(qn("a:t"))]
        assert texts == [f"Discover {index}", f"Deliver {index}"]
        drawing = next(iter(data_part.rels.values())).target_part
        assert [node.text for node in etree.fromstring(drawing.blob).iter(qn("a:t"))] == texts
    assert data_parts[0] is not data_parts[1]
    assert template.read_bytes() == original


def test_smartart_requires_compatible_exemplar_and_count(tmp_path):
    template, result = tmp_path / "source.pptx", tmp_path / "result.pptx"
    _smartart_template(template)
    with pytest.raises(NativeBuildError, match="число узлов"):
        build_deck(template, _plan(["One", "Two", "Three"]), result)
    with pytest.raises(NativeBuildError, match="нативный образец"):
        build_deck(template, _plan(["One", "Two"], source_id=999), result)
    assert not result.exists()
