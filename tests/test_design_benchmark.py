from pathlib import Path

import pytest
from pptx import Presentation

from exposlides.design_content import extractive_plan, source_excerpts, validate_story
from scripts.benchmark_designer import CASES, main, make_request, make_template


@pytest.mark.parametrize("case,count", CASES)
def test_benchmark_sources_are_valid_editable_offline_fixtures(tmp_path, case, count):
    path = tmp_path / "template.pptx"
    make_template(path, case)
    presentation = Presentation(path)
    assert len(presentation.slides) == 1
    assert not list(presentation.slides[0].placeholders)
    request = make_request(count)
    excerpts = source_excerpts(request.script)
    story = extractive_plan(request, excerpts)
    assert len(story.slides) == count
    assert not validate_story(story, request, excerpts)
    assert request.mode == "extractive"
    assert not request.contextual_audit
    assert not request.generated_image.enabled


def test_benchmark_never_overwrites_an_existing_artifact(tmp_path):
    marker = tmp_path / "report.json"
    marker.write_text("keep this result", encoding="utf-8")
    with pytest.raises(SystemExit) as error:
        main(["--output-dir", str(tmp_path)])
    assert error.value.code == 2
    assert marker.read_text(encoding="utf-8") == "keep this result"
    assert list(tmp_path.iterdir()) == [Path(marker)]



def test_verifier_accepts_preserved_native_text_and_page_numbers(tmp_path, service_importer):
    import asyncio
    import json

    from exposlides.design_builder import build_deck
    from exposlides.design_layout import create_variants
    from exposlides.design_models import VisualRequest
    from exposlides.design_pipeline import save_model
    from exposlides.template_profile import profile_from_json
    from scripts.benchmark_designer import verify_variant

    template = tmp_path / "template.pptx"
    make_template(template, "ordinary_textboxes")
    request = make_request(3)
    parser = service_importer(
        Path(__file__).resolve().parents[1] / "services/parsing-service",
        "app.parsers.pptx_parser",
    ).PPTXParser
    data = asyncio.run(parser.parse(template)).model_dump(mode="json")
    profile = profile_from_json(data, template)
    story = extractive_plan(request, source_excerpts(request.script))
    story.slides[0].visual = VisualRequest(kind="bar", dataset_id="revenue")
    story.slides[1].visual = VisualRequest(kind="process", labels=["Анализ", "Проверка"])
    story.slides[2].visual = VisualRequest(kind="table", dataset_id="revenue")
    plan = create_variants(profile, story, request.datasets)[0]
    build_deck(template, plan, tmp_path / "presentation.pptx")
    save_model(tmp_path / "plan.json", plan)
    save_model(tmp_path / "audit.json", {"issues": [], "checks": [], "limitations": [],
                                       "contextual_status": "not_run"})
    # Offline renderer stand-ins: native objects and text are read from a real reopened PPTX.
    (tmp_path / "presentation.pdf").write_bytes(b"%PDF-test")
    (tmp_path / "presentation.html").write_text("test", encoding="utf-8")
    images = []
    for index in range(3):
        name = f"slide-{index}.png"
        (tmp_path / name).write_bytes(b"test-preview")
        images.append(name)
    (tmp_path / "exports.json").write_text(
        json.dumps({"images": images, "html_visual": "png"}), encoding="utf-8",
    )
    assert verify_variant(tmp_path, story, request)["native_objects"]["text"] >= 3
    deck = Presentation(tmp_path / "presentation.pptx")
    title = next(b for b in plan.slides[0].blocks if b.kind == "title")
    shape = next(s for s in deck.slides[0].shapes if s.shape_id == title.source_shape_id)
    shape.text = "Повреждённый заголовок"
    deck.save(tmp_path / "presentation.pptx")
    with pytest.raises(ValueError, match="Изменился заголовок"):
        verify_variant(tmp_path, story, request)


def test_visual_difference_is_required_between_decks_not_every_shared_slide():
    from scripts.benchmark_designer import distinct_deck_count

    variants = [{"png_sha256": ["shared-cover", content]} for content in ("a", "b", "c")]
    assert distinct_deck_count(variants) == 3
    assert distinct_deck_count([variants[0], variants[1], variants[0]]) == 2
    assert distinct_deck_count([variants[0]] * 3) == 1
