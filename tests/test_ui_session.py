from __future__ import annotations

import gc
import importlib
import io
import json
import sys
import weakref
from pathlib import Path
from unittest.mock import Mock

import pytest
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
PresentationSession = importlib.import_module("exposlides.ui_session").PresentationSession
web = importlib.import_module("exposlides.web")
APIError, CONTENT_ERRORS, Workspace = web.APIError, web.CONTENT_ERRORS, web.Workspace


class FakeRenderer:
    available = False

    def close(self):
        pass


def make_template(title="Шаблон", slide_count=3):
    presentation = Presentation()
    for index in range(1, slide_count + 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"{title} {index}"
        slide.placeholders[1].text = "Исходный текст"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


@pytest.fixture
def session(monkeypatch):
    workspace = Workspace(renderer=FakeRenderer())
    # Контролируем завершение задания сами; никакие CLI или LLM не вызываются.
    monkeypatch.setattr(workspace, "_run_job", lambda job: None)
    session = PresentationSession(workspace=workspace)
    yield session, workspace
    session.close()


def complete(session, workspace):
    job = workspace.get_job(session.snapshot()["job"]["id"])
    result = make_template("Готовый слайд", slide_count=2)
    (job.directory / "result.pptx").write_bytes(result)
    workspace._complete_job(job)
    return job, result


def publish_preview(workspace, namespace, metadata, image):
    directory = workspace.root / namespace / metadata["id"]
    directory.mkdir(parents=True)
    paths = []
    for index in range(metadata["slide_count"]):
        path = directory / f"{index}.png"
        path.write_bytes(image + bytes([index]))
        paths.append(path)
    with workspace.lock:
        workspace.preview_files[(namespace, metadata["id"])] = tuple(paths)
        metadata["preview"] = {"status": "ready"}


def test_rerun_reuses_upload_and_preserves_running_job(session):
    ui, workspace = session
    data = make_template()
    template = ui.select_template("../../Шаблон.pptx", data)
    job = ui.start("Текст доклада", 3)

    repeated = ui.select_template("Другое имя.pptx", data)

    assert repeated["id"] == template["id"]
    assert repeated["name"] == "Шаблон.pptx"
    assert ui.snapshot()["job"] == job
    assert len(workspace.templates) == len(workspace.jobs) == 1
    assert str(workspace.root) not in json.dumps(ui.snapshot(), ensure_ascii=False)


def test_cannot_change_template_or_start_again_during_generation(session):
    ui, workspace = session
    template = ui.select_template("first.pptx", make_template())
    job = ui.start("Текст")

    for action in (
        lambda: ui.select_template("second.pptx", make_template("Другой")),
        ui.load_example,
        ui.clear_template,
        lambda: ui.start("Ещё один текст"),
    ):
        with pytest.raises(APIError, match="Дождитесь"):
            action()

    assert ui.snapshot()["template"]["id"] == template["id"]
    assert ui.snapshot()["job"]["id"] == job["id"]
    assert len(workspace.templates) == len(workspace.jobs) == 1


def test_completed_job_displays_result_even_without_rendered_images(session):
    ui, workspace = session
    data = make_template()
    template = ui.select_template("template.pptx", data)
    publish_preview(workspace, "templates", workspace.templates[template["id"]], b"template")
    assert ui.get_slide(1) == b"template\x00"
    ui.start("Текст")
    job, result = complete(ui, workspace)
    ui.select_template("template.pptx", data)

    snapshot = ui.snapshot()

    assert snapshot["view"] == "result"
    assert snapshot["job"]["id"] == job.id
    assert snapshot["presentation"]["slide_count"] == 2
    assert snapshot["presentation"]["slides"][0]["title"] == "Готовый слайд 1"
    assert ui.get_slide(1) is None
    assert ui.download() == result
    publish_preview(workspace, "jobs", job.presentation, b"result")
    assert ui.get_slide(1) == b"result\x00"
    assert ui.get_slide(2) == b"result\x01"
    assert ui.get_slide(3) is None


def test_template_change_clears_old_result_and_reuses_cached_upload(session):
    ui, workspace = session
    original_data = make_template()
    original = ui.select_template("first.pptx", original_data)
    ui.start("Текст")
    complete(ui, workspace)

    second = ui.select_template("second.pptx", make_template("Другой"))

    assert ui.snapshot()["job"] is None
    assert ui.snapshot()["view"] == "template"
    assert ui.snapshot()["presentation"]["id"] == second["id"]
    with pytest.raises(APIError, match="ещё не готова"):
        ui.download()
    assert ui.select_template("first.pptx", original_data)["id"] == original["id"]
    assert len(workspace.templates) == 2


def test_example_can_be_reloaded_without_clearing_current_generation(session):
    ui, workspace = session
    example = ui.load_example()
    job = ui.start(example["script"])

    assert ui.load_example()["id"] == example["id"]
    assert ui.snapshot()["job"]["id"] == job["id"]
    assert len(workspace.templates) == 1


def test_snapshots_are_copies_and_include_safe_failure_progress(session):
    ui, workspace = session
    template = ui.select_template("template.pptx", make_template())
    started = ui.start("Текст")
    job = workspace.get_job(started["id"])
    with workspace.lock:
        job.stage = "content"
        job.status = "failed"
        job.error_code = "network"
        job.error = CONTENT_ERRORS["network"]
        job.progress = {"phase": "slides", "retrying": True}

    snapshot = ui.snapshot()
    assert snapshot["job"]["error"] == CONTENT_ERRORS["network"]
    assert snapshot["job"]["progress"] == {"phase": "slides", "retrying": True}
    assert "directory" not in snapshot["job"]
    assert "process" not in snapshot["job"]
    assert snapshot["view"] == "template"
    snapshot["template"]["slides"][0]["title"] = "Подменённый текст"
    snapshot["presentation"]["preview"]["status"] = "ready"
    snapshot["job"]["progress"]["phase"] = "changed"
    template["slides"].clear()
    fresh = ui.snapshot()
    assert fresh["presentation"]["slides"][0]["title"] == "Шаблон 1"
    assert fresh["presentation"]["preview"]["status"] == "unavailable"
    assert fresh["job"]["progress"]["phase"] == "slides"
    assert ui.start("Повтор")["id"] != job.id


def test_invalid_inputs_leave_selected_template_unchanged(session):
    ui, workspace = session
    original = ui.select_template("valid.pptx", make_template())

    for name, data in (("file.txt", b"no"), ("broken.pptx", b"no")):
        with pytest.raises(APIError):
            ui.select_template(name, data)
    for script, count in (("", None), ("Текст", 4), ("Текст", True)):
        with pytest.raises(APIError):
            ui.start(script, count)

    assert ui.snapshot()["template"]["id"] == original["id"]
    assert ui.snapshot()["job"] is None
    assert not workspace.jobs


def test_clear_template_removes_result_and_can_be_repeated(session):
    ui, workspace = session
    data = make_template()
    original = ui.select_template("template.pptx", data)
    ui.start("Текст")
    complete(ui, workspace)

    ui.clear_template()
    ui.clear_template()

    assert ui.snapshot() == {"template": None, "job": None, "presentation": None, "view": None}
    assert ui.get_slide(1) is None
    with pytest.raises(APIError, match="Выберите"):
        ui.start("Текст")
    with pytest.raises(APIError, match="ещё не готова"):
        ui.download()
    assert ui.select_template("template.pptx", data)["id"] == original["id"]


def test_download_and_preview_handle_missing_files_without_exposing_paths(session):
    ui, workspace = session
    ui.select_template("template.pptx", make_template())
    with pytest.raises(APIError, match="ещё не готова"):
        ui.download()
    ui.start("Текст")
    job, _ = complete(ui, workspace)
    publish_preview(workspace, "jobs", job.presentation, b"result")
    workspace.preview_files[("jobs", job.id)][0].unlink()
    (job.directory / "result.pptx").unlink()

    assert ui.get_slide(1) is None
    with pytest.raises(APIError) as error:
        ui.download()
    assert str(workspace.root) not in str(error.value)


def test_sessions_keep_files_and_jobs_isolated_and_close_once(monkeypatch):
    workspaces = [Workspace(renderer=FakeRenderer()) for _ in range(2)]
    for workspace in workspaces:
        monkeypatch.setattr(workspace, "_run_job", lambda job: None)
        monkeypatch.setattr(workspace, "close", Mock(wraps=workspace.close))
    first, second = [PresentationSession(workspace=workspace) for workspace in workspaces]
    try:
        data = make_template()
        first.select_template("one.pptx", data)
        second.select_template("two.pptx", data)
        first.start("Первая сессия")
        second.start("Вторая сессия")
        assert first.snapshot()["job"]["id"] != second.snapshot()["job"]["id"]
        first.close()
        first.close()
        assert not workspaces[0].root.exists()
        assert workspaces[1].root.exists()
        assert second.snapshot()["job"]["status"] == "running"
        with pytest.raises(APIError, match="Сессия завершена"):
            first.snapshot()
        assert workspaces[0].close.call_count == 1
    finally:
        first.close()
        second.close()
    assert not workspaces[1].root.exists()
    assert workspaces[1].close.call_count == 1


def test_garbage_collection_releases_workspace_without_registry_leak(monkeypatch):
    workspace = Workspace(renderer=FakeRenderer())
    monkeypatch.setattr(workspace, "close", Mock(wraps=workspace.close))
    session = PresentationSession(workspace=workspace)
    reference = weakref.ref(session)

    del session
    gc.collect()

    assert reference() is None
    assert workspace.closed
    assert not workspace.root.exists()
    assert workspace.close.call_count == 1


def test_shutdown_closes_every_live_session_and_is_repeatable(monkeypatch):
    module = importlib.import_module("exposlides.ui_session")
    monkeypatch.setattr(module, "_sessions", weakref.WeakSet())
    workspaces = [Workspace(renderer=FakeRenderer()) for _ in range(2)]
    sessions = [PresentationSession(workspace=workspace) for workspace in workspaces]

    module._close_sessions()
    module._close_sessions()

    assert all(workspace.closed and not workspace.root.exists() for workspace in workspaces)
    assert not module._sessions
    for session in sessions:
        with pytest.raises(APIError, match="Сессия завершена"):
            session.snapshot()
