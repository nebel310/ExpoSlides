from __future__ import annotations

import asyncio
import importlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
PARSING_SERVICE_ROOT = REPOSITORY_ROOT / "services" / "parsing-service"
CONTENT_SERVICE_ROOT = REPOSITORY_ROOT / "services" / "content-service"
BUILDER_SERVICE_ROOT = REPOSITORY_ROOT / "services" / "builder-service"
SOURCE = (
    "LangChain помогает создавать агентов и подключать инструменты. "
    "LangGraph управляет состоянием и переходами между шагами."
)
TITLE = "LangChain и LangGraph"
INITIAL_LABEL = "Разработка и управление интеллектуальными агентами"
REPAIRED_LABEL = "Инструменты разработки интеллектуальных агентов"
ALTERNATIVES = [
    "Создание интеллектуальных агентов",
    "Разработка агентных систем",
    "Управление агентными сценариями",
]
MICRO_LABEL = "Архитектура агентных решений"


def test_validated_title_term_recovers_small_speaker_slot_through_real_graph_and_builder(
    tmp_path, monkeypatch, service_importer,
):
    template = tmp_path / "template.pptx"
    parsed = tmp_path / "template.json"
    script = tmp_path / "script.txt"
    content = tmp_path / "generated_content.json"
    output = tmp_path / "result.pptx"
    original = Presentation()
    slide = original.slides.add_slide(original.slide_layouts[3])
    slide.shapes.title.text = "Заголовок презентации"
    slide.placeholders[1].text = "Описание разработки и управления интеллектуальными агентами"
    speaker = slide.placeholders[2]
    speaker.text = "Имя докладчика"
    speaker.width = Inches(1.6)
    speaker.height = Inches(0.5)
    speaker_run = speaker.text_frame.paragraphs[0].runs[0]
    speaker_run.font.name = "Arial"
    speaker_run.font.size = Pt(18)
    speaker_run.font.bold = True
    speaker_run.font.color.rgb = RGBColor(0x24, 0x36, 0x68)
    original.save(template)
    script.write_text(SOURCE, encoding="utf-8")

    parsing_main = service_importer(PARSING_SERVICE_ROOT, "app.main")
    asyncio.run(parsing_main.run(template, parsed))

    content_main = service_importer(CONTENT_SERVICE_ROOT, "app.main")
    fast = importlib.import_module("app.graph.fast")
    request = importlib.import_module("app.models.request")
    original_parse = content_main.PresentationParser.parse

    def parse_with_fixed_capacity(payload):
        presentation = original_parse(payload)
        # Изолируем восстановление короткого поля от эвристики вместимости.
        limits = {0: 80, 1: 200, 2: 16}
        for placeholder in presentation.slides[0].placeholders:
            placeholder.max_length = limits[placeholder.idx]
        return presentation

    monkeypatch.setattr(content_main.PresentationParser, "parse", parse_with_fixed_capacity)
    calls = []

    class FakeLLM:
        async def generate_json(self, prompt, model, strict=True):
            calls.append(("outline", None))
            return model.model_validate({
                "analysis": {
                    "topic": TITLE, "audience": "", "objective": "",
                    "blocks": [{
                        "index": 1, "heading": TITLE, "summary": SOURCE,
                        "key_points": [SOURCE], "facts": [],
                    }],
                    "key_messages": [SOURCE], "facts": [],
                },
                "plan": {"slides": [{
                    "template_slide_index": 1, "title": TITLE, "content": SOURCE,
                    "purpose": "Описать разработку агентов", "key_message": SOURCE,
                    "source_block_indices": [1],
                }]},
            })

        async def generate_json_object(self, prompt, schema, strict=True, *, model=None):
            calls.append(("content", schema))
            if len(calls) == 2:
                assert schema["required"] == ["field_0000", "field_0001", "field_0002"]
                return {"field_0000": TITLE, "field_0001": SOURCE, "field_0002": INITIAL_LABEL}
            assert schema["required"] == ["field_0002"]
            if len(calls) == 3:
                return {"field_0002": REPAIRED_LABEL}
            if len(calls) == 4:
                assert schema["properties"]["field_0002"]["type"] == "array"
                return {"field_0002": ALTERNATIVES}
            assert len(calls) == 5
            return {"field_0002": MICRO_LABEL}

    assert all(len(value) > 16 for value in (
        INITIAL_LABEL, REPAIRED_LABEL, *ALTERNATIVES, MICRO_LABEL,
    ))
    assert all("LangChain" not in value and "LangGraph" not in value for value in (
        INITIAL_LABEL, REPAIRED_LABEL, *ALTERNATIVES, MICRO_LABEL,
    ))
    monkeypatch.setattr(fast, "fast_llm_client", FakeLLM())
    asyncio.run(content_main.run(
        parsed, script, content,
        request.GenerationSettings(max_slides=None, generation_mode="fast"),
    ))

    generated = json.loads(content.read_text(encoding="utf-8"))
    assert generated["validation_report"] == {"ok": True, "issues": []}
    assert list(generated["content"]) == ["1"]
    assert generated["content"]["1"]["placeholders"] == {
        "0": TITLE, "1": SOURCE, "2": "LangChain",
    }
    assert [kind for kind, _ in calls] == ["outline"] + ["content"] * 4

    builder_main = service_importer(BUILDER_SERVICE_ROOT, "app.main")
    asyncio.run(builder_main.run(template, parsed, content, output))
    reopened = Presentation(output)
    assert len(reopened.slides) == 1
    assert reopened.slides[0].shapes.title.text == TITLE
    assert reopened.slides[0].placeholders[1].text == SOURCE
    restored_speaker = reopened.slides[0].placeholders[2]
    assert restored_speaker.text == "LangChain"
    assert len(restored_speaker.text) <= 16
    restored_run = restored_speaker.text_frame.paragraphs[0].runs[0]
    assert restored_run.font.name == "Arial"
    assert restored_run.font.size == Pt(18)
    assert restored_run.font.bold is True
    assert restored_run.font.color.rgb == RGBColor(0x24, 0x36, 0x68)
    assert restored_speaker.width == Inches(1.6)
    assert restored_speaker.height == Inches(0.5)
