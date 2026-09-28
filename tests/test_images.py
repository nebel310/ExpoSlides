from __future__ import annotations

import asyncio
import base64
import importlib
import io
import json
import sys
from pathlib import Path

import httpx
import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
image_api = importlib.import_module("exposlides.image_api")
images = importlib.import_module("exposlides.images")
find_image_slots = importlib.import_module("exposlides.image_layout").find_image_slots
ImageGenerationError = image_api.ImageGenerationError
ImageSettings = image_api.ImageSettings
generate_image = image_api.generate_image


def png(color="green", mode="RGB") -> bytes:
    stream = io.BytesIO()
    Image.new(mode, (512, 512), color).save(stream, format="PNG")
    return stream.getvalue()


def payload() -> dict:
    return {"images": [{"url": "data:image/png;base64," + base64.b64encode(png()).decode()}]}


@pytest.fixture
def settings() -> ImageSettings:
    return ImageSettings(_env_file=None, llm_api_key="hf_offline-test",
                         llm_base_url="https://router.huggingface.co/v1")


def write_deck(tmp_path: Path, *, placeholder=False) -> tuple[Path, Path]:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[8] if placeholder else deck.slide_layouts[5])
    slide.shapes.title.text = "Команда разрабатывает новый продукт"
    run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(28)
    run.font.bold = True
    if not placeholder:
        picture = slide.shapes.add_picture(
            io.BytesIO(png("red")), Inches(1), Inches(2), Inches(5), Inches(3)
        )
        picture.name = "Illustration"
        picture.click_action.hyperlink.address = "https://example.com/reference"
        picture.line.color.rgb = RGBColor(10, 20, 30)
    path = tmp_path / "deck.pptx"
    deck.save(path)
    template = tmp_path / "template.json"
    template.write_text(json.dumps({"theme": {"colors": {"accent1": "123456"}}}), encoding="utf-8")
    return path, template


def test_api_uses_hf_key_expected_route_and_inline_output(settings):
    def handler(request):
        assert str(request.url) == image_api.IMAGE_ENDPOINT
        assert request.headers["Authorization"] == "Bearer hf_offline-test"
        body = json.loads(request.content)
        assert body["image_size"] == {"width": 1024, "height": 768}
        assert body["sync_mode"] is True
        assert body["enable_safety_checker"] is True
        assert body["enable_prompt_expansion"] is False
        return httpx.Response(200, json=payload())

    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await generate_image(client, settings, prompt="test", width=1024,
                                        height=768, seed=42)
    assert asyncio.run(run()) == png()


@pytest.mark.parametrize("status", [401, 403, 402, 422, 500, 503])
def test_api_errors_are_safe_and_do_not_retry_ambiguous_failure(status, settings):
    calls = []

    def handler(request):
        calls.append(request)
        return httpx.Response(status, text="secret-token and private prompt")

    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            await generate_image(client, settings, prompt="private prompt", width=512,
                                 height=512, seed=0)
    with pytest.raises(ImageGenerationError) as caught:
        asyncio.run(run())
    assert "secret" not in str(caught.value)
    assert "private" not in str(caught.value)
    assert len(calls) == 1


@pytest.mark.parametrize("result", [
    {}, [], {"images": []}, {"images": [{"url": "https://evil.example/image.png"}]},
    {"images": [{"url": "data:image/png;base64,invalid"}]},
    {"images": [{"url": "data:image/png;base64," + base64.b64encode(b"bad").decode()}]},
    {"has_nsfw_concepts": [True]},
])
def test_bad_or_filtered_images_rejected(result):
    with pytest.raises(ImageGenerationError):
        image_api.decode_image(result)


def test_wrong_provider_never_receives_key(settings):
    settings.llm_base_url = "https://other.example/v1"
    with pytest.raises(ImageGenerationError, match="LLM_BASE_URL"):
        settings.token()
    assert "hf_offline-test" not in repr(settings)


def test_scene_prompt_keeps_slide_text_out_of_image_prompt(settings):
    def handler(request):
        body = json.loads(request.content)
        assert body["model"] == settings.llm_fast_model
        assert body["response_format"]["type"] == "json_schema"
        return httpx.Response(200, json={"choices": [{"message": {
            "content": json.dumps({"scene": "A team arranging colorful abstract blocks on a desk."}),
        }}]})

    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await image_api.describe_scene(client, settings, "PRIVATE SLIDE TITLE")
    prompt = asyncio.run(run())
    assert "PRIVATE SLIDE TITLE" not in prompt
    assert "Absolutely no text" in prompt


@pytest.mark.parametrize("content", ["not json", '{"scene": ""}', '{"scene": 7}'])
def test_invalid_scene_fails_before_image_generation(settings, content):
    async def run():
        transport = httpx.MockTransport(lambda r: httpx.Response(200, json={
            "choices": [{"message": {"content": content}}],
        }))
        async with httpx.AsyncClient(transport=transport) as client:
            await image_api.describe_scene(client, settings, "topic")
    with pytest.raises(ImageGenerationError):
        asyncio.run(run())


@pytest.mark.parametrize("placeholder", [False, True])
def test_generated_image_reopens_and_preserves_geometry_style_layer_and_links(
    tmp_path, monkeypatch, settings, placeholder,
):
    path, template = write_deck(tmp_path, placeholder=placeholder)
    before = Presentation(path)
    slots = find_image_slots(before)
    assert len(slots) == 1
    before_ids = [s.shape_id for s in before.slides[0].shapes]
    calls = []

    async def fake(jobs, config):
        calls.extend(jobs)
        return [png() for _ in jobs]
    monkeypatch.setattr(images, "_generate_all", fake)
    report = images.illustrate_presentation(
        path, template, report_path=tmp_path / "report.json", settings=settings,
    )
    after = Presentation(path)
    assert len(after.slides) == 1
    assert [s.shape_id for s in after.slides[0].shapes] == before_ids
    shape = next(s for s in after.slides[0].shapes if s.shape_id == slots[0].shape_id)
    assert (shape.left, shape.top, shape.width, shape.height) == (
        slots[0].left, slots[0].top, slots[0].width, slots[0].height,
    )
    with Image.open(io.BytesIO(shape.image.blob)) as result:
        assert result.getpixel((10, 10)) == (0, 128, 0)
        assert abs(result.width / result.height - shape.width / shape.height) < 0.002
    title = after.slides[0].shapes.title
    run = title.text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Arial" and run.font.bold and run.font.size == Pt(28)
    if not placeholder:
        assert shape.click_action.hyperlink.address == "https://example.com/reference"
        assert shape.line.color.rgb == RGBColor(10, 20, 30)
    assert report["status"] == "generated"
    assert "#123456" in calls[0]["prompt"]
    assert title.text in calls[0]["prompt"]
    assert "hf_offline-test" not in (tmp_path / "report.json").read_text()


def test_api_failure_keeps_original_file(tmp_path, monkeypatch, settings):
    path, template = write_deck(tmp_path)
    original = path.read_bytes()

    async def fail(*args):
        raise ImageGenerationError("API unavailable")
    monkeypatch.setattr(images, "_generate_all", fail)
    with pytest.raises(ImageGenerationError):
        images.illustrate_presentation(path, template, report_path=tmp_path / "report.json",
                                       settings=settings)
    assert path.read_bytes() == original


@pytest.mark.parametrize("mode", ["off", "disabled"])
def test_disable_never_calls_api(tmp_path, monkeypatch, settings, mode):
    path, template = write_deck(tmp_path)
    original = path.read_bytes()
    settings.image_generation_enabled = False
    monkeypatch.setattr(images, "_generate_all", lambda *a: pytest.fail("Unexpected API call"))
    result = images.illustrate_presentation(
        path, template, report_path=tmp_path / "report.json", settings=settings,
        mode="off" if mode == "off" else "auto",
    )
    assert result["status"] == "disabled"
    assert path.read_bytes() == original


def test_no_slots_does_not_load_credentials(tmp_path, monkeypatch):
    path = tmp_path / "deck.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[5])
    deck.save(path)
    monkeypatch.setattr(images, "ImageSettings", lambda **kw: pytest.fail("Read credentials"))
    assert images.illustrate_presentation(
        path, tmp_path / "unused.json", report_path=tmp_path / "report.json",
    )["status"] == "no_slots"


@pytest.mark.parametrize("protection", ["logo", "keep", "repeat", "overlap", "outside", "small"])
def test_template_protected_images_remain_untouched(tmp_path, protection):
    path, _ = write_deck(tmp_path)
    deck = Presentation(path)
    slide = deck.slides[0]
    picture = slide.shapes[-1]
    if protection in {"logo", "keep"}:
        picture.name = protection
    elif protection == "repeat":
        other = deck.slides.add_slide(deck.slide_layouts[6])
        other.shapes.add_picture(io.BytesIO(png("red")), Inches(1), Inches(2), Inches(5), Inches(3))
    elif protection == "overlap":
        slide.shapes.add_textbox(Inches(1), Inches(2), Inches(1), Inches(1)).text = "Подпись"
    elif protection == "outside":
        picture.left = -1
    else:
        picture.width = Inches(0.3)
    assert find_image_slots(deck) == []


def test_generation_deadline_cancels_pending_requests(monkeypatch, settings):
    settings.image_generation_timeout = 0.02
    cancelled = []

    async def slow(*args, **kwargs):
        try:
            await asyncio.sleep(5)
        finally:
            cancelled.append(True)
    async def scene(*args):
        return "An editorial scene with no text"
    monkeypatch.setattr(images, "describe_scene", scene)
    monkeypatch.setattr(images, "generate_image", slow)
    with pytest.raises(ImageGenerationError, match="общее время"):
        asyncio.run(images._generate_all([{"prompt": "topic"}] * 3, settings))
    assert len(cancelled) == 2


def test_invalid_image_does_not_publish_partial_deck(tmp_path, monkeypatch, settings):
    path, template = write_deck(tmp_path)
    original = path.read_bytes()

    async def invalid(*args):
        return [b"not an image"]
    monkeypatch.setattr(images, "_generate_all", invalid)
    with pytest.raises(ImageGenerationError):
        images.illustrate_presentation(path, template, report_path=tmp_path / "report.json",
                                       settings=settings)
    assert path.read_bytes() == original


def test_429_has_one_bounded_retry(monkeypatch, settings):
    calls = []

    def handler(request):
        calls.append(request)
        return httpx.Response(429) if len(calls) == 1 else httpx.Response(200, json=payload())

    async def no_wait(*args):
        return None
    monkeypatch.setattr(image_api.asyncio, "sleep", no_wait)

    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await generate_image(client, settings, prompt="test", width=512,
                                        height=512, seed=0)
    assert asyncio.run(run()) == png()
    assert len(calls) == 2


def test_settings_reuses_llm_key_and_environment_wins(tmp_path, monkeypatch):
    env = tmp_path / "test.env"
    env.write_text("LLM_API_KEY=hf_file_dummy\nIMAGE_CONCURRENCY=1\n", encoding="utf-8")
    monkeypatch.setenv("LLM_API_KEY", "hf_process_dummy")
    assert ImageSettings(_env_file=env).llm_api_key.get_secret_value() == "hf_process_dummy"
    assert ImageSettings(_env_file=env).image_concurrency == 1


def test_image_limit_fails_before_external_call(tmp_path, monkeypatch, settings):
    path, template = write_deck(tmp_path, placeholder=True)
    deck = Presentation(path)
    slide = deck.slides.add_slide(deck.slide_layouts[8])
    slide.shapes.title.text = "Второй слайд"
    deck.save(path)
    settings.image_max_count = 1
    monkeypatch.setattr(images, "_generate_all", lambda *a: pytest.fail("Unexpected API call"))
    with pytest.raises(ImageGenerationError, match="IMAGE_MAX_COUNT"):
        images.illustrate_presentation(path, template, report_path=tmp_path / "report.json",
                                       settings=settings)


def test_existing_picture_is_not_changed_in_other_slide(tmp_path, monkeypatch, settings):
    path, template = write_deck(tmp_path)
    deck = Presentation(path)
    deck.slides[0].shapes[-1].name = "exposlides:image"
    second = deck.slides.add_slide(deck.slide_layouts[5])
    picture = second.shapes.add_picture(io.BytesIO(png("red")), Inches(1), Inches(2), Inches(5))
    picture.name = "keep reference"
    deck.save(path)

    async def fake(*args):
        return [png()]
    monkeypatch.setattr(images, "_generate_all", fake)
    images.illustrate_presentation(path, template, report_path=tmp_path / "report.json",
                                   settings=settings)
    reopened = Presentation(path)
    assert reopened.slides[1].shapes[-1].image.blob == png("red")
    assert reopened.slides[0].shapes[-1].image.blob != png("red")
