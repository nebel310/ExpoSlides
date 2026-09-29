"""Продолжения обычного текста выровнены с первой строкой; списки сохранены."""

from copy import deepcopy

import pytest
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from exposlides.design_builder import build_deck
from exposlides.design_models import Box, DeckPlan, PlacedBlock, SlideInstance, TextStyle
from exposlides.design_native_text import native_plain_text_margin
from exposlides.design_saved_audit import _native_text_matches


def _fixture(tmp_path, layer, *, bullet="buNone", indent=-228600, alignment="l", kind="text"):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    body = slide.placeholders[1]
    body.text = "Исходный текст"
    paragraph = body.text_frame.paragraphs[0]
    paragraph.font.size = Pt(18)
    if layer == "paragraph":
        properties = paragraph._p.get_or_add_pPr()
    else:
        if layer == "shape":
            style = body.text_frame._txBody.find(qn("a:lstStyle"))
        elif layer == "layout":
            style = slide.slide_layout.placeholders[1].text_frame._txBody.find(qn("a:lstStyle"))
        else:
            style = slide.slide_layout.slide_master._element.find(qn("p:txStyles")).find(qn("p:bodyStyle"))
        properties = style.find(qn("a:lvl1pPr"))
        if properties is None:
            properties = OxmlElement("a:lvl1pPr")
            style.append(properties)
    properties.set("marL", "457200")
    properties.set("indent", str(indent))
    properties.set("algn", alignment)
    for child in list(properties):
        if child.tag in {qn("a:buNone"), qn("a:buChar"), qn("a:buAutoNum"), qn("a:buBlip")}:
            properties.remove(child)
    marker = OxmlElement("a:" + bullet)
    if bullet == "buChar":
        marker.set("char", "•")
    elif bullet == "buAutoNum":
        marker.set("type", "arabicPeriod")
    properties.insert(0, marker)
    block = PlacedBlock(
        id="body", kind=kind, source_shape_id=body.shape_id,
        box=Box(left=body.left, top=body.top, width=body.width, height=body.height),
        style=TextStyle(font="Arial", size=18, color="000000"),
        items=["Определение задачи и желаемого результата", "Выбор архитектуры агента"],
    )
    plan = DeckPlan(variant_id="story", name="Текст", description="Проверка переносов",
                    width=prs.slide_width, height=prs.slide_height,
                    slides=[SlideInstance(id="s1", story_slide_id="s1", source_slide_index=1,
                                          blocks=[block])])
    source = tmp_path / "source.pptx"
    prs.save(source)
    return source, prs, body, block, plan


@pytest.mark.parametrize("layer", ["paragraph", "shape", "layout", "master"])
def test_plain_hanging_indent_aligns_wrapped_lines_without_changing_template(tmp_path, layer):
    source, prs, body, block, plan = _fixture(tmp_path, layer)
    source_bytes = source.read_bytes()
    layout_xml = etree.tostring(prs.slides[0].slide_layout._element)
    master_xml = etree.tostring(prs.slides[0].slide_layout.slide_master._element)
    assert native_plain_text_margin(body, body.text_frame.paragraphs[0]) == 228600
    assert etree.tostring(prs.slides[0].slide_layout._element) == layout_xml
    assert etree.tostring(prs.slides[0].slide_layout.slide_master._element) == master_xml
    output = build_deck(source, plan, tmp_path / "result.pptx")
    result = Presentation(output)
    assert len(result.slides) == 1
    filled = result.slides[0].placeholders[1]
    assert filled.text == "\n".join(block.items)
    for paragraph in filled.text_frame.paragraphs:
        assert paragraph._p.pPr.get("marL") == "228600"
        assert paragraph._p.pPr.get("indent") == "0"
    assert _native_text_matches(filled, body, block)
    assert source.read_bytes() == source_bytes
    # Незаявленные изменения оформления по-прежнему обнаруживаются аудитом.
    filled.text_frame.paragraphs[0]._p.pPr.set("marL", "0")
    assert not _native_text_matches(filled, body, block)


@pytest.mark.parametrize("options", [
    {"bullet": "buChar"}, {"bullet": "buAutoNum"}, {"indent": 228600},
    {"alignment": "ctr"}, {"alignment": "r"}, {"kind": "title"},
])
def test_real_lists_and_other_paragraph_styles_keep_their_indents(tmp_path, options):
    source, _, body, block, plan = _fixture(tmp_path, "paragraph", **options)
    before = deepcopy(body.text_frame.paragraphs[0]._p.pPr)
    output = build_deck(source, plan, tmp_path / "result.pptx")
    filled = Presentation(output).slides[0].placeholders[1]
    assert etree.tostring(filled.text_frame.paragraphs[0]._p.pPr, method="c14n", exclusive=True) == etree.tostring(before, method="c14n", exclusive=True)
    assert _native_text_matches(filled, body, block)


def test_explicit_bullet_overrides_inherited_plain_hanging_indent(tmp_path):
    source, prs, body, block, plan = _fixture(tmp_path, "layout")
    properties = body.text_frame.paragraphs[0]._p.get_or_add_pPr()
    bullet = OxmlElement("a:buChar")
    bullet.set("char", "•")
    properties.insert(0, bullet)
    assert native_plain_text_margin(body, body.text_frame.paragraphs[0]) is None
    prs.save(source)
    output = build_deck(source, plan, tmp_path / "result.pptx")
    filled = Presentation(output).slides[0].placeholders[1]
    assert filled.text_frame.paragraphs[0]._p.pPr.get("indent") is None
    assert _native_text_matches(filled, body, block)
