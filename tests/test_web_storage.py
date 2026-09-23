from __future__ import annotations

import base64
import io
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation
from test_web import FakeRenderer, request, web

from exposlides.file_storage import StorageError, StoredFile
from exposlides.web_catalog import WebCatalog


class MemoryStorage:
    """Байты переживают закрытие клиента; сеть и LLM не используются."""

    def __init__(self, files=None):
        self.files = files if files is not None else {}
        self.uploads = []
        self.downloads = []
        self.fail_upload = False
        self.fail_download = False
        self.closed = False

    def upload(self, name, data, *, task_id=""):
        if self.fail_upload:
            raise StorageError("private storage address and credentials")
        file = StoredFile(f"file-{len(self.files) + 1}", 7)
        self.files[file] = data
        self.uploads.append((file, name, task_id))
        return file

    def download(self, file):
        self.downloads.append(file)
        if self.fail_download:
            raise StorageError("private storage address and credentials")
        return self.files[file]

    def close(self):
        self.closed = True


@pytest.fixture
def storage_workspace(tmp_path):
    storage = MemoryStorage()
    workspace = web.Workspace(
        storage=storage, data_dir=tmp_path / "library", renderer=FakeRenderer(),
    )
    yield workspace, storage
    workspace.close()


@pytest.fixture
def fake_pipeline(monkeypatch):
    commands = []

    class Process:
        def __init__(self, command, **kwargs):
            commands.append(command)
            template = Path(command[command.index("--template") + 1])
            presentation = Presentation(template)
            presentation.slides[0].shapes.title.text = "Сохранённый результат"
            presentation.save(command[command.index("--output") + 1])
            self.stdout = io.StringIO("[1/3] Парсинг\n[2/3] Текст\n[3/3] Сборка\n")

        def wait(self):
            return 0

    monkeypatch.setattr(web.subprocess, "Popen", Process)
    return commands


def complete(workspace, template_id):
    result = workspace.create_job({"template_id": template_id, "script": "Приватный исходный текст"})
    workspace.workers[-1].join(timeout=5)
    assert not workspace.workers[-1].is_alive()
    return workspace.get_job(result["id"])


def server(workspace):
    return SimpleNamespace(workspace=workspace, server_port=8765)


def test_saved_template_and_result_survive_restart_without_local_pptx(
    storage_workspace, fake_pipeline,
):
    first, storage = storage_workspace
    template = first.add_template("Русский шаблон.pptx", web.make_example())
    job = complete(first, template["id"])
    assert job.status == "completed"
    assert [name for _, name, _ in storage.uploads] == ["Русский шаблон.pptx", "Презентация.pptx"]
    assert storage.uploads[-1][2] == job.id
    for data in storage.files.values():
        assert len(Presentation(io.BytesIO(data)).slides) == 5
    catalog_dir = first.catalog.path.parent
    catalog_text = first.catalog.path.read_text(encoding="utf-8")
    assert "Приватный исходный текст" not in catalog_text
    assert str(first.root) not in catalog_text
    first.close()
    assert not first.root.exists()
    assert storage.closed

    restored_storage = MemoryStorage(storage.files)
    second = web.Workspace(
        storage=restored_storage, data_dir=catalog_dir, renderer=FakeRenderer(),
    )
    try:
        assert restored_storage.downloads == []
        status, body, _ = request(server(second), "GET", "/api/library")
        library = json.loads(body)
        assert status == 200
        assert library["persistent"] is True
        assert library["templates"][0]["id"] == template["id"]
        assert library["jobs"] == [{
            "id": job.id, "status": "completed", "stage": "complete", "slide_count": 5,
            "name": "Русский шаблон.pptx",
        }]
        assert "file_id" not in library["templates"][0]
        status, data, _ = request(server(second), "GET", f"/api/jobs/{job.id}/download")
        assert status == 200
        assert Presentation(io.BytesIO(data)).slides[0].shapes.title.text == "Сохранённый результат"
        assert restored_storage.downloads == [storage.uploads[1][0]]
        next_job = complete(second, template["id"])
        assert next_job.status == "completed"
        assert restored_storage.downloads[-1] == storage.uploads[0][0]
        assert len(fake_pipeline) == 2
    finally:
        second.close()


def test_upload_failure_is_503_and_not_a_successful_local_upload(storage_workspace):
    workspace, storage = storage_workspace
    storage.fail_upload = True
    status, data, _ = request(server(workspace), "POST", "/api/templates", {
        "name": "Шаблон.pptx", "data": base64.b64encode(web.make_example()).decode(),
    })
    assert status == 503
    assert json.loads(data) == {"error": web.STORAGE_UNAVAILABLE}
    assert workspace.library()["templates"] == []
    assert not list(workspace.root.glob("*.pptx"))
    assert not workspace.catalog.path.exists()


def test_failed_catalog_write_preserves_previous_library(storage_workspace, monkeypatch):
    workspace, _ = storage_workspace
    template = workspace.add_template("first.pptx", web.make_example())
    original = workspace.catalog.path.read_bytes()
    replace = web.os.replace

    def no_space(source, destination):
        if destination == workspace.catalog.path:
            raise OSError("private disk detail")
        return replace(source, destination)

    monkeypatch.setattr("exposlides.web_catalog.os.replace", no_space)
    with pytest.raises(StorageError, match="сохранить каталог"):
        workspace.add_template("second.pptx", web.make_example())
    assert workspace.catalog.path.read_bytes() == original
    assert list(workspace.templates) == [template["id"]]
    assert len(list(workspace.root.glob("*.pptx"))) == 1
    assert not list(workspace.catalog.path.parent.glob(".library-*"))


@pytest.mark.parametrize("failure", ["upload", "catalog"])
def test_result_not_completed_when_durable_save_fails(
    storage_workspace, fake_pipeline, monkeypatch, failure,
):
    workspace, storage = storage_workspace
    template = workspace.add_template("first.pptx", web.make_example())
    if failure == "upload":
        storage.fail_upload = True
    else:
        def fail(*args, **kwargs):
            raise StorageError("private disk detail")
        monkeypatch.setattr(workspace.catalog, "save_result", fail)
    job = complete(workspace, template["id"])
    assert job.status == "failed"
    assert job.error == web.STORAGE_SAVE_FAILED
    assert workspace.library()["jobs"] == []
    assert workspace.catalog.library.jobs == {}
    assert request(server(workspace), "GET", f"/api/jobs/{job.id}/download")[0] == 409


def test_restored_download_outage_returns_safe_error_then_recovers(
    storage_workspace, fake_pipeline,
):
    first, storage = storage_workspace
    template = first.add_template("first.pptx", web.make_example())
    job = complete(first, template["id"])
    directory = first.catalog.path.parent
    first.close()
    next_storage = MemoryStorage(storage.files)
    next_storage.fail_download = True
    second = web.Workspace(storage=next_storage, data_dir=directory, renderer=FakeRenderer())
    try:
        url = f"/api/jobs/{job.id}/download"
        status, data, _ = request(server(second), "GET", url)
        assert status == 503
        assert json.loads(data) == {"error": web.STORAGE_UNAVAILABLE}
        assert not (second.jobs[job.id].directory / "result.pptx").exists()
        status, data, _ = request(server(second), "POST", "/api/jobs", {
            "template_id": template["id"], "script": "Текст",
        })
        assert status == 503
        assert len(second.jobs) == 1
        next_storage.fail_download = False
        assert request(server(second), "GET", url)[0] == 200
    finally:
        second.close()


def test_restored_previews_render_from_saved_files(storage_workspace, fake_pipeline):
    first, storage = storage_workspace
    template = first.add_template("first.pptx", web.make_example())
    job = complete(first, template["id"])
    directory = first.catalog.path.parent
    first.close()
    renderer = FakeRenderer(available=True)
    second = web.Workspace(
        storage=MemoryStorage(storage.files), data_dir=directory, renderer=renderer,
    )
    try:
        assert second.get_preview(template["id"])["status"] == "pending"
        assert second.get_presentation(job.id)["preview"]["status"] == "pending"
        second.preview_queue.join()
        assert second.get_preview(template["id"])["status"] == "ready"
        assert second.get_presentation(job.id)["preview"]["status"] == "ready"
        assert len(renderer.calls) == 2
    finally:
        second.close()


@pytest.mark.parametrize("corruption", ["json", "version", "extra", "identifier", "reference"])
def test_invalid_catalog_is_rejected_without_overwrite(storage_workspace, corruption):
    workspace, storage = storage_workspace
    template = workspace.add_template("first.pptx", web.make_example())
    path = workspace.catalog.path
    workspace.close()
    data = json.loads(path.read_text(encoding="utf-8"))
    if corruption == "version":
        data["version"] = 2
    elif corruption == "extra":
        data["unknown"] = True
    elif corruption == "identifier":
        data["templates"]["../../escape"] = data["templates"].pop(template["id"])
    elif corruption == "reference":
        record = data["templates"][template["id"]]
        data["jobs"][template["id"]] = {**record, "template_id": "f" * 32}
    text = "{invalid" if corruption == "json" else json.dumps(data)
    path.write_text(text, encoding="utf-8")
    with pytest.raises(StorageError, match="открыть каталог"):
        WebCatalog(path.parent)
    assert path.read_text(encoding="utf-8") == text
    # Не оставляет блокировку после неудачного открытия.
    path.write_text('{"version":1,"templates":{},"jobs":{}}', encoding="utf-8")
    WebCatalog(path.parent).close()


def test_catalog_lock_prevents_concurrent_servers(storage_workspace):
    workspace, _ = storage_workspace
    directory = workspace.catalog.path.parent
    with pytest.raises(StorageError, match="другим процессом"):
        WebCatalog(directory)
    workspace.close()
    WebCatalog(directory).close()


def test_temporary_mode_stays_explicit_and_offline():
    workspace = web.Workspace(renderer=FakeRenderer())
    try:
        assert workspace.library() == {"persistent": False, "templates": [], "jobs": []}
    finally:
        workspace.close()


@pytest.mark.parametrize("arguments", [["--file-service", "localhost:50051"], ["--data-dir", "/tmp"]])
def test_web_requires_paired_storage_settings(arguments, monkeypatch):
    monkeypatch.delenv("EXPOSLIDES_FILE_SERVICE", raising=False)
    monkeypatch.delenv("EXPOSLIDES_DATA_DIR", raising=False)
    with pytest.raises(SystemExit) as error:
        web.main(arguments)
    assert error.value.code == 2


def test_web_reads_storage_environment_and_closes_resources(tmp_path, monkeypatch):
    created = []
    storage = MemoryStorage()
    monkeypatch.setenv("EXPOSLIDES_FILE_SERVICE", "localhost:50051")
    monkeypatch.setenv("EXPOSLIDES_DATA_DIR", str(tmp_path / "catalog"))
    monkeypatch.setattr(web, "PreviewRenderer", FakeRenderer)

    def connect(target):
        assert target == "localhost:50051"
        return storage

    class Server:
        def __init__(self, address, workspace):
            created.append(workspace)
            assert workspace.catalog.path.parent == tmp_path / "catalog"
            self.workspace = workspace

        def serve_forever(self):
            raise KeyboardInterrupt

        def server_close(self):
            self.workspace.close()

    monkeypatch.setattr(web, "FileServiceStorage", connect)
    monkeypatch.setattr(web, "WebServer", Server)
    assert web.main([]) == 0
    assert storage.closed
    assert not created[0].root.exists()


def test_web_bind_failure_releases_catalog_and_channel(tmp_path, monkeypatch):
    storage = MemoryStorage()
    monkeypatch.setattr(web, "FileServiceStorage", lambda target: storage)
    monkeypatch.setattr(web, "PreviewRenderer", FakeRenderer)

    def unavailable(*args, **kwargs):
        raise OSError("port in use")

    monkeypatch.setattr(web, "WebServer", unavailable)
    assert web.main(["--file-service", "localhost:50051", "--data-dir", str(tmp_path)]) == 1
    assert storage.closed
    WebCatalog(tmp_path).close()


def test_long_upload_filename_keeps_extension(storage_workspace):
    workspace, storage = storage_workspace
    status, data, _ = request(server(workspace), "POST", "/api/templates", {
        "name": "а" * 200 + ".PPTX", "data": base64.b64encode(web.make_example()).decode(),
    })
    assert status == 201
    assert json.loads(data)["name"] == "а" * 175 + ".pptx"
    assert storage.uploads[0][1] == "а" * 175 + ".pptx"


def test_partial_cache_write_is_discarded_and_download_can_retry(
    storage_workspace, fake_pipeline, monkeypatch,
):
    first, storage = storage_workspace
    template = first.add_template("first.pptx", web.make_example())
    job = complete(first, template["id"])
    directory = first.catalog.path.parent
    first.close()
    next_storage = MemoryStorage(storage.files)
    second = web.Workspace(storage=next_storage, data_dir=directory, renderer=FakeRenderer())
    fdopen = web.os.fdopen

    class PartialWriter:
        def __init__(self, descriptor, mode):
            self.stream = fdopen(descriptor, mode)

        def __enter__(self):
            return self

        def __exit__(self, *args):
            self.stream.close()

        def write(self, data):
            self.stream.write(data[:10])
            raise OSError("disk full")

    try:
        with monkeypatch.context() as patch:
            patch.setattr(web.os, "fdopen", PartialWriter)
            with pytest.raises(StorageError, match="локальную копию"):
                second.download_result(job.id)
        assert not (second.jobs[job.id].directory / "result.pptx").exists()
        assert not list(second.root.rglob(".download-*"))
        data = second.download_result(job.id)
        assert Presentation(io.BytesIO(data)).slides[0].shapes.title.text == "Сохранённый результат"
        assert len(next_storage.downloads) == 2
    finally:
        second.close()


def test_initialization_failure_releases_catalog_lock(tmp_path, monkeypatch):
    def unavailable(*args, **kwargs):
        raise OSError("no space for temp folder")

    monkeypatch.setattr(web.tempfile, "TemporaryDirectory", unavailable)
    with pytest.raises(OSError):
        web.Workspace(storage=MemoryStorage(), data_dir=tmp_path, renderer=FakeRenderer())
    WebCatalog(tmp_path).close()


def test_invalid_utf8_catalog_releases_lock_without_modification(tmp_path):
    path = tmp_path / "library.json"
    path.write_bytes(b"\xff")
    with pytest.raises(StorageError):
        WebCatalog(tmp_path)
    assert path.read_bytes() == b"\xff"
    path.write_text('{"version":1,"templates":{},"jobs":{}}', encoding="utf-8")
    WebCatalog(tmp_path).close()
