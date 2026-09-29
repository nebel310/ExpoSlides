"""Контраст итогового PPTX: фон, наследование, подложки и исходные текстовые фигуры."""

from io import BytesIO

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

from exposlides.design_builder import _contrast, build_deck
from exposlides.design_models import (
    Box,
    DeckPlan,
    PlacedBlock,
    SlideInstance,
    SlidePattern,
    TemplateProfile,
    TextStyle,
)
from exposlides.design_saved_audit import audit_saved_pptx


def _solid(fill, color):
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(color)


@pytest.mark.parametrize('kind', ['title', 'text', 'page_number'])
@pytest.mark.parametrize('surface', ['slide', 'layout', 'master', 'theme', 'rectangle',
                                     'picture', 'partial', 'own', 'native', 'light'])
def test_saved_text_has_contrast(tmp_path, surface, kind):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = Box(left=Inches(1), top=Inches(1), width=Inches(3), height=Inches(1))
    dark = '111111'
    expected_background = dark
    if surface in {'slide', 'layout', 'master', 'theme'}:
        owner = {'slide': slide, 'layout': slide.slide_layout,
                 'master': slide.slide_layout.slide_master, 'theme': slide}[surface]
        _solid(owner.background.fill, dark)
        if surface == 'theme':
            fill = owner._element.find('.//' + qn('a:solidFill'))
            fill.remove(fill[0])
            color = OxmlElement('a:schemeClr')
            color.set('val', 'tx1')
            fill.append(color)
            expected_background = '000000'
    elif surface in {'rectangle', 'partial'}:
        rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                      Inches(5 if surface == 'rectangle' else 2), Inches(3))
        _solid(rect.fill, dark)
        if surface == 'partial':
            expected_background = 'FFFFFF'
    elif surface == 'picture':
        picture = BytesIO()
        Image.new('RGB', (40, 20), (17, 17, 17)).save(picture, format='PNG')
        picture.seek(0)
        slide.shapes.add_picture(picture, 0, 0, width=Inches(5), height=Inches(3))
        expected_background = 'FFFFFF'
    elif surface == 'light':
        expected_background = 'FFFFFF'
    native = None
    if surface == 'native':
        native = slide.shapes.add_textbox(box.left, box.top, box.width, box.height)
        native.text = 'Исходный текст'
        _solid(native.fill, dark)
        run = native.text_frame.paragraphs[0].runs[0]
        run.font.color.rgb = RGBColor.from_string(dark)
        run.font.bold = True
        run.hyperlink.address = 'https://example.org'
    block = PlacedBlock(id='body', kind=kind, box=box,
                        style=TextStyle(color='FFFFFF' if surface == 'light' else dark),
                        text='Новый текст', fill=dark if surface == 'own' else None,
                        source_shape_id=native.shape_id if native is not None else None)
    plan = DeckPlan(variant_id='story', name='Test', description='',
                    width=prs.slide_width, height=prs.slide_height,
                    slides=[SlideInstance(id='s1', story_slide_id='story1',
                                          source_slide_index=1, blocks=[block])])
    source, output = tmp_path / 'source.pptx', tmp_path / 'result.pptx'
    prs.save(source)
    original = source.read_bytes()
    snapshot = plan.model_dump()
    build_deck(source, plan, output)
    reopened = Presentation(output)
    assert len(reopened.slides) == 1
    shape = reopened.slides[0].shapes[-1]
    assert shape.text == 'Новый текст'
    run = shape.text_frame.paragraphs[0].runs[0]
    assert _contrast(str(run.font.color.rgb), expected_background) >= 4.5
    if surface in {'partial', 'picture'}:
        assert shape.fill.type == MSO_FILL_TYPE.SOLID
        assert str(shape.fill.fore_color.rgb) == 'FFFFFF'
    if surface == 'native':
        assert run.font.bold
        assert run.hyperlink.address == 'https://example.org'
    assert source.read_bytes() == original
    assert plan.model_dump() == snapshot
    profile = TemplateProfile(
        template_sha256="0"*64, width=prs.slide_width, height=prs.slide_height,
        patterns=[SlidePattern(source_slide_index=1, name="Test", slots=[], content_box=box,
                               palette=[dark, "FFFFFF"], font="Arial",
                               title_style=TextStyle(), body_style=TextStyle())],
    )
    report = audit_saved_pptx(output, source, plan, profile)
    assert not [issue for issue in report.issues if issue.rule == "saved_text_style"]
    run.font.color.rgb = RGBColor.from_string("FF0000")
    reopened.save(output)
    report = audit_saved_pptx(output, source, plan, profile)
    assert any(issue.rule == "saved_text_style" for issue in report.issues)


def test_native_readable_accent_and_inherited_color(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    _solid(slide.background.fill, '111111')
    body = slide.placeholders[1]
    body.text = 'Текст'
    paragraph = body.text_frame.paragraphs[0]
    paragraph.font.color.rgb = RGBColor.from_string('111111')
    accent = paragraph.add_run()
    accent.text = ' Акцент'
    accent.font.color.rgb = RGBColor.from_string('00FFCC')
    before = accent._r.rPr.xml
    block = PlacedBlock(id='body', kind='text', source_shape_id=body.shape_id,
                        box=Box(left=body.left, top=body.top, width=body.width, height=body.height),
                        style=TextStyle(size=40), text='Основной текст с ярким акцентом')
    plan = DeckPlan(variant_id='story', name='Test', description='',
                    width=prs.slide_width, height=prs.slide_height,
                    slides=[SlideInstance(id='s1', story_slide_id='story1',
                                          source_slide_index=1, blocks=[block])])
    source, output = tmp_path / 'source.pptx', tmp_path / 'result.pptx'
    prs.save(source)
    build_deck(source, plan, output)
    runs = Presentation(output).slides[0].placeholders[1].text_frame.paragraphs[0].runs
    assert _contrast(str(runs[0].font.color.rgb), '111111') >= 4.5
    assert runs[1]._r.rPr.xml == before
