from __future__ import annotations

import asyncio
import importlib
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation as PPTXPresentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt
from pydantic import ValidationError

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
BUILDER_ROOT = REPOSITORY_ROOT / "services" / "builder-service"
PARSER_ROOT = REPOSITORY_ROOT / "services" / "parsing-service"


def _payload() -> dict:
    return {
        "slide_width": 9144000,
        "slide_height": 6858000,
        "slides": [{
            "index": 1,
            "elements": [{
                "id": "title",
                "type": "text",
                "bbox": {"left": 0, "top": 0, "width": 100, "height": 100},
                "placeholder_idx": 0,
                "placeholder_type": "CENTER_TITLE",
                "text": {
                    "full_text": "Legacy full text",
                    "paragraphs": [{
                        "text": "Legacy paragraph",
                        "runs": [{"text": "Run", "style": {"size_pt": 32}}],
                    }],
                },
            }],
        }],
    }


@pytest.mark.parametrize("version", [None, "1.0.0"])
def test_builder_keeps_explicit_legacy_text_and_placeholder_type(service_importer, version):
    models = service_importer(BUILDER_ROOT, "app.models.presentation")
    data = _payload()
    if version is not None:
        data["schema_version"] = version
    parsed = models.Presentation.model_validate(data)
    element = parsed.slides[0].elements[0]
    assert element.placeholder_type == "CENTER_TITLE"
    assert element.text.full_text == "Legacy full text"
    assert element.text.paragraphs[0].text == "Legacy paragraph"


@pytest.mark.parametrize("kind, expected", [
    ("title", "TITLE"),
    ("section_header", "TITLE"),
    ("body", "BODY"),
    ("content", "OBJECT"),
    ("subtitle", "SUBTITLE"),
    ("other", "OTHER"),
])
def test_builder_reads_v2_text_placeholder_and_theme(service_importer, kind, expected):
    models = service_importer(BUILDER_ROOT, "app.models.presentation")
    data = _payload()
    data["schema_version"] = "2.0.0"
    data["source_path"] = None
    data["tokens"] = {"theme": {"colors": {"accent1": "123456"}, "fonts": {"major": "Arial"}}}
    element = data["slides"][0]["elements"][0]
    del element["placeholder_type"]
    element["placeholder_kind"] = kind
    element["text"] = {"paragraphs": [
        {"runs": [{"text": "First ", "style": {}}, {"text": "line", "style": {}}]},
        {"runs": [{"text": "Second line", "style": {}}]},
    ]}

    parsed = models.Presentation.model_validate(data)
    assert parsed.source_path is None
    assert parsed.theme.colors == {"accent1": "123456"}
    assert parsed.theme.fonts == {"major": "Arial"}
    assert parsed.slides[0].elements[0].placeholder_type == expected
    assert parsed.slides[0].elements[0].text.full_text == "First line\nSecond line"
    assert parsed.slides[0].elements[0].text.paragraphs[0].text == "First line"


@pytest.mark.parametrize("version", ["3.0.0", 2, ""])
def test_builder_rejects_unknown_schema_versions(service_importer, version):
    models = service_importer(BUILDER_ROOT, "app.models.presentation")
    data = _payload()
    data["schema_version"] = version
    with pytest.raises(ValidationError, match="schema_version"):
        models.Presentation.model_validate(data)


def test_builder_rejects_unknown_v2_placeholder_kind(service_importer):
    models = service_importer(BUILDER_ROOT, "app.models.presentation")
    data = _payload()
    data["schema_version"] = "2.0.0"
    data["slides"][0]["elements"][0]["placeholder_kind"] = "unsupported_kind"
    with pytest.raises(ValidationError, match="Неизвестный тип placeholder v2"):
        models.Presentation.model_validate(data)


def test_builder_accepts_legacy_and_v2_table_cells_without_ignoring_invalid_text(service_importer):
    models = service_importer(BUILDER_ROOT, "app.models.presentation")
    table = models.TableElement.model_validate({
        "rows": 1, "cols": 2,
        "cells": [["Legacy", {"text": "Modern", "row_span": 2, "is_merged_origin": True}]],
    })
    assert table.cells == [["Legacy", "Modern"]]
    with pytest.raises(ValidationError):
        models.TableElement.model_validate({"rows": 1, "cols": 1, "cells": [[{"text": []}]]})


def test_v2_parser_to_builder_preserves_images_table_style_and_slide_order(tmp_path, service_importer):
    source = PPTXPresentation()
    for index in range(1, 3):
        slide = source.slides.add_slide(source.slide_layouts[5])
        title = slide.shapes.title
        title.text = f"Title {index}"
        run = title.text_frame.paragraphs[0].runs[0]
        run.font.name = "Arial"
        run.font.size = Pt(31)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string("123456")

    slide = source.slides[1]
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "First cell"
    table.cell(0, 1).text = "Second cell"
    table.cell(0, 0).fill.solid()
    table.cell(0, 0).fill.fore_color.rgb = RGBColor.from_string("ABCDEF")
    image_stream = BytesIO()
    Image.new("RGB", (8, 8), color=(20, 100, 180)).save(image_stream, "PNG")
    image_blob = image_stream.getvalue()
    slide.shapes.add_picture(BytesIO(image_blob), Inches(6), Inches(2), Inches(1), Inches(1))
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "result.pptx"
    source.save(template_path)

    parser = service_importer(PARSER_ROOT, "app.parsers.pptx")
    parsed = asyncio.run(parser.PPTXParser.parse(template_path))
    payload = parsed.presentation.model_dump(mode="json", exclude_none=True)
    assert payload["schema_version"] == "2.0.0"

    builder = service_importer(BUILDER_ROOT, "app.builder")
    models = importlib.import_module("app.models.presentation")
    content = importlib.import_module("app.models.content")
    metadata = models.Presentation.model_validate(payload)
    assert any(element.table for element in metadata.slides[1].elements)
    assert any(element.image and element.image.asset_id for element in metadata.slides[1].elements)
    generated = content.GeneratedContent(content={
        2: content.SlideContent(placeholders={"0": "Second generated"}),
        1: content.SlideContent(placeholders={"0": "First generated"}),
    })
    asyncio.run(builder.PPTXBuilder.build(template_path, metadata, generated, output_path))

    reopened = PPTXPresentation(output_path)
    assert len(reopened.slides) == 2
    assert [slide.shapes.title.text for slide in reopened.slides] == [
        "Second generated", "First generated",
    ]
    first = reopened.slides[0]
    run = first.shapes.title.text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Arial"
    assert run.font.size.pt == 31
    assert run.font.bold is True
    assert run.font.color.rgb == RGBColor.from_string("123456")
    table = next(shape.table for shape in first.shapes if shape.has_table)
    assert [cell.text for cell in table.rows[0].cells] == ["First cell", "Second cell"]
    assert table.cell(0, 0).fill.fore_color.rgb == RGBColor.from_string("ABCDEF")
    picture = next(shape for shape in first.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert picture.image.blob == image_blob
