"""Повторная сборка обновляет старые текстовые слоты без потери содержимого."""

from copy import deepcopy

import pytest
from PIL import Image
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from exposlides.design_models import DesignRequest, TemplateProfile
from exposlides.design_pipeline import DesignPipeline, save_model
from exposlides.template_profile import profile_from_json
from tests.test_design_integration import OfflineRenderer
from tests.test_template_design import _data, _element, _story


@pytest.mark.parametrize("changed_template", [False, True])
def test_build_refreshes_old_picture_text_slot(tmp_path, monkeypatch, changed_template):
    monkeypatch.setattr("exposlides.design_pipeline.PreviewRenderer", OfflineRenderer)
    monkeypatch.setattr("exposlides.design_export.shutil.which", lambda name: None)
    template = tmp_path / "template.pptx"
    image = tmp_path / "photo.png"
    Image.new("RGB", (20, 20), "blue").save(image)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(12), Inches(8)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = _data()
    elements = []
    for text, top, height, kind, size in [("Title", .4, .85, "title", 32),
                                           ("Body", 2, 4, "content", 20)]:
        shape = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(10), Inches(height))
        shape.text = text
        shape.text_frame.paragraphs[0].font.size = Pt(size)
        elements.append(_element(shape.shape_id, text, 1, top, 10, height, kind, size=size))
    picture = slide.shapes.add_picture(str(image), Inches(10), Inches(7), Inches(1), Inches(.5))
    ph = OxmlElement("p:ph")
    ph.set("type", "obj")
    picture._element.nvPicPr.nvPr.append(ph)
    fake = _element(picture.shape_id, "", 10, 7, 1, .5, "content")
    fake.update(type="shape", text=None, has_text_frame=False)
    elements.append(fake)
    data["slides"][0]["elements"] = elements
    prs.save(template)
    stale_data = deepcopy(data)
    stale_data["slides"][0]["elements"][-1].pop("has_text_frame")
    stale = profile_from_json(stale_data, template)
    assert picture.shape_id in {s.shape_id for s in stale.patterns[0].slots}
    pipeline = DesignPipeline(tmp_path / "job")
    save_model(pipeline.directory / "template.json", data)
    save_model(pipeline.directory / "profile.json", stale)
    before = (pipeline.directory / "profile.json").read_bytes()
    story = _story()
    request = DesignRequest(script=story.slides[0].paragraphs[0], mode="llm", slide_count=1)
    if changed_template:
        template.write_bytes(template.read_bytes() + b"changed")
        with pytest.raises(ValueError, match="не соответствует"):
            pipeline.build(template, request, stale, story)
        assert (pipeline.directory / "profile.json").read_bytes() == before
        return
    try:
        variants = pipeline.build(template, request, stale, story)
    finally:
        pipeline.close()
    assert len(variants) == 3
    fresh = TemplateProfile.model_validate_json((pipeline.directory / "profile.json").read_text())
    assert picture.shape_id not in {s.shape_id for s in fresh.patterns[0].slots}
    for variant in variants:
        result = Presentation(pipeline.directory / "variants" / variant["id"] / "1" / "presentation.pptx")
        assert len(result.slides) == 1
        shapes = list(result.slides[0].shapes)
        assert any(s.has_text_frame and story.slides[0].paragraphs[0] in s.text for s in shapes)
        saved_picture = next(s for s in shapes if s.shape_id == picture.shape_id)
        assert not saved_picture.has_text_frame
        assert saved_picture.image.blob == image.read_bytes()
