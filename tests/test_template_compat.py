"""Совместимость при загрузке не должна менять оригинал или скрывать другие ошибки."""

import base64
import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from exposlides.studio import PlanInput, Studio, create_app
from exposlides.template_compat import prepare_template
from exposlides.web import APIError


def template(*, broken=False, extra_damage=False, references=None):
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "Первый слайд"
    deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = "Второй слайд"
    deck.slide_layouts[0]._element.set("matchingName", 'Слайд "Спасибо!"')
    deck.slide_layouts[0]._element.cSld.set("name", 'Слайд "Спасибо!"')
    if references == "custom_show":
        etree.SubElement(deck._element, qn("p:custShowLst"))
    elif references == "section":
        extensions = etree.SubElement(deck._element, qn("p:extLst"))
        extension = etree.SubElement(extensions, qn("p:ext"), uri="sections")
        etree.SubElement(
            extension, "{http://schemas.microsoft.com/office/powerpoint/2010/main}sectionLst",
        )
    output = io.BytesIO()
    deck.save(output)
    if not broken:
        return output.getvalue()
    rewritten = io.BytesIO()
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(rewritten, "w") as target:
        for member in source.infolist():
            data = source.read(member)
            if member.filename == "ppt/slideLayouts/slideLayout1.xml":
                data = data.replace(
                    'matchingName="Слайд &quot;Спасибо!&quot;"'.encode(),
                    'matchingName="Слайд "Спасибо!""'.encode(),
                )
                if extra_damage:
                    data = data.replace(b"</p:sldLayout>", b"</p:wrong>")
            target.writestr(member, data)
    return rewritten.getvalue()


def test_valid_template_is_byte_identical():
    original = template()
    working, metadata = prepare_template(original, "template.pptx")
    assert working == original
    assert metadata["slide_count"] == 2


def test_only_known_attribute_is_changed_and_output_reopens(tmp_path):
    original = template(broken=True)
    working, metadata = prepare_template(original, "template.pptx")
    with zipfile.ZipFile(io.BytesIO(original)) as before, zipfile.ZipFile(io.BytesIO(working)) as after:
        assert before.namelist() == after.namelist()
        changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
        assert changed == ["ppt/slideLayouts/slideLayout1.xml"]
    assert working != original
    assert metadata["slide_count"] == 2
    deck = Presentation(io.BytesIO(working))
    assert deck.slide_layouts[0]._element.get("matchingName") == 'Слайд "Спасибо!"'
    path = tmp_path / "result.pptx"
    deck.save(path)
    reopened = Presentation(path)
    assert [s.shapes.title.text for s in reopened.slides] == ["Первый слайд", "Второй слайд"]


@pytest.mark.parametrize("data", [b"not a pptx", template(broken=True, extra_damage=True)])
def test_other_damage_is_rejected(data):
    with pytest.raises(APIError):
        prepare_template(data, "broken.pptx")


def test_upload_preserves_original_and_job_uses_compatible_copy_after_restart(tmp_path, monkeypatch):
    original = template(broken=True)
    app = create_app(tmp_path)
    with TestClient(app) as client:
        token = client.get("/api/session").json()["token"]
        response = client.post("/api/templates", headers={"X-Session-Token": token}, json={
            "name": "template.pptx", "data": base64.b64encode(original).decode(),
        })
        assert response.status_code == 201
        identifier = response.json()["id"]
        path = app.state.studio.template(identifier)
        assert path.read_bytes() == original
    studio = Studio(tmp_path)
    monkeypatch.setattr(studio, "_submit", lambda *args: None)
    try:
        job = studio.start(PlanInput.model_validate({
            "template_id": identifier,
            "request": {"script": "Текст презентации", "slide_count": 2, "mode": "extractive"},
        }))
        working = studio.directory(job["id"]) / "template.pptx"
        assert len(Presentation(working).slides) == 2
        assert path.read_bytes() == original
        assert working.read_bytes() == path.with_suffix(".compatible.pptx").read_bytes()
    finally:
        studio.close()


@pytest.mark.parametrize("references", ["custom_show", "section"])
@pytest.mark.parametrize("broken", [False, True])
def test_presentation_navigation_accepted_at_upload_and_after_restart(
    tmp_path, monkeypatch, references, broken,
):
    original = template(references=references, broken=broken)
    working, metadata = prepare_template(original, "template.pptx")
    assert metadata["slide_count"] == 2
    assert (working != original) is broken
    assert len(Presentation(io.BytesIO(working)).slides) == 2

    app = create_app(tmp_path / "studio")
    with TestClient(app) as client:
        token = client.get("/api/session").json()["token"]
        response = client.post("/api/templates", headers={"X-Session-Token": token}, json={
            "name": "sections.pptx", "data": base64.b64encode(original).decode(),
        })
        assert response.status_code == 201
        identifier = response.json()["id"]
        assert app.state.studio.template(identifier).read_bytes() == original

    studio = Studio(tmp_path / "studio")
    submitted = []
    monkeypatch.setattr(studio, "_submit", lambda *args: submitted.append(args))
    try:
        job = studio.start(PlanInput.model_validate({
            "template_id": identifier,
            "request": {"script": "Текст презентации", "slide_count": 2, "mode": "extractive"},
        }))
        assert submitted
        assert (studio.directory(job["id"]) / "template.pptx").read_bytes() == working
        assert studio.template(identifier).read_bytes() == original
    finally:
        studio.close()
