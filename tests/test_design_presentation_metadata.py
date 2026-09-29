"""Навигация исходного PPTX не должна мешать сборке новой последовательности."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.oxml import serialize_part_xml
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from exposlides.design_builder import build_deck
from exposlides.design_models import Box, DeckPlan, PlacedBlock, SlideInstance, TextStyle
from exposlides.design_pptx_parts import NativeBuildError

P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
UNRELATED = "urn:exposlides:metadata-regression"


def _template(path: Path, selector: str) -> DeckPlan:
    deck = Presentation()
    for index in range(1, 4):
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = f"Образец {index}"
        run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
        run.font.name, run.font.size, run.font.bold = "Arial", Pt(31), True
        run.font.color.rgb = RGBColor.from_string("123456")
        run.hyperlink.address = "https://example.com/kept"
        etree.SubElement(
            slide.shapes.title._element.find(".//" + qn("p:cNvPr")),
            qn("a:hlinkClick"), action="ppaction://customshow?id=42&return=true",
        )
        notes = slide.notes_slide.notes_text_frame
        notes.text = f"Заметки {index}"
        etree.SubElement(
            slide.notes_slide.shapes[0]._element.find(".//" + qn("p:cNvPr")),
            qn("a:hlinkHover"), action="ppaction://customshow?id=42",
        )
    etree.SubElement(
        deck.slide_layouts[0].shapes[0]._element.find(".//" + qn("p:cNvPr")),
        qn("a:hlinkClick"), action="ppaction://customshow?id=42",
    )
    ids = list(deck.slides._sldIdLst)
    shows = etree.Element(qn("p:custShowLst"))
    show = etree.SubElement(shows, qn("p:custShow"), name="Исходный показ", id="42")
    listing = etree.SubElement(show, qn("p:sldLst"))
    for identifier in ids:
        etree.SubElement(listing, qn("p:sld"), {qn("r:id"): identifier.rId})
    deck._element.insert_element_before(shows, "p:defaultTextStyle", "p:extLst")
    extensions = etree.SubElement(deck._element, qn("p:extLst"))
    extension = etree.SubElement(extensions, qn("p:ext"), uri="sections")
    sections = etree.SubElement(extension, f"{{{P14}}}sectionLst")
    section = etree.SubElement(
        sections, f"{{{P14}}}section", name="Исходный раздел",
        id="{67F85625-23A0-410C-80E6-77DA30B6F99A}",
    )
    listing = etree.SubElement(section, f"{{{P14}}}sldIdLst")
    for identifier in ids:
        etree.SubElement(listing, f"{{{P14}}}sldId", id=str(identifier.id))
    # Соседние и чужие расширения не являются навигацией PowerPoint.
    etree.SubElement(extension, f"{{{UNRELATED}}}keep", value="same-extension")
    extra = etree.SubElement(extensions, qn("p:ext"), uri="unrelated")
    etree.SubElement(extra, f"{{{UNRELATED}}}sectionLst", value="unrelated-extension")

    properties = deck.part.part_related_by(RT.PRES_PROPS)
    xml = parse_xml(properties.blob)
    show_properties = etree.SubElement(xml, qn("p:showPr"), loop="1", useTimings="0")
    etree.SubElement(show_properties, qn("p:present"))
    if selector == "custom_show":
        etree.SubElement(show_properties, qn("p:custShow"), id="42")
    else:
        etree.SubElement(show_properties, qn("p:sldRg"), st="2", end="3")
    etree.SubElement(
        etree.SubElement(show_properties, qn("p:penClr")), qn("a:srgbClr"), val="FF0000",
    )
    properties.blob = serialize_part_xml(xml)
    deck.save(path)
    return DeckPlan(
        variant_id="story", name="Собранная колода", description="Регрессия навигации",
        width=deck.slide_width, height=deck.slide_height,
        slides=[
            SlideInstance(
                id=f"slide-{position}", story_slide_id=f"story-{position}",
                source_slide_index=index, notes=f"Новые заметки {position}",
                blocks=[PlacedBlock(
                    id="generated", kind="text", text=f"Текст {position}",
                    box=Box(left=Inches(1), top=Inches(5), width=Inches(5), height=Inches(1)),
                    style=TextStyle(font="Arial", size=18, color="202124"),
                )],
            )
            for position, index in enumerate([3, 1, 3], start=1)
        ],
    )


@pytest.mark.parametrize("selector", ["custom_show", "range"])
def test_populated_navigation_is_pruned_without_changing_source_or_styles(tmp_path, selector):
    source, result = tmp_path / "source.pptx", tmp_path / "result.pptx"
    plan = _template(source, selector)
    original = source.read_bytes()
    build_deck(source, plan, result)

    reopened = Presentation(result)
    assert [slide.shapes.title.text for slide in reopened.slides] == [
        "Образец 3", "Образец 1", "Образец 3",
    ]
    for index, slide in enumerate(reopened.slides, start=1):
        run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
        assert (run.font.name, run.font.size, run.font.bold) == ("Arial", Pt(31), True)
        assert str(run.font.color.rgb) == "123456"
        assert run.hyperlink.address == "https://example.com/kept"
        assert slide.notes_slide.notes_text_frame.text == f"Новые заметки {index}"
        assert any(shape.has_text_frame and shape.text == f"Текст {index}"
                   for shape in slide.shapes)
    assert reopened._element.find(qn("p:custShowLst")) is None
    assert list(reopened._element.iter(f"{{{P14}}}sectionLst")) == []
    assert next(reopened._element.iter(f"{{{UNRELATED}}}keep")).get("value") == "same-extension"
    assert next(reopened._element.iter(f"{{{UNRELATED}}}sectionLst")).get("value") == (
        "unrelated-extension"
    )
    properties = parse_xml(reopened.part.part_related_by(RT.PRES_PROPS).blob)
    show = properties.find(qn("p:showPr"))
    assert [child.tag for child in show] == [qn("p:present"), qn("p:sldAll"), qn("p:penClr")]
    assert show.attrib == {"loop": "1", "useTimings": "0"}
    with ZipFile(result) as archive:
        assert archive.testzip() is None
        names = archive.namelist()
        assert len(names) == len(set(names))
        slides = [name for name in names if name.startswith("ppt/slides/") and name.endswith(".xml")]
        assert len(slides) == 3
        for name in names:
            if name.endswith(".xml"):
                assert b"ppaction://customshow" not in archive.read(name)
    assert source.read_bytes() == original


def test_unknown_presentation_reference_still_fails_without_overwriting_output(tmp_path):
    source, result = tmp_path / "source.pptx", tmp_path / "result.pptx"
    plan = _template(source, "custom_show")
    deck = Presentation(source)
    extension = etree.SubElement(deck._element.find(qn("p:extLst")), qn("p:ext"), uri="unknown")
    etree.SubElement(
        extension, f"{{{UNRELATED}}}slide", {qn("r:id"): deck.slides._sldIdLst[1].rId},
    )
    deck.save(source)
    original = source.read_bytes()
    result.write_bytes(b"Existing result")
    with pytest.raises(NativeBuildError, match="дополнительные ссылки"):
        build_deck(source, plan, result)
    assert result.read_bytes() == b"Existing result"
    assert source.read_bytes() == original
