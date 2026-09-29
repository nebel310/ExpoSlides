"""Переполненные подписи скрываются без потери точных данных нативного графика."""

import pytest
from pptx import Presentation
from pptx.util import Inches

from exposlides.design_builder import build_deck
from exposlides.design_models import Dataset
from tests.test_design_builder import _block, _plan, _template


def chart_case(tmp_path, values, *, kind="bar", width=10, height=4, font_size=18):
    template, output = tmp_path / "template.pptx", tmp_path / "result.pptx"
    removals = _template(template)
    dataset = Dataset(id="data", name="Значение", columns=["Категория", "Значение"],
                      rows=[[str(index), value] for index, value in enumerate(values)], source="CSV")
    block = _block("visual", "chart", dataset_id=dataset.id, chart_type=kind)
    block.box.width, block.box.height = Inches(width), Inches(height)
    block.style.size = font_size
    build_deck(template, _plan(removals, [block], count=1, datasets=[dataset]), output)
    reopened = Presentation(output)
    assert len(reopened.slides) == 1
    chart = next(shape.chart for shape in reopened.slides[0].shapes
                 if shape.name == "exposlides:visual")
    assert list(chart.series[0].values) == values
    assert [category.label for category in chart.plots[0].categories] == [
        str(index) for index in range(len(values))
    ]
    return chart


@pytest.mark.parametrize("kind", ["bar", "line"])
@pytest.mark.parametrize("width,height", [(10, 3), (5, 4)])
def test_precise_long_labels_are_hidden_instead_of_wrapping(tmp_path, kind, width, height):
    chart = chart_case(tmp_path, [1234567.89, 1234568.25, 1234569.75],
                       kind=kind, width=width, height=height)
    assert not chart.plots[0].has_data_labels
    assert chart.category_axis.tick_labels.font.size.pt >= 14


@pytest.mark.parametrize("kind", ["bar", "line"])
def test_short_labels_remain_visible_when_geometry_allows(tmp_path, kind):
    chart = chart_case(tmp_path, [12, 18, 24], kind=kind)
    assert chart.plots[0].has_data_labels
    assert chart.plots[0].data_labels.font.size.pt == 18


@pytest.mark.parametrize("kind", ["bar", "line", "pie"])
def test_dense_categories_preserve_every_value_without_overlapping_labels(tmp_path, kind):
    values = [index + 0.123456789 for index in range(1001)]
    chart = chart_case(tmp_path, values, kind=kind)
    assert not chart.plots[0].has_data_labels
    if kind == "pie":
        assert not chart.has_legend
    else:
        assert chart.category_axis.tick_labels.font.size.pt >= 14
        stride = chart.category_axis._element.xpath("./c:tickLblSkip")[0]
        assert int(stride.get("val")) > 1


@pytest.mark.parametrize("values,minimum,maximum", [
    ([1234567.89, 1234568.25], 0, None),
    ([-1234567.89, -1234568.25], None, 0),
    ([-1234567.89, 1234568.25], None, None),
    ([0, 0], 0, None),
])
def test_bar_axis_preserves_a_zero_baseline(tmp_path, values, minimum, maximum):
    chart = chart_case(tmp_path, values)
    assert chart.value_axis.minimum_scale == minimum
    assert chart.value_axis.maximum_scale == maximum


def test_million_value_ticks_reserve_space_for_long_axis_numbers(tmp_path):
    chart = chart_case(tmp_path, [1234567.89, 1234568.25, 1234569.75], width=5)
    assert 1234569.75 / chart.value_axis.major_unit <= 3


def test_short_bar_keeps_all_three_category_labels_when_each_line_fits(tmp_path):
    chart = chart_case(tmp_path, [1234567.89, 1234568.25, 1234569.75], height=2.1)
    assert not chart.plots[0].has_data_labels
    assert chart.category_axis.tick_labels.font.size.pt == 14
    for name in ("tickLblSkip", "tickMarkSkip"):
        element = chart.category_axis._element.xpath(f"./c:{name}")[0]
        assert element.get("val") == "1"


@pytest.mark.parametrize("kind", ["bar", "line", "pie"])
def test_generated_chart_disables_libreoffice_automatic_series_title(tmp_path, kind):
    chart = chart_case(tmp_path, [12, 18, 24], kind=kind)
    assert not chart.has_title
    assert chart._chartSpace.chart.xpath("./c:autoTitleDeleted")[0].get("val") == "1"


def test_rendered_story_bar_reserves_the_area_for_all_three_categories(tmp_path):
    # Размеры проблемного файла: без autoTitleDeleted LibreOffice рисовал
    # заголовок «Значение», сжимал график и пропускал B при формальном stride=1.
    chart = chart_case(tmp_path, [1234567.89, 1234568.25, 1234569.75],
                       width=849.6 / 72, height=208.90086614173228 / 72, font_size=22)
    assert chart.category_axis.tick_labels.font.size.pt == 16
    assert chart.category_axis._element.xpath("./c:tickLblSkip")[0].get("val") == "1"
    assert chart._chartSpace.chart.xpath("./c:autoTitleDeleted")[0].get("val") == "1"
