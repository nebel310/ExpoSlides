"""Привязанный текст заполняет фигуры шаблона, сохраняя их нативное оформление."""

from copy import deepcopy

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from exposlides.design_builder import build_deck
from exposlides.design_models import Box, DeckPlan, PlacedBlock, SlideInstance, TextStyle
from exposlides.design_native_text import native_reference_size, native_shapes
from exposlides.design_pptx_parts import NativeBuildError


def _xml(element):
    return etree.tostring(element, method="c14n", exclusive=True) if element is not None else b""


def _shape_design(shape):
    xml = deepcopy(shape._element)
    body = xml.find(qn("p:txBody"))
    if body is not None:
        xml.remove(body)
    return _xml(xml)


def _plan(prs, blocks, removals, count=1):
    return DeckPlan(variant_id="story", name="Native", description="Template preservation",
                    width=prs.slide_width, height=prs.slide_height,
                    slides=[SlideInstance(id=f"s{index}", story_slide_id=f"s{index}",
                                          source_slide_index=1, blocks=deepcopy(blocks),
                                          remove_shape_ids=removals) for index in range(count)])


def _block(shape, *, kind="text", size=26, items=None, text=""):
    return PlacedBlock(id=f"bound-{shape.shape_id}", kind=kind, source_shape_id=shape.shape_id,
                       # Намеренно другая геометрия и стиль: привязанный слот должен
                       # сохранить исходные свойства, кроме разрешённого уменьшения размера.
                       box=Box(left=1, top=1, width=100, height=100),
                       style=TextStyle(font="Arial", color="000000", size=size),
                       fill="FFFFFF", items=items or [], text=text)


def _template(tmp_path):
    path = tmp_path / "source.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    body = group.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
    body.name, body.rotation = "VK Education body", 5
    body.fill.solid()
    body.fill.fore_color.rgb = RGBColor.from_string("112244")
    body.line.color.rgb = RGBColor.from_string("0088FF")
    body.line.width = Pt(2)
    frame = body.text_frame
    frame.margin_left, frame.margin_right = Pt(12), Pt(9)
    frame.margin_top, frame.margin_bottom = Pt(8), Pt(6)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment, paragraph.level = PP_ALIGN.RIGHT, 1
    paragraph.space_before, paragraph.space_after = Pt(9), Pt(14)
    paragraph.line_spacing = 1.4
    paragraph.font.size = Pt(26)
    bullet = OxmlElement("a:buChar")
    bullet.set("char", "◆")
    paragraph._p.get_or_add_pPr().append(bullet)
    for text, color, size, bold in [("Исходный выделенный фрагмент ", "FFFFFF", 26, True),
                                   ("деталь", "00AACC", 18, False)]:
        run = paragraph.add_run()
        run.text = text
        run.font.name, run.font.size = "Verdana", Pt(size)
        run.font.bold, run.font.italic = bold, True
        run.font.color.rgb = RGBColor.from_string(color)
    paragraph.runs[0].hyperlink.address = "https://example.org/reference"
    paragraph = frame.add_paragraph()
    paragraph.level, paragraph.alignment = 2, PP_ALIGN.CENTER
    paragraph.font.size = Pt(22)
    paragraph.font.name = "Georgia"
    run = paragraph.add_run()
    run.text = "Второй стиль"
    run.font.size, run.font.underline = Pt(22), True
    list_style = frame._txBody.find(qn("a:lstStyle"))
    level = OxmlElement("a:lvl3pPr")
    level.set("marL", "250000")
    list_style.append(level)
    group.left += Inches(0.4)
    group.width += Inches(1)
    sibling = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8), Inches(1), Inches(1), Inches(1))
    page = slide.shapes.add_textbox(Inches(8), Inches(6), Inches(1), Inches(0.4))
    page.name = "Original page number"
    paragraph = page.text_frame.paragraphs[0]
    paragraph.font.name, paragraph.font.size = "Cambria", Pt(11)
    field = OxmlElement("a:fld")
    field.set("id", "{12345678-1234-1234-1234-123456789ABC}")
    field.set("type", "slidenum")
    text = OxmlElement("a:t")
    text.text = "1"
    field.append(text)
    paragraph._p.append(field)
    prs.save(path)
    return path, prs, group, body, sibling, page


def test_bound_native_text_preserves_groups_shape_design_and_all_style_prototypes(tmp_path):
    path, prs, group, body, sibling, page = _template(tmp_path)
    original_bytes = path.read_bytes()
    paragraphs = list(body.text_frame.paragraphs)
    source_runs = [[_xml(run._r.rPr) for run in paragraph.runs] for paragraph in paragraphs]
    texts = ["Новая мысль в фирменном оформлении", "Вторая мысль", "Дополнительная мысль"]
    plan = _plan(prs, [_block(body, items=texts), _block(page, kind="page_number", size=11, text="3")],
                 [group.shape_id, body.shape_id, page.shape_id, sibling.shape_id], count=2)
    output = build_deck(path, plan, tmp_path / "result.pptx")
    result = Presentation(output)
    for slide in result.slides:
        shapes = {shape.shape_id: shape for shape in native_shapes(slide.shapes)}
        filled = shapes[body.shape_id]
        assert filled.name == "VK Education body"
        assert len(shapes) == 3  # original group, original body and original page number
        assert sibling.shape_id not in shapes
        assert _shape_design(filled) == _shape_design(body)
        assert _xml(shapes[group.shape_id]._element.find(qn("p:grpSpPr"))) == _xml(group._element.find(qn("p:grpSpPr")))
        for tag in ("a:bodyPr", "a:lstStyle"):
            assert _xml(filled.text_frame._txBody.find(qn(tag))) == _xml(body.text_frame._txBody.find(qn(tag)))
        assert filled.text == "\n".join(texts)
        for index, paragraph in enumerate(filled.text_frame.paragraphs):
            prototype = paragraphs[min(index, 1)]
            assert _xml(paragraph._p.pPr) == _xml(prototype._p.pPr)
            if index == 0:
                assert [_xml(run._r.rPr) for run in paragraph.runs] == source_runs[0]
            else:
                # Наследуемый тёмный текст теперь исправляется на тёмной заливке.
                run = paragraph.runs[0]
                assert str(run.font.color.rgb) == "FFFFFF"
                properties = deepcopy(run._r.rPr)
                properties.remove(properties.find(qn("a:solidFill")))
                assert _xml(properties) == source_runs[1][0]
        assert filled.text_frame.paragraphs[0].runs[0].hyperlink.address == "https://example.org/reference"
        field = shapes[page.shape_id]._element.find(".//"+qn("a:fld"))
        assert field.get("type") == "slidenum"
        assert field.find(qn("a:t")).text == "3"
        assert shapes[page.shape_id].name == page.name
    assert path.read_bytes() == original_bytes


def test_native_shrink_scales_run_sizes_without_replacing_font_or_color(tmp_path):
    path, prs, _, body, _, _ = _template(tmp_path)
    assert native_reference_size(body) == 26
    plan = _plan(prs, [_block(body, size=20.8, items=["Коротко", "Ещё короче"])], [body.shape_id])
    output = build_deck(path, plan, tmp_path / "smaller.pptx")
    filled = next(shape for shape in native_shapes(Presentation(output).slides[0].shapes)
                  if shape.shape_id == body.shape_id)
    first, second = filled.text_frame.paragraphs
    assert [run.font.size.pt for run in first.runs] == [20.8, 14.4]
    assert [run.font.name for run in first.runs] == ["Verdana", "Verdana"]
    assert [str(run.font.color.rgb) for run in first.runs] == ["FFFFFF", "00AACC"]
    assert first.runs[0].font.bold and first.runs[0].font.italic
    assert first.font.size.pt == 20.8
    assert second.font.size.pt == second.runs[0].font.size.pt == 17.6
    assert second.runs[0].font.underline
    assert _shape_design(filled) == _shape_design(body)


def test_empty_placeholder_inherits_layout_and_reduces_only_font_size(tmp_path):
    path = tmp_path / "placeholder.pptx"
    prs = Presentation()
    layout = prs.slide_layouts[1]
    prototype = layout.placeholders[1].text_frame.paragraphs[0]
    prototype.font.name, prototype.font.size = "Georgia", Pt(32)
    slide = prs.slides.add_slide(layout)
    body = slide.placeholders[1]
    assert native_reference_size(body) == 32
    body_design = _shape_design(body)
    prs.save(path)
    output = build_deck(path, _plan(prs, [_block(body, size=24, text="Заполненный слот")], [body.shape_id]),
                        tmp_path / "inherited.pptx")
    filled = Presentation(output).slides[0].placeholders[1]
    run = filled.text_frame.paragraphs[0].runs[0]
    assert run.text == "Заполненный слот"
    assert run.font.size.pt == 24
    assert run.font.name is None  # Georgia remains inherited from the unchanged layout.
    assert _shape_design(filled) == body_design


@pytest.mark.parametrize("invalid", ["missing", "nontext", "duplicate"])
def test_invalid_native_binding_fails_without_publishing(tmp_path, invalid):
    path, prs, group, body, _, _ = _template(tmp_path)
    block = _block(body, text="Новая мысль")
    if invalid == "missing":
        block.source_shape_id = 99999
    if invalid == "nontext":
        block.source_shape_id = group.shape_id
    blocks = [block] if invalid != "duplicate" else [block, block.model_copy(update={"id": "second"})]
    output = tmp_path / "result.pptx"
    output.write_bytes(b"Previous output")
    with pytest.raises(NativeBuildError, match="текстовая фигура|текстовой|текстовая"):
        build_deck(path, _plan(prs, blocks, []), output)
    assert output.read_bytes() == b"Previous output"


@pytest.mark.parametrize("new_paragraphs", [["Одна мысль"], ["Первая мысль", "Вторая мысль"]])
def test_shrink_uses_font_of_emitted_prototypes_not_unused_long_small_paragraph(tmp_path, new_paragraphs):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = "Краткий заголовок"
    paragraph.runs[0].font.name, paragraph.runs[0].font.size = "Verdana", Pt(18)
    paragraph = shape.text_frame.add_paragraph()
    paragraph.text = "Подробный исходный комментарий мелким шрифтом. " * 8
    paragraph.runs[0].font.name, paragraph.runs[0].font.size = "Georgia", Pt(14)
    assert native_reference_size(shape) == 14
    assert native_reference_size(shape, paragraph_count=1) == 18
    assert native_reference_size(shape, paragraph_count=2) == 18
    assert native_reference_size(shape, paragraph_count=8) == 18
    source = tmp_path / "mixed.pptx"
    prs.save(source)
    block = _block(shape, size=14.4, items=new_paragraphs)
    output = build_deck(source, _plan(prs, [block], [shape.shape_id]), tmp_path / "filled.pptx")
    filled = Presentation(output).slides[0].shapes[0]
    assert filled.text == "\n".join(new_paragraphs)
    assert filled.text_frame.paragraphs[0].runs[0].font.size.pt == 14.4
    assert filled.text_frame.paragraphs[0].runs[0].font.name == "Verdana"
    if len(new_paragraphs) > 1:
        assert filled.text_frame.paragraphs[1].runs[0].font.size.pt == 11.2
        assert filled.text_frame.paragraphs[1].runs[0].font.name == "Georgia"
