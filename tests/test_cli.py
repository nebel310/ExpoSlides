from __future__ import annotations

import importlib
import json
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation as PPTXPresentation

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPOSITORY_ROOT))
cli = importlib.import_module("exposlides.cli")


def _write_presentation(path: Path, slide_count: int = 1) -> None:
    presentation = PPTXPresentation()
    for _ in range(slide_count):
        presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)


def _option(command: list[str], name: str) -> str:
    return command[command.index(name) + 1]


def _successful_subprocess(
    calls: list[tuple[list[str], Path, dict[str, str] | None]],
    *,
    content: dict | None = None,
    built_slide_count: int = 1,
):
    response = content or {
        "content": {"1": {"placeholders": {"0": "Готовый заголовок"}}},
        "validation_report": {"ok": True, "issues": []},
        "error": None,
    }

    def fake_run(command, *, cwd, env, check):
        assert check is False
        command = list(command)
        calls.append((command, Path(cwd), env))
        if Path(cwd).name == "parsing-service":
            Path(_option(command, "--output-json")).write_text(
                json.dumps({"slides": [{"index": 1}]}),
                encoding="utf-8",
            )
        elif Path(cwd).name == "content-service":
            Path(_option(command, "--output-json")).write_text(
                json.dumps(response, ensure_ascii=False),
                encoding="utf-8",
            )
            Path(env["LOG_FILE"]).write_text("test log", encoding="utf-8")
        elif Path(cwd).name == "builder-service":
            _write_presentation(
                Path(_option(command, "--output-pptx")),
                slide_count=built_slide_count,
            )
        else:
            raise AssertionError(f"Unexpected service cwd: {cwd}")
        return subprocess.CompletedProcess(command, 0)

    return fake_run


def _inputs(tmp_path: Path) -> tuple[Path, Path]:
    template = tmp_path / "template.pptx"
    script = tmp_path / "script.txt"
    _write_presentation(template)
    script.write_text("Текст для презентации", encoding="utf-8")
    return template, script


def test_run_pipeline_invokes_isolated_services_and_publishes_valid_pptx(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    template, script = _inputs(tmp_path)
    output = tmp_path / "nested" / "result.pptx"
    calls: list[tuple[list[str], Path, dict[str, str] | None]] = []
    monkeypatch.setattr(cli.subprocess, "run", _successful_subprocess(calls))

    result = cli.run_pipeline(
        template=template,
        script=script,
        output=output,
        max_slides=1,
    )

    assert result.output_path == output.resolve()
    assert result.artifacts_path is None
    assert len(PPTXPresentation(output).slides) == 1
    assert [call[1].name for call in calls] == [
        "parsing-service",
        "content-service",
        "builder-service",
    ]
    assert all(call[0][:3] == [sys.executable, "-m", "app.main"] for call in calls)

    parser_command, _, _ = calls[0]
    assert parser_command[3:] == [
        "--input-pptx",
        str(template.resolve()),
        "--output-json",
        _option(parser_command, "--output-json"),
    ]

    content_command, _, content_env = calls[1]
    assert content_command[3:] == [
        "--cli",
        "--template-json",
        _option(content_command, "--template-json"),
        "--script",
        str(script.resolve()),
        "--output-json",
        _option(content_command, "--output-json"),
        "--language",
        "ru",
        "--tone",
        "professional",
        "--complexity",
        "medium",
        "--max-slides",
        "1",
    ]
    assert content_env is not None
    temporary_work_directory = Path(_option(content_command, "--output-json")).parent
    assert Path(content_env["LOG_FILE"]).parent == temporary_work_directory
    assert content_env["LOG_LEVEL"] == "INFO"

    builder_command, _, _ = calls[2]
    assert builder_command[3:] == [
        "--template-pptx",
        str(template.resolve()),
        "--template-json",
        _option(parser_command, "--output-json"),
        "--content-json",
        _option(content_command, "--output-json"),
        "--output-pptx",
        _option(builder_command, "--output-pptx"),
    ]
    assert not temporary_work_directory.exists()


def test_artifacts_dir_keeps_unique_intermediate_directory(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    template, script = _inputs(tmp_path)
    artifacts_root = tmp_path / "artifacts"
    calls: list[tuple[list[str], Path, dict[str, str] | None]] = []
    monkeypatch.setattr(cli.subprocess, "run", _successful_subprocess(calls))

    result = cli.run_pipeline(
        template=template,
        script=script,
        output=tmp_path / "result.pptx",
        artifacts_dir=artifacts_root,
    )

    assert result.artifacts_path is not None
    assert result.artifacts_path.parent == artifacts_root.resolve()
    assert result.artifacts_path.name.startswith("exposlides-")
    assert (result.artifacts_path / "template.json").is_file()
    assert (result.artifacts_path / "generated_content.json").is_file()
    assert (result.artifacts_path / "content-service.log").is_file()


@pytest.mark.parametrize(
    ("content_response", "expected_message"),
    [
        ({"content": {}, "error": None}, "пустой content"),
        ({"content": {}, "error": "LLM недоступна"}, "LLM недоступна"),
    ],
)
def test_content_contract_error_stops_before_builder_and_preserves_output(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    content_response: dict,
    expected_message: str,
) -> None:
    template, script = _inputs(tmp_path)
    output = tmp_path / "result.pptx"
    output.write_bytes(b"previous result")
    calls: list[tuple[list[str], Path, dict[str, str] | None]] = []
    monkeypatch.setattr(
        cli.subprocess,
        "run",
        _successful_subprocess(calls, content=content_response),
    )

    with pytest.raises(cli.PipelineError, match=expected_message) as error:
        cli.run_pipeline(
            template=template,
            script=script,
            output=output,
            force=True,
        )

    assert error.value.stage == "content"
    assert output.read_bytes() == b"previous result"
    assert [call[1].name for call in calls] == ["parsing-service", "content-service"]


def test_nonzero_stage_exit_is_pipeline_error_and_preserves_output(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    template, script = _inputs(tmp_path)
    output = tmp_path / "result.pptx"
    output.write_bytes(b"previous result")
    calls: list[tuple[list[str], Path, dict[str, str] | None]] = []

    def fake_run(command, *, cwd, env, check):
        command = list(command)
        calls.append((command, Path(cwd), env))
        if Path(cwd).name == "parsing-service":
            Path(_option(command, "--output-json")).write_text("{}", encoding="utf-8")
            return subprocess.CompletedProcess(command, 0)
        return subprocess.CompletedProcess(command, 7)

    monkeypatch.setattr(cli.subprocess, "run", fake_run)

    with pytest.raises(cli.PipelineError, match="кодом 7") as error:
        cli.run_pipeline(
            template=template,
            script=script,
            output=output,
            force=True,
        )

    assert error.value.stage == "content"
    assert output.read_bytes() == b"previous result"
    assert [call[1].name for call in calls] == ["parsing-service", "content-service"]


@pytest.mark.parametrize("built_slide_count", [0, 2])
def test_invalid_built_slide_count_does_not_replace_existing_output(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    built_slide_count: int,
) -> None:
    template, script = _inputs(tmp_path)
    output = tmp_path / "result.pptx"
    output.write_bytes(b"previous result")
    calls: list[tuple[list[str], Path, dict[str, str] | None]] = []
    monkeypatch.setattr(
        cli.subprocess,
        "run",
        _successful_subprocess(calls, built_slide_count=built_slide_count),
    )

    with pytest.raises(cli.PipelineError, match="число слайдов") as error:
        cli.run_pipeline(
            template=template,
            script=script,
            output=output,
            force=True,
        )

    assert error.value.stage == "builder"
    assert output.read_bytes() == b"previous result"
    staged_output = Path(_option(calls[-1][0], "--output-pptx"))
    assert not staged_output.exists()


def test_corrupt_built_pptx_does_not_replace_existing_output(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    template, script = _inputs(tmp_path)
    output = tmp_path / "result.pptx"
    output.write_bytes(b"previous result")
    calls: list[tuple[list[str], Path, dict[str, str] | None]] = []

    def fake_run(command, *, cwd, env, check):
        command = list(command)
        calls.append((command, Path(cwd), env))
        if Path(cwd).name == "parsing-service":
            Path(_option(command, "--output-json")).write_text("{}", encoding="utf-8")
        elif Path(cwd).name == "content-service":
            Path(_option(command, "--output-json")).write_text(
                json.dumps({"content": {"1": {"placeholders": {}}}}),
                encoding="utf-8",
            )
        else:
            Path(_option(command, "--output-pptx")).write_bytes(b"not a pptx")
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(cli.subprocess, "run", fake_run)

    with pytest.raises(cli.PipelineError, match="PPTX/ZIP"):
        cli.run_pipeline(
            template=template,
            script=script,
            output=output,
            force=True,
        )

    assert output.read_bytes() == b"previous result"


def test_existing_output_requires_force_without_starting_services(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    template, script = _inputs(tmp_path)
    output = tmp_path / "result.pptx"
    output.write_bytes(b"existing")

    def unexpected_run(*args, **kwargs):
        raise AssertionError("subprocess must not start")

    monkeypatch.setattr(cli.subprocess, "run", unexpected_run)

    with pytest.raises(cli.PipelineError, match="--force"):
        cli.run_pipeline(template=template, script=script, output=output)

    assert output.read_bytes() == b"existing"


def test_force_replaces_existing_output_after_successful_validation(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    template, script = _inputs(tmp_path)
    output = tmp_path / "result.pptx"
    output.write_bytes(b"existing")
    calls: list[tuple[list[str], Path, dict[str, str] | None]] = []
    monkeypatch.setattr(cli.subprocess, "run", _successful_subprocess(calls))

    result = cli.run_pipeline(
        template=template,
        script=script,
        output=output,
        force=True,
    )

    assert result.output_path == output.resolve()
    assert output.read_bytes() != b"existing"
    assert len(PPTXPresentation(output).slides) == 1


def test_concurrent_output_is_not_overwritten_without_force(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    template, script = _inputs(tmp_path)
    output = tmp_path / "result.pptx"
    calls: list[tuple[list[str], Path, dict[str, str] | None]] = []
    monkeypatch.setattr(cli.subprocess, "run", _successful_subprocess(calls))
    original_link = cli.os.link

    def racing_link(source, destination):
        Path(destination).write_bytes(b"concurrent result")
        return original_link(source, destination)

    monkeypatch.setattr(cli.os, "link", racing_link)

    with pytest.raises(cli.PipelineError, match="появился во время выполнения"):
        cli.run_pipeline(template=template, script=script, output=output)

    assert output.read_bytes() == b"concurrent result"
    staged_output = Path(_option(calls[-1][0], "--output-pptx"))
    assert not staged_output.exists()


def test_output_cannot_equal_template(tmp_path: Path) -> None:
    template, script = _inputs(tmp_path)

    with pytest.raises(cli.PipelineError, match="не должен совпадать"):
        cli.run_pipeline(
            template=template,
            script=script,
            output=template,
            force=True,
        )


def test_empty_script_is_rejected_before_starting_services(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    template, script = _inputs(tmp_path)
    script.write_text("  \n", encoding="utf-8")

    def unexpected_run(*args, **kwargs):
        raise AssertionError("subprocess must not start")

    monkeypatch.setattr(cli.subprocess, "run", unexpected_run)

    with pytest.raises(cli.PipelineError, match="пуст"):
        cli.run_pipeline(
            template=template,
            script=script,
            output=tmp_path / "result.pptx",
        )


def test_main_reports_pipeline_error_without_traceback(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    missing_template = tmp_path / "missing.pptx"
    script = tmp_path / "script.txt"
    script.write_text("Текст", encoding="utf-8")

    exit_code = cli.main(
        [
            "--template",
            str(missing_template),
            "--script",
            str(script),
            "--output",
            str(tmp_path / "result.pptx"),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 1
    assert "Ошибка:" in captured.err
    assert "не найден" in captured.err
