"""Точные числа одинаково отображаются и проверяются после сохранения PPTX."""

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.util import Inches

from exposlides.design_builder import build_deck
from exposlides.design_models import (
    Box,
    Dataset,
    DeckPlan,
    PlacedBlock,
    SlideInstance,
    SlidePattern,
    TemplateProfile,
    TextStyle,
)
from exposlides.design_saved_audit import audit_saved_pptx


@pytest.mark.parametrize("kind", ["table", "chart"])
def test_saved_audit_accepts_precise_values_and_detects_rounded_data(tmp_path, kind):
    template, output = tmp_path / "template.pptx", tmp_path / "result.pptx"
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(12), Inches(8)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(template)
    box = Box(left=Inches(1), top=Inches(2), width=Inches(10), height=Inches(4))
    style = TextStyle(size=20)
    profile = TemplateProfile(
        template_sha256="0" * 64, width=presentation.slide_width,
        height=presentation.slide_height, patterns=[SlidePattern(
            source_slide_index=1, name="Data", slots=[], font="Arial",
            content_box=box, palette=["202124", "FFFFFF"],
            title_style=style, body_style=style,
        )],
    )
    dataset = Dataset(id="data", name="Выручка", source="sales.csv",
                      columns=["Период", "Значение"], rows=[
                          [1234567.89, 0.123456789012345],
                          [1e20, 1234567.89],
                          [0.123456789012345, 20.0],
                      ])
    plan = DeckPlan(
        variant_id="story", name="Data", description="Data",
        width=profile.width, height=profile.height, datasets=[dataset],
        slides=[SlideInstance(id="s1", story_slide_id="s1", source_slide_index=1,
                              blocks=[PlacedBlock(
                                  id="visual", kind=kind, dataset_id=dataset.id,
                                  chart_type="line" if kind == "chart" else None,
                                  box=box, style=style,
                              )])],
    )
    build_deck(template, plan, output)
    report = audit_saved_pptx(output, template, plan, profile)
    assert report.ok, [(issue.rule, issue.message) for issue in report.issues]
    saved = Presentation(output)
    assert len(saved.slides) == 1
    shape = saved.slides[0].shapes[0]
    if kind == "table":
        assert shape.table.cell(1, 0).text == "1234567.89"
        assert shape.table.cell(2, 0).text == "100000000000000000000"
        shape.table.cell(1, 0).text = "1.23457e+06"
    else:
        categories = [category.label for category in shape.chart.plots[0].categories]
        assert categories == ["1234567.89", "100000000000000000000", "0.123456789012345"]
        replacement = CategoryChartData()
        replacement.categories = ["1.23457e+06", *categories[1:]]
        replacement.add_series(dataset.columns[1], [row[1] for row in dataset.rows])
        shape.chart.replace_data(replacement)
    saved.save(output)
    assert "saved_data" in {
        issue.rule for issue in audit_saved_pptx(output, template, plan, profile).issues
    }
