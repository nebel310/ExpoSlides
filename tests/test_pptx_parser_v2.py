from __future__ import annotations

import asyncio
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation as PPTXPresentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

PARSING_SERVICE_ROOT = Path(__file__).resolve().parents[1] / "services" / "parsing-service"


def _parse(presentation, tmp_path: Path, service_importer):
    path = tmp_path / "styles.pptx"
    presentation.save(path)
    parser = service_importer(PARSING_SERVICE_ROOT, "app.parsers.pptx").PPTXParser
    return asyncio.run(parser.parse(path))


def _styles(result) -> dict:
    def visit(elements):
        for element in elements:
            if element.text:
                for paragraph in element.text.paragraphs:
                    for run in paragraph.runs:
                        yield run.text, run.style
            if element.group:
                yield from visit(element.group.children)

    return dict(visit(result.presentation.slides[0].elements))


def _list_font(parent, level: int, size: float, name: str) -> None:
    properties = parent.find(qn(f"a:lvl{level + 1}pPr"))
    if properties is None:
        properties = OxmlElement(f"a:lvl{level + 1}pPr")
        parent.append(properties)
    defaults = properties.find(qn("a:defRPr"))
    if defaults is None:
        defaults = OxmlElement("a:defRPr")
        properties.append(defaults)
    defaults.set("sz", str(round(size * 100)))
    latin = defaults.find(qn("a:latin"))
    if latin is None:
        latin = OxmlElement("a:latin")
        defaults.append(latin)
    latin.set("typeface", name)


def test_v2_preserves_contract_assets_and_stable_element_ids(tmp_path: Path, service_importer):
    presentation = PPTXPresentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Заголовок"
    slide.placeholders[1].text = "Содержание"
    image = BytesIO()
    Image.new("RGB", (2, 2), "blue").save(image, format="PNG")
    slide.shapes.add_picture(image, Inches(1), Inches(1), Inches(1), Inches(1))

    first = _parse(presentation, tmp_path, service_importer)
    second = _parse(presentation, tmp_path, service_importer)

    payload = first.presentation.model_dump(mode="json")
    assert payload["schema_version"] == "2.0.0"
    assert payload["tokens"]["theme"]["fonts"]
    assert payload["patterns"]
    assert payload["masters"]
    assert payload["slides"][0]["layout_index"] == 2
    assert payload["slides"][0]["elements"][0]["placeholder_kind"] == "title"
    assert "placeholder_type" not in payload["slides"][0]["elements"][0]
    assert len(first.assets) == 1
    asset = next(iter(first.assets.values()))
    assert asset.data == image.getvalue()
    assert payload["assets"][0]["asset_id"] == asset.asset_id
    assert payload == second.presentation.model_dump(mode="json")


def test_v2_explicit_fonts_override_paragraph_and_list_defaults(tmp_path: Path, service_importer):
    presentation = PPTXPresentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.font.size = Pt(32)
    paragraph.font.name = "Verdana"
    paragraph.add_run().text = "Paragraph"
    explicit = paragraph.add_run()
    explicit.text = "Explicit"
    explicit.font.size = Pt(18)
    explicit.font.name = "Calibri"
    nested = shape.text_frame.add_paragraph()
    nested.text = "List level"
    nested.level = 2
    list_style = shape._element.find(qn("p:txBody")).find(qn("a:lstStyle"))
    _list_font(list_style, 0, 50, "Arial")
    _list_font(list_style, 2, 15, "Cambria")

    styles = _styles(_parse(presentation, tmp_path, service_importer))

    assert (styles["Paragraph"].size_pt, styles["Paragraph"].font_name) == (32, "Verdana")
    assert (styles["Explicit"].size_pt, styles["Explicit"].font_name) == (18, "Calibri")
    assert (styles["List level"].size_pt, styles["List level"].font_name) == (15, "Cambria")


def test_v2_layout_font_matches_placeholder_idx_and_paragraph_level(
    tmp_path: Path, service_importer,
):
    presentation = PPTXPresentation()
    layout = presentation.slide_layouts[3]
    columns = [
        shape for shape in layout.placeholders
        if shape.placeholder_format.type == PP_PLACEHOLDER.OBJECT
    ]
    assert len(columns) == 2
    for shape, size in zip(columns, (16, 24), strict=True):
        shape.text_frame.paragraphs[0].font.size = Pt(size)
        shape.text_frame.paragraphs[0].font.name = "Arial Black"
        # Стандартный образец уже содержит пример абзаца второго уровня.
        nested = next(p for p in shape.text_frame.paragraphs if p.level == 1)
        nested.font.size = Pt(size - 2)
        nested.font.name = "Georgia"
    slide = presentation.slides.add_slide(layout)
    for index, shape in enumerate(columns):
        target = slide.placeholders[shape.placeholder_format.idx]
        target.text = f"Column {index}"
        nested = target.text_frame.add_paragraph()
        nested.text = f"Nested {index}"
        nested.level = 1

    styles = _styles(_parse(presentation, tmp_path, service_importer))

    assert styles["Column 0"].size_pt == 16
    assert styles["Column 1"].size_pt == 24
    assert styles["Column 1"].font_name == "Arial Black"
    assert styles["Nested 0"].size_pt == 14
    assert styles["Nested 1"].size_pt == 22
    assert styles["Nested 1"].font_name == "Georgia"


@pytest.mark.parametrize(
    ("target", "style_name", "size"),
    [("title", "p:titleStyle", 44), ("body", "p:bodyStyle", 27),
     ("group", "p:otherStyle", 19)],
)
def test_v2_inherits_master_fonts_including_grouped_text(
    tmp_path: Path, service_importer, target: str, style_name: str, size: int,
):
    presentation = PPTXPresentation()
    layout = presentation.slide_layouts[1]
    text_styles = layout.slide_master._element.find(qn("p:txStyles"))
    _list_font(text_styles.find(qn(style_name)), 0, size, "Arial Black")
    slide = presentation.slides.add_slide(layout)
    if target == "title":
        shape = slide.shapes.title
    elif target == "body":
        shape = slide.placeholders[1]
    else:
        group = slide.shapes.add_group_shape()
        shape = group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.text = "Inherited master font"

    style = _styles(_parse(presentation, tmp_path, service_importer))["Inherited master font"]

    assert style.size_pt == size
    assert style.font_name == "Arial Black"


def test_v2_master_placeholder_precedes_master_text_style(tmp_path: Path, service_importer):
    presentation = PPTXPresentation()
    layout = presentation.slide_layouts[1]
    master_title = next(
        shape for shape in layout.slide_master.placeholders
        if shape.placeholder_format.type == PP_PLACEHOLDER.TITLE
    )
    master_title.text_frame.paragraphs[0].font.size = Pt(39)
    master_title.text_frame.paragraphs[0].font.name = "Georgia"
    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "Master placeholder"

    style = _styles(_parse(presentation, tmp_path, service_importer))["Master placeholder"]

    assert style.size_pt == 39
    assert style.font_name == "Georgia"


def test_v2_inherited_theme_font_is_resolved(tmp_path: Path, service_importer):
    presentation = PPTXPresentation()
    layout = presentation.slide_layouts[0]
    title_style = layout.slide_master._element.find(qn("p:txStyles")).find(qn("p:titleStyle"))
    _list_font(title_style, 0, 44, "+mj-lt")
    for paragraph in layout.placeholders.get(idx=0).text_frame.paragraphs:
        paragraph.font.size = None
    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "Center title"

    result = _parse(presentation, tmp_path, service_importer)
    style = _styles(result)["Center title"]

    assert style.size_pt == 44
    assert style.font_name == result.presentation.tokens.theme.fonts["major"]


def test_v2_presentation_font_used_when_master_has_no_value(tmp_path: Path, service_importer):
    presentation = PPTXPresentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.text = "Presentation default"
    master_styles = slide.slide_layout.slide_master._element.find(qn("p:txStyles"))
    master_styles.remove(master_styles.find(qn("p:otherStyle")))
    _list_font(presentation._element.find(qn("p:defaultTextStyle")), 0, 22, "Verdana")

    style = _styles(_parse(presentation, tmp_path, service_importer))["Presentation default"]

    assert style.size_pt == 22
    assert style.font_name == "Verdana"
