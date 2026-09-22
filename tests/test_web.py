from __future__ import annotations

import base64
import http.client
import importlib
import io
import json
import sys
import threading
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
web = importlib.import_module("exposlides.web")


class FakeRenderer:
    """Ни один HTTP-тест не запускает офисные программы или дочерние процессы."""

    def __init__(self, *, available=False):
        self.available = available
        self.closed = False
        self.calls = []

    def render(self, source, directory, slide_count):
        self.calls.append(source)
        paths = []
        for index in range(1, slide_count + 1):
            image = directory / f"slide-{index}.png"
            image.write_bytes(b"\x89PNG\r\n\x1a\n" + str(index).encode())
            paths.append(image)
        return paths

    def close(self):
        self.closed = True


@pytest.fixture
def workspace():
    workspace = web.Workspace(renderer=FakeRenderer())
    yield workspace
    workspace.close()


@pytest.fixture
def server(workspace):
    return SimpleNamespace(workspace=workspace, server_port=8765)


class MemoryConnection:
    """HTTP-запрос и ответ в памяти: тесты не открывают сетевые порты."""

    def __init__(self, data: bytes):
        self.data = data
        self.output = bytearray()

    def makefile(self, mode, *args):
        return io.BytesIO(self.data)

    def sendall(self, data):
        self.output.extend(data)


def request(server, method, path, payload=None, *, token=True, headers=None):
    request_headers = dict(headers or {})
    request_headers.setdefault("Host", f"127.0.0.1:{server.server_port}")
    if token:
        request_headers["X-ExpoSlides-Token"] = server.workspace.token
    if payload is not None:
        request_headers["Content-Type"] = "application/json"
    body = json.dumps(payload).encode() if payload is not None else b""
    request_headers.setdefault("Content-Length", str(len(body)))
    head = f"{method} {path} HTTP/1.1\r\n"
    head += "".join(f"{key}: {value}\r\n" for key, value in request_headers.items())
    connection = MemoryConnection(head.encode() + b"\r\n" + body)
    web.WebHandler(connection, ("127.0.0.1", 12345), server)
    response = http.client.HTTPResponse(MemoryConnection(bytes(connection.output)))
    response.begin()
    data = response.read()
    result_headers = dict(response.getheaders())
    status = response.status
    return status, data, result_headers


def test_example_is_real_editable_pptx_and_reused_in_session(workspace):
    metadata = workspace.example()

    assert metadata == workspace.example()
    assert metadata["slide_count"] == 5
    assert metadata["width"] / metadata["height"] == pytest.approx(16 / 9, rel=0.001)
    assert metadata["slides"][0]["title"] == "Квартальный обзор"
    assert all(slide["placeholder_count"] == 2 for slide in metadata["slides"])
    presentation = Presentation(workspace.root / f'{metadata["id"]}.pptx')
    assert len(presentation.slides) == 5
    assert presentation.slides[0].shapes.title.text == "Квартальный обзор"
    assert "обратной связи" in metadata["script"]


def test_upload_returns_inspected_metadata_without_exposing_paths(server):
    status, body, _ = request(
        server,
        "POST",
        "/api/templates",
        {"name": "../../Шаблон.pptx", "data": base64.b64encode(web.make_example()).decode()},
    )

    metadata = json.loads(body)
    assert status == 201
    assert metadata["name"] == "Шаблон.pptx"
    assert metadata["slide_count"] == 5
    assert str(server.workspace.root).encode() not in body
    assert (server.workspace.root / f'{metadata["id"]}.pptx').is_file()
    assert metadata["preview"]["status"] == "unavailable"


def test_preview_returns_every_image_in_slide_order_and_guards_paths():
    renderer = FakeRenderer(available=True)
    workspace = web.Workspace(renderer=renderer)
    server = SimpleNamespace(workspace=workspace, server_port=8765)
    try:
        metadata = workspace.example()
        workspace.preview_queue.join()
        preview_url = f'/api/templates/{metadata["id"]}/preview'
        status, body, _ = request(server, "GET", preview_url)
        preview = json.loads(body)
        assert status == 200
        assert preview["status"] == "ready"
        assert len(preview["slides"]) == metadata["slide_count"]
        assert str(workspace.root).encode() not in body
        for index, url in enumerate(preview["slides"], start=1):
            status, body, headers = request(server, "GET", url)
            assert status == 200
            assert headers["Content-Type"] == "image/png"
            assert body == b"\x89PNG\r\n\x1a\n" + str(index).encode()
        assert workspace.example()["preview"] == preview
        assert len(renderer.calls) == 1
        for suffix in ("0.png", "6.png", "-1.png", "1.png/extra", "../1.png", "01.png"):
            assert request(server, "GET", preview_url.replace("preview", f"slides/{suffix}"))[0] == 404
        assert request(server, "GET", f'/api/templates/{"0" * 32}/preview')[0] == 404
        assert request(server, "GET", "/api/templates/../../.env/preview")[0] == 404
        assert request(
            server, "GET", preview["slides"][0], headers={"Origin": "https://example.org"},
        )[0] == 403
    finally:
        workspace.close()
    assert renderer.closed
    assert not workspace.root.exists()


def test_preview_upload_returns_while_renderer_is_running_and_queue_is_serial():
    started, release = threading.Event(), threading.Event()

    class BlockingRenderer(FakeRenderer):
        def render(self, source, directory, slide_count):
            started.set()
            assert release.wait(timeout=5)
            return super().render(source, directory, slide_count)

        def close(self):
            release.set()
            super().close()

    renderer = BlockingRenderer(available=True)
    workspace = web.Workspace(renderer=renderer)
    try:
        metadata = workspace.add_template("first.pptx", web.make_example())
        assert metadata["preview"] == {"status": "pending"}
        assert started.wait(timeout=3)
        second = workspace.add_template("second.pptx", web.make_example())
        assert second["preview"] == {"status": "pending"}
        assert workspace.preview_queue.qsize() == 1
        with pytest.raises(web.APIError, match="недоступен"):
            workspace.get_preview_image(metadata["id"], 1)
        release.set()
        workspace.preview_queue.join()
        assert renderer.calls == [
            workspace.root / f'{metadata["id"]}.pptx',
            workspace.root / f'{second["id"]}.pptx',
        ]
        assert workspace.get_preview(second["id"])["status"] == "ready"
    finally:
        workspace.close()


@pytest.mark.parametrize("failure", ["error", "missing", "outside"])
def test_preview_failure_hides_details_and_keeps_template_usable(failure, monkeypatch, tmp_path):
    class BrokenRenderer(FakeRenderer):
        def render(self, source, directory, slide_count):
            if failure == "error":
                raise RuntimeError("secret /private/file.pptx")
            paths = super().render(source, directory, slide_count)
            if failure == "missing":
                return paths[:-1]
            outside = tmp_path / "private.png"
            outside.write_bytes(b"private")
            return [outside, *paths[1:]]

    monkeypatch.setattr(web.Workspace, "_run_job", lambda *args: None)
    workspace = web.Workspace(renderer=BrokenRenderer(available=True))
    try:
        metadata = workspace.example()
        workspace.preview_queue.join()
        preview = workspace.get_preview(metadata["id"])
        assert preview == {"status": "failed", "message": web.PREVIEW_FAILED}
        assert not (workspace.root / "previews" / "templates" / metadata["id"]).exists()
        assert (workspace.root / f'{metadata["id"]}.pptx').exists()
        assert workspace.create_job({
            "template_id": metadata["id"], "script": "Текст",
        })["status"] == "running"
    finally:
        workspace.close()


def test_close_cancels_preview_and_skips_queued_templates():
    started, cancelled = threading.Event(), threading.Event()

    class CancelledRenderer(FakeRenderer):
        def render(self, source, directory, slide_count):
            self.calls.append(source)
            started.set()
            assert cancelled.wait(timeout=5)
            raise RuntimeError("cancelled")

        def close(self):
            cancelled.set()
            super().close()

    renderer = CancelledRenderer(available=True)
    workspace = web.Workspace(renderer=renderer)
    try:
        workspace.example()
        assert started.wait(timeout=3)
        workspace.add_template("second.pptx", web.make_example())
    finally:
        workspace.close()
    assert len(renderer.calls) == 1
    assert not workspace.preview_worker.is_alive()
    assert not workspace.root.exists()


@pytest.mark.parametrize(
    "payload",
    [
        {"name": "test.ppt", "data": "aGVsbG8="},
        {"name": "test.pptx", "data": "not base64"},
        {"name": "test.pptx", "data": "aGVsbG8="},
        {"name": "test.pptx", "data": 42},
    ],
)
def test_bad_upload_is_rejected(server, payload):
    status, body, _ = request(server, "POST", "/api/templates", payload)
    assert status == 400
    assert json.loads(body)["error"]
    assert server.workspace.templates == {}


def test_zip_limits_and_missing_placeholders_are_rejected(monkeypatch):
    data = web.make_example()
    monkeypatch.setattr(web, "MAX_UNCOMPRESSED_BYTES", 32)
    with pytest.raises(web.APIError, match="Распакованный"):
        web.inspect_template(data, "test.pptx")
    monkeypatch.setattr(web, "MAX_UNCOMPRESSED_BYTES", 150 * 1024 * 1024)
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    output = io.BytesIO()
    presentation.save(output)
    with pytest.raises(web.APIError, match="нет текстовых заполнителей"):
        web.inspect_template(output.getvalue(), "test.pptx")


def test_zip_path_traversal_is_rejected():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("[Content_Types].xml", "contents")
        archive.writestr("ppt/presentation.xml", "presentation")
        archive.writestr("../unsafe", "content")
    with pytest.raises(web.APIError, match="неподдерживаемые записи"):
        web.inspect_template(output.getvalue(), "test.pptx")


def test_session_token_origin_and_host_are_guarded(server):
    status, body, headers = request(server, "GET", "/api/session", token=False)
    assert status == 200
    assert json.loads(body)["token"] == server.workspace.token
    assert headers["Cache-Control"] == "no-store"
    assert "Access-Control-Allow-Origin" not in headers

    assert request(server, "POST", "/api/jobs", {}, token=False)[0] == 403
    assert request(server, "GET", "/api/session", headers={"Host": "example.org"})[0] == 403
    assert request(
        server, "POST", "/api/jobs", {}, headers={"Origin": "https://example.org"}
    )[0] == 403
    assert request(server, "GET", "/api/session", headers={"Origin": "null"})[0] == 403


def test_static_routes_are_explicit_and_cannot_expose_local_files(server, monkeypatch, tmp_path):
    (tmp_path / "index.html").write_text("<h1>ExpoSlides</h1>", encoding="utf-8")
    (tmp_path / "private.txt").write_text("private", encoding="utf-8")
    monkeypatch.setattr(web, "STATIC_ROOT", tmp_path)

    assert request(server, "GET", "/")[1] == b"<h1>ExpoSlides</h1>"
    for path in ("/private.txt", "/../private.txt", "/.env", "/api/jobs/../download"):
        assert request(server, "GET", path)[0] == 404


@pytest.mark.parametrize("max_slides", [0, -1, 251, "5", 2.5, True])
def test_invalid_slide_limits_never_start_process(workspace, monkeypatch, max_slides):
    monkeypatch.setattr(web.subprocess, "Popen", lambda *a, **kw: pytest.fail("unexpected process"))
    metadata = workspace.example()
    with pytest.raises(web.APIError, match="целым числом"):
        workspace.create_job({"template_id": metadata["id"], "script": "Текст", "max_slides": max_slides})


def test_blank_script_and_unknown_template_are_rejected(workspace, monkeypatch):
    monkeypatch.setattr(web.subprocess, "Popen", lambda *a, **kw: pytest.fail("unexpected process"))
    metadata = workspace.example()
    with pytest.raises(web.APIError, match="Добавьте текст"):
        workspace.create_job({"template_id": metadata["id"], "script": "  \n"})
    with pytest.raises(web.APIError, match="Шаблон не найден"):
        workspace.create_job({"template_id": "../../unknown", "script": "Текст"})


def test_slide_count_cannot_exceed_template_and_script_length_is_limited(workspace, monkeypatch):
    monkeypatch.setattr(web.subprocess, "Popen", lambda *a, **kw: pytest.fail("unexpected process"))
    metadata = workspace.example()
    with pytest.raises(web.APIError, match="В шаблоне только 5 слайдов"):
        workspace.create_job({"template_id": metadata["id"], "script": "Текст", "max_slides": 6})
    with pytest.raises(web.APIError, match="100 000"):
        workspace.create_job({"template_id": metadata["id"], "script": "а" * 100_001})


def test_shutdown_kills_descendants_even_if_parent_exited(workspace, monkeypatch):
    signals = []
    alive = True

    class Worker:
        def join(self, timeout):
            assert timeout == 5

        def is_alive(self):
            return alive

    def kill_group(pid, signal):
        nonlocal alive
        signals.append((pid, signal))
        if signal == web.signal.SIGKILL:
            alive = False

    monkeypatch.setattr(web.os, "killpg", kill_group)
    job = web.Job("id", workspace.root, "template", None)
    job.process = SimpleNamespace(pid=12345, poll=lambda: 0)
    workspace.jobs[job.id] = job
    workspace.workers.append(Worker())

    workspace.close()

    assert signals == [(12345, web.signal.SIGTERM), (12345, web.signal.SIGKILL)]
    assert not workspace.root.exists()
    job.process = None


def test_only_one_job_can_run_and_unfinished_result_cannot_download(server, monkeypatch):
    monkeypatch.setattr(web.Workspace, "_run_job", lambda *args: None)
    metadata = server.workspace.example()
    payload = {"template_id": metadata["id"], "script": "Текст"}

    status, body, _ = request(server, "POST", "/api/jobs", payload)
    job_id = json.loads(body)["id"]
    assert status == 202
    assert request(server, "POST", "/api/jobs", payload)[0] == 409
    assert request(server, "GET", f"/api/jobs/{job_id}/download")[0] == 409
    assert request(server, "GET", "/api/jobs/missing/download")[0] == 404


def test_successful_job_invokes_cli_and_downloads_verified_pptx(server, monkeypatch):
    commands = []

    class FakeProcess:
        def __init__(self, command, **kwargs):
            commands.append((command, kwargs))
            self.stdout = io.StringIO("[1/3] Парсинг\nprivate log\n[2/3] Контент\n[3/3] Сборка\n")
            Path(command[command.index("--output") + 1]).write_bytes(web.make_example())

        def wait(self):
            return 0

    monkeypatch.setattr(web.subprocess, "Popen", FakeProcess)
    metadata = server.workspace.example()
    status, body, _ = request(
        server, "POST", "/api/jobs",
        {"template_id": metadata["id"], "script": "Текст", "max_slides": 5},
    )
    assert status == 202
    job_id = json.loads(body)["id"]
    server.workspace.workers[-1].join(timeout=3)
    status, body, _ = request(server, "GET", f"/api/jobs/{job_id}")
    assert json.loads(body) == {
        "id": job_id, "status": "completed", "stage": "complete", "slide_count": 5,
    }
    assert b"private log" not in body
    command, kwargs = commands[0]
    assert command[:3] == [web.sys.executable, "-m", "exposlides"]
    assert command[command.index("--generation-mode") + 1] == "fast"
    assert command[-2:] == ["--max-slides", "5"]
    assert kwargs["cwd"] == web.REPOSITORY_ROOT
    assert kwargs["start_new_session"] is True
    status, body, headers = request(server, "GET", f"/api/jobs/{job_id}/download")
    assert status == 200
    assert len(Presentation(io.BytesIO(body)).slides) == 5
    assert "attachment" in headers["Content-Disposition"]


def changed_result() -> bytes:
    """Итог отличается от примера числом, порядком, текстом и размером слайдов."""
    presentation = Presentation()
    for title, body in (
        ("Готово: Главное", "Вывод из последнего слайда шаблона"),
        ("Готово: Что изменилось", "Новый текст для второго слайда шаблона"),
    ):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def stub_successful_cli(monkeypatch) -> None:
    result = changed_result()

    class FakeProcess:
        def __init__(self, command, **kwargs):
            self.stdout = io.StringIO("[3/3] Сборка\n")
            Path(command[command.index("--output") + 1]).write_bytes(result)

        def wait(self):
            return 0

    monkeypatch.setattr(web.subprocess, "Popen", FakeProcess)


def complete_result(workspace, template_id) -> str:
    created = workspace.create_job({"template_id": template_id, "script": "Новый исходный текст"})
    workspace.workers[-1].join(timeout=3)
    assert not workspace.workers[-1].is_alive()
    return created["id"]


class ResultRenderer(FakeRenderer):
    def render(self, source, directory, slide_count):
        paths = super().render(source, directory, slide_count)
        presentation = Presentation(source)
        assert len(presentation.slides) == slide_count
        for slide, path in zip(presentation.slides, paths, strict=True):
            path.write_bytes(b"\x89PNG\r\n\x1a\n" + slide.shapes.title.text.encode())
        return paths


def test_result_metadata_without_renderer_comes_from_output_not_template(server, monkeypatch):
    stub_successful_cli(monkeypatch)
    template = server.workspace.example()
    job_id = complete_result(server.workspace, template["id"])
    base = f"/api/jobs/{job_id}"
    status, body, _ = request(server, "GET", f"{base}/presentation")
    result = json.loads(body)

    assert status == 200
    assert result == {
        **web.inspect_template(changed_result(), "Презентация.pptx"),
        "id": job_id,
        "preview": {"status": "unavailable", "message": web.RESULT_PREVIEW_UNAVAILABLE},
    }
    assert result["slide_count"] == 2 != template["slide_count"]
    assert result["width"] != template["width"]
    assert [slide["title"] for slide in result["slides"]] == [
        "Готово: Главное", "Готово: Что изменилось",
    ]
    assert [slide["index"] for slide in result["slides"]] == [1, 2]
    assert result["slides"][0]["texts"][1] == "Вывод из последнего слайда шаблона"
    assert str(server.workspace.root).encode() not in body
    assert json.loads(request(server, "GET", f"{base}/preview")[1]) == result["preview"]
    assert request(server, "GET", f"{base}/slides/1.png")[0] == 409
    assert json.loads(request(server, "GET", base)[1]) == {
        "id": job_id, "status": "completed", "stage": "complete", "slide_count": 2,
    }
    status, data, _ = request(server, "GET", f"{base}/download")
    reopened = Presentation(io.BytesIO(data))
    assert status == 200
    assert [slide.shapes.title.text for slide in reopened.slides] == [
        slide["title"] for slide in result["slides"]
    ]


def test_completed_result_images_follow_result_order_and_never_reuse_template(monkeypatch):
    stub_successful_cli(monkeypatch)
    renderer = ResultRenderer(available=True)
    workspace = web.Workspace(renderer=renderer)
    server = SimpleNamespace(workspace=workspace, server_port=8765)
    try:
        template = workspace.example()
        workspace.preview_queue.join()
        template_preview = workspace.get_preview(template["id"])
        job_id = complete_result(workspace, template["id"])
        workspace.preview_queue.join()
        result = workspace.get_presentation(job_id)
        assert result["preview"] == {
            "status": "ready",
            "slides": [f"/api/jobs/{job_id}/slides/{index}.png" for index in (1, 2)],
        }
        for slide, url in zip(result["slides"], result["preview"]["slides"], strict=True):
            status, data, headers = request(server, "GET", url)
            assert status == 200
            assert headers["Content-Type"] == "image/png"
            assert data == b"\x89PNG\r\n\x1a\n" + slide["title"].encode()
        assert workspace.get_preview(template["id"]) == template_preview
        assert workspace.get_presentation(job_id) == result
        assert renderer.calls == [
            workspace.root / f'{template["id"]}.pptx',
            workspace.get_job(job_id).directory / "result.pptx",
        ]
        for suffix in ("0.png", "3.png", "-1.png", "1.png/extra", "../1.png", "01.png"):
            assert request(server, "GET", f"/api/jobs/{job_id}/slides/{suffix}")[0] == 404
        for route in ("presentation", "preview", "slides/1.png"):
            assert request(server, "GET", f'/api/jobs/{"0" * 32}/{route}')[0] == 404
            assert request(server, "GET", f"/api/jobs/../../.env/{route}")[0] == 404
            assert request(
                server, "GET", f"/api/jobs/{job_id}/{route}",
                headers={"Origin": "https://example.org"},
            )[0] == 403
    finally:
        workspace.close()


def test_completed_result_downloads_while_preview_waits_in_shared_serial_queue(monkeypatch):
    stub_successful_cli(monkeypatch)
    started, release = threading.Event(), threading.Event()

    class BlockingRenderer(ResultRenderer):
        def render(self, source, directory, slide_count):
            started.set()
            assert release.wait(timeout=5)
            return super().render(source, directory, slide_count)

        def close(self):
            release.set()
            super().close()

    renderer = BlockingRenderer(available=True)
    workspace = web.Workspace(renderer=renderer)
    server = SimpleNamespace(workspace=workspace, server_port=8765)
    try:
        template = workspace.example()
        assert started.wait(timeout=3)
        job_id = complete_result(workspace, template["id"])
        assert workspace.preview_queue.qsize() == 1
        base = f"/api/jobs/{job_id}"
        status, body, _ = request(server, "GET", f"{base}/presentation")
        assert status == 200
        assert json.loads(body)["preview"] == {"status": "pending"}
        assert json.loads(body)["slides"][0]["title"] == "Готово: Главное"
        assert json.loads(request(server, "GET", base)[1])["status"] == "completed"
        assert request(server, "GET", f"{base}/download")[0] == 200
        assert request(server, "GET", f"{base}/slides/1.png")[0] == 409
        release.set()
        workspace.preview_queue.join()
        assert workspace.get_job_preview(job_id)["status"] == "ready"
        assert renderer.calls == [
            workspace.root / f'{template["id"]}.pptx',
            workspace.get_job(job_id).directory / "result.pptx",
        ]
    finally:
        workspace.close()


def test_close_cancels_result_preview_and_skips_next_completed_result(monkeypatch):
    stub_successful_cli(monkeypatch)
    started, cancelled = threading.Event(), threading.Event()

    class CancelledResultRenderer(ResultRenderer):
        def render(self, source, directory, slide_count):
            if source.name != "result.pptx":
                return super().render(source, directory, slide_count)
            self.calls.append(source)
            started.set()
            assert cancelled.wait(timeout=5)
            raise RuntimeError("cancelled")

        def close(self):
            cancelled.set()
            super().close()

    renderer = CancelledResultRenderer(available=True)
    workspace = web.Workspace(renderer=renderer)
    try:
        template = workspace.example()
        workspace.preview_queue.join()
        first = complete_result(workspace, template["id"])
        assert started.wait(timeout=3)
        second = complete_result(workspace, template["id"])
        assert workspace.get_job(first).status == "completed"
        assert workspace.get_job(second).status == "completed"
        assert workspace.preview_queue.qsize() == 1
        assert workspace.get_job_preview(second) == {"status": "pending"}
        expected_sources = [
            workspace.root / f'{template["id"]}.pptx',
            workspace.get_job(first).directory / "result.pptx",
        ]
    finally:
        workspace.close()
    assert renderer.calls == expected_sources
    assert not workspace.preview_worker.is_alive()
    assert not workspace.root.exists()


@pytest.mark.parametrize("failure", ["error", "missing", "outside", "duplicate"])
def test_result_preview_failure_keeps_completed_metadata_and_download(
    monkeypatch, tmp_path, failure,
):
    stub_successful_cli(monkeypatch)

    class BrokenResultRenderer(ResultRenderer):
        def render(self, source, directory, slide_count):
            paths = super().render(source, directory, slide_count)
            if source.name != "result.pptx":
                return paths
            if failure == "error":
                raise RuntimeError("private result text /secret/path.pptx")
            if failure == "missing":
                return paths[:-1]
            if failure == "duplicate":
                return [paths[0]] * slide_count
            outside = tmp_path / "private.png"
            outside.write_bytes(b"private")
            return [outside, *paths[1:]]

    workspace = web.Workspace(renderer=BrokenResultRenderer(available=True))
    server = SimpleNamespace(workspace=workspace, server_port=8765)
    try:
        template = workspace.example()
        job_id = complete_result(workspace, template["id"])
        workspace.preview_queue.join()
        base = f"/api/jobs/{job_id}"
        status, body, _ = request(server, "GET", f"{base}/presentation")
        result = json.loads(body)
        assert status == 200
        assert result["preview"] == {"status": "failed", "message": web.RESULT_PREVIEW_FAILED}
        assert result["slide_count"] == 2
        assert result["slides"][0]["title"] == "Готово: Главное"
        assert b"private" not in body and b"/secret" not in body
        assert json.loads(request(server, "GET", base)[1])["status"] == "completed"
        assert request(server, "GET", f"{base}/download")[0] == 200
        assert request(server, "GET", f"{base}/slides/1.png")[0] == 409
        assert workspace.get_preview(template["id"])["status"] == "ready"
        assert not (workspace.root / "previews" / "jobs" / job_id).exists()
    finally:
        workspace.close()


@pytest.mark.parametrize("job_status", ["running", "failed"])
def test_incomplete_job_does_not_expose_template_as_result(server, monkeypatch, job_status):
    monkeypatch.setattr(web.Workspace, "_run_job", lambda *args: None)
    template = server.workspace.example()
    created = server.workspace.create_job({"template_id": template["id"], "script": "Текст"})
    job = server.workspace.get_job(created["id"])
    with server.workspace.lock:
        job.status = job_status
    for route in ("presentation", "preview", "slides/1.png"):
        status, body, _ = request(server, "GET", f"/api/jobs/{job.id}/{route}")
        assert status == 409
        assert "Квартальный обзор" not in body.decode()
        assert job.presentation is None


@pytest.mark.parametrize("returncode,invalid_output", [(1, False), (0, True)])
def test_failed_job_hides_raw_logs_and_does_not_allow_download(
    server, monkeypatch, returncode, invalid_output,
):
    class FakeProcess:
        def __init__(self, command, **kwargs):
            self.stdout = io.StringIO("[2/3] Контент\nsecret-api-key /private/path\n")
            if invalid_output:
                Path(command[command.index("--output") + 1]).write_bytes(b"not PPTX")

        def wait(self):
            return returncode

    monkeypatch.setattr(web.subprocess, "Popen", FakeProcess)
    metadata = server.workspace.example()
    job = server.workspace.create_job({"template_id": metadata["id"], "script": "Текст"})
    server.workspace.workers[-1].join(timeout=3)
    status, body, _ = request(server, "GET", f'/api/jobs/{job["id"]}')
    assert status == 200
    assert json.loads(body)["status"] == "failed"
    assert "error" in json.loads(body)
    assert b"secret-api-key" not in body
    assert b"/private/path" not in body
    assert request(server, "GET", f'/api/jobs/{job["id"]}/download')[0] == 409
    saved_log = server.workspace.get_job(job["id"]).directory / "job.log"
    assert "secret-api-key" in saved_log.read_text(encoding="utf-8")
    assert request(server, "GET", f'/api/jobs/{job["id"]}/job.log')[0] == 404


@pytest.mark.parametrize("code", list(web.CONTENT_ERRORS))
def test_failed_job_reports_specific_safe_message(server, monkeypatch, code):
    class FakeProcess:
        def __init__(self, command, **kwargs):
            self.stdout = io.StringIO(
                "[2/3] Контент\nprivate credentials and source\n"
                f"{web.CONTENT_ERROR_MARKER}{code}\n"
            )

        def wait(self):
            return 1

    monkeypatch.setattr(web.subprocess, "Popen", FakeProcess)
    template = server.workspace.example()
    created = server.workspace.create_job({"template_id": template["id"], "script": "Текст"})
    server.workspace.workers[-1].join(timeout=3)
    _, body, _ = request(server, "GET", f'/api/jobs/{created["id"]}')
    assert json.loads(body)["error"] == web.CONTENT_ERRORS[code]
    assert json.loads(body)["error_code"] == code
    assert b"private" not in body


def test_error_marker_accepts_only_exact_codes_in_content_stage(tmp_path):
    job = web.Job("job", tmp_path, "template", None)
    job.observe_output(f"{web.CONTENT_ERROR_MARKER}auth\n")
    assert job.error_code is None
    job.observe_output("[2/3] Контент\n")
    for line in (
        f"prefix {web.CONTENT_ERROR_MARKER}auth\n",
        f"{web.CONTENT_ERROR_MARKER}auth private response\n",
        f"{web.CONTENT_ERROR_MARKER}private-source\n",
    ):
        job.observe_output(line)
        assert job.error_code is None
    job.observe_output(f"{web.CONTENT_ERROR_MARKER}auth\n")
    assert job.error_code == "auth"
    assert "error_code" not in job.public()
    job.observe_output("[3/3] Сборка\n")
    assert job.error_code is None


def test_request_limit_rejects_before_reading_body(server):
    status, _, _ = request(
        server, "POST", "/api/templates", {},
        headers={
            "Content-Length": str(web.MAX_REQUEST_BYTES + 1),
        },
    )
    assert status == 413


def content_log(message, *, logger="app.graph.nodes", level="INFO"):
    return f"2026-09-22 15:10:00 - {logger} - {level} - {message}\n"


def test_job_reports_content_phases_and_resets_progress_between_stages(tmp_path):
    job = web.Job("job", tmp_path, "template", None)
    assert "progress" not in job.public()
    job.observe_output(content_log("=== Запуск узла analyze_script ==="))
    assert "progress" not in job.public()
    job.observe_output("[2/3] Генерация контента\n")
    for node, phase in (
        ("analyze_script", "analysis"), ("plan_slides", "planning"),
        ("generate_content", "slides"), ("validate_content", "validation"),
    ):
        job.observe_output(content_log(f"=== Запуск узла {node} ==="))
        assert job.public()["progress"] == {"phase": phase}
    job.observe_output("[3/3] Сборка итогового PPTX\n")
    assert job.stage == "builder"
    assert "progress" not in job.public()
    job.observe_output(content_log("--- Генерация контента для слайда 1/5 ---"))
    assert "progress" not in job.public()


def test_fast_job_reports_batch_phases_and_retry_without_fake_slide_counter(tmp_path):
    job = web.Job("job", tmp_path, "template", None, stage="content")
    for node, phase in (
        ("analyze_and_plan", "planning"),
        ("generate_batch", "slides"),
        ("validate_content", "validation"),
    ):
        job.observe_output(content_log(f"=== Запуск узла {node} ===", logger="app.graph.fast"))
        assert job.public()["progress"] == {"phase": phase}
        job.observe_output(content_log(
            "Повтор пакетного запроса после проверки", logger="app.graph.fast", level="WARNING",
        ))
        assert job.public()["progress"] == {"phase": phase, "retrying": True}
    job.observe_output("[3/3] Сборка итогового PPTX\n")
    assert "progress" not in job.public()


def test_fast_progress_requires_exact_message_and_logger(tmp_path):
    job = web.Job("job", tmp_path, "template", None, stage="content")
    rejected = [
        content_log("=== Запуск узла analyze_and_plan ==="),
        content_log("=== Запуск узла generate_batch ===", logger="app.chains.llm"),
        content_log("=== Запуск узла analyze_script ===", logger="app.graph.fast"),
        content_log("--- Генерация контента для слайда 1/5 ---", logger="app.graph.fast"),
        content_log(
            "Повтор пакетного запроса после проверки", logger="app.graph.fast", level="WARNING",
        ),
    ]
    for line in rejected:
        job.observe_output(line)
        assert "progress" not in job.public()
    job.observe_output(content_log("=== Запуск узла generate_batch ===", logger="app.graph.fast"))
    for line in (
        content_log("Повтор пакетного запроса после проверки", level="WARNING"),
        content_log("Повтор пакетного запроса после проверки", logger="app.graph.fast"),
        content_log(
            "Повтор пакетного запроса после проверки: private content",
            logger="app.graph.fast", level="WARNING",
        ),
        content_log(
            "Некорректный JSON-ответ LLM, повтор 1/2: private content",
            logger="app.graph.fast", level="WARNING",
        ),
    ):
        job.observe_output(line)
        assert job.public()["progress"] == {"phase": "slides"}


@pytest.mark.parametrize(
    "phase,logger,level,message",
    [
        ("analysis", "app.graph.nodes", "WARNING", "Анализ источника отклонён, повтор 1: secret"),
        ("planning", "app.graph.nodes", "WARNING", "План отклонён, повтор 2: secret"),
        ("slides", "app.graph.nodes", "WARNING", "Контент слайда 15 отклонён, повтор 1: secret"),
        ("validation", "app.graph.nodes", "INFO", "Будет повторная генерация (retry 1)"),
        (
            "analysis", "app.chains.llm", "WARNING",
            "Некорректный ответ LLM для ScriptAnalysis, повтор 1/2: secret",
        ),
        (
            "planning", "app.chains.llm", "WARNING",
            "Некорректный ответ LLM для SlidePlan, повтор 2/2: secret",
        ),
        (
            "slides", "app.chains.llm", "WARNING",
            "Некорректный JSON-ответ LLM, повтор 1/2: secret /private/path",
        ),
        *[
            (
                "slides", "app.chains.llm", "WARNING",
                f"Временная ошибка GigaChat HTTP {status}, повтор 1/2 через 1 с",
            ) for status in (500, 502, 503, 504)
        ],
    ],
)
def test_job_reports_retries_without_raw_messages(tmp_path, phase, logger, level, message):
    job = web.Job("job", tmp_path, "template", None, stage="content", progress={"phase": phase})
    if phase == "slides":
        job.progress.update(current=2, total=5)
    expected = {**job.progress, "retrying": True}
    job.observe_output(content_log(message, logger=logger, level=level))
    assert job.public()["progress"] == expected
    assert "secret" not in json.dumps(job.public())
    assert "/private/path" not in json.dumps(job.public())


def test_slide_progress_keeps_ordinal_and_total_through_retry_then_resets(tmp_path):
    job = web.Job("job", tmp_path, "template", None, stage="content")
    job.observe_output(content_log("--- Генерация контента для слайда 2/5 ---"))
    job.observe_output(content_log(
        "Контент слайда 15 отклонён, повтор 1: private content", level="WARNING",
    ))
    assert job.public()["progress"] == {
        "phase": "slides", "current": 2, "total": 5, "retrying": True,
    }
    job.observe_output(content_log("--- Генерация контента для слайда 3/5 ---"))
    assert job.public()["progress"] == {"phase": "slides", "current": 3, "total": 5}
    job.observe_output(content_log("=== Запуск узла validate_content ==="))
    assert job.public()["progress"] == {"phase": "validation"}


def test_unrelated_or_injected_lines_do_not_change_public_progress(tmp_path):
    job = web.Job(
        "job", tmp_path, "template", None, stage="content",
        progress={"phase": "slides", "current": 2, "total": 5},
    )
    expected = job.public()
    marker = "=== Запуск узла analyze_script ==="
    rejected = [
        marker,
        f"source text: {content_log(marker)}",
        content_log(marker, logger="app.main"),
        content_log(marker, logger="app.chains.llm"),
        content_log(marker, level="DEBUG"),
        content_log(marker, level="WARNING"),
        content_log(marker + " extra source text"),
        content_log("Результат валидации: ok=False, issues=['private data']"),
        content_log("Некорректный ответ: " + marker, level="WARNING"),
        content_log("=== Запуск узла unexpected ==="),
        content_log("--- Генерация контента для слайда 0/5 ---"),
        content_log("--- Генерация контента для слайда 6/5 ---"),
        content_log("--- Генерация контента для слайда 1/251 ---"),
        content_log("--- Генерация контента для слайда 01/5 ---"),
        content_log("Анализ источника отклонён, повтор 1: private", level="WARNING"),
        content_log("Контент слайда 999 отклонён, повтор 1: private", level="WARNING"),
        content_log(
            "Некорректный JSON-ответ LLM, повтор 3/2: private",
            logger="app.chains.llm", level="WARNING",
        ),
        content_log(
            "Временная ошибка GigaChat HTTP 429, повтор 1/2 через 1 с",
            logger="app.chains.llm", level="WARNING",
        ),
        content_log(
            "Некорректный JSON-ответ LLM, повтор 1/2: private",
            logger="app.graph.nodes", level="WARNING",
        ),
        "[2/3] Генерация контента\n",
    ]
    for line in rejected:
        job.observe_output(line)
        assert job.public() == expected, repr(line)


def test_retry_does_not_invent_a_phase_before_known_marker(tmp_path):
    job = web.Job("job", tmp_path, "template", None, stage="content")
    job.observe_output(content_log(
        "Временная ошибка GigaChat HTTP 503, повтор 1/2 через 1 с",
        logger="app.chains.llm", level="WARNING",
    ))
    assert "progress" not in job.public()


def test_job_worker_reads_live_progress_and_clears_it_after_success(workspace, monkeypatch):
    snapshots = []

    class Output:
        def __iter__(self):
            for line in (
                "[2/3] Генерация контента\n",
                content_log("=== Запуск узла analyze_script ==="),
                content_log("--- Генерация контента для слайда 2/5 ---"),
                content_log(
                    "Некорректный JSON-ответ LLM, повтор 1/2: private value",
                    logger="app.chains.llm", level="WARNING",
                ),
                "[3/3] Сборка итогового PPTX\n",
            ):
                yield line
                snapshots.append(next(iter(workspace.jobs.values())).public())

        def close(self):
            pass

    class FakeProcess:
        def __init__(self, command, **kwargs):
            self.stdout = Output()
            Path(command[command.index("--output") + 1]).write_bytes(web.make_example())

        def wait(self):
            return 0

    monkeypatch.setattr(web.subprocess, "Popen", FakeProcess)
    metadata = workspace.example()
    job = workspace.create_job({"template_id": metadata["id"], "script": "Текст"})
    workspace.workers[-1].join(timeout=3)
    assert [snapshot.get("progress") for snapshot in snapshots] == [
        None,
        {"phase": "analysis"},
        {"phase": "slides", "current": 2, "total": 5},
        {"phase": "slides", "current": 2, "total": 5, "retrying": True},
        None,
    ]
    completed = workspace.get_job(job["id"]).public()
    assert completed["status"] == "completed"
    assert "progress" not in completed
