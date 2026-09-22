from __future__ import annotations

import importlib
import io
import sys
from pathlib import Path
from unittest.mock import Mock

import pytest
from pptx import Presentation
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
ui = importlib.import_module("exposlides.streamlit_ui")
web = importlib.import_module("exposlides.web")
PresentationSession = importlib.import_module("exposlides.ui_session").PresentationSession


class NoRenderer:
    available = False

    def close(self):
        pass


def pptx_bytes(title="Исходный шаблон", count=3):
    deck = Presentation()
    for index in range(1, count + 1):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"{title} {index}"
        slide.placeholders[1].text = f"Содержание слайда {index}"
    data = io.BytesIO()
    deck.save(data)
    return data.getvalue()


@pytest.fixture
def app(monkeypatch):
    workspace = web.Workspace(renderer=NoRenderer())
    # Настоящие загрузки и состояние; процессы генерации и конвертации отключены.
    monkeypatch.setattr(workspace, "_run_job", lambda job: None)
    session = PresentationSession(workspace=workspace)
    monkeypatch.setattr(ui, "get_session", lambda: session)
    test = AppTest.from_file(ROOT / "streamlit_app.py", default_timeout=5).run()
    assert not test.exception
    yield test, session, workspace
    session.close()


def html_body(test):
    return "\n".join(element.proto.body for element in test.get("html"))


def upload_template(test, data=None):
    test.file_uploader[0].set_value((
        "Шаблон.pptx", data if data is not None else pptx_bytes(),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )).run()
    assert not test.exception


def start_example(test, workspace):
    test.button(key="example").click().run()
    test.button(key="generate").click().run()
    assert not test.exception
    return next(reversed(workspace.jobs.values()))


def complete(workspace, job, *, preview="pending"):
    result = pptx_bytes("Готовая презентация", count=2)
    (job.directory / "result.pptx").write_bytes(result)
    workspace._complete_job(job)
    with workspace.lock:
        job.presentation["preview"] = {"status": preview}
    return result


def test_empty_form_requires_template_and_nonempty_text(app):
    test, _, workspace = app
    assert test.button(key="generate").disabled
    assert test.selectbox[0].disabled
    assert "Здесь будут слайды" in html_body(test)

    test.text_area(key="script").set_value("Доклад").run()
    assert test.button(key="generate").disabled
    upload_template(test)
    assert not test.button(key="generate").disabled
    test.text_area(key="script").set_value(" \n ").run()
    assert test.button(key="generate").disabled
    assert not workspace.jobs


def test_example_loads_materials_without_starting_generation(app):
    test, session, workspace = app
    test.button(key="example").click().run()

    snapshot = session.snapshot()
    assert not test.exception
    assert test.text_area(key="script").value == web.EXAMPLE_SCRIPT
    assert snapshot["view"] == "template"
    assert len(workspace.templates) == 1
    assert not workspace.jobs
    assert not test.button(key="generate").disabled
    assert not test.download_button
    assert 'aria-label="Сборка слайдов"' not in html_body(test)


def test_upload_and_text_reruns_preserve_edited_materials_without_jobs(app):
    test, session, workspace = app
    upload_template(test)
    template_id = session.snapshot()["template"]["id"]
    test.file_uploader[1].set_value(("Доклад.txt", b"\xef\xbb\xbfImported", "text/plain"))
    test.run()
    assert test.text_area(key="script").value == "Imported"
    test.text_area(key="script").set_value("Отредактированный текст").run()
    test.selectbox[0].select(2).run()
    test.run()

    assert not test.exception
    assert test.text_area(key="script").value == "Отредактированный текст"
    assert session.snapshot()["template"]["id"] == template_id
    assert len(workspace.templates) == 1
    assert not workspace.jobs


def test_generate_uses_selected_limit_once_and_locks_materials(app):
    test, _, workspace = app
    upload_template(test)
    test.text_area(key="script").set_value("Текст для двух слайдов").run()
    test.selectbox[0].select(2).run()
    test.button(key="generate").click().run()
    test.run()
    test.run()

    assert not test.exception
    assert len(workspace.jobs) == 1
    job = next(iter(workspace.jobs.values()))
    assert job.max_slides == 2
    assert (job.directory / "script.txt").read_text(encoding="utf-8") == "Текст для двух слайдов"
    assert test.button(key="generate").disabled
    assert test.button(key="example").disabled
    assert test.text_area(key="script").disabled
    assert all(upload.disabled for upload in test.file_uploader)
    assert test.selectbox[0].disabled
    assert 'aria-label="Сборка слайдов"' in html_body(test)
    assert not test.download_button


def test_running_preview_can_show_template_without_restarting(app):
    test, _, workspace = app
    job = start_example(test, workspace)
    with workspace.lock:
        job.stage = "content"
        job.progress = {"phase": "slides", "retrying": True}
    test.run()
    assert any("Уточняем текст после проверки" in caption.value for caption in test.caption)

    test.toggle[0].set_value(True).run()
    assert 'aria-label="Сборка слайдов"' not in html_body(test)
    assert "ТЕКСТ СЛАЙДА" in html_body(test)
    test.toggle[0].set_value(False).run()
    assert 'aria-label="Сборка слайдов"' in html_body(test)
    assert len(workspace.jobs) == 1
    assert not test.exception


def test_failed_job_keeps_materials_and_allows_explicit_retry(app):
    test, _, workspace = app
    failed = start_example(test, workspace)
    script = test.text_area(key="script").value
    with workspace.lock:
        failed.status = "failed"
        failed.error = web.CONTENT_ERRORS["network"]
    test.run()

    assert not test.exception
    assert [error.value for error in test.error] == [web.CONTENT_ERRORS["network"]]
    assert test.text_area(key="script").value == script
    assert not test.button(key="generate").disabled
    assert not test.download_button
    assert str(workspace.root) not in html_body(test)
    test.button(key="generate").click().run()
    assert not test.error
    assert not test.exception
    assert len(workspace.jobs) == 2
    assert next(reversed(workspace.jobs.values())).id != failed.id


@pytest.mark.parametrize("status", ["completed", "failed"])
def test_job_finishing_during_materials_render_unlocks_form(app, monkeypatch, status):
    test, session, workspace = app
    job = start_example(test, workspace)
    original = ui._materials
    transitioned = False

    def finish_after_materials(session, snapshot):
        nonlocal transitioned
        original(session, snapshot)
        if transitioned:
            return
        transitioned = True
        if status == "completed":
            complete(workspace, job, preview="unavailable")
        else:
            with workspace.lock:
                job.status = "failed"
                job.error = web.CONTENT_ERRORS["network"]

    monkeypatch.setattr(ui, "_materials", finish_after_materials)
    test.run()

    assert not test.exception
    assert session.snapshot()["job"]["status"] == status
    assert not test.button(key="generate").disabled
    assert not test.button(key="example").disabled
    assert not test.text_area(key="script").disabled
    assert all(not upload.disabled for upload in test.file_uploader)
    assert not test.selectbox[0].disabled
    assert len(workspace.jobs) == 1
    assert bool(test.download_button) == (status == "completed")


def test_start_validation_error_stays_in_form_and_can_be_retried(app, monkeypatch):
    test, session, workspace = app
    test.button(key="example").click().run()
    original = session.start
    monkeypatch.setattr(session, "start", Mock(side_effect=web.APIError("Попробуйте ещё раз.")))
    test.button(key="generate").click().run()
    assert [error.value for error in test.error] == ["Попробуйте ещё раз."]
    assert not test.exception
    assert not workspace.jobs

    monkeypatch.setattr(session, "start", original)
    test.button(key="generate").click().run()
    assert not test.error
    assert len(workspace.jobs) == 1


def test_completed_result_is_available_before_images_and_replaces_template(app, monkeypatch):
    test, session, workspace = app
    job = start_example(test, workspace)
    result = complete(workspace, job)
    download = Mock(wraps=session.download)
    monkeypatch.setattr(session, "download", download)
    test.run()

    assert not test.exception
    assert "Готовая презентация 1" in html_body(test)
    assert any(heading.value == "Готовая презентация" for heading in test.subheader)
    assert any("Готовим изображения" in caption.value for caption in test.caption)
    assert test.selectbox[-1].options == ["01 / 02", "02 / 02"]
    assert len(test.download_button) == 1
    assert test.download_button[0].proto.url
    download.assert_called_once_with()
    assert session.download() == result
    test.download_button[0].click().run()
    assert not test.exception
    assert len(workspace.jobs) == 1


def test_result_navigation_stays_inside_result_and_keeps_download(app):
    test, session, workspace = app
    job = start_example(test, workspace)
    complete(workspace, job, preview="failed")
    test.run()
    key = f"slide_result_{job.id}"
    assert test.button(key=f"prev_{key}").disabled
    assert not test.button(key=f"next_{key}").disabled

    test.button(key=f"next_{key}").click().run()
    assert test.selectbox(key=key).value == 2
    assert "Готовая презентация 2" in html_body(test)
    assert test.button(key=f"next_{key}").disabled
    test.button(key=f"prev_{key}").click().run()
    assert test.selectbox(key=key).value == 1
    test.selectbox(key=key).select(2).run()
    assert "Готовая презентация 2" in html_body(test)
    assert session.snapshot()["view"] == "result"
    assert len(test.download_button) == 1
    assert len(workspace.jobs) == 1
    assert not test.exception


def test_missing_slide_image_uses_result_text_and_keeps_download(app, monkeypatch):
    test, session, workspace = app
    job = start_example(test, workspace)
    complete(workspace, job, preview="ready")
    monkeypatch.setattr(session, "get_slide", Mock(side_effect=web.APIError("Недоступно")))
    test.run()

    assert not test.exception
    assert "Готовая презентация 1" in html_body(test)
    assert any("Изображение слайда недоступно" in caption.value for caption in test.caption)
    assert len(test.download_button) == 1


def test_missing_result_file_shows_safe_error_and_keeps_text_preview(app):
    test, _, workspace = app
    job = start_example(test, workspace)
    complete(workspace, job)
    (job.directory / "result.pptx").unlink()
    test.run()

    assert not test.exception
    assert any("Не удалось открыть презентацию" in error.value for error in test.error)
    assert all(str(workspace.root) not in error.value for error in test.error)
    assert "Готовая презентация 1" in html_body(test)
    assert not test.download_button
    assert not test.button(key="generate").disabled
    assert len(workspace.jobs) == 1


def test_invalid_template_disables_generation_until_removed(app):
    test, session, workspace = app
    upload_template(test)
    test.text_area(key="script").set_value("Доклад").run()
    original_id = session.snapshot()["template"]["id"]
    upload_template(test, b"not a pptx")
    assert test.error
    assert test.button(key="generate").disabled
    assert session.snapshot()["template"]["id"] == original_id

    test.file_uploader[0].clear().run()
    assert not test.error
    assert session.snapshot()["template"] is None
    assert test.button(key="generate").disabled
    assert not workspace.jobs


@pytest.mark.parametrize("data, message", [
    (b"\xff\xfe\x80", "Сохраните файл в UTF-8"),
    (("я" * (web.MAX_SCRIPT_LENGTH + 1)).encode("utf-8"), "Текст слишком длинный"),
    (b"x" * (web.MAX_SCRIPT_LENGTH * 4 + 1), "Текст слишком длинный"),
])
def test_invalid_text_import_preserves_existing_text(app, data, message):
    test, _, workspace = app
    test.text_area(key="script").set_value("Сохранённый доклад").run()
    test.file_uploader[1].set_value(("Доклад.txt", data, "text/plain")).run()

    assert not test.exception
    assert test.text_area(key="script").value == "Сохранённый доклад"
    assert any(message in error.value for error in test.error)
    assert not workspace.jobs
    test.file_uploader[1].set_value(("Новый.txt", "Новый доклад".encode(), "text/plain")).run()
    assert not test.error
    assert test.text_area(key="script").value == "Новый доклад"
