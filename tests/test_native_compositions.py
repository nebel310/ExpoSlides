"""Регрессии: весь набор пригодных макетов и независимые фигурные фото в layout."""

from copy import deepcopy

import pytest
from PIL import Image
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

from exposlides.design_audit import audit_deck, audit_saved_pptx
from exposlides.design_builder import build_deck
from exposlides.design_layout import create_variants
from exposlides.design_models import ContentPlan, StorySlide, TemplateProfile
from exposlides.template_layout import native_layout
from exposlides.template_profile import profile_from_json
from tests.test_design_images_batch import _photo_template
from tests.test_template_design import _profile, _story


def _layout_photo_template(tmp_path):
    template, _, originals = _photo_template(tmp_path)
    presentation = Presentation(template)
    slide = presentation.slides[0]
    layout = slide.slide_layout
    for shape in list(layout.shapes):
        shape._element.getparent().remove(shape._element)
    photo_id, logo_id = originals["Photo"][0], originals["Logo"][0]
    photo = next(shape for shape in slide.shapes if shape.shape_id == photo_id)
    for shape in list(slide.shapes):
        if shape.shape_id not in {photo_id, logo_id}:
            continue
        # IDs могут совпадать с фигурами слайда, но область источника однозначна.
        image_part, rid = layout.part.get_or_add_image_part(str(tmp_path / f'{shape.name}.png'))
        element = deepcopy(shape._element)
        element.find('.//' + qn('a:blip')).set(qn('r:embed'), rid)
        if shape.shape_id == photo_id:
            properties = element.find(qn('p:spPr'))
            geometry = properties.find(qn('a:prstGeom'))
            if geometry is not None:
                properties.remove(geometry)
            properties.append(parse_xml(
                '<a:custGeom ' + nsdecls('a') + '><a:avLst/><a:gdLst/><a:ahLst/>'
                '<a:cxnLst/><a:rect l="l" t="t" r="r" b="b"/><a:pathLst>'
                '<a:path w="100" h="100"><a:moveTo><a:pt x="0" y="0"/></a:moveTo>'
                '<a:lnTo><a:pt x="100" y="0"/></a:lnTo>'
                '<a:lnTo><a:pt x="75" y="100"/></a:lnTo><a:close/>'
                '</a:path></a:pathLst></a:custGeom>'
            ))
        layout.shapes._spTree.insert_element_before(element, 'p:extLst')
        shape._element.getparent().remove(shape._element)
    presentation.save(template)
    # Реальный parser проверяется отдельно; здесь минимальный сериализованный контракт.
    from tests.test_template_design import _element
    elements = [
        _element(s.shape_id, s.text, s.left/914400, s.top/914400,
                 s.width/914400, s.height/914400,
                 'title' if index == 0 else None, size=32 if index == 0 else 20)
        for index, s in enumerate(slide.shapes)
    ]
    inherited = [dict(id=f'layout-shape-{s.shape_id}', shape_id=s.shape_id,
                      type='image', shape_name=s.name,
                      bbox=dict(left=s.left, top=s.top, width=s.width, height=s.height))
                 for s in layout.shapes if s.shape_id in {photo_id, logo_id}]
    profile = profile_from_json(dict(
        schema_version='2.0.0', slide_width=presentation.slide_width,
        slide_height=presentation.slide_height,
        slides=[dict(index=1, layout_index=1, elements=elements)],
        layouts=[dict(index=1, elements=inherited)],
    ), template)
    assert set(profile.patterns[0].layout_images) == {photo.shape_id}
    assert not profile.patterns[0].replaceable_images
    return template, profile, originals


def _build_masked(tmp_path):
    template, profile, originals = _layout_photo_template(tmp_path)
    story = ContentPlan(title='Reading', slides=[
        StorySlide(id=f'slide-{i}', title=f'Topic {i}', paragraphs=[f'Point {i}'],
                   source_ids=['source-1']) for i in range(2)
    ])
    images = {}
    for index, slide in enumerate(story.slides):
        path = tmp_path / f'new-{index}.png'
        Image.new('RGB', (200, 100), ['red', 'blue'][index]).save(path)
        images[slide.id] = path
    plan = create_variants(profile, story, generated_images=images, template_images_only=True)[0]
    output = build_deck(template, plan, tmp_path / 'result.pptx')
    return template, profile, originals, plan, output, images


def test_layout_photo_keeps_mask_and_logo_with_independent_images(tmp_path):
    template, profile, originals, plan, output, images = _build_masked(tmp_path)
    source, result = Presentation(template), Presentation(output)
    photo_id, logo_id = originals['Photo'][0], originals['Logo'][0]
    original = next(s for s in source.slides[0].slide_layout.shapes if s.shape_id == photo_id)
    assert len(result.slides) == 2
    assert result.slides[0].slide_layout.part is not result.slides[1].slide_layout.part
    for slide, planned in zip(result.slides, plan.slides, strict=True):
        photos = {s.shape_id: s for s in slide.slide_layout.shapes if s.shape_type == 13}
        photo = photos[photo_id]
        assert photo.image.blob == images[planned.story_slide_id].read_bytes()
        assert photo.image.blob != originals['Photo'][1]
        assert photos[logo_id].image.blob == originals['Logo'][1]
        assert photo._element.find(qn('p:spPr')).xml == original._element.find(qn('p:spPr')).xml
        assert photo.crop_left == photo.crop_right == .25
        assert (photo.left, photo.top, photo.width, photo.height) == (
            original.left, original.top, original.width, original.height,
        )
        assert slide.slide_layout in list(slide.slide_layout.slide_master.slide_layouts)
    assert original.image.blob == originals['Photo'][1]
    assert audit_saved_pptx(output, template, plan, profile).ok
    assert audit_deck(plan, profile).ok
    payload = profile.model_dump()
    payload['patterns'][0].pop('layout_images')
    assert TemplateProfile.model_validate(payload).patterns[0].layout_images == {}


@pytest.mark.parametrize('damage,rule', [('mask','native_image_mask'), ('photo','image_content'),
                                        ('crop','image_aspect'), ('logo','template_layout')])
def test_audit_detects_native_layout_image_corruption(tmp_path, damage, rule):
    template, profile, originals, plan, output, _ = _build_masked(tmp_path)
    result = Presentation(output)
    layout = result.slides[0].slide_layout
    photo = next(s for s in layout.shapes if s.shape_id == originals['Photo'][0])
    if damage == 'mask':
        point = photo._element.find('.//' + qn('a:pt'))
        point.set('x', '50')
    elif damage == 'crop':
        photo.crop_left = 0
    else:
        target = photo if damage == 'photo' else next(s for s in layout.shapes if s.shape_id == originals['Logo'][0])
        _, rid = layout.part.get_or_add_image_part(str(tmp_path/'Photo.png'))
        target._element.find('.//' + qn('a:blip')).set(qn('r:embed'), rid)
    result.save(output)
    assert rule in {issue.rule for issue in audit_saved_pptx(output, template, plan, profile).issues}


def test_unused_fitting_compositions_win_before_cost_driven_repeats(tmp_path):
    profile = _profile(tmp_path)
    for index in range(2, 7):
        pattern = profile.patterns[0].model_copy(deep=True)
        pattern.source_slide_index = index
        pattern.slots[1].box.width -= index * 200000
        # Несколько маленьких пустых служебных областей не означают переполнение.
        extra = pattern.slots[1].model_copy(deep=True)
        extra.shape_id = 90 + index
        extra.box.top = 6600000
        extra.box.height = 50000
        pattern.slots.append(extra)
        profile.patterns.append(pattern)
    story = _story()
    story.slides = [story.slides[0].model_copy(update={'id': f'slide-{i}'}) for i in range(6)]
    for plan in create_variants(profile, story):
        assert len({s.source_slide_index for s in plan.slides}) == 6
    # Нельзя ради уникальности переходить на переполняющийся слайд.
    profile.patterns[-1].slots[1].box.height = 50000
    chosen = native_layout(profile, story.slides[0], 'story', 6, tuple(range(1,6)))
    assert chosen[0].source_slide_index != 6


def test_audit_detects_four_layout_cycle_with_different_text(tmp_path):
    profile = _profile(tmp_path)
    for index in range(2,5):
        pattern = profile.patterns[0].model_copy(deep=True)
        pattern.source_slide_index = index
        profile.patterns.append(pattern)
    plan = create_variants(profile, _story())[0]
    sample = plan.slides[0]
    plan.slides = []
    for index, source in enumerate([1,2,3,4,1,2,3,4]):
        slide = sample.model_copy(deep=True)
        slide.id = f'slide-{index}'
        slide.story_slide_id = slide.id
        slide.source_slide_index = source
        slide.blocks[1].items = [f'Different content {index}']
        plan.slides.append(slide)
    assert len([i for i in audit_deck(plan,profile).issues if i.rule == 'layout_cycle']) == 1


def test_model_receives_template_block_counts_without_sample_content(tmp_path, service_importer):
    import asyncio
    from pathlib import Path

    from exposlides.design_models import DesignRequest

    module = service_importer(Path(__file__).resolve().parents[1]/'services/content-service', 'app.design_main')
    profile = _profile(tmp_path)
    pattern = profile.patterns[0].model_copy(deep=True)
    pattern.source_slide_index = 2
    extra = pattern.slots[1].model_copy(deep=True)
    extra.shape_id = 90
    extra.box.left += extra.box.width // 2
    extra.box.width //= 2
    pattern.slots[1].box.width //= 2
    pattern.slots.append(extra)
    profile.patterns.append(pattern)
    profile.patterns[0].slots[1].text = 'SECRET SAMPLE 999%'
    source = 'Команда развивает платформу. Пользователи получают поддержку.'

    class Client:
        calls = 0

        async def generate_json(self, prompt, model):
            self.calls += 1
            assert 'paragraph_counts' in prompt and 'SECRET SAMPLE' not in prompt
            assert '999%' not in prompt
            return model.model_validate(dict(title='Платформа', slides=[
                dict(id='slide-1',title='Платформа',paragraphs=[source],source_ids=['source-1']),
                dict(id='slide-2',title='Поддержка',paragraphs=['Пользователи получают поддержку.'],source_ids=['source-1']),
            ]))
    client = Client()
    result = asyncio.run(module.generate(DesignRequest(script=source,slide_count=2),client,profile=profile))
    assert len(result.slides) == 2 and client.calls == 1


def test_complete_layout_wins_over_unused_half_empty_cards(tmp_path):
    profile = _profile(tmp_path)
    empty = profile.patterns[0].model_copy(deep=True)
    empty.source_slide_index = 2
    second = empty.slots[1].model_copy(deep=True)
    second.shape_id = 90
    second.box.top = 5700000
    second.box.height = 1000000
    empty.slots[1].box.height = 3000000
    empty.slots.append(second)
    profile.patterns.append(empty)
    selected = native_layout(profile, _story().slides[0], 'story', 1, (1,))
    assert selected[0].source_slide_index == 1


def test_small_template_does_not_repeat_the_same_three_layout_cycle(tmp_path):
    profile = _profile(tmp_path)
    for index in (2,3):
        pattern = profile.patterns[0].model_copy(deep=True)
        pattern.source_slide_index = index
        profile.patterns.append(pattern)
    story = _story()
    story.slides = [story.slides[0].model_copy(update={'id': f'slide-{i}'}) for i in range(12)]
    for plan in create_variants(profile,story):
        assert not [i for i in audit_deck(plan,profile).issues if i.rule == 'layout_cycle']
