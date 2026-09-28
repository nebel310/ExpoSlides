"""Integration boundaries use actual PPTX and fake external render/model services."""

import json
import threading

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from exposlides.design_audit import apply_fixes, audit_deck
from exposlides.design_content import extractive_plan, source_excerpts, validate_story
from exposlides.design_models import (
    Box,
    ContentPlan,
    Dataset,
    DeckPlan,
    DesignRequest,
    PlacedBlock,
    SlideInstance,
    SlidePattern,
    StorySlide,
    TemplateProfile,
    TextStyle,
    VisualRequest,
)
from exposlides.design_pipeline import CancelledError, DesignPipeline, save_model
from exposlides.studio import FixInput, Studio, create_app


def profile():
    return TemplateProfile(
        template_sha256="a"*64, width=9144000, height=5143500,
        patterns=[SlidePattern(
            source_slide_index=1, layout_index=1, name="Basic", slots=[],
            palette=["202124", "FFFFFF", "4472C4"], font="Arial",
            title_style=TextStyle(size=32), body_style=TextStyle(size=20),
            content_box=Box(left=400000, top=1000000, width=8200000, height=3500000),
        )],
    )


def test_extractive_keeps_signed_facts_and_all_sources():
    request = DesignRequest(script="Выручка составила 12 млн рублей. Прибыль составила −2 млн рублей.\n\n"
                            "Клиенты выбирают обучение.", mode="extractive", slide_count=3)
    excerpts = source_excerpts(request.script)
    plan = extractive_plan(request, excerpts)
    assert len(plan.slides) == 3
    assert validate_story(plan, request, excerpts) == []
    broken = plan.model_copy(deep=True)
    broken.slides[1].paragraphs = ["Прибыль составила 2 млн рублей."]
    broken.slides[1].title = "Прибыль составила 2 млн рублей."
    assert validate_story(broken, request, excerpts)


def test_false_source_references_do_not_satisfy_coverage():
    request = DesignRequest(script="Поддержка клиентов доступна круглосуточно.\n\n"
                            "Обучение сотрудников проходит очно в офисе.", slide_count=1)
    story = ContentPlan(title="Поддержка", slides=[StorySlide(
        id="s1", title="Поддержка", paragraphs=["Поддержка клиентов доступна круглосуточно."],
        source_ids=["source-1", "source-2"],
    )])
    assert any("source-2" in e for e in validate_story(story, request, source_excerpts(request.script)))


def test_numbers_from_referenced_dataset_are_grounded_but_new_numbers_fail():
    request = DesignRequest(script="Итоги продаж за год.", slide_count=1, datasets=[Dataset(
        id="sales", name="Продажи", columns=["Период", "Выручка"],
        rows=[["2025", 42]], source="Демонстрационные данные", unit="млн рублей",
    )])
    story = ContentPlan(title="Продажи", slides=[StorySlide(
        id="s1", title="Итоги продаж за год", paragraphs=["Выручка (2025): 42 млн рублей."],
        source_ids=["source-1"], visual=VisualRequest(kind="bar", dataset_id="sales"),
    )])
    assert validate_story(story, request, source_excerpts(request.script)) == []
    story.slides[0].paragraphs = ["Выручка (2025): 43 млн рублей."]
    assert validate_story(story, request, source_excerpts(request.script))


def test_speaker_notes_cannot_hide_invented_facts_or_replace_visible_coverage():
    request = DesignRequest(script="Команда развивает платформу анализа данных.", slide_count=1)
    story = ContentPlan(title="Платформа", slides=[StorySlide(
        id="s1", title="Платформа", paragraphs=[request.script], source_ids=["source-1"],
        notes="Выручка составила 999 млн рублей. Компания стала лидером рынка.",
    )])
    assert validate_story(story, request, source_excerpts(request.script))
    story.slides[0].notes = request.script
    story.slides[0].paragraphs = ["Обсудим вопросы."]
    assert any("раскрыто не полностью" in e for e in
               validate_story(story, request, source_excerpts(request.script)))


def test_selected_fixes_preserve_unselected_objects_and_original():
    template = profile()
    block = PlacedBlock(id="a", kind="text", text="Короткий текст", style=TextStyle(),
                        box=Box(left=-200, top=1000000, width=2500000, height=1000000))
    untouched = PlacedBlock(id="b", kind="text", text="Другой текст", style=TextStyle(),
                            box=Box(left=4000000, top=1000000, width=2500000, height=1000000))
    plan = DeckPlan(variant_id="story", name="Story", description="Desc", width=template.width,
                    height=template.height, slides=[SlideInstance(
                        id="slide", story_slide_id="s1", source_slide_index=1,
                        blocks=[block, untouched],
                    )])
    report = audit_deck(plan, template)
    outside = next(i for i in report.issues if i.rule == "outside_slide")
    result = apply_fixes(plan, template, report, [outside.id])
    assert result.slides[0].blocks[0].box.left == 0
    assert plan.slides[0].blocks[0].box.left == -200
    assert result.slides[0].blocks[1] == untouched
    assert not any(i.rule == "outside_slide" for i in audit_deck(result, template).issues)
    with pytest.raises(ValueError, match="существующие"):
        apply_fixes(plan, template, report, ["stale"])


def test_title_fix_keeps_title_scale_and_uses_free_space():
    template = profile()
    title = PlacedBlock(id="title", kind="title", text="Длинный заголовок "*7,
                        style=TextStyle(size=32),
                        box=Box(left=300000, top=200000, width=8000000, height=450000))
    body = PlacedBlock(id="body", kind="text", text="Содержание", style=TextStyle(),
                       box=Box(left=300000, top=2300000, width=8000000, height=1500000))
    plan = DeckPlan(variant_id="story", name="Story", description="Desc", width=template.width,
                    height=template.height, slides=[SlideInstance(
                        id="slide", story_slide_id="s1", source_slide_index=1, blocks=[title, body],
                    )])
    report = audit_deck(plan, template)
    issue = next(i for i in report.issues if i.rule == "text_capacity")
    fixed = apply_fixes(plan, template, report, [issue.id])
    assert fixed.slides[0].blocks[0].style.size >= 32*0.8
    assert fixed.slides[0].blocks[1] == body
    assert not {"type_scale", "text_capacity", "overlap"}.intersection(
        i.rule for i in audit_deck(fixed, template).issues
    )


def test_pipeline_cancellation_does_not_launch_command(tmp_path):
    cancel = threading.Event()
    cancel.set()
    pipeline = DesignPipeline(tmp_path, cancel=cancel)
    with pytest.raises(CancelledError):
        pipeline.command(["command-must-not-run"], tmp_path)
    pipeline.close()


class FakePipeline:
    def __init__(self, directory, **kwargs):
        self.directory = directory

    def plan(self, template, request):
        story = extractive_plan(request, source_excerpts(request.script))
        result = profile()
        save_model(self.directory / "profile.json", result)
        return result, story

    def build(self, *args):
        return [{"id": "story", "revision": 1, "issues": [], "preview_urls": [], "exports": {}}]

    def fix(self, *args):
        raise ValueError("Исправление не применимо")

    def close(self):
        pass


def wait_for(studio, job_id, status):
    # Futures are waited by executor shutdown, avoiding sleeps/time-dependent polling in tests.
    studio.pool.shutdown(wait=True)
    assert studio.get(job_id)["status"] == status


def test_studio_upload_security_and_textbox_templates(tmp_path):
    import base64
    import io

    from pptx.util import Inches

    app = create_app(tmp_path, pipeline_factory=FakePipeline)
    with TestClient(app) as client:
        token = client.get("/api/session").json()["token"]
        assert client.post("/api/templates", json={}).status_code == 403
        headers = {"X-Session-Token": token}
        assert client.post("/api/templates", headers=headers,
                           json={"name": "bad.pptx", "data": "AAAA"}).status_code == 422
        assert client.post("/api/templates", headers=headers | {"Origin": "https://evil.test"},
                           json={}).status_code == 403
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Заголовок"
        data = io.BytesIO()
        presentation.save(data)
        response = client.post("/api/templates", headers=headers, json={
            "name": "textboxes.pptx", "data": base64.b64encode(data.getvalue()).decode(),
        })
        assert response.status_code == 201
        request = {"script": "Первый результат. Следующий шаг.", "mode": "extractive", "slide_count": 2}
        response = client.post("/api/design/plan", headers=headers, json={
            "template_id": response.json()["id"], "request": request,
        })
        assert response.status_code == 202
        job_id = response.json()["id"]
        wait_for(app.state.studio, job_id, "awaiting_review")
        job = client.get(f"/api/design/jobs/{job_id}").json()
        assert job["request"]["script"] == request["script"]
        assert job["template"]["name"] == "textboxes.pptx"
        assert len(job["story"]["slides"]) == 2
        assert client.get(f"/api/design/jobs/{job_id}/variants/story/files/1/request.json").status_code == 404
        owner_cookie = client.cookies.get("studio_owner")
    restored = create_app(tmp_path, pipeline_factory=FakePipeline)
    with TestClient(restored, cookies={"studio_owner": owner_cookie}) as client:
        assert client.get(f"/api/design/jobs/{job_id}").json()["status"] == "awaiting_review"


def test_studio_recovers_interrupted_jobs(tmp_path):
    identifier = "a"*32
    save_model(tmp_path / "jobs" / identifier / "job.json", {
        "id": identifier, "status": "running", "variants": [], "_owner_id": "f"*32,
    })
    with TestClient(create_app(tmp_path, pipeline_factory=FakePipeline),
                    cookies={"studio_owner": "f"*32}) as client:
        result = client.get(f"/api/design/jobs/{identifier}").json()
        assert result["status"] == "failed"
        assert "перезапущен" in result["error"]


def test_save_model_is_utf8_and_replaces_atomically(tmp_path):
    target = tmp_path / "data.json"
    save_model(target, {"title": "Привет"})
    save_model(target, {"title": "Мир"})
    assert json.loads(target.read_text(encoding="utf-8")) == {"title": "Мир"}
    assert list(tmp_path.iterdir()) == [target]


def prepared_studio(tmp_path, pipeline_factory):
    studio = Studio(tmp_path, pipeline_factory=pipeline_factory)
    job_id = "b"*32
    directory = studio.directory(job_id)
    directory.mkdir()
    save_model(directory / "profile.json", profile())
    save_model(directory / "request.json", DesignRequest(script="Исходные материалы.", slide_count=1))
    studio.jobs[job_id] = {
        "id": job_id, "status": "completed", "stage": "review", "active_seconds": 0,
        "variants": [{"id": "story", "revision": 1, "preview_urls": [], "exports": {}}],
    }
    return studio, job_id


def test_cancelled_fix_preserves_usable_previous_revision(tmp_path):
    class CancelledFix(FakePipeline):
        def fix(self, *args):
            raise CancelledError()

    studio, job_id = prepared_studio(tmp_path, CancelledFix)
    studio.fix(job_id, "story", FixInput(revision=1, issue_ids=["issue"]))
    studio.pool.shutdown(wait=True)
    result = studio.get(job_id)
    assert result["status"] == "completed"
    assert result["variants"][0]["revision"] == 1
    assert "отменена" in result["error"]
    studio.close()


def test_late_cancel_does_not_hide_committed_revision(tmp_path):
    class CommittedFix(FakePipeline):
        def __init__(self, directory, **kwargs):
            super().__init__(directory, **kwargs)
            self.cancel = kwargs["cancel"]

        def fix(self, *args):
            self.cancel.set()  # Cancellation loses to the publication commit point.
            return {"id": "story", "revision": 2, "preview_urls": [], "exports": {}}

    studio, job_id = prepared_studio(tmp_path, CommittedFix)
    studio.fix(job_id, "story", FixInput(revision=1, issue_ids=["issue"]))
    studio.pool.shutdown(wait=True)
    result = studio.get(job_id)
    assert result["status"] == "completed"
    assert result["variants"][0]["revision"] == 2
    assert result["error"] is None
    studio.close()


def test_build_failure_exposes_only_already_published_variants(tmp_path):
    class PartialBuild(FakePipeline):
        def build(self, *args):
            save_model(self.directory / "variants.json", [
                {"id": "story", "revision": 1, "preview_urls": [], "exports": {}},
            ])
            raise TimeoutError("Превышен лимит")

    studio, job_id = prepared_studio(tmp_path, PartialBuild)
    studio.jobs[job_id].update(status="awaiting_review", variants=[])
    story = ContentPlan(title="Исходные материалы", slides=[StorySlide(
        id="s1", title="Исходные материалы", paragraphs=["Исходные материалы."],
        source_ids=["source-1"],
    )])
    studio.build(job_id, story)
    studio.pool.shutdown(wait=True)
    result = studio.get(job_id)
    assert result["status"] == "failed"
    assert result["error"] == "Превышен лимит"
    assert len(result["variants"]) == 1
    studio.close()


def test_invalid_story_keeps_review_editable(tmp_path):
    from fastapi import HTTPException

    studio, job_id = prepared_studio(tmp_path, FakePipeline)
    studio.jobs[job_id].update(status="awaiting_review", variants=[])
    story = ContentPlan(title="Ошибка", slides=[StorySlide(
        id="s1", title="Ошибка", paragraphs=["Исходные материалы содержат 999 фактов."],
        source_ids=["source-1"],
    )])
    with pytest.raises(HTTPException) as failure:
        studio.build(job_id, story)
    assert failure.value.status_code == 422
    assert studio.get(job_id)["status"] == "awaiting_review"
    studio.close()


def test_studio_jobs_templates_and_exports_are_scoped_to_browser_owner(tmp_path):
    app = create_app(tmp_path, pipeline_factory=FakePipeline)
    with TestClient(app) as owner, TestClient(app) as stranger:
        owner_token = owner.get("/api/session").json()["token"]
        other_token = stranger.get("/api/session").json()["token"]
        assert owner.cookies.get("studio_owner") != stranger.cookies.get("studio_owner")
        template = owner.get("/api/example").json()["template_id"]
        payload = {"template_id": template, "request": {
            "script": "Первый вывод. Второй вывод.", "slide_count": 2, "mode": "extractive",
        }}
        result = owner.post("/api/design/plan", json=payload,
                            headers={"X-Session-Token": owner_token})
        job_id = result.json()["id"]
        wait_for(app.state.studio, job_id, "awaiting_review")
        assert len(owner.get("/api/design/jobs").json()["jobs"]) == 1
        assert stranger.get("/api/design/jobs").json() == {"jobs": []}
        assert stranger.get(f"/api/design/jobs/{job_id}").status_code == 404
        assert stranger.post("/api/design/plan", json=payload,
                             headers={"X-Session-Token": other_token}).status_code == 404
        assert stranger.post(f"/api/design/jobs/{job_id}/cancel",
                             headers={"X-Session-Token": other_token}).status_code == 404
        directory = app.state.studio.directory(job_id)/"variants/story/1"
        directory.mkdir(parents=True)
        (directory/"presentation.pdf").write_bytes(b"%PDF-synthetic")
        url = f"/api/design/jobs/{job_id}/variants/story/files/1/presentation.pdf"
        assert owner.get(url).status_code == 200
        assert stranger.get(url).status_code == 404
        assert "_owner_id" not in owner.get(f"/api/design/jobs/{job_id}").json()
