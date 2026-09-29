from __future__ import annotations

import asyncio
import hashlib
import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from exposlides.design_audit import audit_deck, audit_saved_pptx
from exposlides.design_builder import build_deck
from exposlides.design_layout import create_variants
from exposlides.design_models import ContentPlan, DesignRequest, StorySlide, TemplateProfile
from exposlides.design_pipeline import DesignPipeline
from exposlides.template_profile import profile_from_json
from tests.test_template_design import _element, _profile

SERVICE = Path(__file__).resolve().parents[1] / "services/content-service"


def _story():
    return ContentPlan(title="Тема", slides=[
        StorySlide(id=key, title=title, paragraphs=[title], source_ids=["source-1"])
        for key, title in [("cover", "Тема"), ("one", "Разработка"), ("two", "Обучение")]
    ])


def test_pipeline_auto_generates_distinct_requests_and_validates_binding(tmp_path, monkeypatch):
    pipeline = DesignPipeline(tmp_path)
    request = DesignRequest(script="Разработка и обучение.", slide_count=3)
    story = _story()
    _, profile, _ = _photo_template(tmp_path)
    captured = []

    def command(args, cwd, **kwargs):
        payload = json.loads(Path(args[args.index("--request") + 1]).read_text())
        captured.append(payload)
        assert 0 < float(args[args.index("--timeout") + 1]) <= 90
        results = {}
        for index, item in enumerate(payload["images"]):
            path = tmp_path / f"image-{index}.png"
            Image.new("RGB", (64, 64), (index * 80, 100, 200)).save(path)
            results[item["slide_id"]] = {
                "status": "completed", "path": str(path),
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "source_ids": item["request"]["source_ids"],
            }
        Path(args[args.index("--output") + 1]).write_text(json.dumps({
            "status": "completed", "images": results,
        }))

    monkeypatch.setattr(pipeline, "command", command)
    monkeypatch.setenv("IMAGE_GENERATION_ENABLED", "true")
    try:
        result = pipeline.generate_images(request, profile, story)
        assert set(result) == {"one", "two"}
        assert len({item["request"]["seed"] for item in captured[0]["images"]}) == 2
        assert [item["request"]["prompt"] for item in captured[0]["images"]] == [
            "Разработка\nРазработка", "Обучение\nОбучение",
        ]
        request.generated_image.auto = False
        assert pipeline.generate_images(request, profile, story) == {}
        assert len(captured) == 1
    finally:
        pipeline.close()


def _batch(module, count=5):
    return module.BatchImageRequest(images=[module.SlideImageRequest(
        slide_id=f"slide-{index}", request=module.ImageRequest(enabled=True, prompt=f"Topic {index}"),
    ) for index in range(count)])


def test_image_batch_limits_concurrency_and_returns_distinct_results(service_importer, monkeypatch, tmp_path):
    module = service_importer(SERVICE, "app.design_images")
    active = maximum = 0

    async def generate(request, output):
        nonlocal active, maximum
        active += 1
        maximum = max(active, maximum)
        try:
            await asyncio.sleep(0.01)
            return module.ImageResult(status="completed", path=str(output), sha256=request.prompt)
        finally:
            active -= 1

    monkeypatch.setattr(module, "generate_image", generate)
    result = asyncio.run(module.generate_images(_batch(module), tmp_path / "result.json"))
    assert result["status"] == "completed"
    assert len(result["images"]) == 5
    assert maximum == 3 and active == 0


@pytest.mark.parametrize("failure", ["timeout", "provider", "duplicate"])
def test_image_batch_fails_atomically_and_cancels_children(service_importer, monkeypatch, tmp_path, failure):
    module = service_importer(SERVICE, "app.design_images")
    active = 0

    async def generate(request, output):
        nonlocal active
        active += 1
        try:
            if failure == "timeout":
                await asyncio.Future()
            if failure == "provider" and request.prompt == "Topic 0":
                return module.ImageResult(status="failed")
            await asyncio.sleep(0.01)
            return module.ImageResult(status="completed", path=str(output), sha256="same")
        finally:
            active -= 1

    monkeypatch.setattr(module, "generate_image", generate)
    result = asyncio.run(module.generate_images(_batch(module), tmp_path / "result.json", timeout=0.03))
    assert result["status"] == "failed" and result["images"] == {}
    assert active == 0


def _photo_template(tmp_path):
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(12), Inches(8)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    elements = []
    for label, x, y, width, height, kind, size in [
        ("Тема", 1, .4, 10, .85, "title", 32), ("Содержание", 1, 2, 4, 4, None, 20),
    ]:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        shape.text = label
        shape.text_frame.paragraphs[0].runs[0].font.name = "Arial"
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
        elements.append(_element(shape.shape_id, label, x, y, width, height, kind, size=size))
    originals = {}
    for label, color, x, y, width, height in [
        ("Photo", "green", 6, 2, 4, 4), ("Logo", "yellow", .2, .2, .5, .5),
    ]:
        path = tmp_path / f"{label}.png"
        Image.new("RGB", (128, 128), color).save(path)
        shape = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(width), Inches(height))
        shape.name = label
        originals[label] = (shape.shape_id, path.read_bytes())
        elements.append({
            "id": f"slide-1-shape-{shape.shape_id}", "shape_id": shape.shape_id,
            "type": "image", "shape_name": label,
            "bbox": dict(left=shape.left, top=shape.top, width=shape.width, height=shape.height),
        })
    template = tmp_path / "template.pptx"
    presentation.save(template)
    profile = profile_from_json({
        "schema_version": "2.0.0", "slide_width": presentation.slide_width,
        "slide_height": presentation.slide_height, "slides": [{"index": 1, "elements": elements}],
    }, template)
    return template, profile, originals


def test_reopened_variants_have_different_images_and_preserve_logo(tmp_path):
    template, profile, originals = _photo_template(tmp_path)
    assert set(profile.patterns[0].replaceable_images) == {originals["Photo"][0]}
    restored = TemplateProfile.model_validate_json(profile.model_dump_json())
    assert restored.patterns[0].replaceable_images == profile.patterns[0].replaceable_images
    story = _story().model_copy(update={"slides": _story().slides[1:]})
    images = {}
    for slide, color in zip(story.slides, ["red", "blue"], strict=True):
        path = tmp_path / f"{slide.id}.png"
        Image.new("RGB", (128, 128), color).save(path)
        images[slide.id] = path
    for plan in create_variants(profile, story, generated_images=images):
        output = build_deck(template, plan, tmp_path / f"{plan.variant_id}.pptx")
        reopened = Presentation(output)
        assert len(reopened.slides) == 2
        for slide, content in zip(reopened.slides, story.slides, strict=True):
            blobs = [s.image.blob for s in slide.shapes if s.shape_type == 13]
            assert originals["Logo"][1] in blobs
            assert originals["Photo"][1] not in blobs
            assert images[content.id].read_bytes() in blobs
        assert not [i for i in audit_deck(plan, profile).issues if i.rule == "protected_overlap"]
        assert not [i for i in audit_saved_pptx(output, template, plan, profile).issues
                    if i.rule in {"protected_changed", "old_image_retained", "image_content"}]


def test_old_profile_still_loads_and_unknown_image_slide_is_rejected(tmp_path):
    profile = _profile(tmp_path)
    payload = profile.model_dump()
    for pattern in payload["patterns"]:
        pattern.pop("replaceable_images", None)
    assert TemplateProfile.model_validate(payload).patterns[0].replaceable_images == {}
    with pytest.raises(ValueError, match="отсутствующий"):
        create_variants(profile, _story(), generated_images={"unknown": tmp_path / "x.png"})


def _mixed_template(tmp_path):
    template, profile, originals = _photo_template(tmp_path)
    presentation = Presentation(template)
    slides = []
    for index, columns in [(2, 1), (3, 2)]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        elements = []
        positions = [("Title", 1, .4, 10, .85, "title", 32)]
        positions += [("Body", 1 + column * 5, 2, 10 if columns == 1 else 4, 4, None, 20)
                      for column in range(columns)]
        for label, x, y, width, height, kind, size in positions:
            shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
            shape.text = label
            shape.text_frame.paragraphs[0].runs[0].font.name = "Arial"
            shape.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
            elements.append(_element(shape.shape_id, label, x, y, width, height, kind, size=size))
        slides.append({"index": index, "elements": elements})
    presentation.save(template)
    text_profile = profile_from_json({
        "schema_version": "2.0.0", "slide_width": presentation.slide_width,
        "slide_height": presentation.slide_height, "slides": slides,
    }, template)
    profile.patterns.extend(text_profile.patterns)
    return template, profile, originals


def test_auto_images_preserve_varied_native_compositions_in_reopened_pptx(tmp_path):
    template, profile, originals = _mixed_template(tmp_path)
    story = _story()
    story.slides = [StorySlide(id=f"slide-{index}", title=f"Topic {index}",
                              paragraphs=[f"First point {index}", f"Second point {index}"],
                              source_ids=["source-1"]) for index in range(8)]
    images = {}
    for index, slide in enumerate(story.slides):
        path = tmp_path / f"{slide.id}.png"
        Image.new("RGB", (128, 128), (index * 30, 30, 160)).save(path)
        images[slide.id] = path
    reference = create_variants(profile, story)
    plans = create_variants(profile, story, generated_images=images, template_images_only=True)
    source = Presentation(template)
    for plan, expected in zip(plans, reference, strict=True):
        indices = [slide.source_slide_index for slide in plan.slides]
        assert indices == [slide.source_slide_index for slide in expected.slides]
        assert set(indices) == {1, 2, 3}
        output = build_deck(template, plan, tmp_path / f"{plan.variant_id}.pptx")
        reopened = Presentation(output)
        assert len(reopened.slides) == len(story.slides)
        for actual, selected, original in zip(reopened.slides, plan.slides, expected.slides, strict=True):
            assert [b for b in selected.blocks if b.kind != "image"] == original.blocks
            native_shapes = {shape.shape_id: shape for shape in source.slides[selected.source_slide_index-1].shapes}
            saved_shapes = {shape.shape_id: shape for shape in actual.shapes}
            for block in selected.blocks:
                if block.kind in {"title", "text"}:
                    before, after = native_shapes[block.source_shape_id], saved_shapes[block.source_shape_id]
                    assert (before.left, before.top, before.width, before.height) == (
                        after.left, after.top, after.width, after.height,
                    )
                    assert after.text == "\n".join(block.items or [block.text])
            pictures = [shape.image.blob for shape in actual.shapes if shape.shape_type == 13]
            if selected.source_slide_index == 1:
                assert images[selected.story_slide_id].read_bytes() in pictures
                assert originals["Photo"][1] not in pictures
                assert originals["Logo"][1] in pictures
            else:
                assert not pictures
        assert audit_saved_pptx(output, template, plan, profile).ok
        assert not [issue for issue in audit_deck(plan, profile).issues if issue.rule == "repeated_layout"]


def test_auto_images_skip_text_only_templates_without_calling_provider(tmp_path, monkeypatch):
    pipeline = DesignPipeline(tmp_path)
    monkeypatch.setenv("IMAGE_GENERATION_ENABLED", "true")
    monkeypatch.setattr(pipeline, "command", lambda *a, **kw: pytest.fail("Unexpected API call"))
    try:
        assert pipeline.generate_images(DesignRequest(script="Тема"), _profile(tmp_path), _story()) == {}
    finally:
        pipeline.close()


def test_auto_images_request_only_selected_photo_slots(tmp_path, monkeypatch):
    _, profile, _ = _mixed_template(tmp_path)
    pipeline = DesignPipeline(tmp_path)
    story = _story()
    story.slides = [StorySlide(id=f"slide-{index}", title=f"Topic {index}",
                              paragraphs=[f"Point {index}", "Other point"], source_ids=["source-1"])
                    for index in range(9)]
    monkeypatch.setenv("IMAGE_GENERATION_ENABLED", "true")
    requested = []

    def command(args, cwd, **kwargs):
        payload = json.loads(Path(args[args.index("--request") + 1]).read_text())
        results = {}
        for index, item in enumerate(payload["images"]):
            requested.append(item["slide_id"])
            path = tmp_path / f"auto-{index}.png"
            Image.new("RGB", (64, 64), (index * 20, 70, 90)).save(path)
            results[item["slide_id"]] = {
                "status": "completed", "path": str(path),
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "source_ids": ["source-1"],
            }
        Path(args[args.index("--output") + 1]).write_text(json.dumps({
            "status": "completed", "images": results,
        }), encoding="utf-8")

    monkeypatch.setattr(pipeline, "command", command)
    try:
        images = pipeline.generate_images(DesignRequest(script="Тема"), profile, story)
        assert "slide-0" not in requested
        plans = create_variants(profile, story, generated_images=images, template_images_only=True)
        used = {slide.story_slide_id for plan in plans for slide in plan.slides
                if any(block.kind == "image" for block in slide.blocks)}
        assert used == set(requested)
        assert used and len(used) < len(story.slides)
    finally:
        pipeline.close()


def test_explicit_image_searches_native_photo_layout_before_free_layout(tmp_path):
    _, profile, _ = _mixed_template(tmp_path)
    for plan in create_variants(profile, _story(), generated_image=tmp_path / "explicit.png"):
        slide = plan.slides[0]
        assert slide.source_slide_index == 1
        assert all(block.source_shape_id is not None for block in slide.blocks)
