"""Регрессия: выбор пустоты между слотами не должен стирать композицию шаблона."""

from copy import deepcopy

from exposlides.design_layout import create_variants
from exposlides.design_models import Box
from tests.test_template_design import _data, _element, _profile, _story


def test_native_columns_keep_geometry_order_and_all_source_paragraphs(tmp_path):
    data = _data()
    data["slides"][0]["elements"] = [
        _element(2, "Title", 1, 0.4, 10, 1, "title", size=32),
        _element(4, "Right", 6.2, 2, 4.8, 4),
        _element(3, "Left", 1, 2, 4.8, 4),
    ]
    profile = _profile(tmp_path, data)
    paragraphs = ["First argument.", "Second argument.", "Third argument."]
    for plan in create_variants(profile, _story(paragraphs)):
        slide = plan.slides[0]
        body = [block for block in slide.blocks if block.kind == "text"]
        assert [block.source_shape_id for block in body] == [3, 4]
        assert [text for block in body for text in block.items] == paragraphs
        assert not slide.remove_shape_ids
        for block in slide.blocks:
            slot = next(slot for slot in profile.patterns[0].slots
                        if slot.shape_id == block.source_shape_id)
            assert block.box == slot.box


def test_many_tiny_slots_do_not_win_by_their_large_union_area(tmp_path):
    data = _data()
    tiny = deepcopy(data["slides"][0])
    tiny["index"] = 2
    tiny["elements"] = [_element(2, "Title", 1, 0.4, 10, 1, "title", size=32)]
    tiny["elements"] += [
        _element(3 + i, "Sample", 0.5 + (i % 4) * 3, 2 + (i // 4) * 4, 1, 0.5, size=10)
        for i in range(8)
    ]
    data["slides"].append(tiny)
    profile = _profile(tmp_path, data)
    story = _story(["A substantial grounded explanation. " * 5] * 3)
    assert all(plan.slides[0].source_slide_index == 1 for plan in create_variants(profile, story))


def test_unused_native_text_is_cleared_without_removing_its_shape(tmp_path):
    data = _data()
    data["slides"][0]["elements"][1:] = [
        _element(3, "Left example", 1, 2, 4, 4),
        _element(4, "Right example", 6, 2, 4, 4),
    ]
    for plan in create_variants(_profile(tmp_path, data), _story()):
        body = [block for block in plan.slides[0].blocks if block.kind == "text"]
        assert len(body) == 2
        assert body[1].items == [] and body[1].text == ""
        assert body[1].source_shape_id == 4
        assert not plan.slides[0].remove_shape_ids


def test_capacity_uses_the_paragraph_style_that_will_actually_be_filled(tmp_path):
    data = _data()
    body = data["slides"][0]["elements"][1]
    body["text"]["paragraphs"] = [
        {"runs": [{"text": "Short heading", "style": {"size_pt": 24}}]},
        {"runs": [{"text": "Long small explanation. " * 20, "style": {"size_pt": 14}}]},
    ]
    profile = _profile(tmp_path, data)
    slot = profile.patterns[0].slots[1]
    assert slot.style.size == 14
    assert slot.paragraph_font_sizes == [24, 14]
    for plan in create_variants(profile, _story()):
        body = next(block for block in plan.slides[0].blocks if block.kind == "text")
        assert body.style.size == 24


def test_columns_balance_text_when_multiple_assignments_fit(tmp_path):
    data = _data()
    data["slides"][0]["elements"][1:] = [
        _element(3, "Left", 1, 2, 4, 4),
        _element(4, "Right", 6, 2, 4, 4),
    ]
    paragraphs = ["A grounded explanation of the result."] * 4
    for plan in create_variants(_profile(tmp_path, data), _story(paragraphs)):
        body = [block for block in plan.slides[0].blocks if block.kind == "text"]
        assert [len(block.items) for block in body] == [2, 2]


def test_native_sequence_uses_cover_photos_and_varied_fitting_compositions(tmp_path):
    profile = _profile(tmp_path)
    plain = profile.patterns[0]
    cover = plain.model_copy(deep=True)
    cover.source_slide_index, cover.name = 2, "Титульный слайд"
    photo = plain.model_copy(deep=True)
    photo.source_slide_index, photo.name = 3, "Паттерн + фото"
    photo.slots[1].box.width //= 2
    photo.protected_regions = [Box(left=int(profile.width * .55), top=int(profile.height * .25),
                                  width=int(profile.width * .45), height=int(profile.height * .75))]
    photo.protected_shape_ids = [9]
    profile.patterns.extend([cover, photo])
    story = _story()
    story.slides = [story.slides[0].model_copy(update={"id": f"slide-{i}"}) for i in range(8)]
    for plan in create_variants(profile, story):
        ids = [slide.source_slide_index for slide in plan.slides]
        assert ids[0] == 2
        assert 3 in ids[1:] and 1 in ids[1:]
        assert 2 not in ids[1:]
        assert all(9 not in slide.remove_shape_ids for slide in plan.slides)
        assert len(set(ids)) == 3


def test_diversity_never_selects_overflow_when_a_fitting_layout_exists(tmp_path):
    from exposlides.template_layout import native_layout

    profile = _profile(tmp_path)
    tiny = profile.patterns[0].model_copy(deep=True)
    tiny.source_slide_index = 2
    tiny.slots[1].box.height = 12700 * 20
    tiny.slots[1].box.width = 12700 * 90
    profile.patterns.append(tiny)
    result = native_layout(profile, _story().slides[0], "story", 20, (1,) * 20)
    assert result[0].source_slide_index == 1


def test_selected_cover_background_and_photo_survive_actual_pptx_build(tmp_path):
    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    from exposlides.design_builder import build_deck
    from exposlides.template_profile import profile_from_json

    template, picture = tmp_path / "source.pptx", tmp_path / "photo.png"
    Image.new("RGB", (32, 32), "#eeaa22").save(picture)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(12), Inches(8)
    data = _data()
    data["slides"] = []
    for index, name in enumerate(("Text", "Титульный слайд", "Паттерн + фото"), 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(
            "0077FF" if index == 2 else "FFFFFF",
        )
        elements = []
        for text, left, top, width, height, role, size in [
            ("Title", 1, .4, 10, .85, "title", 32),
            ("Body", 1, 2, 5 if index == 3 else 10, 4, None, 20),
        ]:
            shape = slide.shapes.add_textbox(Inches(left), Inches(top),
                                            Inches(width), Inches(height))
            shape.text = text
            shape.text_frame.paragraphs[0].font.size = Pt(size)
            elements.append(_element(shape.shape_id, text, left, top, width, height, role,
                                     color="FFFFFF" if index == 2 else "202124", size=size))
        if index == 3:
            photo = slide.shapes.add_picture(str(picture), Inches(6.5), Inches(2),
                                            Inches(5.5), Inches(6))
            elements.append(dict(id="photo", shape_id=photo.shape_id, type="picture",
                                 bbox=dict(left=photo.left, top=photo.top,
                                           width=photo.width, height=photo.height)))
        data["slides"].append(dict(index=index, layout_name=name, elements=elements))
    prs.save(template)
    profile = profile_from_json(data, template)
    story = _story()
    story.slides = [story.slides[0].model_copy(update={"id": f"slide-{i}"}) for i in range(6)]
    for plan in create_variants(profile, story):
        output = tmp_path / f"{plan.variant_id}.pptx"
        build_deck(template, plan, output)
        result = Presentation(output)
        assert len(result.slides) == 6
        assert str(result.slides[0].background.fill.fore_color.rgb) == "0077FF"
        assert any(slide.source_slide_index == 3 for slide in plan.slides)
        for slide, selected in zip(result.slides, plan.slides, strict=True):
            if selected.source_slide_index == 3:
                assert any(getattr(shape, "image", None) and shape.image.blob == picture.read_bytes()
                           for shape in slide.shapes)
            assert any(shape.has_text_frame and "grounded" in shape.text for shape in slide.shapes)


def test_unusable_extra_slot_does_not_discard_fitting_template(tmp_path):
    profile = _profile(tmp_path)
    pattern = profile.patterns[0]
    extra = pattern.slots[1].model_copy(deep=True)
    extra.shape_id = 99
    extra.box = Box(left=0, top=0, width=10000, height=10000)
    pattern.slots.append(extra)
    pattern.mutable_shape_ids.append(99)
    pattern.protected_regions.append(extra.box.model_copy())
    plans = create_variants(profile, _story())
    for plan in plans:
        assert plan.slides[0].source_slide_index == 1
        assert 99 in plan.slides[0].remove_shape_ids
        assert all(block.source_shape_id != 99 for block in plan.slides[0].blocks)


def test_variants_choose_different_fitting_original_compositions(tmp_path):
    profile = _profile(tmp_path)
    for index in (2, 3):
        pattern = profile.patterns[0].model_copy(deep=True)
        pattern.source_slide_index = index
        pattern.slots[1].box.width -= index * 100000
        profile.patterns.append(pattern)
    plans = create_variants(profile, _story())
    assert len({plan.slides[0].source_slide_index for plan in plans}) == 3
    for plan in plans:
        pattern = profile.patterns[plan.slides[0].source_slide_index - 1]
        assert all(block.box == next(slot.box for slot in pattern.slots
                                    if slot.shape_id == block.source_shape_id)
                   for block in plan.slides[0].blocks)


def test_audit_reports_repeated_composition_with_different_content(tmp_path):
    from exposlides.design_audit import audit_deck

    profile = _profile(tmp_path)
    story = _story()
    story.slides = [story.slides[0].model_copy(update={
        "id": f"slide-{index}", "title": f"Title {index}", "paragraphs": [f"Point {index}"],
    }) for index in range(4)]
    plan = create_variants(profile, story)[0]
    report = audit_deck(plan, profile)
    assert "repeated_layout" in report.checks
    assert len([issue for issue in report.issues if issue.rule == "repeated_layout"]) == 1
    assert not [issue for issue in report.issues if issue.rule == "duplicate_slide"]
    plan.slides[1].blocks[1].box.width //= 2
    assert not [issue for issue in audit_deck(plan, profile).issues if issue.rule == "repeated_layout"]


def test_single_slide_uses_other_native_compositions_beyond_shared_cover(tmp_path):
    from exposlides.design_diversity import variant_diversity

    profile = _profile(tmp_path)
    plain = profile.patterns[0]
    cover = plain.model_copy(deep=True)
    cover.source_slide_index, cover.name = 2, "Титульный слайд"
    cover.slots[1].box.width = int(plain.slots[1].box.width * .8)
    narrow = plain.model_copy(deep=True)
    narrow.source_slide_index, narrow.name = 3, "Узкий текст"
    narrow.slots[1].box.width = int(plain.slots[1].box.width * .6)
    profile.patterns.extend([cover, narrow])
    original = profile.model_dump()
    story = _story()

    variants = create_variants(profile, story)

    assert len({plan.slides[0].source_slide_index for plan in variants}) == 3
    assert variant_diversity(profile, variants)["sufficient"]
    assert profile.model_dump() == original


def test_content_placeholder_without_text_frame_is_not_a_text_slot(tmp_path):
    data = _data()
    fake_text = _element(77, "", 1, 2, 10, 4, "content")
    fake_text.update(has_text_frame=False, type="shape", text=None)
    data["slides"][0]["elements"].append(fake_text)
    profile = _profile(tmp_path, data)
    assert all(slot.shape_id != 77 for slot in profile.patterns[0].slots)
    assert any(slot.role == "body" for slot in profile.patterns[0].slots)
