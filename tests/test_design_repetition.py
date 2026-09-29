"""Повторы фотографий и соседних макетов проверяются на настоящих PPTX."""

import asyncio
import hashlib
import json
from pathlib import Path

import pytest
from PIL import Image, PngImagePlugin
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

from exposlides.design_audit import audit_saved_pptx
from exposlides.design_builder import build_deck
from exposlides.design_layout import create_variants
from exposlides.design_models import Box, DesignRequest
from exposlides.design_pipeline import DesignPipeline
from exposlides.design_repetition import image_identity, repeated_content_images
from exposlides.template_layout import composition_key, native_layout
from tests.test_design_images_batch import _photo_template, _story
from tests.test_layout_selection_quality import _photo
from tests.test_template_design import _profile
from tests.test_template_design import _story as text_story


def _two_photos(tmp_path):
    template, profile, originals = _photo_template(tmp_path)
    prs = Presentation(template)
    photo = next(s for s in prs.slides[0].shapes if s.shape_id == originals["Photo"][0])
    photo.height = Inches(2)
    ph = OxmlElement("p:ph")
    ph.set("type", "obj")
    photo._element.nvPicPr.nvPr.append(ph)
    other = prs.slides[0].shapes.add_picture(str(tmp_path / "Photo.png"),
                                            Inches(6), Inches(4.2), Inches(4), Inches(2))
    pattern = profile.patterns[0]
    pattern.replaceable_images = {
        s.shape_id: Box(left=s.left, top=s.top, width=s.width, height=s.height)
        for s in (photo, other)
    }
    pattern.protected_regions = list(pattern.replaceable_images.values())
    pattern.protected_shape_ids.append(other.shape_id)
    prs.save(template)
    profile.template_sha256 = hashlib.sha256(template.read_bytes()).hexdigest()
    return template, profile, originals


def test_parser_recognizes_picture_inside_object_placeholder(tmp_path, service_importer):
    template, _, originals = _two_photos(tmp_path)
    service = Path(__file__).resolve().parents[1] / "services/parsing-service"
    parser = service_importer(service, "app.parsers.pptx").PPTXParser
    parsed = asyncio.run(parser.parse(template)).presentation.model_dump(mode="json")
    picture = next(e for e in parsed["slides"][0]["elements"]
                   if e["shape_id"] == originals["Photo"][0])
    assert picture["type"] == "image"
    assert picture["placeholder_kind"] == "content"
    assert picture["has_text_frame"] is False
    assert picture["image"]["asset_id"]


def test_each_native_photo_slot_gets_distinct_image_in_saved_pptx(tmp_path, monkeypatch):
    template, profile, originals = _two_photos(tmp_path)
    pipeline = DesignPipeline(tmp_path)
    monkeypatch.setenv("IMAGE_GENERATION_ENABLED", "true")
    captured = []

    def command(args, cwd, **kwargs):
        payload = json.loads(Path(args[args.index("--request") + 1]).read_text())
        captured.extend(payload["images"])
        entries = {}
        for i, item in enumerate(payload["images"]):
            path = tmp_path / f"generated-{i}.png"
            Image.new("RGB", (80, 80), (i * 40, 80, 120)).save(path)
            entries[item["slide_id"]] = dict(status="completed", path=str(path),
                sha256=hashlib.sha256(path.read_bytes()).hexdigest(), source_ids=["source-1"])
        Path(args[args.index("--output") + 1]).write_text(json.dumps(
            dict(status="completed", images=entries)))

    monkeypatch.setattr(pipeline, "command", command)
    story = _story()
    try:
        images = pipeline.generate_images(DesignRequest(script="Тема"), profile, story)
    finally:
        pipeline.close()
    assert len(captured) == 4  # two content slides, two photo slots per slide
    assert len({i["request"]["prompt"] for i in captured}) == 4
    assert len({i["request"]["seed"] for i in captured}) == 4
    story.slides = story.slides[1:]
    for plan in create_variants(profile, story, generated_images=images, template_images_only=True):
        output = build_deck(template, plan, tmp_path / f"{plan.variant_id}.pptx")
        assert not repeated_content_images(output, plan, profile)
        assert audit_saved_pptx(output, template, plan, profile).ok
        for slide in Presentation(output).slides:
            pictures = [s.image.blob for s in slide.shapes if hasattr(s, "image")]
            assert originals["Logo"][1] in pictures
            assert originals["Photo"][1] not in pictures
            assert len({image_identity(blob) for blob in pictures}) == len(pictures)


def test_legacy_single_image_cannot_be_duplicated_into_two_slots(tmp_path):
    _, profile, _ = _two_photos(tmp_path)
    with pytest.raises(ValueError, match="каждой фотообласти"):
        create_variants(profile, _story(), generated_images={"one": tmp_path / "Photo.png"},
                        template_images_only=True)


def test_repeated_template_photos_block_publication_but_logo_is_ignored(tmp_path):
    template, profile, _ = _photo_template(tmp_path)
    story = _story()
    plan = create_variants(profile, story)[0]
    pipeline = DesignPipeline(tmp_path / "job")
    try:
        with pytest.raises(ValueError, match="повторяются содержательные"):
            pipeline.publish_variant(template, profile, plan, revision=1)
        assert not (pipeline.directory / "variants/story/1").exists()
        assert not list(pipeline.directory.rglob("*.pptx"))
    finally:
        pipeline.close()


def test_different_png_metadata_does_not_hide_identical_pixels(tmp_path):
    first, second = tmp_path / "a.png", tmp_path / "b.png"
    image = Image.new("RGB", (30, 30), "red")
    image.save(first)
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("description", "different metadata")
    image.save(second, pnginfo=metadata)
    assert first.read_bytes() != second.read_bytes()
    assert image_identity(first.read_bytes()) == image_identity(second.read_bytes())


def test_neighbor_composition_beats_family_quota_and_duplicate_source_id(tmp_path):
    profile = _profile(tmp_path)
    plain = profile.patterns[0]
    photo = _photo(plain, 2)
    duplicate_photo = photo.model_copy(update={"source_slide_index": 3}, deep=True)
    profile.patterns.extend([photo, duplicate_photo])
    selected, _ = native_layout(profile, text_story().slides[0], "story", 4,
                                previous=(1, 1, 1, 2))
    assert composition_key(selected) != composition_key(photo)
