from __future__ import annotations

import asyncio
import importlib
import json
from pathlib import Path

from pptx import Presentation as PPTXPresentation
from pptx.util import Pt

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
PARSING_SERVICE_ROOT = REPOSITORY_ROOT / "services" / "parsing-service"
CONTENT_SERVICE_ROOT = REPOSITORY_ROOT / "services" / "content-service"
BUILDER_SERVICE_ROOT = REPOSITORY_ROOT / "services" / "builder-service"


def _write_template(path: Path) -> None:
    presentation = PPTXPresentation()
    for title in ("Первый исходный заголовок презентации", "Второй исходный заголовок"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = (
            "Исходный текст достаточно длинный для безопасной оценки нового содержимого"
        )
        title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
        title_run.font.name = "Arial"
        title_run.font.size = Pt(28)
        title_run.font.bold = True
    presentation.save(path)


class FakeLLMClient:
    """Deterministic structured responses for the real content graph."""

    async def generate_json(self, _prompt, model, strict=True):
        if model.__name__ == "ScriptAnalysis":
            return model(
                topic="Запуск продукта",
                audience="",
                objective="",
                blocks=[
                    {
                        "index": 1,
                        "heading": "Подготовка запуска",
                        "summary": "Команда готовит запуск продукта и описывает стратегию",
                        "key_points": ["Команда готовит запуск продукта"],
                        "facts": [],
                    }
                ],
                key_messages=["Команда готовит запуск продукта"],
                facts=[],
            )
        if model.__name__ == "SlidePlan":
            return model(
                slides=[
                    {
                        "template_slide_index": 2,
                        "layout_type": "bullets",
                        "title": "Стратегия запуска продукта",
                        "content": "Команда готовит запуск продукта",
                        "purpose": "Описать стратегию",
                        "key_message": "Команда готовит запуск продукта",
                        "source_block_indices": [1],
                    },
                    {
                        "template_slide_index": 1,
                        "layout_type": "bullets",
                        "title": "Команда готовит запуск",
                        "content": "Стратегия продукта описана командой",
                        "purpose": "Завершить рассказ",
                        "key_message": "Команда готовит запуск продукта",
                        "source_block_indices": [1],
                    },
                ]
            )
        raise AssertionError(f"Unexpected model: {model.__name__}")

    async def generate_json_with_schema(self, prompt, schema, strict=True):
        if "Слайд 2:" in prompt:
            texts = ["Стратегия запуска продукта", "Команда готовит запуск продукта"]
        else:
            texts = ["Команда готовит запуск", "Стратегия продукта описана командой"]
        return {
            key: texts[index % len(texts)]
            for index, key in enumerate(schema["required"])
        }


def test_full_offline_pipeline_creates_ordered_editable_pptx(
    tmp_path: Path,
    monkeypatch,
    service_importer,
) -> None:
    template_pptx = tmp_path / "template.pptx"
    template_json = tmp_path / "template.json"
    script = tmp_path / "script.txt"
    content_json = tmp_path / "generated_content.json"
    output_pptx = tmp_path / "result.pptx"
    _write_template(template_pptx)
    script.write_text(
        "Команда готовит запуск продукта и описывает стратегию.",
        encoding="utf-8",
    )

    parsing_main = service_importer(PARSING_SERVICE_ROOT, "app.main")
    asyncio.run(parsing_main.run(template_pptx, template_json))

    content_main = service_importer(CONTENT_SERVICE_ROOT, "app.main")
    nodes = importlib.import_module("app.graph.nodes")
    request = importlib.import_module("app.models.request")
    monkeypatch.setattr(nodes, "llm_client", FakeLLMClient())
    asyncio.run(
        content_main.run(
            template_json,
            script,
            content_json,
            request.GenerationSettings(max_slides=2),
        )
    )

    builder_main = service_importer(BUILDER_SERVICE_ROOT, "app.main")
    asyncio.run(builder_main.run(template_pptx, template_json, content_json, output_pptx))

    parsed_payload = json.loads(template_json.read_text(encoding="utf-8"))
    generated_payload = json.loads(content_json.read_text(encoding="utf-8"))
    result = PPTXPresentation(output_pptx)

    assert parsed_payload["schema_version"] == "2.0.0"
    assert len(parsed_payload["slides"]) == 2
    assert list(generated_payload["content"]) == ["2", "1"]
    assert generated_payload["validation_report"] == {"ok": True, "issues": []}
    assert len(result.slides) == 2
    assert [slide.shapes.title.text for slide in result.slides] == [
        "Стратегия запуска продукта",
        "Команда готовит запуск",
    ]
    for slide in result.slides:
        title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
        assert title_run.font.name == "Arial"
        assert title_run.font.size == Pt(28)
        assert title_run.font.bold is True
