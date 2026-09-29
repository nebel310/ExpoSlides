"""Нативные PPTX: независимые экземпляры, editable визуализации и атомарность."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

from exposlides.design_builder import build_deck
from exposlides.design_models import Box, Dataset, DeckPlan, PlacedBlock, SlideInstance, TextStyle
from exposlides.design_pptx_parts import NativeBuildError


def _template(path: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(12), Inches(8)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(0.2), Inches(10), Inches(1))
    title.text = "Replace me"
    group = slide.shapes.add_group_shape()
    editable = group.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    editable.text = "Old grouped text"
    protected = group.shapes.add_textbox(Inches(8), Inches(7), Inches(3), Inches(0.5))
    protected.text = "Corporate footer"
    protected.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.com/company"
    picture = BytesIO()
    Image.new("RGB", (20, 10), (12, 34, 56)).save(picture, format="PNG")
    picture.seek(0)
    slide.shapes.add_picture(picture, Inches(10.5), Inches(0.2), width=Inches(1))
    data = CategoryChartData()
    data.categories = ["Before", "After"]
    data.add_series("Source", [10, 20])
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7), Inches(2), Inches(4), Inches(3), data
    )
    slide.notes_slide.notes_text_frame.text = "Old source notes"
    prs.save(path)
    return [title.shape_id, editable.shape_id]


def _block(identifier="title", kind="title", **kwargs):
    return PlacedBlock(
        id=identifier,
        kind=kind,
        box=Box(left=Inches(1), top=Inches(1), width=Inches(5), height=Inches(4)),
        style=TextStyle(font="Arial", size=18, color="202124"),
        **kwargs,
    )


def _plan(removals, blocks=None, count=2, datasets=None):
    return DeckPlan(
        variant_id="story",
        name="Story",
        description="Test",
        width=Inches(12),
        height=Inches(8),
        datasets=datasets or [],
        slides=[
            SlideInstance(
                id=f"slide-{index}",
                story_slide_id=f"story-{index}",
                source_slide_index=1,
                remove_shape_ids=removals,
                notes=f"Notes {index}",
                blocks=blocks or [_block(text=f"Title {index}")],
            )
            for index in range(count)
        ],
    )


def test_twelve_copies_preserve_native_parts_and_nested_neighbors(tmp_path):
    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    source_bytes = template.read_bytes()
    build_deck(template, _plan(removals, count=12), result)
    prs = Presentation(result)
    assert len(prs.slides) == 12
    chart_parts, workbook_parts, notes_parts = [], [], []
    for index, slide in enumerate(prs.slides):
        assert (
            next(shape for shape in slide.shapes if shape.name == "exposlides:title").text
            == f"Title {index}"
        )
        group = next(shape for shape in slide.shapes if hasattr(shape, "shapes"))
        assert [shape.text for shape in group.shapes] == ["Corporate footer"]
        assert (
            group.shapes[0].text_frame.paragraphs[0].runs[0].hyperlink.address
            == "https://example.com/company"
        )
        assert any(shape.shape_type == 13 for shape in slide.shapes)
        chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
        chart_parts.append(str(chart.part.partname))
        workbook_parts.append(str(chart.part.part_related_by(RT.PACKAGE).partname))
        notes_parts.append(str(slide.notes_slide.part.partname))
        assert slide.notes_slide.notes_text_frame.text == f"Notes {index}"
    assert len(set(chart_parts)) == len(set(workbook_parts)) == len(set(notes_parts)) == 12
    first_chart = next(shape.chart for shape in prs.slides[0].shapes if shape.has_chart)
    changed = CategoryChartData()
    changed.categories = ["Before", "After"]
    changed.add_series("Source", [99, 100])
    first_chart.replace_data(changed)
    prs.save(result)
    reopened = Presentation(result)
    second_chart = next(shape.chart for shape in reopened.slides[1].shapes if shape.has_chart)
    assert list(second_chart.series[0].values) == [10, 20]
    assert template.read_bytes() == source_bytes
    with ZipFile(result) as archive:
        assert len(archive.namelist()) == len(set(archive.namelist()))


@pytest.mark.parametrize(
    "kind,chart_type", [("table", None), ("chart", "bar"), ("chart", "line"), ("chart", "pie")]
)
def test_data_visuals_remain_native_editable(tmp_path, kind, chart_type):
    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    data = Dataset(
        id="data",
        name="Revenue",
        columns=["Quarter", "Value"],
        rows=[["Q1", 10], ["Q2", 20]],
        source="Test brief",
        unit="mln",
    )
    block = _block("visual", kind, dataset_id="data", chart_type=chart_type, fill="224488")
    build_deck(template, _plan(removals, [block], count=1, datasets=[data]), result)
    shape = next(s for s in Presentation(result).slides[0].shapes if s.name == "exposlides:visual")
    if kind == "table":
        assert shape.has_table
        assert shape.table.cell(1, 0).text == "Q1"
        assert shape.table.cell(2, 1).text == "20"
    else:
        assert shape.has_chart
        assert list(shape.chart.series[0].values) == [10, 20]
        assert shape.chart.part.part_related_by(RT.PACKAGE).blob.startswith(b"PK")
        if chart_type != "pie":
            assert shape.chart.value_axis.axis_title.text_frame.text == "mln"
            assert shape.chart.category_axis.axis_title.text_frame.text == "Quarter"


@pytest.mark.parametrize("kind", ["process", "comparison"])
def test_diagram_is_a_group_of_editable_objects(tmp_path, kind):
    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    block = _block(
        "diagram", kind, items=["Discover", "Plan", "Build", "Audit", "Export"], fill="224488"
    )
    build_deck(template, _plan(removals, [block], count=1), result)
    group = next(s for s in Presentation(result).slides[0].shapes if s.name == "exposlides:diagram")
    assert hasattr(group, "shapes")
    assert [s.text for s in group.shapes if s.has_text_frame] == block.items
    assert len(group.shapes) == len(block.items) + (
        len(block.items) - 1 if kind == "process" else 0
    )


def test_invalid_data_preserves_existing_output_and_template(tmp_path):
    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    original = template.read_bytes()
    result.write_bytes(b"Previous valid result")
    data = Dataset(
        id="data",
        name="Revenue",
        columns=["Quarter", "Value"],
        rows=[["Q1", "unconfirmed"]],
        source="Test",
    )
    block = _block("visual", "chart", dataset_id="data", chart_type="bar")
    with pytest.raises(NativeBuildError, match="нечисловые"):
        build_deck(template, _plan(removals, [block], count=1, datasets=[data]), result)
    assert result.read_bytes() == b"Previous valid result"
    assert template.read_bytes() == original


def test_failure_after_save_does_not_publish(tmp_path, monkeypatch):
    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    result.write_bytes(b"Previous valid result")

    def fail(*args):
        raise NativeBuildError("test validation failure")

    monkeypatch.setattr("exposlides.design_builder._validate_saved", fail)
    with pytest.raises(NativeBuildError, match="validation failure"):
        build_deck(template, _plan(removals, count=1), result)
    assert result.read_bytes() == b"Previous valid result"
    assert not list(tmp_path.glob(".output.pptx.*"))


def test_rejects_output_overwriting_template(tmp_path):
    template = tmp_path / "source.pptx"
    removals = _template(template)
    with pytest.raises(NativeBuildError, match="перезаписывать"):
        build_deck(template, _plan(removals), template)


@pytest.mark.parametrize("icon", ["arrow", "check", "info"])
def test_icons_are_native_shapes(tmp_path, icon):
    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    block = _block("icon", "icon", icon=icon, fill="224488")
    build_deck(template, _plan(removals, [block], count=1), result)
    shape = next(s for s in Presentation(result).slides[0].shapes if s.name == "exposlides:icon")
    assert shape.shape_type != 13
    if icon == "check":
        assert len(shape.shapes) == 2
    elif icon == "info":
        assert shape.text == "i"


def test_internal_slide_links_point_to_output_instances(tmp_path):
    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    prs = Presentation(template)
    source = prs.slides[0]
    target = prs.slides.add_slide(prs.slide_layouts[6])
    link = source.shapes.add_textbox(Inches(1), Inches(6), Inches(2), Inches(1))
    link.text = "Go to target"
    link.click_action.target_slide = target
    prs.save(template)
    plan = _plan(removals, count=3)
    plan.slides[1].source_slide_index = 2
    plan.slides[1].remove_shape_ids = []
    build_deck(template, plan, result)
    reopened = Presentation(result)
    assert len(reopened.slides) == 3
    for index in [0, 2]:
        link = next(
            s
            for s in reopened.slides[index].shapes
            if s.has_text_frame and s.text == "Go to target"
        )
        assert link.click_action.target_slide == reopened.slides[1]
    with pytest.raises(NativeBuildError, match="внутреннюю ссылку"):
        build_deck(template, _plan(removals, count=1), result)
    assert len(Presentation(result).slides) == 3


def test_generated_picture_is_embedded_without_distortion(tmp_path):
    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    picture = tmp_path / "generated.png"
    Image.new("RGB", (320, 160), (22, 44, 88)).save(picture)
    block = _block("generated", "image", image_path=str(picture))
    build_deck(template, _plan(removals, [block], count=1), result)
    shape = next(
        s for s in Presentation(result).slides[0].shapes if s.name == "exposlides:generated"
    )
    assert shape.shape_type == 13
    assert shape.width / shape.height == pytest.approx(2)
    assert shape.left >= block.box.left and shape.top >= block.box.top
    assert shape.left + shape.width <= block.box.left + block.box.width
    assert shape.top + shape.height <= block.box.top + block.box.height
    assert shape.image.blob == picture.read_bytes()


def test_invalid_image_does_not_publish(tmp_path):
    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    picture = tmp_path / "invalid.png"
    picture.write_bytes(b"not an image")
    block = _block("generated", "image", image_path=str(picture))
    with pytest.raises(NativeBuildError, match="прочитать изображение"):
        build_deck(template, _plan(removals, [block], count=1), result)
    assert not result.exists()


def test_dark_table_and_chart_use_readable_native_colors(tmp_path):
    from pptx.oxml.ns import qn

    from exposlides.design_builder import _contrast

    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    data = Dataset(
        id="data", name="Value", columns=["Period", "Total"], rows=[["Q1", 10]], source="Test"
    )
    for kind in ["table", "chart"]:
        block = _block(
            "visual",
            kind,
            dataset_id="data",
            chart_type="bar" if kind == "chart" else None,
            fill="224488",
        )
        block.style.color = "FFFFFF"
        build_deck(template, _plan(removals, [block], count=1, datasets=[data]), result)
        shape = next(
            s for s in Presentation(result).slides[0].shapes if s.name == "exposlides:visual"
        )
        if kind == "table":
            cell = shape.table.cell(1, 0)
            color = str(cell.text_frame.paragraphs[0].runs[0].font.color.rgb)
            assert _contrast(color, str(cell.fill.fore_color.rgb)) >= 4.5
        else:
            assert shape.chart._chartSpace.find(qn("c:spPr")).find(qn("a:noFill")) is not None
            assert str(shape.chart.value_axis.tick_labels.font.color.rgb) == "FFFFFF"


def test_links_from_shared_layout_cannot_leave_hidden_source_slides(tmp_path):
    from copy import deepcopy

    from pptx.oxml.ns import qn

    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    prs = Presentation(template)
    source = prs.slides[0]
    target = prs.slides.add_slide(prs.slide_layouts[6])
    shape = source.shapes.add_textbox(Inches(1), Inches(6), Inches(2), Inches(1))
    shape.text = "Layout navigation"
    shape.click_action.target_slide = target
    element = deepcopy(shape._element)
    shape._element.getparent().remove(shape._element)
    # Исходная ссылка слайда больше не используется и не должна влиять на проверку.
    unused = next(
        rel.rId
        for rel in source.part.rels.values()
        if rel.reltype == RT.SLIDE and rel.target_part is target.part
    )
    source.part.drop_rel(unused)
    layout = source.slide_layout
    rid = layout.part.relate_to(target.part, RT.SLIDE)
    element.find(".//" + qn("a:hlinkClick")).set(qn("r:id"), rid)
    layout.shapes._spTree.insert_element_before(element, "p:extLst")
    prs.save(template)
    with pytest.raises(NativeBuildError, match="скрытые исходные слайды"):
        build_deck(template, _plan(removals, count=1), result)
    assert not result.exists()


@pytest.mark.parametrize("tag", ["custom_show", "section"])
def test_empty_presentation_navigation_does_not_block_build(tmp_path, tag):
    from lxml import etree
    from pptx.oxml.ns import qn

    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    prs = Presentation(template)
    if tag == "custom_show":
        etree.SubElement(prs._element, qn("p:custShowLst"))
    else:
        extension_list = etree.SubElement(prs._element, qn("p:extLst"))
        extension = etree.SubElement(
            extension_list, qn("p:ext"), uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
        )
        etree.SubElement(
            extension, "{http://schemas.microsoft.com/office/powerpoint/2010/main}sectionLst"
        )
    prs.save(template)
    original = template.read_bytes()
    build_deck(template, _plan(removals, count=1), result)
    reopened = Presentation(result)
    assert len(reopened.slides) == 1
    assert reopened._element.find(qn("p:custShowLst")) is None
    assert not any(element.tag.endswith("}sectionLst") for element in reopened._element.iter())
    assert template.read_bytes() == original


@pytest.mark.parametrize("values", [[12, 18], [-12, 18], [0.001, 0.003], [10_000_000, 30_000_000]])
def test_narrow_chart_limits_tick_density_without_altering_source_values(tmp_path, values):
    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    dataset = Dataset(id="data", name="Value", columns=["Name", "Value"],
                      rows=[["A", values[0]], ["B", values[1]]], source="CSV")
    block = _block("visual", "chart", dataset_id="data", chart_type="bar")
    block.box.width = Inches(2)
    build_deck(template, _plan(removals, [block], count=1, datasets=[dataset]), result)
    chart = next(s.chart for s in Presentation(result).slides[0].shapes
                 if s.name == "exposlides:visual")
    assert list(chart.series[0].values) == values
    span = max(0, *values) - min(0, *values)
    assert span / chart.value_axis.major_unit <= 2
    assert chart.value_axis.tick_labels.number_format_is_linked is False


def test_browser_chart_long_categories_reduce_ticks_without_truncating_labels(tmp_path):
    template = tmp_path / "source.pptx"
    removals = _template(template)
    units = []
    for index, labels in enumerate([["A", "B", "C"], ["Первый месяц", "Второй месяц", "Третий месяц"]]):
        data = Dataset(id="tasks", name="Готовые задачи", source="CSV", unit="задач",
                       columns=["Период", "Готовые задачи"],
                       rows=[[label, value] for label, value in zip(labels, [12, 18, 24], strict=True)])
        block = _block("visual", "chart", dataset_id="tasks", chart_type="bar")
        block.box.width, block.box.height, block.style.size = 5537004, 2971800, 25
        result = build_deck(template, _plan(removals, [block], count=1, datasets=[data]),
                            tmp_path / f"chart-{index}.pptx")
        chart = next(s.chart for s in Presentation(result).slides[0].shapes
                     if s.name == "exposlides:visual")
        assert list(chart.series[0].values) == [12, 18, 24]
        assert [category.label for category in chart.plots[0].categories] == labels
        # Повторное открытие должно сохранять явное отображение средней категории.
        for element_name in ("tickLblSkip", "tickMarkSkip"):
            settings = chart.category_axis._element.xpath(f"./c:{element_name}")
            assert len(settings) == 1 and settings[0].get("val") == "1"
        assert chart.category_axis.tick_labels.font.size.pt == 18
        units.append(chart.value_axis.major_unit)
    assert units[1] > units[0]
    assert 24 / units[1] <= 3


@pytest.mark.parametrize("value", [1234567.89, 0.123456789012345, -1234567.89, 1e-20, 20.0])
def test_table_preserves_all_significant_digits_after_save_and_reopen(tmp_path, value):
    template, result = tmp_path / "source.pptx", tmp_path / "output.pptx"
    removals = _template(template)
    data = Dataset(id="data", name="Exact", columns=["Category", "Value"],
                   rows=[["Original", value]], source="input.csv")
    block = _block("visual", "table", dataset_id=data.id)
    build_deck(template, _plan(removals, [block], count=1, datasets=[data]), result)
    presentation = Presentation(result)
    assert len(presentation.slides) == 1
    shape = next(s for s in presentation.slides[0].shapes if s.name == "exposlides:visual")
    cell = shape.table.cell(1, 1)
    assert float(cell.text) == value
    assert cell.text == (str(int(value)) if value.is_integer() else str(value))
    assert cell.text_frame.paragraphs[0].runs[0].font.size.pt == block.style.size
