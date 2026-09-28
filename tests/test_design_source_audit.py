"""Аудит повторно используемых фигур сравнивает оформление с исходным PPTX."""

from copy import deepcopy

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from exposlides.design_audit import apply_fixes, audit_deck, capacity_risk
from exposlides.design_builder import build_deck
from exposlides.design_models import (
    Box,
    DeckPlan,
    PlacedBlock,
    SlideInstance,
    SlidePattern,
    Slot,
    TemplateProfile,
    TextStyle,
)
from exposlides.design_saved_audit import audit_saved_pptx


@pytest.fixture
def bound_case(tmp_path):
    template = tmp_path / "native.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    shape = group.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(6), Inches(3),
    )
    shape.name = "Original branded card"
    shape.rotation = 5
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string("AACCFF")
    frame = shape.text_frame
    frame.margin_left = Pt(13)
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    paragraph.font.name = "Georgia"
    paragraph.font.size = Pt(24)
    paragraph.font.color.rgb = RGBColor.from_string("112244")
    bullet = OxmlElement("a:buChar")
    bullet.set("char", "◆")
    paragraph._p.get_or_add_pPr().append(bullet)
    first = paragraph.add_run()
    first.text = "Heading phrase "
    first.font.bold = True
    second = paragraph.add_run()
    second.text = "and its detailed explanation"
    second.font.italic = True
    second.font.size = Pt(20)
    end = OxmlElement("a:endParaRPr")
    end.set("sz", "2400")
    paragraph._p.append(end)
    # Размер и позиция родителя отличаются от координат дочерней фигуры.
    group.width = int(group.width * 1.2)
    presentation.save(template)
    box = Box(left=shape.left, top=shape.top, width=shape.width, height=shape.height)
    style = TextStyle(font="Arial", size=24, color="000000")
    block = PlacedBlock(
        id="content", kind="text", source_shape_id=shape.shape_id, box=box, style=style,
        items=["Новая основная мысль и её объяснение", "Второй абзац с тем же оформлением"],
    )
    plan = DeckPlan(
        variant_id="story", name="История", description="Исходные фигуры",
        width=presentation.slide_width, height=presentation.slide_height,
        slides=[SlideInstance(
            id="slide-1", story_slide_id="s1", source_slide_index=1, blocks=[block],
        )],
    )
    profile = TemplateProfile(
        template_sha256="0" * 64, width=plan.width, height=plan.height,
        patterns=[SlidePattern(
            source_slide_index=1, name="Original", slots=[], content_box=box,
            font=style.font, palette=[style.color], title_style=style, body_style=style,
            mutable_shape_ids=[shape.shape_id],
        )],
    )
    return template, plan, profile


def _shape(presentation):
    return presentation.slides[0].shapes[0].shapes[0]


def test_native_text_audit_accepts_inherited_mixed_styles_and_scaled_parent(tmp_path, bound_case):
    template, plan, profile = bound_case
    output = build_deck(template, plan, tmp_path / "result.pptx")
    shape = _shape(Presentation(output))
    assert shape.name == "Original branded card"
    assert shape.text_frame.paragraphs[0].runs[0].font.name is None
    assert audit_saved_pptx(output, template, plan, profile).ok


def test_native_text_audit_accepts_only_requested_proportional_font_shrink(tmp_path, bound_case):
    template, plan, profile = bound_case
    plan.slides[0].blocks[0].style.size = 12
    output = build_deck(template, plan, tmp_path / "smaller.pptx")
    assert audit_saved_pptx(output, template, plan, profile).ok
    presentation = Presentation(output)
    _shape(presentation).text_frame.paragraphs[0].runs[0].font.size = Pt(6)
    presentation.save(output)
    assert "saved_text_style" in {
        item.rule for item in audit_saved_pptx(output, template, plan, profile).issues
    }


@pytest.mark.parametrize("prototype", ["empty", "field", "multiple"])
def test_native_text_audit_handles_original_paragraph_prototypes(tmp_path, bound_case, prototype):
    template, plan, profile = bound_case
    presentation = Presentation(template)
    frame = _shape(presentation).text_frame
    if prototype == "empty":
        frame.clear()
    elif prototype == "field":
        paragraph = frame.paragraphs[0]
        field = OxmlElement("a:fld")
        field.set("id", "{B442EA87-EE1A-4DE8-848F-B51F52B8A2E4}")
        field.set("type", "slidenum")
        text = OxmlElement("a:t")
        text.text = "1"
        field.append(text)
        paragraph._p.insert_element_before(field, "a:endParaRPr")
    else:
        paragraph = frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.size = Pt(16)
        paragraph.font.name = "Times New Roman"
        paragraph.add_run().text = "Другой образец абзаца"
        plan.slides[0].blocks[0].items.append("Третий абзац повторяет последний образец")
    presentation.save(template)
    plan.slides[0].blocks[0].style.size = 12
    output = build_deck(template, plan, tmp_path / "result.pptx")
    assert audit_saved_pptx(output, template, plan, profile).ok


def test_native_audit_checks_font_for_used_prototype_not_unused_body_paragraph(tmp_path, bound_case):
    template, plan, profile = bound_case
    presentation = Presentation(template)
    frame = _shape(presentation).text_frame
    frame.clear()
    frame.paragraphs[0].font.size = Pt(18)
    frame.paragraphs[0].add_run().text = "Короткий заголовок карточки"
    body = frame.add_paragraph()
    body.font.size = Pt(14)
    body.add_run().text = "Подробное описание в исходном образце. " * 20
    presentation.save(template)
    block = plan.slides[0].blocks[0]
    block.items = ["Один новый абзац использует только первый исходный абзац."]
    block.style.size = 14.4
    output = build_deck(template, plan, tmp_path / "result.pptx")
    changed = Presentation(output)
    run = _shape(changed).text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == 14.4
    assert audit_saved_pptx(output, template, plan, profile).ok
    run.font.size = Pt(18)
    changed.save(output)
    assert "saved_text_style" in {
        item.rule for item in audit_saved_pptx(output, template, plan, profile).issues
    }


@pytest.mark.parametrize("mutation,rule", [
    ("margin", "saved_text_style"), ("bullet", "saved_text_style"),
    ("alignment", "saved_text_style"), ("font", "saved_text_style"),
    ("color", "saved_text_style"), ("bold", "saved_text_style"),
    ("size", "saved_text_style"), ("fill", "saved_text_style"),
    ("name", "saved_text_style"), ("position", "saved_geometry"),
    ("parent", "saved_geometry"), ("text", "saved_text"),
])
def test_native_text_audit_rejects_changed_original_format(
    tmp_path, bound_case, mutation, rule,
):
    template, plan, profile = bound_case
    output = build_deck(template, plan, tmp_path / "result.pptx")
    presentation = Presentation(output)
    shape = _shape(presentation)
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.runs[0]
    if mutation == "margin":
        shape.text_frame.margin_left += Pt(1)
    elif mutation == "bullet":
        paragraph._p.get_or_add_pPr().find(qn("a:buChar")).set("char", "•")
    elif mutation == "alignment":
        paragraph.alignment = PP_ALIGN.LEFT
    elif mutation == "font":
        run.font.name = "Arial"
    elif mutation == "color":
        run.font.color.rgb = RGBColor.from_string("FF0000")
    elif mutation == "bold":
        run.font.bold = False
    elif mutation == "size":
        run.font.size = Pt(8)
    elif mutation == "fill":
        shape.fill.fore_color.rgb = RGBColor.from_string("FFAAAA")
    elif mutation == "name":
        shape.name = "replacement box"
    elif mutation == "position":
        shape.left += Pt(1)
    elif mutation == "parent":
        presentation.slides[0].shapes[0].width += Pt(1)
    elif mutation == "text":
        run.text = "Подмена текста"
    presentation.save(output)
    assert rule in {item.rule for item in audit_saved_pptx(output, template, plan, profile).issues}


def test_native_text_audit_accepts_legacy_removal_list_when_bound_shape_is_kept(tmp_path, bound_case):
    template, plan, profile = bound_case
    original = Presentation(template)
    plan.slides[0].remove_shape_ids = [
        plan.slides[0].blocks[0].source_shape_id, original.slides[0].shapes[0].shape_id,
    ]
    output = build_deck(template, plan, tmp_path / "result.pptx")
    assert audit_saved_pptx(output, template, plan, profile).ok


@pytest.mark.parametrize("invalid", ["duplicate", "absent"])
def test_native_text_audit_rejects_invalid_source_mapping(tmp_path, bound_case, invalid):
    template, plan, profile = bound_case
    output = build_deck(template, plan, tmp_path / "result.pptx")
    block = plan.slides[0].blocks[0]
    if invalid == "duplicate":
        other = deepcopy(block)
        other.id = "duplicate"
        plan.slides[0].blocks.append(other)
    else:
        block.source_shape_id = 999999
    expected = "native_object_missing" if invalid == "absent" else "source_shape_binding"
    assert expected in {item.rule for item in audit_saved_pptx(output, template, plan, profile).issues}


def test_native_fit_rejects_overflow_instead_of_expanding_source_geometry(bound_case):
    _, plan, profile = bound_case
    block = plan.slides[0].blocks[0]
    block.box.height = int(Pt(60))
    block.items = ["Содержательный тезис. " * 12]
    original = plan.model_dump()
    report = audit_deck(plan, profile)
    issue = next(item for item in report.issues if item.rule == "text_capacity")
    with pytest.raises(ValueError, match="Сократите план или выберите другой макет"):
        apply_fixes(plan, profile, report, [issue.id])
    assert plan.model_dump() == original
    # Та же геометрия позволяет расширить новый textbox, у которого нет исходного слота.
    block.source_shape_id = None
    revised = apply_fixes(plan, profile, report, [issue.id])
    assert revised.slides[0].blocks[0].box.height > block.box.height
    assert not capacity_risk(revised.slides[0].blocks[0])


def test_native_fit_allows_font_shrink_without_moving_or_resizing_shape(bound_case):
    _, plan, profile = bound_case
    block = plan.slides[0].blocks[0]
    block.box.height = int(Pt(90))
    block.items = ["Текст " * 20]
    report = audit_deck(plan, profile)
    issue = next(item for item in report.issues if item.rule == "text_capacity")
    revised = apply_fixes(plan, profile, report, [issue.id]).slides[0].blocks[0]
    assert revised.style.size < block.style.size
    assert revised.box == block.box
    assert not capacity_risk(revised)


def test_native_move_inside_requires_different_layout_and_preserves_plan(bound_case):
    _, plan, profile = bound_case
    plan.slides[0].blocks[0].box.left = -Pt(20)
    original = plan.model_dump()
    report = audit_deck(plan, profile)
    issue = next(item for item in report.issues if item.rule == "outside_slide")
    with pytest.raises(ValueError, match="Исходную фигуру шаблона нельзя переместить"):
        apply_fixes(plan, profile, report, [issue.id])
    assert plan.model_dump() == original


@pytest.mark.parametrize("paragraph_sizes", [[18, 14, 22], []])
def test_native_reference_uses_selected_slot_prototypes_for_scale_and_minimum(
    bound_case, paragraph_sizes,
):
    _, plan, profile = bound_case
    block = plan.slides[0].blocks[0]
    block.style = block.style.model_copy(update={"size": 18})
    block.box.height = int(Pt(84))
    block.items = ["Тезис " * 28]
    pattern = profile.patterns[0]
    pattern.body_style = pattern.body_style.model_copy(update={"size": 14})
    pattern.slots = [Slot(
        element_id="original", shape_id=block.source_shape_id, box=block.box, role="body",
        style=block.style.model_copy(update={"size": 14 if paragraph_sizes else 18}),
        paragraph_font_sizes=paragraph_sizes,
    )]
    report = audit_deck(plan, profile)
    assert "type_scale" not in {item.rule for item in report.issues}
    issue = next(item for item in report.issues if item.rule == "text_capacity")
    revised = apply_fixes(plan, profile, report, [issue.id]).slides[0].blocks[0]
    assert revised.style.size == pytest.approx(14.4)
    assert revised.box == block.box
    assert not capacity_risk(revised)
