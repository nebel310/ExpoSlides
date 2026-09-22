from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
CLI_WITHOUT_NETWORK_MODULES = """
import importlib.abc
import runpy
import sys

class BlockNetworkModules(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if (
            fullname.split(".")[0] in {"grpc", "aiokafka"}
            or fullname.startswith(("app.grpc", "app.kafka"))
            or fullname.endswith(("_pb2", "_pb2_grpc"))
        ):
            raise ModuleNotFoundError(f"Network module is unavailable: {fullname}")

sys.meta_path.insert(0, BlockNetworkModules())
sys.argv = ["app.main", *sys.argv[1:]]
runpy.run_module("app.main", run_name="__main__")
"""


def _run_cli(service: str, arguments: list[str], cwd: Path) -> subprocess.CompletedProcess[str]:
    env = os.environ.copy()
    env["PYTHONPATH"] = str(REPOSITORY_ROOT / "services" / service)
    env["LLM_API_KEY"] = ""
    return subprocess.run(
        [sys.executable, "-c", CLI_WITHOUT_NETWORK_MODULES, *arguments],
        cwd=cwd,
        env=env,
        text=True,
        capture_output=True,
        check=False,
        timeout=30,
    )


@pytest.mark.parametrize("service", ["parsing-service", "content-service"])
def test_service_help_does_not_import_network_modules(service: str, tmp_path: Path) -> None:
    result = _run_cli(service, ["--help"], tmp_path)

    assert result.returncode == 0, result.stderr
    assert "--output-json" in result.stdout


def test_parser_cli_creates_json_without_network_modules(tmp_path: Path) -> None:
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "template.json"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Проверка локального запуска"
    presentation.save(template_path)

    result = _run_cli(
        "parsing-service",
        ["--input-pptx", str(template_path), "--output-json", str(output_path)],
        tmp_path,
    )

    assert result.returncode == 0, result.stderr
    payload = json.loads(output_path.read_text(encoding="utf-8"))
    assert len(payload["slides"]) == 1
    assert "Проверка локального запуска" in output_path.read_text(encoding="utf-8")


def test_content_cli_validates_inputs_without_network_modules(tmp_path: Path) -> None:
    result = _run_cli(
        "content-service",
        [
            "--cli",
            "--template-json", str(tmp_path / "missing.json"),
            "--script", str(tmp_path / "missing.txt"),
            "--output-json", str(tmp_path / "result.json"),
        ],
        tmp_path,
    )

    assert result.returncode == 1, result.stderr
    assert "Presentation JSON не найден" in result.stderr
    assert "Traceback" not in result.stderr
