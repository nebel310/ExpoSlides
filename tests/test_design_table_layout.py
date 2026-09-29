"""Большие наборы не превращаются в нечитабельные таблицы и не обрезаются."""

import pytest
from pptx import Presentation
from pptx.util import Inches

from exposlides.design_builder import build_deck
from exposlides.design_layout import _table_fits, create_variants, story_layout_issues
from exposlides.design_models import Box, Dataset, VisualRequest
from tests.test_template_design import _profile, _story


def data(rows):
    return Dataset(id="data", name="Продажи", columns=["Период", "Выручка"],
                   rows=rows, source="sales.csv")


def visuals(plans):
    return [next(block for block in plan.slides[0].blocks
                 if block.kind in {"table", "chart"}) for plan in plans]


def test_compact_dataset_remains_a_readable_evidence_table(tmp_path):
    dataset = data([["Q1", 10], ["Q2", 20]])
    story = _story(visual=VisualRequest(kind="bar", dataset_id=dataset.id))
    plans = create_variants(_profile(tmp_path), story, [dataset])
    assert [block.kind for block in visuals(plans)] == ["chart", "table", "chart"]
    assert all(plan.datasets[0].rows == dataset.rows for plan in plans)


@pytest.mark.parametrize("kind", ["bar", "line", "pie"])
def test_large_dataset_keeps_original_chart_and_every_row(tmp_path, kind):
    dataset = data([[str(index), index] for index in range(1001)])
    story = _story(visual=VisualRequest(kind=kind, dataset_id=dataset.id))
    plans = create_variants(_profile(tmp_path), story, [dataset])
    assert all(block.kind == "chart" and block.chart_type == kind for block in visuals(plans))
    assert all(plan.datasets[0].rows == dataset.rows for plan in plans)


def test_long_cells_do_not_overflow_the_evidence_table(tmp_path):
    dataset = data([["Длинная подпись " * 100, 1]])
    story = _story(visual=VisualRequest(kind="bar", dataset_id=dataset.id))
    assert all(block.kind == "chart" for block in visuals(
        create_variants(_profile(tmp_path), story, [dataset]),
    ))


def test_explicit_overflowing_table_fails_before_build_without_truncation(tmp_path):
    dataset = data([[str(index), index] for index in range(1001)])
    story = _story(visual=VisualRequest(kind="table", dataset_id=dataset.id))
    profile = _profile(tmp_path)
    with pytest.raises(ValueError, match="таблица.*не помещается"):
        create_variants(profile, story, [dataset])
    errors = story_layout_issues(profile, story, [dataset])
    assert len(errors) == 1 and "Выберите график" in errors[0]
    assert len(dataset.rows) == 1001


def test_large_evidence_chart_keeps_all_values_in_reopened_pptx(tmp_path):
    template, output = tmp_path / "template.pptx", tmp_path / "result.pptx"
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(12), Inches(8)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(0.4), Inches(10), Inches(0.85)).text = "Title"
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(10), Inches(4)).text = "Body"
    presentation.save(template)
    dataset = data([[str(index), index + 0.123456789] for index in range(1001)])
    story = _story(visual=VisualRequest(kind="line", dataset_id=dataset.id))
    evidence = create_variants(_profile(tmp_path), story, [dataset])[1]
    build_deck(template, evidence, output)
    reopened = Presentation(output)
    assert len(reopened.slides) == 1
    chart = next(shape.chart for shape in reopened.slides[0].shapes if shape.has_chart)
    assert len(chart.series[0].values) == 1001
    assert list(chart.series[0].values) == [row[1] for row in dataset.rows]
    assert [category.label for category in chart.plots[0].categories] == [
        row[0] for row in dataset.rows
    ]


def test_table_capacity_uses_the_expanded_integer_float_text():
    box = Box(left=0, top=0, width=Inches(4), height=Inches(1))
    assert _table_fits(data([["Q1", 1000.0]]), box, 20)
    # В PPTX это 21 цифра, хотя str(1e20) занимает только пять знаков.
    assert not _table_fits(data([["Q1", 1e20]]), box, 20)
