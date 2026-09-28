"""Профили и три компоновки на непредварительно размеченных шаблонах."""

from __future__ import annotations

from copy import deepcopy

import pytest
from pptx.util import Inches

from exposlides.design_audit import audit_deck, capacity_risk
from exposlides.design_geometry import overlap
from exposlides.design_layout import create_variants
from exposlides.design_models import Box, ContentPlan, Dataset, StorySlide, VisualRequest
from exposlides.template_profile import profile_from_json


def _element(identifier, text, left, top, width, height, kind=None, color="202124", size=20):
    return dict(
        id=f"slide-1-shape-{identifier}",
        shape_id=identifier,
        shape_name=f"TextBox {identifier}",
        type="text",
        placeholder_kind=kind,
        bbox=dict(left=Inches(left), top=Inches(top), width=Inches(width), height=Inches(height)),
        text=dict(
            paragraphs=[
                dict(
                    runs=[
                        dict(
                            text=text, style=dict(font_name="Arial", size_pt=size, color_hex=color)
                        )
                    ]
                )
            ]
        ),
    )


def _data():
    return dict(
        schema_version="2.0.0",
        slide_width=Inches(12),
        slide_height=Inches(8),
        tokens=dict(
            theme=dict(
                colors={"dk1": "202124", "lt1": "FFFFFF", "accent1": "224488"},
                fonts={"major": "Arial", "minor": "Arial"},
            )
        ),
        slides=[
            dict(
                index=1,
                layout_index=1,
                master_index=1,
                elements=[
                    _element(2, "Title", 1, 0.4, 10, 0.85, "title", size=32),
                    _element(3, "Body", 1, 2, 10, 4),
                ],
            )
        ],
    )


def _profile(tmp_path, data=None):
    template = tmp_path / "profile-source.pptx"
    template.write_bytes(b"profile fixture")
    return profile_from_json(data or _data(), template)


def _story(paragraphs=None, visual=None, title="One conclusion"):
    return ContentPlan(
        title="Story",
        slides=[
            StorySlide(
                id="one",
                title=title,
                paragraphs=paragraphs or ["A single grounded statement."],
                source_ids=["source-1"],
                visual=visual,
            )
        ],
    )


def test_empty_placeholders_use_inherited_default_styles(tmp_path):
    data = _data()
    for element in data["slides"][0]["elements"]:
        element["text"] = None
        element["default_text_style"] = dict(
            font_name="+mj-lt", size_pt=36, color_hex="ABCDEF", bold=True
        )
    profile = _profile(tmp_path, data)
    assert profile.patterns[0].title_style.font == "Arial"
    assert profile.patterns[0].title_style.size == 36
    assert profile.patterns[0].title_style.color == "ABCDEF"
    assert "ABCDEF" in profile.patterns[0].palette
    assert profile.patterns[0].mutable_shape_ids == [2, 3]


def test_dominant_run_and_nested_world_geometry_are_retained(tmp_path):
    data = _data()
    body = data["slides"][0]["elements"][1]
    body["id"] = "slide-1-shape-8-child-0"
    body["shape_id"] = 9
    body["slide_bbox"] = dict(left=Inches(2), top=Inches(2), width=Inches(8), height=Inches(3))
    body["text"]["paragraphs"][0]["runs"].insert(
        0, dict(text="*", style=dict(font_name="Wingdings", size_pt=8))
    )
    data["slides"][0]["elements"][1] = dict(group=dict(children=[body]))
    pattern = _profile(tmp_path, data).patterns[0]
    slot = next(slot for slot in pattern.slots if slot.shape_id == 9)
    assert slot.style.font == "Arial"
    assert slot.style.size == 20
    assert slot.box.left == Inches(2)
    assert pattern.content_box.width == Inches(8)
    assert 8 not in pattern.mutable_shape_ids


def test_old_ambiguous_group_ids_fail_instead_of_deleting_parent(tmp_path):
    data = _data()
    body = data["slides"][0]["elements"][1]
    body.pop("shape_id")
    body["id"] = "slide-1-shape-8-child-0"
    with pytest.raises(ValueError, match="группы нужен shape_id"):
        _profile(tmp_path, data)


def test_protected_picture_and_inherited_logo_are_avoided(tmp_path):
    data = _data()
    image = dict(
        id="slide-1-shape-4",
        shape_id=4,
        type="image",
        bbox=dict(left=Inches(9), top=Inches(2), width=Inches(2), height=Inches(2)),
    )
    data["slides"][0]["elements"].append(image)
    logo = dict(
        id="master-1-shape-2",
        shape_id=2,
        type="image",
        bbox=dict(left=Inches(1), top=Inches(5), width=Inches(2), height=Inches(1)),
    )
    data["masters"] = [dict(index=1, elements=[logo])]
    profile = _profile(tmp_path, data)
    pattern = profile.patterns[0]
    assert 4 in pattern.protected_shape_ids
    assert 2 not in pattern.protected_shape_ids
    assert all(not overlap(pattern.content_box, obstacle) for obstacle in pattern.protected_regions)
    # Нельзя незаметно вырезать новую рамку из перекрытого исходного блока.
    with pytest.raises(ValueError, match="нет подходящих текстовых областей"):
        create_variants(profile, _story())


def test_background_container_does_not_block_its_text(tmp_path):
    data = _data()
    data["slides"][0]["elements"].append(
        dict(
            id="slide-1-shape-4",
            shape_id=4,
            type="shape",
            bbox=dict(left=Inches(0.8), top=Inches(1.8), width=Inches(10.4), height=Inches(4.4)),
            fill=dict(color_hex="334455"),
        )
    )
    pattern = _profile(tmp_path, data).patterns[0]
    assert pattern.content_box.width == Inches(10)
    assert not pattern.protected_regions
    assert 4 in pattern.protected_shape_ids
    assert "334455" in pattern.palette


def test_single_native_layout_is_preserved_without_inventing_alternatives(tmp_path):
    profile = _profile(tmp_path)
    original = profile.model_dump()
    plans = create_variants(profile, _story())
    bodies = [
        next(block for block in plan.slides[0].blocks if block.kind == "text") for plan in plans
    ]
    assert all(block.box == profile.patterns[0].slots[1].box for block in bodies)
    assert all(block.source_shape_id == 3 for block in bodies)
    assert [block.items or [block.text] for block in bodies] == [
        ["A single grounded statement."]
    ] * 3
    assert all(block.fill is None for block in bodies)
    assert profile.model_dump() == original


def test_two_page_numbers_have_unique_ids(tmp_path):
    data = _data()
    data["slides"][0]["elements"].extend(
        [
            _element(4, "01", 10, 7.7, 0.5, 0.2, "slide_number", size=10),
            _element(5, "01", 11, 7.7, 0.5, 0.2, "slide_number", size=10),
        ]
    )
    for plan in create_variants(_profile(tmp_path, data), _story()):
        pages = [block for block in plan.slides[0].blocks if block.kind == "page_number"]
        assert len({block.id for block in pages}) == 2
        assert all(block.text == "01" for block in pages)


def test_smartart_routes_only_to_native_compatible_pattern(tmp_path):
    profile = _profile(tmp_path)
    native = profile.patterns[0].model_copy(deep=True)
    native.source_slide_index = 2
    native.visual_shape_ids = {"smartart": [15]}
    profile.patterns.append(native)
    story = _story(visual=VisualRequest(kind="smartart", labels=["A", "B"]))
    for plan in create_variants(profile, story):
        assert plan.slides[0].source_slide_index == 2
        visual = next(block for block in plan.slides[0].blocks if block.kind == "smartart")
        assert visual.source_shape_id == 15
        assert visual.items == ["A", "B"]
    with pytest.raises(ValueError, match="нет нативного SmartArt"):
        create_variants(_profile(tmp_path), story)


def test_icon_and_generated_image_keep_all_facts_and_distinct_regions(tmp_path):
    profile = _profile(tmp_path)
    story = _story(visual=VisualRequest(kind="icon", icon="check"))
    for plan in create_variants(profile, story, generated_image=tmp_path / "generated.png"):
        blocks = plan.slides[0].blocks
        icon = next(block for block in blocks if block.kind == "icon")
        image = next(block for block in blocks if block.kind == "image")
        text = next(block for block in blocks if block.kind == "text")
        assert icon.icon == "check"
        assert image.image_path.endswith("generated.png")
        assert text.items or text.text
        assert not overlap(icon.box, image.box)
        assert not overlap(text.box, image.box)
        assert not overlap(text.box, icon.box)
    with pytest.raises(ValueError, match="отсутствующий слайд"):
        create_variants(
            profile, story, generated_image=tmp_path / "generated.png", image_slide_id="missing"
        )


def test_data_visuals_share_source_data_across_layouts(tmp_path):
    dataset = Dataset(
        id="revenue",
        name="Revenue",
        columns=["Quarter", "Value"],
        rows=[["Q1", 10]],
        source="Source",
        unit="million",
    )
    story = _story(visual=VisualRequest(kind="bar", dataset_id=dataset.id))
    plans = create_variants(_profile(tmp_path), story, [dataset])
    visuals = [
        next(block for block in plan.slides[0].blocks if block.kind in {"table", "chart"})
        for plan in plans
    ]
    assert [visual.kind for visual in visuals] == ["chart", "table", "chart"]
    assert all(visual.dataset_id == dataset.id for visual in visuals)
    assert len({tuple(visual.box.model_dump().values()) for visual in visuals}) == 3


def test_long_title_fits_without_mutating_template_style(tmp_path):
    data = _data()
    data["slides"][0]["elements"][0]["bbox"]["height"] = Inches(1.2)
    profile = _profile(tmp_path, data)
    story = _story(
        title="A sufficiently long conclusion about the measured product performance and the next delivery stage"
    )
    blocks = [plan.slides[0].blocks[0] for plan in create_variants(profile, story)]
    assert all(not capacity_risk(block) for block in blocks)
    assert all(25.6 <= block.style.size <= 32 for block in blocks)
    assert profile.patterns[0].title_style.size == 32


def test_dark_style_and_actual_colors_survive_profile(tmp_path):
    data = deepcopy(_data())
    for element in data["slides"][0]["elements"]:
        element["text"]["paragraphs"][0]["runs"][0]["style"]["color_hex"] = "FAFAFA"
    data["slides"][0]["background"] = dict(fill=dict(color_hex="101828"))
    profile = _profile(tmp_path, data)
    assert {"FAFAFA", "101828"}.issubset(profile.patterns[0].palette)
    cards = create_variants(profile, _story())[2]
    assert not [issue for issue in audit_deck(cards, profile).issues if issue.rule == "contrast"]


@pytest.mark.parametrize("paragraph_count", [2, 3, 6])
def test_variants_keep_all_paragraphs_in_the_existing_native_body(tmp_path, paragraph_count):
    profile = _profile(tmp_path)
    profile.patterns[0].palette = ["000000", "FFFFFF"]
    profile.patterns[0].body_style.color = "000000"
    story = _story()
    story.slides[0].paragraphs = [f"Grounded statement {i}." for i in range(paragraph_count)]
    plans = create_variants(profile, story)
    evidence, cards = [
        [block for block in plan.slides[0].blocks if block.kind == "text"]
        for plan in plans[1:]
    ]
    assert [text for block in evidence for text in block.items] == story.slides[0].paragraphs
    assert [text for block in cards for text in block.items] == story.slides[0].paragraphs
    assert all(block.source_shape_id == 3 for block in [*cards, *evidence])
    assert [block.box for block in evidence] == [block.box for block in cards]
    assert all(block.fill is None for block in cards)
    assert not any(issue.rule in {"overlap", "outside_slide", "protected_overlap"}
                   for issue in audit_deck(plans[2], profile).issues)


def test_visuals_prefer_actual_brand_accent_over_unused_hyperlink_color(tmp_path):
    data = _data()
    data["tokens"]["theme"]["colors"].update(hlink="0000FF", accent1="4472C4")
    data["slides"][0]["elements"].append(dict(
        id="slide-1-shape-8", shape_id=8, type="shape", fill=dict(color_hex="BB5427"),
        bbox=dict(left=0, top=0, width=Inches(0.15), height=Inches(8)),
    ))
    story = _story(visual=VisualRequest(kind="process", labels=["Discover", "Deliver"]))
    for plan in create_variants(_profile(tmp_path, data), story):
        visual = next(block for block in plan.slides[0].blocks if block.kind == "process")
        assert visual.fill == "BB5427"


def test_browser_case_allocates_rows_by_text_and_keeps_chart_readable(tmp_path):
    profile = _profile(tmp_path)
    profile.width, profile.height = 12191695, 6858000
    pattern = profile.patterns[0]
    pattern.content_box = Box(left=777240, top=2560320, width=9966960, height=2971800)
    pattern.body_style.size = 25
    paragraphs = [
        "Команда выпустила обновлённый первый запуск, подготовила базу знаний и провела интервью с клиентами.",
        "Основной запрос клиентов — понятный путь от знакомства до первого результата.",
    ]
    data = Dataset(id="tasks", name="Готовые задачи", source="CSV", unit="задач",
                   columns=["Период", "Готовые задачи"],
                   rows=[["Первый месяц", 12], ["Второй месяц", 18], ["Третий месяц", 24]])
    story = _story(paragraphs, VisualRequest(kind="bar", dataset_id=data.id))
    plans = create_variants(profile, story, [data])
    for plan in plans:
        assert not [issue for issue in audit_deck(plan, profile).issues
                    if issue.rule in {"text_capacity", "overlap", "outside_slide"}]
        body = [block for block in plan.slides[0].blocks if block.kind == "text"]
        assert [text for block in body for text in (block.items or [block.text])] == paragraphs
    cards = [block for block in plans[2].slides[0].blocks if block.kind == "text"]
    assert cards[0].box.height > cards[1].box.height
    assert all(block.style.size >= 20 for block in cards)
    first = plans[0].slides[0].blocks
    visual, body = next(b for b in first if b.kind == "chart"), next(b for b in first if b.kind == "text")
    assert visual.box.height == pattern.content_box.height
    assert body.box.left+body.box.width < visual.box.left


def test_story_caption_reserves_its_required_height_before_visual(tmp_path):
    profile = _profile(tmp_path)
    profile.patterns[0].content_box.height = Inches(3.25)
    profile.patterns[0].body_style.size = 25
    text = "Исходное подтверждённое содержание для подробного объяснения решения команде. " * 2
    story = _story([text], VisualRequest(kind="process", labels=["Анализ", "Результат"]))
    plan = create_variants(profile, story)[0]
    body = next(block for block in plan.slides[0].blocks if block.kind == "text")
    visual = next(block for block in plan.slides[0].blocks if block.kind == "process")
    assert body.box.height > profile.patterns[0].content_box.height*0.28
    assert not capacity_risk(body)
    assert not overlap(body.box, visual.box)
    assert body.items == [text]
