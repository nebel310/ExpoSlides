"""Аудит допускает очистку custom shows, сохраняя проверку защищённых объектов."""

import hashlib

import pytest
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

from exposlides.design_builder import build_deck
from exposlides.design_models import Box, SlidePattern, TemplateProfile, TextStyle
from exposlides.design_saved_audit import audit_saved_pptx
from tests.test_design_presentation_metadata import _template


def navigation_case(tmp_path):
    source, output = tmp_path / "source.pptx", tmp_path / "result.pptx"
    plan = _template(source, "custom_show")
    profile = TemplateProfile(
        template_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        width=plan.width, height=plan.height, patterns=[SlidePattern(
            source_slide_index=index, name=f"Slide {index}", slots=[], font="Arial",
            palette=["123456", "FFFFFF"], title_style=TextStyle(), body_style=TextStyle(),
            content_box=Box(left=0, top=0, width=plan.width, height=plan.height),
            protected_shape_ids=[2],
        ) for index in range(1, 4)],
    )
    build_deck(source, plan, output)
    return source, output, plan, profile


def test_pruned_navigation_passes_saved_audit_without_changing_the_source(tmp_path):
    source, output, plan, profile = navigation_case(tmp_path)
    report = audit_saved_pptx(output, source, plan, profile)
    assert report.ok, [(issue.rule, issue.message) for issue in report.issues]
    assert hashlib.sha256(source.read_bytes()).hexdigest() == profile.template_sha256
    saved = Presentation(output)
    assert len(saved.slides) == 3
    assert saved.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].hyperlink.address == (
        "https://example.com/kept"
    )


@pytest.mark.parametrize("tampering,rule", [
    ("external_link", "protected_changed"),
    ("font", "protected_changed"),
    ("custom_show_action", "protected_changed"),
    ("layout", "template_layout"),
])
def test_navigation_normalization_does_not_hide_protected_tampering(tmp_path, tampering, rule):
    source, output, plan, profile = navigation_case(tmp_path)
    saved = Presentation(output)
    shape = saved.slides[0].shapes.title
    if tampering == "external_link":
        shape.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.com/changed"
    elif tampering == "font":
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    elif tampering == "custom_show_action":
        etree.SubElement(shape._element.find(".//" + qn("p:cNvPr")), qn("a:hlinkClick"),
                         action="ppaction://customshow?id=42")
    else:
        saved.slides[0].slide_layout.shapes[0].name = "Changed protected layout"
    saved.save(output)
    report = audit_saved_pptx(output, source, plan, profile)
    assert any(issue.rule == rule and issue.severity == "error" for issue in report.issues)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == profile.template_sha256
