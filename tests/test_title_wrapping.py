"""Короткие заголовки читаются цельно, не меняя геометрию шаблона."""

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from exposlides.design_audit import audit_saved_pptx
from exposlides.design_builder import build_deck
from exposlides.design_layout import create_variants
from exposlides.template_layout import native_layout
from exposlides.template_profile import profile_from_json
from tests.test_template_design import _data, _element, _profile, _story


def _title_patterns(tmp_path, *, size=36, narrow=250.5):
    profile = _profile(tmp_path)
    wide = profile.patterns[0]
    wide.slots[0].style.size = size
    wide.slots[0].paragraph_font_sizes = [size]
    wide.slots[0].box.height = int(106.2 * 12700)
    narrow_pattern = wide.model_copy(deep=True)
    narrow_pattern.source_slide_index = 2
    narrow_pattern.slots[0].box.width = int(narrow * 12700)
    profile.patterns.append(narrow_pattern)
    return profile


def test_short_title_uses_limited_shrink_without_changing_native_geometry(tmp_path):
    profile = _title_patterns(tmp_path, narrow=300)
    profile.patterns = profile.patterns[1:]
    original = profile.model_copy(deep=True)
    selected, blocks = native_layout(profile, _story(title="Создание агента").slides[0], "story", 1)
    title = blocks[0]

    assert title.style.size == 29
    assert title.box == selected.slots[0].box
    assert title.source_shape_id == selected.slots[0].shape_id
    assert title.text == "Создание агента"
    assert profile == original


@pytest.mark.parametrize("variant", ["story", "evidence", "cards"])
def test_short_title_readability_precedes_unused_and_alternative_compositions(tmp_path, variant):
    profile = _title_patterns(tmp_path)
    selected, blocks = native_layout(
        profile, _story(title="Создание агента").slides[0], variant, 1,
        previous=(1,), alternatives=(1,),
    )

    assert selected.source_slide_index == 1
    assert blocks[0].style.size == 36


def test_shrinkable_short_title_still_uses_new_composition(tmp_path):
    profile = _title_patterns(tmp_path, narrow=300)
    selected, blocks = native_layout(
        profile, _story(title="Создание агента").slides[0], "story", 1, previous=(1,),
    )

    assert selected.source_slide_index == 2
    assert blocks[0].style.size == 29


def test_only_narrow_title_keeps_fitting_two_lines_and_original_size(tmp_path):
    profile = _title_patterns(tmp_path)
    profile.patterns = profile.patterns[1:]
    selected, blocks = native_layout(profile, _story(title="Создание агента").slides[0], "story", 1)

    assert selected.source_slide_index == 2
    assert blocks[0].style.size == 36


@pytest.mark.parametrize("title", [
    "Создание универсального агента для работы",
    "Создание\nагента",
])
def test_long_or_explicitly_multiline_title_does_not_override_diversity(tmp_path, title):
    profile = _title_patterns(tmp_path)
    for pattern in profile.patterns:
        pattern.slots[0].box.height = int(260 * 12700)
        pattern.slots[1].box.top = int(330 * 12700)
        pattern.slots[1].box.height = int(200 * 12700)
    selected, blocks = native_layout(
        profile, _story(title=title).slides[0], "story", 1, previous=(1,),
    )

    assert selected.source_slide_index == 2
    assert blocks[0].style.size == 36
    assert blocks[0].text == title


def test_short_title_shrink_survives_native_pptx_build(tmp_path):
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(12), Inches(8)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = _data()
    data["slides"][0]["elements"] = []
    for text, left, top, width, height, role, size in [
        ("Title", 1, .4, 300 / 72, 106.2 / 72, "title", 36),
        ("Body", 1, 2, 10, 4, None, 20),
    ]:
        shape = slide.shapes.add_textbox(
            Inches(left), Inches(top), Inches(width), Inches(height),
        )
        shape.text = text
        shape.text_frame.word_wrap = True
        shape.text_frame.paragraphs[0].font.size = Pt(size)
        data["slides"][0]["elements"].append(
            _element(shape.shape_id, text, left, top, width, height, role, size=size),
        )
    presentation.save(template)
    profile = profile_from_json(data, template)
    plan = create_variants(profile, _story(title="Создание агента"))[0]
    output = build_deck(template, plan, tmp_path / "result.pptx")
    result = Presentation(output)
    title = result.slides[0].shapes[0]

    assert title.text == "Создание агента"
    assert title.text_frame.paragraphs[0].runs[0].font.size.pt == 29
    assert (title.left, title.top, title.width, title.height) == (
        slide.shapes[0].left, slide.shapes[0].top,
        slide.shapes[0].width, slide.shapes[0].height,
    )
    assert audit_saved_pptx(output, template, plan, profile).ok
