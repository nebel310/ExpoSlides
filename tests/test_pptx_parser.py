from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from pptx import Presentation as PPTXPresentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
PARSING_SERVICE_ROOT = REPOSITORY_ROOT / "services" / "parsing-service"


def _write_test_presentation(path: Path) -> None:
    presentation = PPTXPresentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    text_box.text = "Stable parser output"
    presentation.save(path)


def _element_ids(presentation) -> list[str]:
    layout_ids = [element.id for layout in presentation.layouts for element in layout.elements]
    slide_ids = [element.id for slide in presentation.slides for element in slide.elements]
    return layout_ids + slide_ids


def test_element_ids_are_stable_and_unique(tmp_path: Path, service_importer) -> None:
    input_path = tmp_path / "input.pptx"
    _write_test_presentation(input_path)
    parser = service_importer(PARSING_SERVICE_ROOT, "app.parsers.pptx_parser").PPTXParser

    first_result = asyncio.run(parser.parse(input_path))
    second_result = asyncio.run(parser.parse(input_path))
    first_ids = _element_ids(first_result)
    second_ids = _element_ids(second_result)

    assert first_ids == second_ids
    assert len(first_ids) == len(set(first_ids))
    assert all(element_id.startswith(("layout-", "slide-")) for element_id in first_ids)


def test_theme_xml_is_read_from_related_part(tmp_path: Path, service_importer) -> None:
    input_path = tmp_path / "input.pptx"
    _write_test_presentation(input_path)
    parser = service_importer(PARSING_SERVICE_ROOT, "app.parsers.pptx_parser").PPTXParser

    result = asyncio.run(parser.parse(input_path))

    assert result.theme is not None
    assert result.theme.colors
    assert result.theme.fonts


def test_slide_layout_index_matches_presentation_layout_order(
    tmp_path: Path,
    service_importer,
) -> None:
    input_path = tmp_path / "input.pptx"
    presentation = PPTXPresentation()
    presentation.slides.add_slide(presentation.slide_layouts[1])
    presentation.slides.add_slide(presentation.slide_layouts[5])
    presentation.save(input_path)
    parser = service_importer(PARSING_SERVICE_ROOT, "app.parsers.pptx_parser").PPTXParser

    result = asyncio.run(parser.parse(input_path))

    assert [slide.layout_index for slide in result.slides] == [2, 6]
    assert result.slides[0].layout_name == result.layouts[1].name
    assert result.slides[1].layout_name == result.layouts[5].name


def test_center_paragraph_alignment_is_preserved(tmp_path: Path, service_importer) -> None:
    input_path = tmp_path / "input.pptx"
    presentation = PPTXPresentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    paragraph = text_box.text_frame.paragraphs[0]
    paragraph.text = "Centered parser output"
    paragraph.alignment = PP_ALIGN.CENTER
    presentation.save(input_path)
    parser = service_importer(PARSING_SERVICE_ROOT, "app.parsers.pptx_parser").PPTXParser

    result = asyncio.run(parser.parse(input_path))

    text_element = next(
        element
        for element in result.slides[0].elements
        if element.text and element.text.full_text == "Centered parser output"
    )
    style = text_element.text.paragraphs[0].runs[0].style
    assert style.alignment.value == "center"


def _list_font(parent, level: int, size: float, name: str) -> None:
    """Задаёт наследуемый шрифт в тестовом XML, не создавая двоичный fixture."""
    from pptx.oxml.ns import qn

    properties = parent.find(qn(f"a:lvl{level + 1}pPr"))
    if properties is None:
        properties = OxmlElement(f"a:lvl{level + 1}pPr")
        parent.append(properties)
    defaults = properties.find(qn("a:defRPr"))
    if defaults is None:
        defaults = OxmlElement("a:defRPr")
        properties.append(defaults)
    defaults.set("sz", str(round(size * 100)))
    latin = defaults.find(qn("a:latin"))
    if latin is None:
        latin = OxmlElement("a:latin")
        defaults.append(latin)
    latin.set("typeface", name)


def _serialized_slide(presentation, tmp_path: Path, service_importer):
    path = tmp_path / "styles.pptx"
    presentation.save(path)
    parser = service_importer(PARSING_SERVICE_ROOT, "app.parsers.pptx_parser").PPTXParser
    return asyncio.run(parser.parse(path)).model_dump(mode="json")["slides"][0]


def _styles_by_text(slide) -> dict:
    return {
        run["text"]: run["style"]
        for element in slide["elements"]
        if element["text"]
        for paragraph in element["text"]["paragraphs"]
        for run in paragraph["runs"]
    }


def test_explicit_run_font_wins_and_paragraph_default_is_inherited(
    tmp_path: Path, service_importer,
) -> None:
    presentation = PPTXPresentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.font.size = Pt(32)
    paragraph.font.name = "Verdana"
    inherited = paragraph.add_run()
    inherited.text = "Inherited"
    explicit = paragraph.add_run()
    explicit.text = "Explicit"
    explicit.font.size = Pt(18)
    explicit.font.name = "Calibri"

    styles = _styles_by_text(_serialized_slide(presentation, tmp_path, service_importer))

    assert styles["Inherited"]["size_pt"] == 32
    assert styles["Inherited"]["font_name"] == "Verdana"
    assert styles["Explicit"]["size_pt"] == 18
    assert styles["Explicit"]["font_name"] == "Calibri"


def test_shape_list_style_uses_paragraph_level(tmp_path: Path, service_importer) -> None:
    from pptx.oxml.ns import qn

    presentation = PPTXPresentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = "Third level"
    paragraph.level = 2
    list_style = shape._element.find(qn("p:txBody")).find(qn("a:lstStyle"))
    _list_font(list_style, 0, 50, "Arial")
    _list_font(list_style, 2, 15, "Cambria")

    styles = _styles_by_text(_serialized_slide(presentation, tmp_path, service_importer))

    assert styles["Third level"]["size_pt"] == 15
    assert styles["Third level"]["font_name"] == "Cambria"


def test_layout_font_matches_placeholder_index_not_first_shape_of_type(
    tmp_path: Path, service_importer,
) -> None:
    presentation = PPTXPresentation()
    layout = presentation.slide_layouts[3]
    objects = [
        shape for shape in layout.placeholders
        if shape.placeholder_format.type == PP_PLACEHOLDER.OBJECT
    ]
    assert len(objects) == 2
    for shape, size in zip(objects, (16, 24), strict=True):
        shape.text_frame.paragraphs[0].font.size = Pt(size)
        shape.text_frame.paragraphs[0].font.name = "Arial Black"
    slide = presentation.slides.add_slide(layout)
    for index, shape in enumerate(objects):
        slide.placeholders[shape.placeholder_format.idx].text = f"Column {index}"

    styles = _styles_by_text(_serialized_slide(presentation, tmp_path, service_importer))

    assert styles["Column 0"]["size_pt"] == 16
    assert styles["Column 1"]["size_pt"] == 24
    assert styles["Column 1"]["font_name"] == "Arial Black"


@pytest.mark.parametrize(
    ("target", "style_name", "size"),
    [("title", "p:titleStyle", 44), ("body", "p:bodyStyle", 27),
     ("textbox", "p:otherStyle", 19)],
)
def test_master_text_style_font_is_serialized(
    tmp_path: Path, service_importer, target: str, style_name: str, size: int,
) -> None:
    from pptx.oxml.ns import qn

    presentation = PPTXPresentation()
    layout = presentation.slide_layouts[1]
    text_styles = layout.slide_master._element.find(qn("p:txStyles"))
    _list_font(text_styles.find(qn(style_name)), 0, size, "Arial Black")
    slide = presentation.slides.add_slide(layout)
    if target == "title":
        shape = slide.shapes.title
    elif target == "body":
        shape = slide.placeholders[1]
    else:
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.text = "Inherited master font"

    styles = _styles_by_text(_serialized_slide(presentation, tmp_path, service_importer))

    assert styles["Inherited master font"]["size_pt"] == size
    assert styles["Inherited master font"]["font_name"] == "Arial Black"


def test_master_placeholder_font_precedes_master_text_style(
    tmp_path: Path, service_importer,
) -> None:
    presentation = PPTXPresentation()
    layout = presentation.slide_layouts[1]
    master_title = next(
        shape for shape in layout.slide_master.placeholders
        if shape.placeholder_format.type == PP_PLACEHOLDER.TITLE
    )
    master_title.text_frame.paragraphs[0].font.size = Pt(39)
    master_title.text_frame.paragraphs[0].font.name = "Georgia"
    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "Master placeholder"

    styles = _styles_by_text(_serialized_slide(presentation, tmp_path, service_importer))

    assert styles["Master placeholder"]["size_pt"] == 39
    assert styles["Master placeholder"]["font_name"] == "Georgia"


def test_center_title_uses_title_master_theme_font(tmp_path: Path, service_importer) -> None:
    from pptx.oxml.ns import qn

    presentation = PPTXPresentation()
    layout = presentation.slide_layouts[0]
    title_style = layout.slide_master._element.find(qn("p:txStyles")).find(qn("p:titleStyle"))
    _list_font(title_style, 0, 44, "+mj-lt")
    slide = presentation.slides.add_slide(layout)
    assert slide.shapes.title.placeholder_format.type == PP_PLACEHOLDER.CENTER_TITLE
    slide.shapes.title.text = "Center title"
    # У стандартного титульного макета есть собственное уменьшение размера.
    for paragraph in layout.placeholders.get(idx=0).text_frame.paragraphs:
        paragraph.font.size = None

    result = _serialized_slide(presentation, tmp_path, service_importer)
    styles = _styles_by_text(result)

    assert styles["Center title"]["size_pt"] == 44
    assert styles["Center title"]["font_name"] == "Calibri"


def test_presentation_default_font_applies_when_master_has_no_value(
    tmp_path: Path, service_importer,
) -> None:
    from pptx.oxml.ns import qn

    presentation = PPTXPresentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.text = "Presentation default"
    other_style = slide.slide_layout.slide_master._element.find(qn("p:txStyles"))
    other_style.remove(other_style.find(qn("p:otherStyle")))
    _list_font(presentation._element.find(qn("p:defaultTextStyle")), 0, 22, "Verdana")

    styles = _styles_by_text(_serialized_slide(presentation, tmp_path, service_importer))

    assert styles["Presentation default"]["size_pt"] == 22
    assert styles["Presentation default"]["font_name"] == "Verdana"
