from __future__ import annotations

import asyncio
import importlib
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation as PPTXPresentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Pt

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
BUILDER_SERVICE_ROOT = REPOSITORY_ROOT / "services" / "builder-service"
P14_NAMESPACE = "http://schemas.microsoft.com/office/powerpoint/2010/main"


def _load_builder_modules(service_importer):
    builder = service_importer(BUILDER_SERVICE_ROOT, "app.builder")
    content = importlib.import_module("app.models.content")
    errors = importlib.import_module("app.errors")
    presentation = importlib.import_module("app.models.presentation")
    return builder, content, errors, presentation


def _write_template(path: Path) -> None:
    presentation = PPTXPresentation()
    for index in range(1, 4):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        title = slide.shapes.title
        title.text = f"Template {index}"
        if index != 3:
            continue

        paragraph = title.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.runs[0]
        run.font.name = "Arial"
        run.font.size = Pt(31)
        run.font.bold = True
        run.font.italic = True
        run.font.underline = True
        run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
        run.hyperlink.address = "https://example.com/template-link"
    presentation.save(path)


def _add_custom_show(path: Path) -> None:
    presentation = PPTXPresentation(path)
    relationship_ids = [slide_id.rId for slide_id in presentation.slides._sldIdLst]
    show_list = etree.Element(qn("p:custShowLst"))
    show = etree.SubElement(show_list, qn("p:custShow"))
    show.set("name", "Regression show")
    show.set("id", "1")
    slide_list = etree.SubElement(show, qn("p:sldLst"))
    for relationship_id in relationship_ids:
        reference = etree.SubElement(slide_list, qn("p:sld"))
        reference.set(qn("r:id"), relationship_id)
    presentation.part._element.insert_element_before(
        show_list,
        "p:photoAlbum",
        "p:custDataLst",
        "p:kinsoku",
        "p:defaultTextStyle",
        "p:modifyVerifier",
        "p:extLst",
    )
    presentation.save(path)


def _add_section(path: Path) -> None:
    presentation = PPTXPresentation(path)
    extension_list = presentation.part._element.find(qn("p:extLst"))
    if extension_list is None:
        extension_list = etree.SubElement(
            presentation.part._element,
            qn("p:extLst"),
        )
    extension = etree.SubElement(extension_list, qn("p:ext"))
    extension.set("uri", "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    section_list = etree.SubElement(
        extension,
        f"{{{P14_NAMESPACE}}}sectionLst",
        nsmap={"p14": P14_NAMESPACE},
    )
    section = etree.SubElement(section_list, f"{{{P14_NAMESPACE}}}section")
    section.set("name", "Regression section")
    section.set("id", "{67F85625-23A0-410C-80E6-77DA30B6F99A}")
    slide_id_list = etree.SubElement(section, f"{{{P14_NAMESPACE}}}sldIdLst")
    for slide_id in presentation.slides._sldIdLst:
        reference = etree.SubElement(slide_id_list, f"{{{P14_NAMESPACE}}}sldId")
        reference.set("id", str(slide_id.id))
    presentation.save(path)


def _template_data(model_module, path: Path):
    source = PPTXPresentation(path)
    slides = []
    for index, slide in enumerate(source.slides, start=1):
        elements = []
        for shape in slide.shapes:
            if not shape.is_placeholder or not shape.has_text_frame:
                continue
            elements.append(
                model_module.SlideElement(
                    id=f"slide-{index}-shape-{shape.shape_id}",
                    type="text",
                    bbox=model_module.BBox(
                        left=shape.left,
                        top=shape.top,
                        width=shape.width,
                        height=shape.height,
                    ),
                    placeholder_type=shape.placeholder_format.type.name,
                    placeholder_idx=shape.placeholder_format.idx,
                    placeholder_name=shape.name,
                )
            )
        slides.append(
            model_module.Slide(
                index=index,
                layout_name=slide.slide_layout.name,
                elements=elements,
            )
        )
    return model_module.Presentation(
        source_path=str(path),
        slide_width=source.slide_width,
        slide_height=source.slide_height,
        slides=slides,
    )


def _valid_content(content_module):
    return content_module.GeneratedContent(
        content={
            3: content_module.SlideContent(placeholders={"0": "Generated 3"}),
            1: content_module.SlideContent(placeholders={"0": "Generated 1"}),
        },
        validation_report={"ok": True},
    )


def test_build_preserves_requested_order_style_and_relationships(
    tmp_path: Path,
    service_importer,
) -> None:
    builder, content, _, presentation = _load_builder_modules(service_importer)
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "nested" / "output" / "result.pptx"
    _write_template(template_path)

    result_path = asyncio.run(
        builder.PPTXBuilder.build(
            template_path,
            _template_data(presentation, template_path),
            _valid_content(content),
            output_path,
        )
    )

    assert result_path == output_path
    reopened = PPTXPresentation(output_path)
    assert len(reopened.slides) == 2
    assert [slide.shapes.title.text for slide in reopened.slides] == [
        "Generated 3",
        "Generated 1",
    ]

    first_paragraph = reopened.slides[0].shapes.title.text_frame.paragraphs[0]
    first_run = first_paragraph.runs[0]
    assert first_paragraph.alignment == PP_ALIGN.CENTER
    assert first_run.font.name == "Arial"
    assert first_run.font.size.pt == 31
    assert first_run.font.bold is True
    assert first_run.font.italic is True
    assert first_run.font.underline is True
    assert first_run.font.color.rgb == RGBColor(0x12, 0x34, 0x56)
    assert first_run.hyperlink.address == "https://example.com/template-link"

    slide_relationships = [
        relationship
        for relationship in reopened.part.rels.values()
        if relationship.reltype == RT.SLIDE
    ]
    assert len(slide_relationships) == 2
    assert {
        relationship.target_part.slide.shapes.title.text
        for relationship in slide_relationships
    } == {"Generated 1", "Generated 3"}


def test_multiline_placeholder_becomes_styled_paragraphs(
    tmp_path: Path,
    service_importer,
) -> None:
    builder, content, _, presentation = _load_builder_modules(service_importer)
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "result.pptx"
    source = PPTXPresentation()
    slide = source.slides.add_slide(source.slide_layouts[1])
    slide.shapes.title.text = "Title"
    body = slide.placeholders[1]
    body.text = "Template body"
    body_run = body.text_frame.paragraphs[0].runs[0]
    body_run.font.name = "Arial"
    body_run.font.size = Pt(20)
    source.save(template_path)
    content_data = content.GeneratedContent(
        content={
            1: content.SlideContent(
                placeholders={"0": "Generated title", "1": "Первый пункт\nВторой пункт"}
            )
        }
    )

    asyncio.run(
        builder.PPTXBuilder.build(
            template_path,
            _template_data(presentation, template_path),
            content_data,
            output_path,
        )
    )

    reopened = PPTXPresentation(output_path)
    paragraphs = reopened.slides[0].placeholders[1].text_frame.paragraphs
    assert [paragraph.text for paragraph in paragraphs] == ["Первый пункт", "Второй пункт"]
    assert all(paragraph.runs[0].font.name == "Arial" for paragraph in paragraphs)
    assert all(paragraph.runs[0].font.size == Pt(20) for paragraph in paragraphs)


@pytest.mark.parametrize(
    "case_name",
    [
        "service_error",
        "empty_content",
        "empty_placeholders",
        "blank_placeholder",
        "unknown_slide",
        "unknown_placeholder",
    ],
)
def test_invalid_input_keeps_existing_output_unchanged(
    case_name: str,
    tmp_path: Path,
    service_importer,
) -> None:
    builder, content, errors, presentation = _load_builder_modules(service_importer)
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "result.pptx"
    _write_template(template_path)
    template_data = _template_data(presentation, template_path)
    sentinel = b"existing output must survive"
    output_path.write_bytes(sentinel)

    if case_name == "service_error":
        content_data = content.GeneratedContent(error="generation failed")
    elif case_name == "empty_content":
        content_data = content.GeneratedContent()
    elif case_name == "empty_placeholders":
        content_data = content.GeneratedContent(
            content={1: content.SlideContent(placeholders={})}
        )
    elif case_name == "blank_placeholder":
        content_data = content.GeneratedContent(
            content={1: content.SlideContent(placeholders={"0": "   "})}
        )
    elif case_name == "unknown_slide":
        content_data = content.GeneratedContent(
            content={99: content.SlideContent(placeholders={"0": "Unknown"})}
        )
    else:
        content_data = content.GeneratedContent(
            content={1: content.SlideContent(placeholders={"999": "Unknown"})}
        )

    with pytest.raises(errors.BuilderError):
        asyncio.run(
            builder.PPTXBuilder.build(
                template_path,
                template_data,
                content_data,
                output_path,
            )
        )

    assert output_path.read_bytes() == sentinel


def test_template_mismatch_keeps_existing_output_unchanged(
    tmp_path: Path,
    service_importer,
) -> None:
    builder, content, errors, presentation = _load_builder_modules(service_importer)
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "result.pptx"
    _write_template(template_path)
    template_data = _template_data(presentation, template_path)
    mismatched_data = template_data.model_copy(update={"slides": template_data.slides[:2]})
    sentinel = b"existing output must survive"
    output_path.write_bytes(sentinel)

    with pytest.raises(errors.BuilderInputError, match="Состав слайдов"):
        asyncio.run(
            builder.PPTXBuilder.build(
                template_path,
                mismatched_data,
                _valid_content(content),
                output_path,
            )
        )

    assert output_path.read_bytes() == sentinel


def test_missing_placeholder_is_rejected(
    tmp_path: Path,
    service_importer,
) -> None:
    builder, content, errors, presentation = _load_builder_modules(service_importer)
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "result.pptx"
    source = PPTXPresentation()
    slide = source.slides.add_slide(source.slide_layouts[1])
    slide.shapes.title.text = "Title"
    slide.placeholders[1].text = "Body"
    source.save(template_path)
    sentinel = b"existing output must survive"
    output_path.write_bytes(sentinel)
    content_data = content.GeneratedContent(
        content={1: content.SlideContent(placeholders={"0": "Generated title"})}
    )

    with pytest.raises(errors.BuilderInputError, match="отсутствуют обязательные"):
        asyncio.run(
            builder.PPTXBuilder.build(
                template_path,
                _template_data(presentation, template_path),
                content_data,
                output_path,
            )
        )

    assert output_path.read_bytes() == sentinel


def test_blank_list_line_is_rejected(
    tmp_path: Path,
    service_importer,
) -> None:
    builder, content, errors, presentation = _load_builder_modules(service_importer)
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "result.pptx"
    source = PPTXPresentation()
    slide = source.slides.add_slide(source.slide_layouts[1])
    slide.shapes.title.text = "Title"
    slide.placeholders[1].text = "Body"
    source.save(template_path)
    sentinel = b"existing output must survive"
    output_path.write_bytes(sentinel)
    content_data = content.GeneratedContent(
        content={
            1: content.SlideContent(
                placeholders={
                    "0": "Generated title",
                    "1": "Первый пункт\n\nВторой пункт",
                }
            )
        }
    )

    with pytest.raises(errors.BuilderInputError, match="пустые строки"):
        asyncio.run(
            builder.PPTXBuilder.build(
                template_path,
                _template_data(presentation, template_path),
                content_data,
                output_path,
            )
        )

    assert output_path.read_bytes() == sentinel


def test_blank_line_in_non_list_placeholder_remains_supported(
    tmp_path: Path,
    service_importer,
) -> None:
    builder, content, _, presentation = _load_builder_modules(service_importer)
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "result.pptx"
    source = PPTXPresentation()
    slide = source.slides.add_slide(source.slide_layouts[5])
    slide.shapes.title.text = "Title"
    source.save(template_path)
    content_data = content.GeneratedContent(
        content={
            1: content.SlideContent(
                placeholders={"0": "Первая строка\n\nВторая строка"}
            )
        }
    )

    asyncio.run(
        builder.PPTXBuilder.build(
            template_path,
            _template_data(presentation, template_path),
            content_data,
            output_path,
        )
    )

    reopened = PPTXPresentation(output_path)
    paragraphs = reopened.slides[0].shapes.title.text_frame.paragraphs
    assert [paragraph.text for paragraph in paragraphs] == [
        "Первая строка",
        "",
        "Вторая строка",
    ]


def test_custom_show_slide_removal_is_rejected_before_output_changes(
    tmp_path: Path,
    service_importer,
) -> None:
    builder, content, errors, presentation = _load_builder_modules(service_importer)
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "result.pptx"
    _write_template(template_path)
    _add_custom_show(template_path)
    original_template = template_path.read_bytes()
    sentinel = b"existing output must survive"
    output_path.write_bytes(sentinel)
    content_data = content.GeneratedContent(
        content={1: content.SlideContent(placeholders={"0": "Generated 1"})}
    )

    with pytest.raises(errors.BuilderInputError, match="custom shows"):
        asyncio.run(
            builder.PPTXBuilder.build(
                template_path,
                _template_data(presentation, template_path),
                content_data,
                output_path,
            )
        )

    assert template_path.read_bytes() == original_template
    assert output_path.read_bytes() == sentinel


def test_section_slide_removal_is_rejected_before_output_changes(
    tmp_path: Path,
    service_importer,
) -> None:
    builder, content, errors, presentation = _load_builder_modules(service_importer)
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "result.pptx"
    _write_template(template_path)
    _add_section(template_path)
    original_template = template_path.read_bytes()
    sentinel = b"existing output must survive"
    output_path.write_bytes(sentinel)
    content_data = content.GeneratedContent(
        content={1: content.SlideContent(placeholders={"0": "Generated 1"})}
    )

    with pytest.raises(errors.BuilderInputError, match="sections"):
        asyncio.run(
            builder.PPTXBuilder.build(
                template_path,
                _template_data(presentation, template_path),
                content_data,
                output_path,
            )
        )

    assert template_path.read_bytes() == original_template
    assert output_path.read_bytes() == sentinel


def test_post_save_validation_failure_is_atomic(
    tmp_path: Path,
    monkeypatch,
    service_importer,
) -> None:
    builder, content, errors, presentation = _load_builder_modules(service_importer)
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "result.pptx"
    _write_template(template_path)
    sentinel = b"existing output must survive"
    output_path.write_bytes(sentinel)

    def fail_validation(path: Path, expected_slide_count: int) -> None:
        del path, expected_slide_count
        raise errors.BuilderOutputError("forced validation failure")

    monkeypatch.setattr(builder, "_validate_saved_presentation", fail_validation)

    with pytest.raises(errors.BuilderOutputError, match="forced validation failure"):
        asyncio.run(
            builder.PPTXBuilder.build(
                template_path,
                _template_data(presentation, template_path),
                _valid_content(content),
                output_path,
            )
        )

    assert output_path.read_bytes() == sentinel
    assert list(tmp_path.glob(f".{output_path.name}.*.tmp")) == []


def test_output_cannot_overwrite_template(
    tmp_path: Path,
    service_importer,
) -> None:
    builder, content, errors, presentation = _load_builder_modules(service_importer)
    template_path = tmp_path / "template.pptx"
    _write_template(template_path)
    original_template = template_path.read_bytes()

    with pytest.raises(errors.BuilderInputError, match="должен отличаться"):
        asyncio.run(
            builder.PPTXBuilder.build(
                template_path,
                _template_data(presentation, template_path),
                _valid_content(content),
                template_path,
            )
        )

    assert template_path.read_bytes() == original_template


def test_repeated_template_index_is_explicitly_rejected(service_importer) -> None:
    builder, _, errors, _ = _load_builder_modules(service_importer)

    with pytest.raises(errors.BuilderInputError, match="Повторное использование"):
        builder._validate_requested_indices([1, 1], {1: object()})


def test_unsafe_clone_is_explicitly_rejected(service_importer) -> None:
    _, _, errors, _ = _load_builder_modules(service_importer)
    clone = importlib.import_module("app.utils.clone")

    with pytest.raises(errors.SlideReuseNotSupportedError):
        clone.clone_slide(object(), object(), object())
