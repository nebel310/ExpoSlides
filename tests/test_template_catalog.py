"""Нативные layouts доступны планировщику без изменения загруженного оригинала."""

import hashlib
import json

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from exposlides.design_diversity import variant_diversity
from exposlides.design_layout import create_variants
from exposlides.design_native_text import native_reference_size
from exposlides.template_catalog import prepare_template_catalog, resolve_template_catalog
from exposlides.template_layout import composition_key, native_layout
from tests.test_template_design import _profile, _story


def test_catalog_preserves_source_and_adds_each_eligible_layout_once(tmp_path):
    source = tmp_path / "original.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(source)
    original = source.read_bytes()
    catalog = prepare_template_catalog(source, tmp_path / "job")
    assert source.read_bytes() == original
    result = Presentation(catalog)
    assert len(result.slides) > 1
    layouts = [s.slide_layout.part.partname for s in result.slides]
    assert len(layouts) == len(set(layouts))
    assert result.slides[0].slide_layout.name == presentation.slides[0].slide_layout.name
    assert not any(s.placeholder_format.type.name == "PICTURE"
                   for slide in result.slides for s in slide.placeholders)
    assert prepare_template_catalog(catalog, tmp_path / "again") == catalog
    metadata = json.loads((catalog.parent / "layout-template.json").read_text())
    assert resolve_template_catalog(source, catalog.parent, metadata["catalog_sha256"]) == catalog
    assert resolve_template_catalog(source, catalog.parent, hashlib.sha256(original).hexdigest()) == source
    catalog.write_bytes(catalog.read_bytes() + b"corruption")
    with pytest.raises(ValueError, match="Каталог макетов"):
        resolve_template_catalog(source, catalog.parent, metadata["catalog_sha256"])


def test_inherited_font_lookup_does_not_mutate_layout_or_master():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    layers = [slide, slide.slide_layout, slide.slide_layout.slide_master]
    before = [layer._element.xml for layer in layers]
    for shape in slide.shapes:
        if shape.has_text_frame:
            native_reference_size(shape)
    assert [layer._element.xml for layer in layers] == before


def test_different_source_numbers_do_not_count_as_different_compositions(tmp_path):
    profile = _profile(tmp_path)
    original = profile.patterns[0]
    duplicate = original.model_copy(deep=True)
    duplicate.source_slide_index = 2
    duplicate.slots[0].shape_id += 100
    duplicate.slots[0].text = "Другой пример"
    distinct = original.model_copy(deep=True)
    distinct.source_slide_index = 3
    distinct.slots[1].box.width //= 2
    profile.patterns.extend([duplicate, distinct])
    assert composition_key(original) == composition_key(duplicate)
    selected = native_layout(profile, _story().slides[0], "cards", 1, alternatives=(1,))
    assert selected[0].source_slide_index == 3
    profile.patterns.pop()
    story = _story()
    story.slides = [story.slides[0], story.slides[0].model_copy(update={"id": "second"})]
    report = variant_diversity(profile, create_variants(profile, story))
    assert not report["sufficient"]
    assert all(p["different"] == 0 for p in report["pairs"])


def test_empty_half_slide_panel_is_not_materialized(tmp_path):
    from copy import deepcopy

    from pptx.enum.shapes import MSO_SHAPE

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    layout = presentation.slide_layouts[1]
    for slot in layout.placeholders:
        slot.width = presentation.slide_width // 3
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, presentation.slide_width // 2,
                                   0, presentation.slide_width // 2, presentation.slide_height)
    layout.shapes._spTree.insert_element_before(deepcopy(panel._element), "p:extLst")
    source = tmp_path / "source.pptx"
    presentation.save(source)
    result = Presentation(prepare_template_catalog(source, tmp_path / "job"))
    assert layout.name not in [s.slide_layout.name for s in result.slides]
    assert result.slides[0].shapes[0]._element.find(".//" + qn("a:prstGeom")) is not None


def test_extractive_duplicate_versions_are_rejected_before_image_generation(tmp_path, monkeypatch):
    from exposlides.design_models import DesignRequest
    from exposlides.design_pipeline import DesignPipeline

    profile = _profile(tmp_path)
    story = _story(title="Grounded statement")
    story.slides.append(story.slides[0].model_copy(update={"id": "second"}))
    pipeline = DesignPipeline(tmp_path / "job")
    monkeypatch.setattr(pipeline, "generate_images", lambda *args: pytest.fail("Внешний вызов"))
    with pytest.raises(ValueError, match="трёх различимых версий"):
        pipeline.build(tmp_path / "unused.pptx", DesignRequest(
            script="A single grounded statement.", slide_count=2, mode="extractive",
        ), profile, story)
    assert not json.loads((pipeline.directory / "variant-diversity.json").read_text())["sufficient"]
    assert not (pipeline.directory / "variants").exists()


def test_title_that_fits_its_box_wins_over_diversity(tmp_path):
    from exposlides.design_audit import capacity_risk

    profile = _profile(tmp_path)
    tight = profile.patterns[0]
    roomy = tight.model_copy(deep=True)
    roomy.source_slide_index = 2
    tight.slots[0].box.width //= 3
    profile.patterns.append(roomy)
    source = _story(title="An important grounded statement").slides[0]
    chosen, blocks = native_layout(profile, source, "cards", 1, alternatives=(2,))
    assert chosen.source_slide_index == 2
    assert not capacity_risk(blocks[0])


def test_large_real_layout_photo_is_kept(tmp_path):
    from PIL import Image

    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    layout = presentation.slide_layouts[1]
    for slot in layout.placeholders:
        slot.width = presentation.slide_width // 3
    photo = tmp_path / "photo.png"
    Image.new("RGB", (400, 300), "blue").save(photo)
    from copy import deepcopy

    slide = presentation.slides[0]
    image = slide.shapes.add_picture(str(photo), presentation.slide_width // 2, 0,
                                    presentation.slide_width // 2, presentation.slide_height)
    _, rid = layout.part.get_or_add_image_part(str(photo))
    element = deepcopy(image._element)
    element.find(".//" + qn("a:blip")).set(qn("r:embed"), rid)
    layout.shapes._spTree.insert_element_before(element, "p:extLst")
    source = tmp_path / "source.pptx"
    presentation.save(source)
    result = Presentation(prepare_template_catalog(source, tmp_path / "job"))
    assert layout.name in [s.slide_layout.name for s in result.slides]


def test_model_receives_title_budget_for_native_compositions(tmp_path, service_importer):
    import asyncio
    from pathlib import Path

    from exposlides.design_models import DesignRequest

    module = service_importer(Path(__file__).resolve().parents[1] / 'services/content-service',
                              'app.design_main')
    profile = _profile(tmp_path)
    source = "A single grounded statement."

    class Client:
        calls = 0

        async def generate_json(self, prompt, model):
            self.calls += 1
            assert 'content_title_max_characters' in prompt
            assert 'узкие фото-макеты' in prompt
            return model.model_validate(_story(title="Grounded statement").model_dump())

    client = Client()
    result = asyncio.run(module.generate(DesignRequest(script=source, slide_count=1),
                                         client, profile=profile))
    assert len(result.slides) == 1 and client.calls == 1


@pytest.mark.parametrize("slide_count", [1, 10])
def test_identical_decks_fail_diversity_even_with_a_single_slide(tmp_path, slide_count):
    profile = _profile(tmp_path)
    story = _story()
    story.slides = [
        story.slides[0].model_copy(deep=True, update={"id": f"slide-{index}"})
        for index in range(slide_count)
    ]
    base = create_variants(profile, story)[0]
    variants = [
        base.model_copy(deep=True, update={"variant_id": variant})
        for variant in ("story", "evidence", "cards")
    ]

    report = variant_diversity(profile, variants)

    assert not report["sufficient"]
    assert all(pair["compared"] == slide_count for pair in report["pairs"])
    assert all(pair["different"] == 0 and not pair["sufficient"] for pair in report["pairs"])


@pytest.mark.parametrize("slide_count,different_slide", [(1, 0), (10, 0), (10, 9)])
def test_one_visual_difference_per_deck_pair_is_sufficient(
    tmp_path, slide_count, different_slide,
):
    profile = _profile(tmp_path)
    story = _story()
    story.slides = [
        story.slides[0].model_copy(deep=True, update={"id": f"slide-{index}"})
        for index in range(slide_count)
    ]
    base = create_variants(profile, story)[0]
    variants = [
        base.model_copy(deep=True, update={"variant_id": variant})
        for variant in ("story", "evidence", "cards")
    ]
    for index, variant in enumerate(variants):
        variant.slides[different_slide].blocks[0].box.left += index * 127000

    report = variant_diversity(profile, variants)

    assert report["sufficient"]
    assert all(pair["compared"] == slide_count for pair in report["pairs"])
    assert all(pair["different"] == 1 and pair["sufficient"] for pair in report["pairs"])
