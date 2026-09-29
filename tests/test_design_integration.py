"""Регрессии границ: настоящий PPTX и persistence, внешние модели/рендер — offline fake."""

from __future__ import annotations

import hashlib
import json
import threading
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from exposlides.design_audit import apply_fixes, audit_deck, audit_saved_pptx, capacity_risk
from exposlides.design_builder import build_deck
from exposlides.design_models import (
    Box,
    ContentPlan,
    Dataset,
    DeckPlan,
    DesignRequest,
    PlacedBlock,
    SlideInstance,
    SlidePattern,
    StorySlide,
    TemplateProfile,
    TextStyle,
)
from exposlides.design_pipeline import CancelledError, DesignPipeline, save_model
from exposlides.studio import FixInput, Studio


@pytest.fixture
def native_case(tmp_path):
    template = tmp_path / "template.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(12), Inches(8)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(0.3), Inches(10), Inches(1))
    title.text = "Заголовок шаблона"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(10), Inches(4))
    body.text = "Содержание шаблона"
    body.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    logo = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(11), 0, Inches(1), Inches(0.2))
    logo.fill.solid()
    logo.fill.fore_color.rgb = RGBColor.from_string("336699")
    group = slide.shapes.add_group_shape()
    footer = group.shapes.add_textbox(Inches(1), Inches(7.5), Inches(3), Inches(0.2))
    footer.name, footer.text = "Corporate footer", "Авторство компании"
    prs.save(template)
    profile = TemplateProfile(
        template_sha256="0"*64, width=prs.slide_width, height=prs.slide_height,
        patterns=[SlidePattern(
            source_slide_index=1, name="Native", slots=[], font="Arial",
            palette=["202124", "FFFFFF", "336699"],
            title_style=TextStyle(size=32), body_style=TextStyle(size=20),
            mutable_shape_ids=[title.shape_id, body.shape_id],
            protected_shape_ids=[logo.shape_id, footer.shape_id],
            content_box=Box(left=Inches(1), top=Inches(2), width=Inches(10), height=Inches(4)),
        )],
    )
    block = PlacedBlock(
        id="body", kind="text", box=profile.patterns[0].content_box,
        style=TextStyle(size=20), text="Команда развивает платформу.", source_ids=["source-1"],
    )
    plan = DeckPlan(
        variant_id="story", name="План", description="Проверка", width=prs.slide_width,
        height=prs.slide_height, slides=[SlideInstance(
            id="s1", story_slide_id="s1", source_slide_index=1, blocks=[block],
            remove_shape_ids=profile.patterns[0].mutable_shape_ids,
        )],
    )
    return template, profile, plan


def _generated(prs):
    return next(shape for shape in prs.slides[0].shapes if shape.name == "exposlides:body")


@pytest.mark.parametrize("corruption,rule", [
    ("text", "saved_text"), ("font", "saved_text_style"), ("size", "saved_text_style"),
    ("bold", "saved_text_style"), ("color", "saved_text_style"),
    ("position", "saved_geometry"), ("protected_fill", "protected_changed"),
    ("protected_parent", "protected_changed"), ("background", "template_background"),
])
def test_saved_audit_detects_actual_pptx_corruption(tmp_path, native_case, corruption, rule):
    template, profile, plan = native_case
    output = build_deck(template, plan, tmp_path / "result.pptx")
    assert audit_saved_pptx(output, template, plan, profile).ok
    prs = Presentation(output)
    shape = _generated(prs)
    run = shape.text_frame.paragraphs[0].runs[0]
    if corruption == "text":
        run.text = "Ошибочно подменённое содержание"
    elif corruption == "font":
        run.font.name = "Comic Sans MS"
    elif corruption == "size":
        run.font.size = Pt(9)
    elif corruption == "bold":
        run.font.bold = True
    elif corruption == "color":
        run.font.color.rgb = RGBColor.from_string("FF0000")
    elif corruption == "position":
        shape.left += Inches(1)
    elif corruption == "protected_fill":
        protected = next(s for s in prs.slides[0].shapes
                         if s.shape_id == profile.patterns[0].protected_shape_ids[0])
        protected.fill.fore_color.rgb = RGBColor.from_string("FF0000")
    elif corruption == "protected_parent":
        group = next(s for s in prs.slides[0].shapes if hasattr(s, "shapes"))
        group.left += Inches(1)
    else:
        prs.slides[0].background.fill.solid()
        prs.slides[0].background.fill.fore_color.rgb = RGBColor.from_string("FF0000")
    prs.save(output)
    report = audit_saved_pptx(output, template, plan, profile)
    assert any(issue.rule == rule and issue.severity == "error" for issue in report.issues)


def test_saved_audit_has_unique_ids_for_multiple_damaged_protected_objects(tmp_path, native_case):
    template, profile, plan = native_case
    output = build_deck(template, plan, tmp_path / "result.pptx")
    prs = Presentation(output)
    for shape in list(prs.slides[0].shapes):
        if shape.shape_id in profile.patterns[0].protected_shape_ids or hasattr(shape, "shapes"):
            shape._element.getparent().remove(shape._element)
    prs.save(output)
    issues = [issue for issue in audit_saved_pptx(output, template, plan, profile).issues
              if issue.rule == "protected_changed"]
    assert len(issues) == len({issue.id for issue in issues}) == 2


@pytest.mark.parametrize("kind", ["table", "chart"])
def test_saved_audit_verifies_visual_data_against_source(tmp_path, native_case, kind):
    template, profile, plan = native_case
    dataset = Dataset(id="data", name="Выручка", source="Данные пользователя", unit="млн руб.",
                      columns=["Регион", "Выручка"], rows=[["Запад", 12], ["Восток", 18]])
    plan.datasets = [dataset]
    block = plan.slides[0].blocks[0]
    block.kind, block.dataset_id = kind, "data"
    block.chart_type = "bar" if kind == "chart" else None
    output = build_deck(template, plan, tmp_path / "result.pptx")
    assert audit_saved_pptx(output, template, plan, profile).ok
    prs = Presentation(output)
    shape = _generated(prs)
    if kind == "table":
        shape.table.cell(1, 1).text = "99"
    else:
        data = CategoryChartData()
        data.categories = ["Запад", "Восток"]
        data.add_series("Выручка", [99, 18])
        shape.chart.replace_data(data)
    prs.save(output)
    assert "saved_data" in {i.rule for i in audit_saved_pptx(output, template, plan, profile).issues}


def test_saved_audit_wrong_chart_type_reports_instead_of_crashing(tmp_path, native_case):
    template, profile, plan = native_case
    plan.datasets = [Dataset(id="data", name="Value", source="CSV", columns=["Name", "Value"],
                             rows=[["A", 12], ["B", 18]])]
    block = plan.slides[0].blocks[0]
    block.kind, block.dataset_id, block.chart_type = "chart", "data", "bar"
    output = build_deck(template, plan, tmp_path / "result.pptx")
    prs = Presentation(output)
    original = _generated(prs)
    original._element.getparent().remove(original._element)
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Value", [12, 18])
    shape = prs.slides[0].shapes.add_chart(XL_CHART_TYPE.PIE, *block.box.model_dump().values(), data)
    shape.name = "exposlides:body"
    prs.save(output)
    assert {"native_chart", "chart_labels"}.issubset(
        {i.rule for i in audit_saved_pptx(output, template, plan, profile).issues}
    )


def test_saved_audit_rejects_textbox_masquerading_as_image(tmp_path, native_case):
    template, profile, plan = native_case
    image = tmp_path / "image.png"
    Image.new("RGB", (20, 10), "blue").save(image)
    block = plan.slides[0].blocks[0]
    block.kind, block.image_path = "image", str(image)
    output = build_deck(template, plan, tmp_path / "result.pptx")
    assert audit_saved_pptx(output, template, plan, profile).ok
    prs = Presentation(output)
    original = _generated(prs)
    original._element.getparent().remove(original._element)
    shape = prs.slides[0].shapes.add_textbox(*block.box.model_dump().values())
    shape.name, shape.text = "exposlides:body", "Fake picture"
    prs.save(output)
    assert "native_image" in {i.rule for i in audit_saved_pptx(output, template, plan, profile).issues}


class OfflineRenderer:
    available = True

    def __init__(self, **kwargs):
        self.timeout = kwargs.get("timeout", 90)

    def render(self, source, directory, count, *, pdf_output):
        assert len(Presentation(source).slides) == count
        directory.mkdir(parents=True, exist_ok=True)
        pdf_output.write_bytes(b"%PDF-1.4\n%Offline render boundary fixture")
        images = []
        for index in range(1, count+1):
            path = directory / f"slide-{index}.png"
            Image.new("RGB", (120, 80), "white").save(path)
            images.append(path)
        return images

    def close(self):
        pass


@pytest.fixture
def offline_pipeline(tmp_path, monkeypatch):
    monkeypatch.setattr("exposlides.design_pipeline.PreviewRenderer", OfflineRenderer)
    monkeypatch.setattr("exposlides.design_export.shutil.which", lambda name: None)
    return DesignPipeline(tmp_path / "job")


def test_offline_pipeline_actual_parse_llm_boundary_native_build_and_exports(native_case, offline_pipeline):
    from copy import deepcopy

    template, _, _ = native_case
    # Три реальные композиции вместо трёх одинаковых копий единственного образца.
    source = Presentation(template)
    original_shapes = list(source.slides[0].shapes)
    for offset in (1, 2):
        slide = source.slides.add_slide(source.slide_layouts[6])
        for shape in original_shapes:
            slide.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
        body = next(shape for shape in slide.shapes
                    if shape.has_text_frame and shape.text == "Содержание шаблона")
        body.left += Inches(offset)
        body.width -= Inches(offset)
    source.save(template)
    assert len(Presentation(template).slides) == 3
    request = DesignRequest(script="Команда развивает платформу. Поддержка помогает клиентам.",
                            slide_count=1, mode="llm")
    story = ContentPlan(title="Платформа", slides=[StorySlide(
        id="s1", title="Платформа", paragraphs=[request.script], source_ids=["source-1"],
    )])
    command = offline_pipeline.command
    calls = []

    def fake_llm(args, cwd, **kwargs):
        calls.append(args[2])
        if args[2] == "app.design_main":
            saved = json.loads((offline_pipeline.directory / "request.json").read_text())
            assert saved["script"] == request.script
            save_model(offline_pipeline.directory / "story.json", story)
            save_model(offline_pipeline.directory / "story-provenance.json", {
                "workflow_id": "designer", "workflow_version": "1.0.0",
                "role": "story", "model": "Qwen/Qwen3.8-27B",
                "endpoint_model": "Qwen/Qwen3.8-27B:deepinfra", "license": "Apache-2.0",
                "parameters_billions": 27, "registry_version": "1.0.0",
                "prompt_sha256": "a"*64, "api_key": "must-not-enter-manifest",
            })
            return 0
        return command(args, cwd, **kwargs)

    offline_pipeline.command = fake_llm
    profile, generated = offline_pipeline.plan(template, request)
    manifest = json.loads((offline_pipeline.directory / "manifest.json").read_text())
    assert manifest["request_sha256"] == hashlib.sha256(
        (offline_pipeline.directory / "request.json").read_bytes()
    ).hexdigest()
    assert manifest["story_model"]["endpoint_model"] == "Qwen/Qwen3.8-27B:deepinfra"
    assert "api_key" not in manifest["story_model"]
    assert manifest["limits"]["seconds"] == offline_pipeline.timeout
    for name in ("story_correction.md", "story_correction_patch.md", "story_correction_full.md",
                 "story_repair_requirements.md", "story_editorial_repair.md"):
        relative = f"prompts/designer/{name}"
        resource = Path(__file__).resolve().parents[1] / relative
        assert manifest["versioned_resources"][relative] == hashlib.sha256(resource.read_bytes()).hexdigest()
    variants = offline_pipeline.build(template, request, profile, generated)
    assert calls == ["app.main", "app.design_main", "app.main"]
    assert len(variants) == 3
    saved_compositions = set()
    for variant in variants:
        revision = offline_pipeline.directory / "variants" / variant["id"] / "1"
        prs = Presentation(revision / "presentation.pptx")
        assert len(prs.slides) == 1
        saved_compositions.add(tuple(
            (shape.left, shape.top, shape.width, shape.height, shape.text)
            for shape in prs.slides[0].shapes if shape.has_text_frame
        ))
        assert any(request.script in shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
        assert not [issue for issue in variant["issues"] if issue["severity"] == "error"]
        assert (revision / "presentation.pdf").is_file()
        assert "data:image/png;base64," in (revision / "presentation.html").read_text()
    assert len(saved_compositions) == 3
    assert not list(offline_pipeline.directory.rglob(".pending-*"))


def test_pipeline_fix_publishes_new_native_revision_without_mutating_previous(native_case, offline_pipeline):
    template, profile, plan = native_case
    plan.slides[0].blocks[0].box.left = -100
    first = offline_pipeline.publish_variant(template, profile, plan, revision=1)
    old_path = offline_pipeline.directory / "variants/story/1/presentation.pptx"
    old_bytes = old_path.read_bytes()
    selected = next(issue["id"] for issue in first["issues"] if issue["rule"] == "outside_slide")
    revised = offline_pipeline.fix(template, profile, "story", 1, [selected])
    assert revised["revision"] == 2
    assert not [issue for issue in revised["issues"] if issue["rule"] == "outside_slide"]
    assert old_path.read_bytes() == old_bytes
    assert _generated(Presentation(offline_pipeline.directory / "variants/story/2/presentation.pptx")).left == 0


def test_renderer_failure_does_not_publish_and_same_revision_can_retry(native_case, offline_pipeline):
    template, profile, plan = native_case
    original = offline_pipeline.renderer.render

    def fail(*args, **kwargs):
        raise RuntimeError("offline renderer failed")

    offline_pipeline.renderer.render = fail
    with pytest.raises(RuntimeError, match="renderer failed"):
        offline_pipeline.publish_variant(template, profile, plan, revision=1)
    parent = offline_pipeline.directory / "variants/story"
    assert not list(parent.iterdir())
    offline_pipeline.renderer.render = original
    assert offline_pipeline.publish_variant(template, profile, plan, revision=1)["revision"] == 1


def test_fit_text_expands_only_selected_block_within_free_space(native_case):
    _, profile, plan = native_case
    selected = plan.slides[0].blocks[0]
    selected.box = Box(left=Inches(1), top=Inches(2), width=Inches(4), height=Inches(0.2))
    selected.text = "Достаточно длинный текст, которому требуется свободное место под блоком."
    untouched = selected.model_copy(deep=True)
    untouched.id, untouched.text = "other", "Другой факт"
    untouched.box = Box(left=Inches(7), top=Inches(2), width=Inches(3), height=Inches(1))
    plan.slides[0].blocks.append(untouched)
    report = audit_deck(plan, profile)
    issue = next(i for i in report.issues if i.rule == "text_capacity" and i.block_id == "body")
    revised = apply_fixes(plan, profile, report, [issue.id])
    assert revised.slides[0].blocks[1] == untouched
    assert revised.slides[0].blocks[0].box.height > selected.box.height
    assert not capacity_risk(revised.slides[0].blocks[0])
    assert selected.box.height == Inches(0.2)


def test_cancelled_fix_preserves_available_versions_after_restart(tmp_path):
    entered, release = threading.Event(), threading.Event()

    class WaitingPipeline:
        def __init__(self, directory, *, cancel, **kwargs):
            self.cancel = cancel

        def check(self):
            if self.cancel.is_set():
                raise CancelledError()

        def fix(self, *args):
            entered.set()
            assert release.wait(5)
            self.check()

        def close(self):
            pass

    identifier = "a"*32
    directory = tmp_path / "jobs" / identifier
    directory.mkdir(parents=True)
    profile = TemplateProfile(template_sha256="a"*64, width=100, height=100, patterns=[SlidePattern(
        source_slide_index=1, name="Basic", slots=[], content_box=Box(left=1, top=1, width=98, height=98),
        palette=["000000"], font="Arial", title_style=TextStyle(), body_style=TextStyle(),
    )])
    save_model(directory / "profile.json", profile)
    variant = dict(id="story", revision=1, preview_urls=[], exports={}, issues=[])
    save_model(directory / "job.json", dict(id=identifier, status="completed", active_seconds=1,
                                           variants=[variant], stage="review"))
    studio = Studio(tmp_path, workers=1, pipeline_factory=WaitingPipeline)
    try:
        studio.fix(identifier, "story", FixInput(revision=1, issue_ids=["issue"]))
        assert entered.wait(5)
        studio.cancel(identifier)
        release.set()
        studio.pool.submit(lambda: None).result(timeout=5)
        job = studio.get(identifier)
        assert job["status"] == "completed"
        assert job["variants"][0]["revision"] == 1
    finally:
        release.set()
        studio.close()
    restored = Studio(tmp_path, pipeline_factory=WaitingPipeline)
    try:
        assert restored.get(identifier)["status"] == "completed"
    finally:
        restored.close()


def test_unverified_story_builds_reopens_and_keeps_audit(native_case, offline_pipeline):
    template, profile, _ = native_case
    request = DesignRequest(script="Команда развивает платформу.", slide_count=1, mode="llm")
    profile, _ = offline_pipeline.plan(
        template, request.model_copy(update={"mode": "extractive"}),
    )
    story = ContentPlan(title="Обзор", slides=[StorySlide(
        id="s1", title="Обзор", paragraphs=["Команда развивает платформу. Рост 99%."],
        source_ids=["source-1"],
    )])
    variants = offline_pipeline.build(template, request, profile, story)
    assert variants
    for variant in variants:
        directory = offline_pipeline.directory / "variants" / variant["id"] / "1"
        saved = Presentation(directory / "presentation.pptx")
        assert len(saved.slides) == 1
        assert any("99%" in shape.text for shape in saved.slides[0].shapes if shape.has_text_frame)
        assert any(issue["rule"] == "story_validation" and "99%" in issue["message"]
                   for issue in variant["issues"])
        assert not json.loads((directory / "audit.json").read_text())["issues"] == []
    story.slides[0].paragraphs = [request.script]
    offline_pipeline.validate_plan(request, profile, story)
    findings = json.loads((offline_pipeline.directory / "story-validation.json").read_text())
    assert not any("99%" in issue for issue in findings["issues"])


def test_unverified_story_cannot_bypass_broken_references(native_case, offline_pipeline):
    _, profile, _ = native_case
    request = DesignRequest(script="Команда развивает платформу.", slide_count=1, mode="llm")
    story = ContentPlan(title="Обзор", slides=[StorySlide(
        id="s1", title="Обзор", paragraphs=[request.script], source_ids=["unknown"],
    )])
    with pytest.raises(ValueError, match="неизвестные ссылки"):
        offline_pipeline.validate_plan(request, profile, story)


def test_saved_audit_accepts_empty_relationship_in_powerpoint_action(tmp_path, native_case):
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    template, profile, plan = native_case
    source = Presentation(template)
    shape = source.slides[0].shapes[2]
    action = OxmlElement("a:hlinkClick")
    action.set(qn("r:id"), "")
    action.set("action", "ppaction://hlinkshowjump?jump=nextslide")
    shape._element.find(".//" + qn("p:cNvPr")).append(action)
    source.save(template)
    output = build_deck(template, plan, tmp_path / "actions.pptx")
    assert len(Presentation(output).slides) == 1
    assert audit_saved_pptx(output, template, plan, profile).ok
