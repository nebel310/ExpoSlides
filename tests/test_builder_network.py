from __future__ import annotations

import asyncio
import importlib
import io
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from pptx import Presentation
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
BUILDER = ROOT / "services/builder-service"


def _inputs(tmp_path, service_importer):
    source = Presentation()
    for number in (1, 2, 3):
        slide = source.slides.add_slide(source.slide_layouts[5])
        slide.shapes.title.text = f"Образец {number}"
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(31)
    path = tmp_path / "template.pptx"
    source.save(path)
    parser = service_importer(ROOT / "services/parsing-service", "app.parsers.pptx")
    result = asyncio.run(parser.PPTXParser.parse(path))
    ids = {name: str(uuid4()) for name in (
        "template_file_id", "structure_file_id", "content_file_id", "script_file_id"
    )}
    files = {
        ids["template_file_id"]: path.read_bytes(),
        ids["structure_file_id"]: result.presentation.model_dump_json().encode(),
        ids["content_file_id"]: json.dumps({
            "content": {
                "3": {"placeholders": {"0": "Итог"}},
                "1": {"placeholders": {"0": "Введение"}},
            }, "validation_report": {"ok": True}, "error": None,
        }).encode(),
    }
    return ids, files


def _envelope(ids):
    return {"task_id": str(uuid4()), "attempt": 2, "payload": ids, "error": None}


def test_network_builder_creates_pptx_and_gateway_completes(tmp_path, service_importer, monkeypatch):
    ids, files = _inputs(tmp_path, service_importer)
    network = service_importer(BUILDER, "app.network")
    client = AsyncMock()
    client.download_file.side_effect = lambda key: files[key]
    result_id = str(uuid4())
    client.upload_file.return_value = result_id
    producer = AsyncMock()
    envelope = _envelope(ids)
    asyncio.run(network.process_message(
        json.dumps(envelope).encode(), network.BuilderPipeline(client), producer, network.Settings()
    ))
    topic, raw = producer.send_and_wait.await_args.args
    assert topic == "task.built"
    event = json.loads(raw)
    assert event == {**envelope, "payload": {**ids, "result_file_id": result_id}}
    upload = client.upload_file.await_args.kwargs
    assert upload["task_id"] == envelope["task_id"]
    assert upload["filename"] == "result.pptx"
    assert upload["content_type"] == network.PPTX_CONTENT_TYPE
    reopened = Presentation(io.BytesIO(upload["content"]))
    assert [s.shapes.title.text for s in reopened.slides] == ["Итог", "Введение"]
    assert all(s.shapes.title.text_frame.paragraphs[0].runs[0].font.size.pt == 31
               for s in reopened.slides)

    gateway = service_importer(ROOT / "services/gateway-service", "app.services.kafka_consumer")
    repository = importlib.import_module("app.repositories.tasks").TaskRepository
    from fakeredis.aioredis import FakeRedis

    emit = AsyncMock()
    monkeypatch.setattr(gateway, "emit_to_session", emit)

    async def consume_result():
        redis = FakeRedis(decode_responses=True)
        try:
            await repository.save(redis, {
                "task_id": envelope["task_id"], "sid": "session", "status": "building"
            })
            await gateway.GatewayKafkaConsumer()._handle(redis, topic, raw)
            task = await repository.get(redis, envelope["task_id"])
            assert task["status"] == "done"
            assert task["result_file_id"] == result_id
            emit.assert_awaited_once()
        finally:
            await redis.aclose()
    asyncio.run(consume_result())


@pytest.mark.parametrize("failure", ["download", "content", "template", "upload", "payload"])
def test_network_builder_failure_publishes_failed(
    tmp_path, service_importer, failure
):
    ids, files = _inputs(tmp_path, service_importer)
    network = service_importer(BUILDER, "app.network")
    client = AsyncMock()
    client.download_file.side_effect = lambda key: files[key]
    client.upload_file.return_value = str(uuid4())
    if failure == "download":
        client.download_file.side_effect = RuntimeError("file service unavailable")
    elif failure == "content":
        files[ids["content_file_id"]] = b'{"content": {}, "error": "invalid"}'
    elif failure == "template":
        files[ids["template_file_id"]] = b"not a pptx"
    elif failure == "upload":
        client.upload_file.side_effect = RuntimeError("storage unavailable")
    else:
        ids.pop("content_file_id")
    envelope = _envelope(ids)
    producer = AsyncMock()
    asyncio.run(network.process_message(
        json.dumps(envelope).encode(), network.BuilderPipeline(client), producer, network.Settings()
    ))
    producer.send_and_wait.assert_awaited_once()
    topic, raw = producer.send_and_wait.await_args.args
    result = json.loads(raw)
    assert topic == "task.failed"
    assert result["payload"]["stage"] == "builder"
    assert result["task_id"] == envelope["task_id"]
    assert result["attempt"] == 2
    if failure != "upload":
        client.upload_file.assert_not_awaited()


@pytest.mark.parametrize("publish_fails", [False, True])
def test_worker_commits_only_after_acknowledged_output(service_importer, publish_fails):
    network = service_importer(BUILDER, "app.network")
    envelope = _envelope({})
    order = []

    class Consumer:
        def __aiter__(self):
            async def messages():
                yield SimpleNamespace(
                    value=json.dumps(envelope).encode(), topic="task.content_ready",
                    partition=2, offset=17,
                )
            return messages()

        async def commit(self, offsets):
            from aiokafka.structs import TopicPartition

            assert offsets == {TopicPartition("task.content_ready", 2): 18}
            order.append("commit")

    async def publish(*args, **kwargs):
        order.append("publish")
        if publish_fails:
            raise RuntimeError("broker unavailable")

    producer = SimpleNamespace(send_and_wait=publish)
    pipeline = SimpleNamespace(process=AsyncMock(return_value={"result_file_id": str(uuid4())}))
    if publish_fails:
        with pytest.raises(RuntimeError, match="broker unavailable"):
            asyncio.run(network.consume(Consumer(), pipeline, producer, network.Settings()))
        assert order == ["publish"]
    else:
        asyncio.run(network.consume(Consumer(), pipeline, producer, network.Settings()))
        assert order == ["publish", "commit"]


def test_invalid_envelope_is_not_sent_to_builder(service_importer):
    network = service_importer(BUILDER, "app.network")
    pipeline = SimpleNamespace(process=AsyncMock())
    producer = AsyncMock()
    asyncio.run(network.process_message(b"{}", pipeline, producer, network.Settings()))
    pipeline.process.assert_not_awaited()
    producer.send_and_wait.assert_not_awaited()


@pytest.mark.parametrize("failure", [None, "producer", "consumer"])
def test_worker_closes_resources_on_stop_and_startup_failure(
    service_importer, monkeypatch, failure
):
    import sys
    from types import ModuleType

    network = service_importer(BUILDER, "app.network")
    started, stopped = [], []

    class Client:
        def __init__(self, name):
            self.name = name

        async def start(self):
            started.append(self.name)
            if failure == self.name:
                raise RuntimeError("startup failed")

        async def stop(self):
            stopped.append(self.name)

        def __aiter__(self):
            async def messages():
                await asyncio.Event().wait()
                yield None
            return messages()

    module = ModuleType("app.file_client")
    module.FileServiceClient = lambda settings: Client("file")
    monkeypatch.setitem(sys.modules, "app.file_client", module)
    monkeypatch.setattr(network, "AIOKafkaProducer", lambda **kwargs: Client("producer"))
    monkeypatch.setattr(network, "AIOKafkaConsumer", lambda *args, **kwargs: Client("consumer"))

    async def run():
        stop = asyncio.Event()
        stop.set()
        await network.serve(network.Settings(), stop)

    if failure:
        with pytest.raises(RuntimeError, match="startup failed"):
            asyncio.run(run())
    else:
        asyncio.run(run())
    assert stopped == list(reversed(started))


def test_file_client_roundtrip_and_cleanup(service_importer, monkeypatch):
    import sys
    from types import ModuleType

    monkeypatch.syspath_prepend(str(ROOT))
    from exposlides._proto import file_service_pb2

    network = service_importer(BUILDER, "app.network")
    stub = AsyncMock()
    stub.DownloadFile.return_value = SimpleNamespace(content=b"pptx")
    stub.UploadFile.return_value = SimpleNamespace(file_id=str(uuid4()))
    generated = ModuleType("file_service_pb2_grpc")
    generated.FileServiceStub = lambda channel: stub
    monkeypatch.setitem(sys.modules, "file_service_pb2", file_service_pb2)
    monkeypatch.setitem(sys.modules, "file_service_pb2_grpc", generated)
    client_module = importlib.import_module("app.file_client")
    channel = AsyncMock()
    monkeypatch.setattr(
        client_module.grpc.aio, "insecure_channel", lambda target, **kwargs: channel
    )
    settings = network.Settings(file_service_timeout=13)
    client = client_module.FileServiceClient(settings)

    async def run():
        with pytest.raises(RuntimeError):
            await client.download_file("file-id")
        await client.start()
        assert await client.download_file("file-id") == b"pptx"
        result = await client.upload_file(
            filename="result.pptx", content=b"pptx", content_type="type", task_id="task"
        )
        assert result == stub.UploadFile.return_value.file_id
        await client.stop()
        assert client.stub is None
        channel.close.assert_awaited_once()
    asyncio.run(run())
    assert stub.DownloadFile.await_args.args[0].file_id == "file-id"
    assert stub.DownloadFile.await_args.kwargs["timeout"] == 13
    upload = stub.UploadFile.await_args.args[0]
    assert (upload.filename, upload.content, upload.task_id) == ("result.pptx", b"pptx", "task")
