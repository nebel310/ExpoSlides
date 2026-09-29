"""Нативная сборка вариантов: повторное применение образцов, текст и визуализации."""

from __future__ import annotations

import math
import os
import posixpath
import tempfile
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlsplit
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.parts.slide import SlidePart
from pptx.util import Pt

from exposlides.design_contrast import repair_text_contrast
from exposlides.design_data import cell_text
from exposlides.design_models import Dataset, DeckPlan, PlacedBlock, TextStyle
from exposlides.design_native_image import fill_native_images
from exposlides.design_native_text import (
    TEXT_KINDS,
    bound_text_shapes,
    fill_native_text,
    native_shapes,
    retained_shape_ids,
)
from exposlides.design_pptx_parts import NativeBuildError, clone_selected_slides, remove_shapes
from exposlides.design_smartart import add_smartart, capture_smartart


def build_deck(template: Path, plan: DeckPlan, output: Path) -> Path:
    """Создаёт и атомарно публикует проверенный PPTX из типизированного плана."""
    template, output = Path(template), Path(output)
    if template.resolve() == output.resolve():
        raise NativeBuildError("Результат не должен перезаписывать исходный шаблон")
    presentation = Presentation(template)
    if (presentation.slide_width, presentation.slide_height) != (plan.width, plan.height):
        raise NativeBuildError("Размер слайда в плане не соответствует шаблону")
    slides = clone_selected_slides(
        presentation, [slide.source_slide_index for slide in plan.slides]
    )
    datasets = {dataset.id: dataset for dataset in plan.datasets}
    for slide, instance in zip(slides, plan.slides, strict=True):
        fill_native_images(presentation, slide, instance.blocks)
        bound_text = bound_text_shapes(slide, instance.blocks)
        retained = retained_shape_ids(bound_text.values())
        smartart = capture_smartart(slide, instance.blocks)
        removals = [identifier for identifier in instance.remove_shape_ids if identifier not in retained]
        remove_shapes(slide, [*removals, *[item[1] for item in smartart.values()]])
        for block in instance.blocks:
            if block.kind == "image" and block.image_fit == "template":
                continue
            if block.id in bound_text:
                fill_native_text(bound_text[block.id], block)
            elif block.kind == "smartart":
                add_smartart(slide, block, smartart[block.id][0])
            else:
                _add_block(slide, block, datasets)
        text_shapes = {**bound_text}
        for shape in slide.shapes:
            for block in instance.blocks:
                if block.kind in TEXT_KINDS and shape.name == "exposlides:" + block.id:
                    text_shapes[block.id] = shape
        palette = _template_palette(slide, "FFFFFF")
        for block in instance.blocks:
            if block.id in text_shapes:
                repair_text_contrast(slide, text_shapes[block.id], block.style.color,
                                     lambda fill, color: _color_on_fill(palette, fill, color))
        if instance.notes or slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None:
                notes.text = instance.notes
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, name = tempfile.mkstemp(
        prefix=f".{output.name}.", suffix=".pptx", dir=output.parent
    )
    os.close(descriptor)
    temporary = Path(name)
    try:
        presentation.save(temporary)
        _validate_saved(temporary, plan)
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    return output


def _add_block(slide: Any, block: PlacedBlock, datasets: dict[str, Dataset]) -> None:
    if block.kind == "image":
        _add_image(slide, block)
    elif block.kind == "icon":
        _add_icon(slide, block)
    elif block.kind == "table":
        _add_table(slide, block, datasets[block.dataset_id])
    elif block.kind == "chart":
        _add_chart(slide, block, datasets[block.dataset_id])
    elif block.kind in {"process", "comparison"}:
        _add_diagram(slide, block)
    else:
        box = block.box
        shape = slide.shapes.add_textbox(box.left, box.top, box.width, box.height)
        shape.name = "exposlides:" + block.id
        if block.fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(block.fill)
        _write_text(shape.text_frame, block.items or block.text.splitlines(), block.style)


def _write_text(frame: Any, lines: list[str], style: TextStyle, *, centered: bool = False) -> None:
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE if centered else MSO_ANCHOR.TOP
    for index, line in enumerate(lines or [""]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
        paragraph.space_after = Pt(style.size * 0.25) if index < len(lines) - 1 else Pt(0)
        paragraph.line_spacing = 1.1
        properties = paragraph._p.get_or_add_pPr()
        for child in list(properties):
            if child.tag in {qn("a:buChar"), qn("a:buAutoNum"), qn("a:buBlip")}:
                properties.remove(child)
        properties.append(OxmlElement("a:buNone"))
        run = paragraph.add_run()
        run.text = line
        _font(run.font, style)


def _font(font: Any, style: TextStyle) -> None:
    font.name = style.font
    font.size = Pt(style.size)
    font.bold = style.bold
    font.color.rgb = RGBColor.from_string(style.color)


def _add_table(slide: Any, block: PlacedBlock, dataset: Dataset) -> None:
    box = block.box
    shape = slide.shapes.add_table(
        len(dataset.rows) + 1, len(dataset.columns), box.left, box.top, box.width, box.height
    )
    shape.name = "exposlides:" + block.id
    table = shape.table
    table.first_row = True
    table.horz_banding = False
    palette = _template_palette(slide, block.style.color)
    body_fill = max(palette, key=lambda color: _contrast(color, block.style.color))
    for row_index, values in enumerate([dataset.columns, *dataset.rows]):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.fill.solid()
            fill = (block.fill or block.style.color) if row_index == 0 else body_fill
            cell.fill.fore_color.rgb = RGBColor.from_string(fill)
            style = block.style.model_copy(
                update={
                    "bold": row_index == 0 or block.style.bold,
                    "color": _color_on_fill(palette, fill, block.style.color)
                    if row_index == 0
                    else block.style.color,
                }
            )
            _write_text(cell.text_frame, [cell_text(value)], style)
            cell.margin_left = cell.margin_right = Pt(6)
            cell.margin_top = cell.margin_bottom = Pt(4)


def _contrasting_text(color: str) -> str:
    values = [int(color[offset : offset + 2], 16) / 255 for offset in (0, 2, 4)]
    linear = [
        value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4 for value in values
    ]
    luminance = sum(value * weight for value, weight in zip(linear, (0.2126, 0.7152, 0.0722)))
    return "000000" if luminance > 0.179 else "FFFFFF"


def _contrast(first: str, second: str) -> float:
    def luminance(color):
        values = [int(color[offset : offset + 2], 16) / 255 for offset in (0, 2, 4)]
        return sum(
            (value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4) * weight
            for value, weight in zip(values, (0.2126, 0.7152, 0.0722))
        )

    a, b = luminance(first), luminance(second)
    return (max(a, b) + 0.05) / (min(a, b) + 0.05)


def _template_palette(slide: Any, fallback: str) -> list[str]:
    try:
        theme = slide.slide_layout.slide_master.part.part_related_by(RT.THEME)
    except KeyError:
        return [fallback, "000000", "FFFFFF"]
    xml = etree.fromstring(theme.blob)
    scheme = xml.find(".//" + qn("a:clrScheme"))
    colors = [fallback]
    if scheme is not None:
        for item in scheme:
            for color in item:
                value = color.get("lastClr") if color.tag == qn("a:sysClr") else color.get("val")
                if (
                    value
                    and len(value) == 6
                    and all(ch in "0123456789abcdefABCDEF" for ch in value)
                ):
                    colors.append(value.upper())
    return list(dict.fromkeys(colors))


def _color_on_fill(palette: list[str], fill: str, preferred: str) -> str:
    if _contrast(fill, preferred) >= 4.5:
        return preferred
    candidate = max(palette, key=lambda color: _contrast(fill, color))
    return candidate if _contrast(fill, candidate) >= 4.5 else _contrasting_text(fill)


def _theme_colors(slide: Any, fallback: str) -> list[str]:
    try:
        theme = slide.slide_layout.slide_master.part.part_related_by(RT.THEME)
    except KeyError:
        return [fallback]
    xml = etree.fromstring(theme.blob)
    colors = [fallback]
    for index in range(1, 7):
        accent = xml.find(".//" + qn(f"a:accent{index}"))
        if accent is None:
            continue
        value = accent.find(qn("a:srgbClr"))
        if value is not None and value.get("val") not in colors:
            colors.append(value.get("val"))
    return colors


def _chart_data_labels_fit(block: PlacedBlock, dataset: Dataset) -> bool:
    """Показывать подписи только с запасом для офисной автоматической верстки.

    Шрифт не уменьшаем; точные данные остаются в workbook и таблице HTML.
    Оценка консервативна: LibreOffice отводит подписи лишь часть ширины графика.
    """
    size = block.style.size
    width = max(0, block.box.width / 12700 - size * 6)
    height = max(0, block.box.height / 12700 - size * 5)
    count = len(dataset.rows)
    series = len(dataset.columns) - 1
    if width <= 0 or height < size * 2 or series > 1:
        return False
    if block.chart_type == "pie":
        values = [float(str(row[1]).replace(",", ".")) for row in dataset.rows]
        total = sum(values)
        if not total or any(value <= 0 for value in values):
            return False
        smallest_arc = math.pi * min(width, height) * min(values) / total
        return smallest_arc >= size * 5
    longest = max(len(cell_text(value)) for row in dataset.rows for value in row[1:])
    label_width = size * (longest * 0.60 + 1)
    available_width = min(width * 0.12, width / (count + 1))
    if label_width > available_width:
        return False
    return block.chart_type != "bar" or height / count >= size * 1.7


def _chart_legend_fits(block: PlacedBlock, labels: list[str]) -> bool:
    size = block.style.size
    label_width = size * (max(len(label) for label in labels) * 0.60 + 2)
    columns = max(1, int(block.box.width / 12700 / label_width))
    lines = math.ceil(len(labels) / columns)
    return lines * size * 1.4 <= block.box.height / 12700 * 0.25


def _add_chart(slide: Any, block: PlacedBlock, dataset: Dataset) -> None:
    if block.chart_type is None:
        raise NativeBuildError("Для диаграммы не указан тип")
    if block.chart_type == "pie" and len(dataset.columns) != 2:
        raise NativeBuildError("Круговая диаграмма требует ровно одну серию данных")
    data = CategoryChartData()
    data.categories = [cell_text(row[0]) for row in dataset.rows]
    for index, column in enumerate(dataset.columns[1:], start=1):
        try:
            values = [float(str(row[index]).replace(",", ".")) for row in dataset.rows]
        except (ValueError, TypeError) as error:
            raise NativeBuildError(f"Столбец {column!r} содержит нечисловые данные") from error
        if not all(math.isfinite(value) for value in values):
            raise NativeBuildError("Диаграмма содержит бесконечные или отсутствующие значения")
        if block.chart_type == "pie" and (any(value < 0 for value in values) or sum(values) <= 0):
            raise NativeBuildError(
                "Круговая диаграмма требует неотрицательные значения с суммой > 0"
            )
        data.add_series(column, values)
    chart_type = {
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
    }[block.chart_type]
    box = block.box
    shape = slide.shapes.add_chart(chart_type, box.left, box.top, box.width, box.height, data)
    shape.name = "exposlides:" + block.id
    chart = shape.chart
    # Явный setter запрещает LibreOffice добавлять заголовок из имени серии.
    # Его лишняя строка сжимала plot и скрывала B даже при tickLblSkip=1.
    chart.has_title = False
    _font(chart.font, block.style)
    for container in (chart._chartSpace, chart._chartSpace.chart.plotArea):
        properties = OxmlElement("c:spPr")
        properties.append(OxmlElement("a:noFill"))
        line = OxmlElement("a:ln")
        line.append(OxmlElement("a:noFill"))
        properties.append(line)
        container.insert_element_before(
            properties, "c:txPr", "c:externalData", "c:printSettings", "c:extLst"
        )
    chart.has_legend = len(dataset.columns) > 2 or (
        block.chart_type == "pie"
        and _chart_legend_fits(block, [cell_text(row[0]) for row in dataset.rows])
    )
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        _font(chart.legend.font, block.style)
    colors = _theme_colors(slide, block.fill or block.style.color)
    for index, series in enumerate(chart.series):
        targets = list(series.points) if block.chart_type == "pie" else [series]
        for target_index, target in enumerate(targets):
            color = colors[(target_index if block.chart_type == "pie" else index) % len(colors)]
            target.format.fill.solid()
            target.format.fill.fore_color.rgb = RGBColor.from_string(color)
            target.format.line.color.rgb = RGBColor.from_string(color)
    if block.chart_type != "pie":
        for axis, title in [
            (chart.category_axis, dataset.columns[0]),
            (chart.value_axis, dataset.unit or dataset.name),
        ]:
            axis.has_title = True
            axis.format.line.color.rgb = RGBColor.from_string(block.style.color)
            _write_text(axis.axis_title.text_frame, [title], block.style)
            _font(axis.tick_labels.font, block.style)
        category_size = block.style.size
        if block.chart_type == "bar":
            # LibreOffice может игнорировать tickLblSkip=1 при столкновении глифов.
            # Резервируем высоту для названий/числовой оси и межстрочных
            # интервалов категорий. Подписи остаются не мельче читаемых 14 pt.
            plot_height_pt = max(0, box.height / 12700 - block.style.size*5)
            label_size = min(block.style.size, max(
                14, math.floor(plot_height_pt / (len(dataset.rows)*2)),
            ))
            category_style = block.style.model_copy(update={"size": label_size})
            _font(chart.category_axis.tick_labels.font, category_style)
            category_size = label_size
            category_length = plot_height_pt
            # Одна строка подписи требует межстрочного запаса, а не 70% пустоты.
            # Избыточный запас скрывал B даже у трёх хорошо разделённых столбцов.
            category_spacing = category_size * 1.25
        else:
            category_length = max(0, box.width / 12700 - block.style.size * 6)
            category_spacing = category_size * (
                max(len(cell_text(row[0])) for row in dataset.rows) * 0.60 + 1
            )
        visible_categories = max(1, int(category_length / category_spacing))
        _set_category_stride(chart, max(1, math.ceil(len(dataset.rows) / visible_categories)))
        values = [float(str(value).replace(",", "."))
                  for row in dataset.rows for value in row[1:]]
        # Столбцы показывают абсолютную величину от нуля, включая отрицательные.
        if block.chart_type == "bar":
            if min(values) >= 0:
                chart.value_axis.minimum_scale = 0
            elif max(values) <= 0:
                chart.value_axis.maximum_scale = 0
        # Число делений зависит от доступной длины оси и кегля; данные не меняются.
        if block.chart_type == "bar":
            # Подписи категорий и название оси занимают существенную часть ширины.
            # Ширина shape целиком завышает реальную длину числовой оси в 1.5–2 раза.
            label_chars = max(len(cell_text(row[0])) for row in dataset.rows)
            categories_width = min(box.width*0.45,
                                   Pt(block.style.size*(label_chars*0.5+1.7)))
            axis_length = max(1, min(box.width*0.75,
                                    box.width-categories_width-Pt(block.style.size*1.2)))
            span = max(0, *values) - min(0, *values)
            tick_chars = (8 if max(abs(value) for value in values) >= 1e7 or span < 1e-4
                          else max(len(format(value, ",.6f").rstrip("0").rstrip("."))
                                   for value in [0, *values]))
            label_spacing = Pt(block.style.size * max(2.5, tick_chars * 0.60 + 1))
        else:
            axis_length = max(1, box.height-Pt(block.style.size*3))
            label_spacing = Pt(block.style.size*3)
        intervals = max(2, min(5, int(axis_length/label_spacing)))
        span = max(0, *values) - min(0, *values)
        if span > 0:
            raw_unit = span / intervals
            magnitude = 10 ** math.floor(math.log10(raw_unit))
            major_unit = next(step * magnitude for step in (1, 2, 2.5, 5, 10)
                              if step * magnitude >= raw_unit)
            chart.value_axis.major_unit = major_unit
            precision = max(0, -math.floor(math.log10(major_unit)))
            if major_unit / magnitude == 2.5:
                precision += 1
            number_format = ("0.##E+0" if max(abs(value) for value in values) >= 1e7
                             or span < 1e-4 else "#,##0" + ("." + "#"*precision if precision else ""))
            chart.value_axis.tick_labels.number_format = number_format
            chart.value_axis.tick_labels.number_format_is_linked = False
    plot = chart.plots[0]
    plot.has_data_labels = _chart_data_labels_fit(block, dataset)
    if plot.has_data_labels:
        plot.data_labels.show_value = block.chart_type != "pie"
        plot.data_labels.show_percentage = block.chart_type == "pie"
        _font(plot.data_labels.font, block.style)


def _set_category_stride(chart: Any, stride: int) -> None:
    """Разрежать только подписи; все исходные категории остаются в диаграмме.

    python-pptx пока не предоставляет setter для этих свойств CategoryAxis.
    Порядок элементов соответствует CT_CatAx; значение 1 означает каждую категорию.
    """
    axis = chart.category_axis._element
    for name, successors in (
        ("c:tickLblSkip", ("c:tickMarkSkip", "c:noMultiLvlLbl", "c:extLst")),
        ("c:tickMarkSkip", ("c:noMultiLvlLbl", "c:extLst")),
    ):
        element = axis.find(qn(name))
        if element is None:
            element = OxmlElement(name)
            axis.insert_element_before(element, *successors)
        element.set("val", str(stride))


def _add_diagram(slide: Any, block: PlacedBlock) -> None:
    items = block.items
    if len(items) < 2:
        raise NativeBuildError("Схема требует хотя бы два элемента")
    box = block.box
    group = slide.shapes.add_group_shape()
    group.name = "exposlides:" + block.id
    columns = min(4 if block.kind == "process" else 3, len(items))
    rows = math.ceil(len(items) / columns)
    gap = min(int(Pt(22)), box.width // (columns * 8), box.height // max(8, rows * 4))
    width = (box.width - gap * (columns - 1)) // columns
    height = (box.height - gap * (rows - 1)) // rows
    positions = []
    for index, _ in enumerate(items):
        row, column = divmod(index, columns)
        if block.kind == "process" and row % 2:
            column = columns - column - 1
        positions.append((box.left + column * (width + gap), box.top + row * (height + gap)))
    if block.kind == "process":
        for index in range(1, len(positions)):
            x1, y1 = positions[index - 1]
            x2, y2 = positions[index]
            if y1 == y2:
                start_x, end_x = (x1 + width, x2) if x2 > x1 else (x1, x2 + width)
                start_y = end_y = y1 + height // 2
            else:
                start_x, start_y = x1 + width // 2, y1 + height
                end_x, end_y = x2 + width // 2, y2
            connector = group.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, start_x, start_y, end_x, end_y
            )
            connector.name = f"exposlides:{block.id}:link:{index}"
            connector.line.color.rgb = RGBColor.from_string(block.fill or block.style.color)
            connector.line.width = Pt(1.5)
            tail = OxmlElement("a:tailEnd")
            tail.set("type", "triangle")
            connector._element.spPr.get_or_add_ln().append(tail)
    palette = _template_palette(slide, block.style.color)
    for index, ((left, top), text) in enumerate(zip(positions, items, strict=True)):
        node = group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        node.name = f"exposlides:{block.id}:node:{index}"
        fill = block.fill or max(palette, key=lambda color: _contrast(color, block.style.color))
        node.fill.solid()
        node.fill.fore_color.rgb = RGBColor.from_string(fill)
        node.line.color.rgb = RGBColor.from_string(block.style.color)
        style = block.style.model_copy(
            update={"color": _color_on_fill(palette, fill, block.style.color)}
        )
        _write_text(node.text_frame, [text], style, centered=True)
        node.text_frame.margin_left = node.text_frame.margin_right = Pt(8)
        node.text_frame.margin_top = node.text_frame.margin_bottom = Pt(6)


def _add_image(slide: Any, block: PlacedBlock) -> None:
    """Помещает отдельное изображение с сохранением пропорций в отведённую область."""
    from PIL import Image

    path = Path(block.image_path)
    if not path.is_file() or path.stat().st_size > 25 * 1024 * 1024:
        raise NativeBuildError("Нужен локальный PNG/JPEG размером не более 25 МБ")
    try:
        with Image.open(path) as image:
            width, height = image.size
            if image.format not in {"PNG", "JPEG"} or width * height > 25_000_000:
                raise NativeBuildError("Поддерживаются PNG/JPEG до 25 миллионов пикселей")
            image.verify()
    except NativeBuildError:
        raise
    except Exception as error:
        raise NativeBuildError("Не удалось прочитать изображение") from error
    box = block.box
    scale = min(box.width / width, box.height / height)
    target_width, target_height = max(1, round(width * scale)), max(1, round(height * scale))
    shape = slide.shapes.add_picture(
        str(path),
        box.left + (box.width - target_width) // 2,
        box.top + (box.height - target_height) // 2,
        width=target_width,
        height=target_height,
    )
    shape.name = "exposlides:" + block.id


def _add_icon(slide: Any, block: PlacedBlock) -> None:
    """Небольшой набор пиктограмм из редактируемых фигур, без растеризации."""
    box = block.box
    color = block.fill or block.style.color
    if block.icon == "check":
        group = slide.shapes.add_group_shape()
        group.name = "exposlides:" + block.id
        points = [
            (box.left, box.top + box.height // 2),
            (box.left + box.width // 3, box.top + box.height),
            (box.left + box.width, box.top),
        ]
        for index in range(2):
            line = group.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, *points[index], *points[index + 1]
            )
            line.name = f"exposlides:{block.id}:stroke:{index}"
            line.line.color.rgb = RGBColor.from_string(color)
            line.line.width = Pt(3)
        return
    kind = MSO_SHAPE.RIGHT_ARROW if block.icon == "arrow" else MSO_SHAPE.OVAL
    shape = slide.shapes.add_shape(kind, box.left, box.top, box.width, box.height)
    shape.name = "exposlides:" + block.id
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    if block.icon == "info":
        style = block.style.model_copy(
            update={
                "color": _color_on_fill(
                    _template_palette(slide, block.style.color), color, block.style.color
                ),
                "bold": True,
            }
        )
        _write_text(shape.text_frame, ["i"], style, centered=True)


def _validate_saved(path: Path, plan: DeckPlan) -> None:
    reopened = Presentation(path)
    if len(reopened.slides) != len(plan.slides):
        raise NativeBuildError("После повторного открытия изменилось число слайдов")
    for slide, instance in zip(reopened.slides, plan.slides, strict=True):
        shapes = {shape.name: shape for shape in native_shapes(slide.shapes)}
        by_id = {shape.shape_id: shape for shape in native_shapes(slide.shapes)}
        layout_ids = {shape.shape_id: shape for shape in native_shapes(slide.slide_layout.shapes)}
        for block in instance.blocks:
            if block.kind == "image" and block.image_fit == "template":
                shape = (layout_ids if block.source_layer == "layout" else by_id).get(block.source_shape_id)
            else:
                shape = (by_id.get(block.source_shape_id)
                     if block.kind in TEXT_KINDS and block.source_shape_id is not None
                     else shapes.get("exposlides:" + block.id))
            if shape is None:
                raise NativeBuildError(f"После сохранения отсутствует объект {block.id}")
            if block.kind == "chart" and not shape.has_chart:
                raise NativeBuildError("Диаграмма не сохранилась нативным объектом")
            if block.kind == "table" and not shape.has_table:
                raise NativeBuildError("Таблица не сохранилась нативным объектом")
    listed_slides = {slide.part for slide in reopened.slides}
    if any(
        isinstance(part, SlidePart) and part not in listed_slides
        for part in reopened.part.package.iter_parts()
    ):
        raise NativeBuildError(
            "Шаблон содержит ссылки из мастеров или макетов на скрытые исходные слайды"
        )
    # Обход всех связей заставляет библиотеку разрешить каждую внутреннюю часть.
    for part in reopened.part.package.iter_parts():
        for relationship in part.rels.values():
            if not relationship.is_external:
                _ = relationship.target_part.blob
    with ZipFile(path) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)) or archive.testzip() is not None:
            raise NativeBuildError("В созданном PPTX найдены повторные или повреждённые части")
        entries = set(names)
        relation_ids = {}
        for name in names:
            if not name.endswith(".rels"):
                continue
            folder, filename = posixpath.split(name)
            base = posixpath.dirname(folder)
            owner = posixpath.join(base, filename[:-5]) if name != "_rels/.rels" else ""
            relations = etree.fromstring(archive.read(name))
            relation_ids[owner] = {relation.get("Id") for relation in relations}
            for relation in relations:
                if relation.get("TargetMode") == "External":
                    continue
                target = unquote(urlsplit(relation.get("Target", "")).path)
                resolved = posixpath.normpath(posixpath.join(base, target)).lstrip("/")
                if resolved not in entries:
                    raise NativeBuildError(f"Связь PPTX ведёт к отсутствующей части: {resolved}")
        namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        for name in names:
            if not name.endswith(".xml"):
                continue
            xml = etree.fromstring(archive.read(name))
            for element in xml.iter():
                for attribute, value in element.attrib.items():
                    if (
                        attribute.startswith(namespace)
                        and value
                        and value not in relation_ids.get(name, set())
                    ):
                        raise NativeBuildError(f"В части {name} отсутствует связь {value}")
