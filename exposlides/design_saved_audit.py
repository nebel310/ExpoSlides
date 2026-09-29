"""Проверки фактического PPTX: данные, стили, геометрия и защищённые части."""

from __future__ import annotations

import hashlib
import math
from collections import Counter
from copy import deepcopy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

from exposlides.design_data import cell_text
from exposlides.design_models import AuditIssue, AuditReport, DeckPlan, TemplateProfile
from exposlides.design_pptx_parts import _prune_custom_show_actions
from exposlides.design_repetition import repeated_content_images

REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
STRUCTURAL_RELATIONS = {RT.SLIDE, RT.SLIDE_LAYOUT, RT.SLIDE_MASTER, RT.NOTES_SLIDE, RT.NOTES_MASTER}


def _xml(element):
    return etree.tostring(element, method="c14n", exclusive=True) if element is not None else b""


def _part_signature(part, visited=None):
    """Имена клонированных parts могут меняться; содержание и ссылки — нет."""
    visited = set() if visited is None else visited
    if id(part) in visited:
        return (part.content_type, "cycle")
    visited = visited | {id(part)}
    relationships = []
    for relation in part.rels.values():
        if relation.reltype in STRUCTURAL_RELATIONS:
            continue
        target = relation.target_ref if relation.is_external else _part_signature(
            relation.target_part, visited,
        )
        relationships.append((relation.rId, relation.reltype, target))
    return (part.content_type, hashlib.sha256(part.blob).hexdigest(), tuple(sorted(relationships)))


def _shape_signature(shape, ancestors):
    links = []
    for element in shape._element.iter():
        for key, value in element.attrib.items():
            if key.startswith("{"+REL_NS+"}") and value:
                # Пустой r:id допустим у встроенных действий PowerPoint.
                # Сам атрибут остаётся в XML-сигнатуре для сравнения.
                relation = shape.part.rels[value]
                target = relation.target_ref if relation.is_external else (
                    relation.reltype if relation.reltype in STRUCTURAL_RELATIONS
                    else _part_signature(relation.target_part)
                )
                links.append((value, relation.reltype, target))
    return _xml(shape._element), ancestors, tuple(sorted(links))


def _shapes(shapes, ancestors=()):
    result = {}
    for shape in shapes:
        result[shape.shape_id] = (shape, ancestors)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            transform = _xml(shape._element.find(qn("p:grpSpPr")))
            result.update(_shapes(shape.shapes, (*ancestors, transform)))
    return result


def _font_matches(font, style):
    try:
        color = str(font.color.rgb).upper()
    except (AttributeError, TypeError):
        return False
    return (font.name == style.font and font.size is not None
            and math.isclose(font.size.pt, style.size, abs_tol=0.01)
            and bool(font.bold) == style.bold and color == style.color.upper())


def _native_shape_format(shape):
    """Исходная фигура сохраняет имя, свойства, заливку и обводку целиком."""
    element = deepcopy(shape._element)
    for body in list(element):
        if body.tag == qn("p:txBody"):
            element.remove(body)
    return _xml(element)


def _paragraph_format(paragraph, ratio=1.0, *, plain_margin=None):
    element = deepcopy(paragraph._p)
    if plain_margin is not None:
        properties = element.get_or_add_pPr()
        properties.set("marL", str(plain_margin))
        properties.set("indent", "0")
    for child in list(element):
        if child.tag in {qn("a:r"), qn("a:fld"), qn("a:br")}:
            element.remove(child)
    if ratio < 1:
        for child in element.iter():
            if child.get("sz") is not None:
                child.set("sz", str(max(100, round(int(child.get("sz")) * ratio))))
    return _xml(element)


def _run_format(run, *, size=None):
    element = deepcopy(run) if run is not None else OxmlElement("a:r")
    for child in list(element):
        if child.tag == qn("a:t"):
            element.remove(child)
    if size is not None:
        properties = element.find(qn("a:rPr"))
        if properties is None:
            properties = OxmlElement("a:rPr")
            element.insert(0, properties)
        properties.set("sz", str(max(100, round(size * 100))))
    return _xml(element)


def _contrast_reference(shape, original, fallback):
    """Ожидаемая коррекция исходного оформления в окружении итогового слайда."""
    from exposlides.design_builder import _color_on_fill, _template_palette
    from exposlides.design_contrast import repair_text_contrast

    expected = type(shape)(deepcopy(original._element), shape._parent)
    for paragraph in expected.text_frame.paragraphs:
        if not any(child.tag in {qn("a:r"), qn("a:fld")} for child in paragraph._p):
            paragraph.add_run()
    parent = shape._element.getparent()
    parent.replace(shape._element, expected._element)
    try:
        palette = _template_palette(shape.part.slide, "FFFFFF")
        repair_text_contrast(shape.part.slide, expected, fallback,
                             lambda fill, color: _color_on_fill(palette, fill, color))
    finally:
        parent.replace(expected._element, shape._element)
    return expected


def _contrast_style(shape, style):
    from exposlides.design_builder import _color_on_fill, _template_palette
    from exposlides.design_contrast import _behind, _theme

    slide = shape.part.slide
    theme, mapping = _theme(slide)
    background = _behind(slide, shape, theme, mapping)
    if background is None:
        return style  # Сборщик обязан был добавить подложку.
    return style.model_copy(update={
        "color": _color_on_fill(_template_palette(slide, "FFFFFF"), background, style.color),
    })


def _native_text_matches(shape, original, block):
    # Разрешение унаследованного кегля общее со сборщиком; само сравнение
    # форматирования независимое и не пересоздаёт текст через его writer.
    from exposlides.design_native_text import (
        native_plain_text_margin,
        native_reference_size,
        native_run_size,
    )

    original = _contrast_reference(shape, original, block.style.color)
    if _native_shape_format(shape) != _native_shape_format(original):
        return False
    body, old_body = shape.text_frame._txBody, original.text_frame._txBody
    if tuple(_xml(child) for child in body if child.tag != qn("a:p")) != tuple(
        _xml(child) for child in old_body if child.tag != qn("a:p")
    ):
        return False
    paragraph_count = sum(len(item.splitlines()) for item in (block.items or [block.text])) or 1
    reference = native_reference_size(original, paragraph_count=paragraph_count)
    ratio = min(1.0, block.style.size / reference) if reference else 1.0
    prototypes = original.text_frame.paragraphs
    for index, paragraph in enumerate(shape.text_frame.paragraphs):
        prototype = prototypes[min(index, len(prototypes)-1)]
        margin = native_plain_text_margin(original, prototype) if block.kind == "text" else None
        if _paragraph_format(paragraph) != _paragraph_format(prototype, ratio, plain_margin=margin):
            return False
        runs = [child for child in paragraph._p if child.tag in {qn("a:r"), qn("a:fld")}]
        old_runs = [child for child in prototype._p if child.tag in {qn("a:r"), qn("a:fld")}]
        old_runs = old_runs or [None]
        if len(runs) != len(old_runs) or any(child.tag == qn("a:br") for child in paragraph._p):
            return False
        for run, old_run in zip(runs, old_runs, strict=True):
            size = native_run_size(original, prototype, old_run)
            size = size * ratio if ratio < 1 and size is not None else None
            if _run_format(run) != _run_format(old_run, size=size):
                return False
    return True


def _photo_format(element):
    element = deepcopy(element)
    for node in element.iter(qn("a:blip")):
        node.set(qn("r:embed"), "native-photo")
    for node in list(element.iter(qn("a:srcRect"))):
        node.getparent().remove(node)
    return element


def _layer_signature(part, image_ids=(), *, master=False):
    """Разрешены только заявленные фото и реестр изолированного layout/master."""
    element = deepcopy(part._element)
    if master:
        listing = element.find(qn("p:sldLayoutIdLst"))
        if listing is not None:
            element.remove(listing)
    for picture in list(element.iter(qn("p:pic"))):
        identity = picture.find("./" + qn("p:nvPicPr") + "/" + qn("p:cNvPr"))
        if identity is not None and int(identity.get("id")) in image_ids:
            picture.getparent().replace(picture, _photo_format(picture))
    links = []
    for node in element.iter():
        for key, value in node.attrib.items():
            if key.startswith("{" + REL_NS + "}") and value != "native-photo":
                relation = part.rels[value]
                target = relation.target_ref if relation.is_external else _part_signature(relation.target_part)
                links.append((value, relation.reltype, target))
    themes = [_part_signature(rel.target_part) for rel in part.rels.values() if rel.reltype == RT.THEME]
    return _xml(element), tuple(sorted(links)), tuple(themes)


def audit_saved_pptx(output: Path, template: Path, plan: DeckPlan,
                     profile: TemplateProfile) -> AuditReport:
    """Сопоставить редактируемые объекты сохранённого файла с подтверждённым планом."""
    result, source = Presentation(output), Presentation(template)
    # Builder удаляет только действия исходных custom shows. Нормализуем
    # эталон в памяти, чтобы разрешённое удаление не считалось порчей оформления.
    # Готовый файл не нормализуем: новые действия и остальные изменения — ошибка.
    _prune_custom_show_actions(source)
    if len(result.slides) != len(plan.slides):
        raise ValueError("Число сохранённых слайдов не совпадает с планом")
    patterns = {pattern.source_slide_index: pattern for pattern in profile.patterns}
    datasets = {dataset.id: dataset for dataset in plan.datasets}
    issues = repeated_content_images(output, plan, profile)

    def issue(rule, instance, block, message, discriminator=""):
        identity = f"{rule}:{instance.id}:{block.id if block else 'slide'}:{discriminator}"
        issues.append(AuditIssue(
            id=hashlib.sha256(identity.encode()).hexdigest()[:16], rule=rule, severity="error",
            slide_id=instance.id, block_id=block.id if block else None,
            box=block.box if block else None, message=message,
            source_ids=block.source_ids if block else [],
        ))

    for index, instance in enumerate(plan.slides):
        slide, original = result.slides[index], source.slides[instance.source_slide_index-1]
        actual, old = _shapes(slide.shapes), _shapes(original.shapes)
        layout, original_layout = slide.slide_layout, original.slide_layout
        pattern = patterns[instance.source_slide_index]
        layout_image_ids = {block.source_shape_id for block in instance.blocks
                            if block.kind == "image" and block.image_fit == "template"
                            and block.source_layer == "layout"}
        if not layout_image_ids.issubset(pattern.layout_images):
            issue("source_shape_binding", instance, None, "Фото макета не разрешено к замене")
        if (_layer_signature(layout.part, layout_image_ids) != _layer_signature(original_layout.part, layout_image_ids)
                or _layer_signature(layout.slide_master.part, master=True)
                != _layer_signature(original_layout.slide_master.part, master=True)):
            issue("template_layout", instance, None, "Изменены макет, мастер или тема шаблона")
        if _xml(slide._element.cSld.find(qn("p:bg"))) != _xml(original._element.cSld.find(qn("p:bg"))):
            issue("template_background", instance, None, "Изменён фон исходного образца")
        pattern = patterns[instance.source_slide_index]
        for shape_id in pattern.protected_shape_ids:
            if shape_id in pattern.replaceable_images and any(
                block.kind == "image" and block.image_fit == "template"
                and block.source_layer == "slide" and block.source_shape_id == shape_id
                for block in instance.blocks
            ):
                continue  # Байты и маска нативного фото проверяются ниже.
            if (shape_id in pattern.replaceable_images and shape_id in instance.remove_shape_ids
                    and any(b.kind == "image" for b in instance.blocks)):
                # Новый image проверяется ниже по байтам и геометрии; старая фотография
                # должна исчезнуть. Логотипы не входят в replaceable_images.
                if (shape_id in actual and shape_id in old
                        and _shape_signature(*actual[shape_id]) == _shape_signature(*old[shape_id])):
                    issue("old_image_retained", instance, None, "Сохранена заменяемая фотография")
                continue
            if (shape_id not in actual or shape_id not in old
                    or _shape_signature(*actual[shape_id]) != _shape_signature(*old[shape_id])):
                issue("protected_changed", instance, None,
                      f"Изменён защищённый объект шаблона {shape_id}", str(shape_id))
        by_name = {}
        for shape, ancestors in actual.values():
            by_name.setdefault(shape.name, []).append((shape, ancestors))
        native_ids = Counter(block.source_shape_id for block in instance.blocks
                             if block.kind in {"title", "text", "page_number"}
                             and block.source_shape_id is not None)
        for block in instance.blocks:
            native_image = block.kind == "image" and block.image_fit == "template"
            native = (block.kind in {"title", "text", "page_number"} and (
                block.source_shape_id is not None
            )) or native_image
            saved_shapes, source_shapes = actual, old
            if native_image and block.source_layer == "layout":
                saved_shapes, source_shapes = _shapes(layout.shapes), _shapes(original_layout.shapes)
            if native and not native_image and native_ids[block.source_shape_id] != 1:
                issue("source_shape_binding", instance, block,
                      "Исходная фигура связана с несколькими блоками")
            source_shape = source_shapes.get(block.source_shape_id) if native else None
            matches = ([saved_shapes[block.source_shape_id]]
                       if native and source_shape and block.source_shape_id in saved_shapes
                       else [] if native else by_name.get("exposlides:"+block.id, []))
            if len(matches) != 1:
                issue("native_object_missing", instance, block,
                      "Объект плана отсутствует или имеет неоднозначное имя")
                continue
            shape, ancestors = matches[0]
            expected_box = (block.box.left, block.box.top, block.box.width, block.box.height)
            if block.kind == "image" and shape._element.tag == qn("p:pic"):
                width, height = shape.image.size
                scale = min(block.box.width/width, block.box.height/height)
                fitted = max(1, round(width*scale)), max(1, round(height*scale))
                expected_box = (block.box.left+(block.box.width-fitted[0])//2,
                                block.box.top+(block.box.height-fitted[1])//2, *fitted)
                visible_width = width * (1 - shape.crop_left - shape.crop_right) if native_image else width
                visible_height = height * (1 - shape.crop_top - shape.crop_bottom) if native_image else height
                target_ratio = block.box.width / block.box.height if native_image else shape.width / shape.height
                if visible_width <= 0 or visible_height <= 0 or abs(
                    target_ratio / (visible_width / visible_height) - 1
                ) > 0.02:
                    issue("image_aspect", instance, block, "Пропорции изображения изменены")
                if native_image and _xml(_photo_format(shape._element)) != _xml(_photo_format(source_shape[0]._element)):
                    issue("native_image_mask", instance, block, "Изменена форма или оформление исходной фотографии")
                image_path = Path(block.image_path)
                if not image_path.is_file() or shape.image.blob != image_path.read_bytes():
                    issue("image_content", instance, block, "Изображение отличается от выбранного файла")
            actual_box = (shape.left, shape.top, shape.width, shape.height)
            if native:
                original_shape, original_ancestors = source_shape
                expected_box = (original_shape.left, original_shape.top,
                                original_shape.width, original_shape.height)
                geometry_changed = (ancestors != original_ancestors
                                    or shape.rotation != original_shape.rotation
                                    or actual_box != expected_box)
            else:
                geometry_changed = (bool(ancestors) or shape.rotation != 0
                                    or any(abs(a-b) > 3 for a, b in zip(
                                        actual_box, expected_box, strict=True)))
            if geometry_changed:
                issue("saved_geometry", instance, block, "Геометрия объекта отличается от плана")
            if block.kind in {"title", "text", "page_number"}:
                if not shape.has_text_frame:
                    issue("native_text", instance, block, "Текстовый блок не является редактируемым текстом")
                    continue
                expected_text = "\n".join(block.items or block.text.splitlines())
                if shape.text != expected_text:
                    issue("saved_text", instance, block, "Текст сохранённого объекта отличается от плана")
                runs = [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
                matches_style = (_native_text_matches(shape, original_shape, block)
                                 if native and original_shape.has_text_frame else
                                 not native and runs and all(
                                     _font_matches(run.font, _contrast_style(shape, block.style)) for run in runs))
                if not matches_style:
                    issue("saved_text_style", instance, block, "Шрифт, размер, начертание или цвет отличается от плана")
            elif block.kind == "table":
                if not shape.has_table:
                    issue("native_table", instance, block, "Таблица не является редактируемой таблицей")
                    continue
                dataset = datasets[block.dataset_id]
                expected = [[cell_text(value) for value in row] for row in [dataset.columns, *dataset.rows]]
                saved = [[cell.text for cell in row.cells] for row in shape.table.rows]
                if expected != saved:
                    issue("saved_data", instance, block, "Ячейки таблицы отличаются от исходных данных")
            elif block.kind == "chart":
                if not shape.has_chart:
                    issue("native_chart", instance, block, "Диаграмма не является нативной диаграммой")
                    continue
                chart, dataset = shape.chart, datasets[block.dataset_id]
                expected_type = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
                                 "pie": XL_CHART_TYPE.PIE}.get(block.chart_type)
                if chart.chart_type != expected_type:
                    issue("native_chart", instance, block, "Тип диаграммы отличается от плана")
                expected_values = [[float(str(row[column]).replace(",", ".")) for row in dataset.rows]
                                   for column in range(1, len(dataset.columns))]
                saved_values = [list(series.values) for series in chart.series]
                categories = [category.label for category in chart.plots[0].categories]
                if (categories != [cell_text(row[0]) for row in dataset.rows]
                        or [series.name for series in chart.series] != dataset.columns[1:]
                        or len(saved_values) != len(expected_values)
                        or any(len(saved) != len(expected) or any(
                            a is None or not math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-9)
                            for a, b in zip(saved, expected)
                        ) for saved, expected in zip(saved_values, expected_values))):
                    issue("saved_data", instance, block, "Диаграмма отличается от исходных данных")
                if block.chart_type != "pie":
                    try:
                        labels = (chart.category_axis, chart.value_axis)
                        expected_labels = (dataset.columns[0], dataset.unit or dataset.name)
                        valid = all(axis.has_title and axis.axis_title.text_frame.text == expected
                                    for axis, expected in zip(labels, expected_labels, strict=True))
                    except ValueError:
                        valid = False
                    if not valid:
                        issue("chart_labels", instance, block, "Названия осей отличаются от исходных данных")
            elif block.kind == "image" and shape._element.tag != qn("p:pic"):
                issue("native_image", instance, block, "Изображение отсутствует в сохранённом файле")
            elif block.kind in {"process", "comparison"}:
                if shape.shape_type != MSO_SHAPE_TYPE.GROUP:
                    issue("native_diagram", instance, block, "Схема не состоит из редактируемых фигур")
                elif [item.text for item in shape.shapes if item.has_text_frame] != block.items:
                    issue("saved_text", instance, block, "Подписи схемы отличаются от плана")
            elif block.kind == "smartart":
                rel_ids = shape._element.find(f".//{{{DGM}}}relIds")
                if rel_ids is None or not rel_ids.get(qn("r:dm")):
                    issue("native_smartart", instance, block, "Объект не является нативным SmartArt")
                else:
                    data = etree.fromstring(slide.part.related_part(rel_ids.get(qn("r:dm"))).blob)
                    labels = ["".join(text.text or "" for text in point.iter(qn("a:t")))
                              for point in data.iter(f"{{{DGM}}}pt")
                              if point.get("type", "node") == "node"
                              and any((text.text or "").strip() for text in point.iter(qn("a:t")))]
                    if labels != block.items:
                        issue("saved_text", instance, block, "Подписи SmartArt отличаются от плана")
            elif block.kind == "icon" and shape.shape_type not in {
                MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.GROUP,
            }:
                issue("native_icon", instance, block, "Иконка не является редактируемой фигурой")
    return AuditReport(issues=issues, checks=[
        "pptx_reopen", "native_text", "native_table", "native_chart", "native_object_missing",
        "template_layout", "template_background", "protected_changed", "chart_labels", "image_aspect",
        "saved_geometry", "saved_text", "saved_text_style", "saved_data", "image_content",
        "native_image", "native_image_mask", "native_diagram", "native_smartart", "native_icon",
        "source_shape_binding", "native_template_text_format",
    ])
