"""Проверка повторов содержательных иллюстраций в уже сохранённой колоде."""

from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

from exposlides.design_models import AuditIssue, DeckPlan, TemplateProfile
from exposlides.design_native_text import native_shapes


def image_identity(data: bytes) -> str:
    """Метаданные и способ упаковки PNG не делают картинку новой."""
    with Image.open(BytesIO(data)) as image:
        pixels = image.convert("RGBA")
        digest = hashlib.sha256(str(pixels.size).encode("ascii"))
        digest.update(pixels.tobytes())
        return digest.hexdigest()


def repeated_content_images(output: Path, plan: DeckPlan,
                            profile: TemplateProfile) -> list[AuditIssue]:
    """Логотипы/фон не считаются контентом; фото в layout и группах учитываются."""
    presentation = Presentation(output)
    patterns = {p.source_slide_index: p for p in profile.patterns}
    seen, identities, issues = {}, {}, []
    for slide_number, (slide, instance) in enumerate(
        zip(presentation.slides, plan.slides, strict=True), 1,
    ):
        pattern = patterns[instance.source_slide_index]
        for layer, shapes, photos in (
            ("slide", slide.shapes, pattern.replaceable_images),
            ("layout", slide.slide_layout.shapes, pattern.layout_images),
        ):
            generated = {"exposlides:" + b.id for b in instance.blocks
                         if b.kind == "image" and b.image_fit != "template"}
            for shape in native_shapes(shapes):
                if shape._element.tag != qn("p:pic"):
                    continue
                if shape.shape_id not in photos and not (layer == "slide" and shape.name in generated):
                    continue
                data = shape.image.blob
                digest = hashlib.sha256(data).hexdigest()
                if digest not in identities:
                    identities[digest] = image_identity(data)
                key = identities[digest]
                if key in seen:
                    identity = f"repeated_image:{instance.id}:{layer}:{shape.shape_id}"
                    issues.append(AuditIssue(
                        id=hashlib.sha256(identity.encode()).hexdigest()[:16],
                        rule="repeated_image", severity="error", slide_id=instance.id,
                        message=f"Иллюстрация повторяет изображение на слайде {seen[key]}",
                    ))
                else:
                    seen[key] = slide_number
    return issues
