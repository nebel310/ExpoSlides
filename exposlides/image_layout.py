"""Поиск областей шаблона и замена изображений без изменения верстки."""

from __future__ import annotations

import hashlib
import io
from collections import Counter
from copy import deepcopy
from dataclasses import dataclass

from PIL import Image, ImageOps
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

from exposlides.image_api import ImageGenerationError

PROTECTED_NAMES = ("logo", "логотип", "brand", "watermark", "фон", "background", "keep")


@dataclass(frozen=True)
class ImageSlot:
    slide_index: int
    shape_id: int
    name: str
    left: int
    top: int
    width: int
    height: int

    @property
    def pixels(self) -> tuple[int, int]:
        ratio = self.width / self.height
        if ratio >= 1:
            return 1024, max(256, round(1024 / ratio / 32) * 32)
        return max(256, round(1024 * ratio / 32) * 32), 1024


def _picture_hash(shape) -> str | None:
    if shape._element.find(qn("p:blipFill")) is not None:
        return hashlib.sha256(shape.image.blob).hexdigest()
    return None


def _overlaps_text(shape, slide) -> bool:
    for other in slide.shapes:
        if other.shape_id == shape.shape_id or not other.has_text_frame or not other.text.strip():
            continue
        overlap_width = min(shape.left + shape.width, other.left + other.width) - max(
            shape.left, other.left
        )
        overlap_height = min(shape.top + shape.height, other.top + other.height) - max(
            shape.top, other.top
        )
        if overlap_width > 0 and overlap_height > 0:
            return True
    return False


def find_image_slots(presentation) -> list[ImageSlot]:
    """PICTURE placeholders, явно помеченные картинки и крупные неповторяющиеся фото.

    Группы, фоны, логотипы и области под текстом не заменяем. Маркер
    exposlides:image в имени позволяет явно выбрать обычную картинку.
    """
    counts = Counter(
        digest for slide in presentation.slides for shape in slide.shapes
        if (digest := _picture_hash(shape)) is not None
    )
    slots = []
    slide_area = presentation.slide_width * presentation.slide_height
    for index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            name = shape.name.lower()
            if any(word in name for word in PROTECTED_NAMES):
                continue
            is_picture_placeholder = (
                shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE
            )
            is_picture = shape.shape_type == MSO_SHAPE_TYPE.PICTURE or is_picture_placeholder
            if not is_picture or shape.width <= 0 or shape.height <= 0:
                continue
            if shape.rotation != 0:
                continue
            if (shape.left < 0 or shape.top < 0
                    or shape.left + shape.width > presentation.slide_width
                    or shape.top + shape.height > presentation.slide_height):
                continue
            if _overlaps_text(shape, slide):
                continue
            explicit = "exposlides:image" in name
            if not explicit and not is_picture_placeholder:
                area = shape.width * shape.height / slide_area
                ratio = shape.width / shape.height
                if not (0.08 <= area <= 0.65 and 0.4 <= ratio <= 2.5):
                    continue
                if counts[_picture_hash(shape)] != 1:
                    continue
                # Прозрачные PNG часто содержат брендинг/иконки, а не фото.
                try:
                    with Image.open(io.BytesIO(shape.image.blob)) as source:
                        if "A" in source.getbands() or "transparency" in source.info:
                            continue
                except (OSError, ValueError):
                    continue
            slots.append(ImageSlot(
                index, shape.shape_id, shape.name,
                shape.left, shape.top, shape.width, shape.height,
            ))
    return slots


def replace_slot(presentation, slot: ImageSlot, data: bytes) -> str:
    """Меняет только картинку, сохраняя shape id, слой, ссылки, рамку и геометрию."""
    slide = presentation.slides[slot.slide_index - 1]
    shape = next(s for s in slide.shapes if s.shape_id == slot.shape_id)
    # Подгоняем по отношению сторон рамки до замены, без растягивания.
    ratio = slot.width / slot.height
    size = (1600, max(1, round(1600 / ratio))) if ratio >= 1 else (
        max(1, round(1600 * ratio)), 1600
    )
    with Image.open(io.BytesIO(data)) as source:
        fitted = ImageOps.fit(source.convert("RGB"), size, method=Image.Resampling.LANCZOS)
        stream = io.BytesIO()
        fitted.save(stream, format="PNG")
    stream.seek(0)
    if shape._element.find(qn("p:blipFill")) is not None:
        _, rid = slide.part.get_or_add_image_part(stream)
        shape._element.blipFill.blip.set(qn("r:embed"), rid)
        shape.crop_left = shape.crop_right = shape.crop_top = shape.crop_bottom = 0
        picture = shape
    else:
        # Публичный API python-pptx заменяет PICTURE placeholder в том же слое.
        original_properties = deepcopy(shape._element.spPr)
        original_nonvisual = deepcopy(shape._element.xpath(".//p:cNvPr")[0])
        picture = shape.insert_picture(stream)
        picture.crop_left = picture.crop_right = picture.crop_top = picture.crop_bottom = 0
        properties = picture._element.spPr
        properties.getparent().replace(properties, original_properties)
        nonvisual = picture._element.xpath(".//p:cNvPr")[0]
        nonvisual.getparent().replace(nonvisual, original_nonvisual)
    picture._element.xpath(".//p:cNvPr")[0].set(
        "descr", "Иллюстрация сгенерирована ExpoSlides / Tongyi-MAI/Z-Image-Turbo"
    )
    if (picture.left, picture.top, picture.width, picture.height) != (
        slot.left, slot.top, slot.width, slot.height
    ):
        raise ImageGenerationError("При вставке изображения изменилась геометрия шаблона.")
    return hashlib.sha256(picture.image.blob).hexdigest()
