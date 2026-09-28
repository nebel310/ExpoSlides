"""Этап иллюстраций общего CLI/web pipeline после нативной сборки PPTX."""

from __future__ import annotations

import asyncio
import hashlib
import json
import os
import tempfile
from pathlib import Path

import httpx
from pptx import Presentation

from exposlides.image_api import (
    IMAGE_MODEL,
    ImageGenerationError,
    ImageSettings,
    describe_scene,
    generate_image,
)
from exposlides.image_layout import find_image_slots, replace_slot

ROOT = Path(__file__).resolve().parent.parent
PROMPT_PATH = Path(__file__).resolve().parent / "prompts" / "slide_image.md"


async def _generate_all(jobs: list[dict], settings: ImageSettings) -> list[bytes]:
    semaphore = asyncio.Semaphore(settings.image_concurrency)
    async with httpx.AsyncClient(timeout=settings.image_api_timeout, follow_redirects=False) as client:
        async def generate(job):
            async with semaphore:
                scene = await describe_scene(client, settings, job["prompt"])
                return await generate_image(client, settings, **(job | {"prompt": scene}))

        tasks = [asyncio.create_task(generate(job)) for job in jobs]
        try:
            async with asyncio.timeout(settings.image_generation_timeout):
                return await asyncio.gather(*tasks)
        except TimeoutError as error:
            raise ImageGenerationError("Истекло общее время генерации изображений.") from error
        finally:
            for task in tasks:
                if not task.done():
                    task.cancel()
            await asyncio.gather(*tasks, return_exceptions=True)


def illustrate_presentation(
    pptx_path: Path,
    template_json: Path,
    *,
    report_path: Path,
    mode: str = "auto",
    settings: ImageSettings | None = None,
) -> dict:
    """Атомарная замена иллюстраций; ошибка не оставляет частично готовую колоду."""
    if mode not in {"auto", "off"}:
        raise ImageGenerationError("Режим изображений должен быть auto или off.")
    report = {"model": IMAGE_MODEL, "provider": "hf/fal-ai", "status": "disabled", "images": []}
    if mode == "off":
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report
    try:
        presentation = Presentation(pptx_path)
        slots = find_image_slots(presentation)
        if not slots:
            report["status"] = "no_slots"
        else:
            # Конфигурация загружается приложением; секрет не сериализуется и не логируется.
            settings = settings or ImageSettings(
                _env_file=ROOT / "services" / "content-service" / ".env"
            )
            if settings.image_generation_enabled:
                settings.token()
                if len(slots) > settings.image_max_count:
                    raise ImageGenerationError(
                        "Число областей изображений превышает IMAGE_MAX_COUNT. "
                        "Увеличьте лимит или пометьте сохраняемые картинки словом keep."
                    )
                print(f"[images] Генерация иллюстраций: {len(slots)}", flush=True)
                template = json.loads(template_json.read_text(encoding="utf-8"))
                colors = (template.get("theme") or {}).get("colors", {})
                palette = ", ".join(
                    f"{key}: #{value}" for key, value in colors.items()
                    if key.startswith("accent") or key in {"dk1", "lt1"}
                ) or "restrained neutral tones"
                prompt_template = PROMPT_PATH.read_text(encoding="utf-8")
                jobs = []
                for slot in slots:
                    slide = presentation.slides[slot.slide_index - 1]
                    content = "\n".join(
                        s.text for s in slide.shapes if s.has_text_frame and s.text.strip()
                    )[:5000]
                    if not content.strip():
                        raise ImageGenerationError(
                            f"Слайд {slot.slide_index}: нет текста для тематической иллюстрации."
                        )
                    prompt = prompt_template.format(
                        palette=palette, ratio=f"{slot.width}:{slot.height}", content=content,
                    )
                    width, height = slot.pixels
                    seed = int(hashlib.sha256(prompt.encode()).hexdigest()[:8], 16) % (2**31)
                    jobs.append(dict(prompt=prompt, width=width, height=height, seed=seed))
                data = asyncio.run(_generate_all(jobs, settings))
                expected = []
                for slot, image in zip(slots, data, strict=True):
                    digest = replace_slot(presentation, slot, image)
                    expected.append((slot, digest))
                    report["images"].append({
                        "slide_index": slot.slide_index, "shape_id": slot.shape_id,
                        "sha256": digest,
                        "bbox": [slot.left, slot.top, slot.width, slot.height],
                    })
                fd, name = tempfile.mkstemp(suffix=".pptx", dir=pptx_path.parent)
                os.close(fd)
                temporary = Path(name)
                try:
                    presentation.save(temporary)
                    reopened = Presentation(temporary)
                    if len(reopened.slides) != len(presentation.slides):
                        raise ImageGenerationError("После вставки изображений изменилось число слайдов.")
                    for slot, digest in expected:
                        shape = next(
                            s for s in reopened.slides[slot.slide_index - 1].shapes
                            if s.shape_id == slot.shape_id
                        )
                        if (hashlib.sha256(shape.image.blob).hexdigest() != digest
                                or (shape.left, shape.top, shape.width, shape.height) != (
                                    slot.left, slot.top, slot.width, slot.height
                                )):
                            raise ImageGenerationError("Проверка вставленных изображений не пройдена.")
                    temporary.replace(pptx_path)
                finally:
                    temporary.unlink(missing_ok=True)
                report["status"] = "generated"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        if report["status"] == "no_slots":
            print("[images] В шаблоне нет подходящих областей для иллюстраций", flush=True)
        return report
    except ImageGenerationError:
        raise
    except Exception as error:
        raise ImageGenerationError("Не удалось подготовить или вставить изображения в PPTX.") from error
