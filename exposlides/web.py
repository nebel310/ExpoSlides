"""Локальный интерфейс ExpoSlides поверх существующего файлового CLI."""

from __future__ import annotations

import argparse
import base64
import binascii
import hmac
import io
import json
import os
import queue
import re
import secrets
import shutil
import signal
import subprocess
import sys
import tempfile
import threading
import zipfile
from dataclasses import dataclass, field
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import urlsplit
from uuid import uuid4

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from exposlides.cli import CONTENT_ERROR_MARKER, REPOSITORY_ROOT
from exposlides.file_storage import FileServiceStorage, StorageError
from exposlides.preview import PreviewRenderer
from exposlides.web_catalog import WebCatalog

STATIC_ROOT = Path(__file__).parent / "web_static"
MAX_TEMPLATE_BYTES = 25 * 1024 * 1024
MAX_REQUEST_BYTES = (MAX_TEMPLATE_BYTES + 2) // 3 * 4 + 8192
MAX_SCRIPT_LENGTH = 100_000
MAX_UNCOMPRESSED_BYTES = 150 * 1024 * 1024
MAX_TEMPLATE_SLIDES = 250
STORAGE_UNAVAILABLE = "Не удалось обратиться к хранилищу файлов. Попробуйте ещё раз позже."
STORAGE_SAVE_FAILED = (
    "Не удалось сохранить презентацию в хранилище. Проверьте доступность хранилища "
    "и запустите сборку ещё раз."
)
PREVIEW_UNAVAILABLE = (
    "Предпросмотр с оформлением недоступен на этом компьютере. "
    "Вы можете продолжить создание презентации."
)
PREVIEW_FAILED = (
    "Не удалось отобразить оформление шаблона. "
    "Вы можете продолжить создание презентации."
)
RESULT_PREVIEW_UNAVAILABLE = (
    "Предпросмотр с оформлением недоступен на этом компьютере. "
    "Готовую презентацию можно скачать."
)
RESULT_PREVIEW_FAILED = (
    "Не удалось отобразить оформление готовой презентации. "
    "Вы можете скачать её в формате PPTX."
)
STATIC_ROUTES = {
    "/": ("index.html", "text/html; charset=utf-8"),
    "/styles.css": ("styles.css", "text/css; charset=utf-8"),
    "/app.js": ("app.js", "text/javascript; charset=utf-8"),
    "/favicon.svg": ("favicon.svg", "image/svg+xml"),
}
STAGE_MARKERS = {"[1/3]": "parser", "[2/3]": "content", "[3/3]": "builder"}
CONTENT_LOG = re.compile(
    r"[0-9]{4}-[0-9]{2}-[0-9]{2} [0-9]{2}:[0-9]{2}:[0-9]{2} - "
    r"(app\.graph\.nodes|app\.graph\.fast|app\.chains\.llm) - (INFO|WARNING) - ([^\r\n]*)"
)
CONTENT_PHASES = {
    "=== Запуск узла analyze_script ===": "analysis",
    "=== Запуск узла plan_slides ===": "planning",
    "=== Запуск узла generate_content ===": "slides",
    "=== Запуск узла validate_content ===": "validation",
}
FAST_CONTENT_PHASES = {
    "=== Запуск узла analyze_and_plan ===": "planning",
    "=== Запуск узла generate_batch ===": "slides",
    "=== Запуск узла validate_content ===": "validation",
}
PHASE_RETRIES = {
    "analysis": r"Анализ источника отклонён, повтор [1-5]: [^\r\n]*",
    "planning": r"План отклонён, повтор [1-5]: [^\r\n]*",
    "slides": r"Контент слайда ([1-9][0-9]{0,2}) отклонён, повтор [1-5]: [^\r\n]*",
}
STAGE_ERRORS = {
    "parser": "Не удалось прочитать шаблон. Проверьте PPTX и попробуйте ещё раз.",
    "content": (
        "Не удалось подготовить текст слайдов. Попробуйте запустить сборку ещё раз."
    ),
    "builder": "Не удалось собрать презентацию. Попробуйте другой PPTX-шаблон.",
}
CONTENT_ERRORS = {
    "invalid_response": (
        "Модель Qwen вернула ответ в неподходящем формате. "
        "Не удалось исправить его автоматически. Запустите сборку ещё раз."
    ),
    "content_validation": (
        "Созданный текст не прошёл проверку: факты или объём текста не соответствуют "
        "исходным материалам и шаблону. Запустите сборку ещё раз."
    ),
    "timeout": (
        "Модель Qwen не успела подготовить текст за отведённое время. "
        "Попробуйте запустить сборку ещё раз."
    ),
    "auth": (
        "Сервис генерации отклонил доступ. Проверьте токен и права доступа к модели."
    ),
    "network": (
        "Не удалось связаться с сервисом генерации или он временно недоступен. "
        "Проверьте соединение и попробуйте ещё раз."
    ),
    "unknown": STAGE_ERRORS["content"],
}
EXAMPLE_SCRIPT = (
    "Квартальный обзор команды продукта. За квартал мы упростили первый запуск, "
    "обновили справочный раздел и собрали обратную связь от клиентов.\n\n"
    "Что изменилось. В продукте появился короткий вводный сценарий. "
    "Инструкции собраны в одном разделе, а заявки поддержки теперь проходят "
    "через единый список приоритетов.\n\n"
    "Ключевые результаты. Команда выпустила обновлённый первый запуск, "
    "подготовила базу знаний и провела интервью с клиентами. "
    "Основной запрос клиентов — понятный путь от знакомства до первого результата.\n\n"
    "Следующий квартал. Проверим новый вводный сценарий с клиентами, "
    "дополним базу знаний и настроим сбор обратной связи внутри продукта.\n\n"
    "Главное. Сосредоточимся на понятном первом опыте и будем выбирать "
    "следующие изменения на основе обратной связи от клиентов."
)


class APIError(Exception):
    """Безопасное сообщение об ошибке для интерфейса."""

    def __init__(self, message: str, status: int = HTTPStatus.BAD_REQUEST) -> None:
        super().__init__(message)
        self.status = status


def inspect_template(
    data: bytes, name: str, *, require_placeholders: bool = True,
) -> dict[str, Any]:
    """Проверить ограниченный по размеру PPTX до чтения python-pptx."""
    if not data or len(data) > MAX_TEMPLATE_BYTES:
        raise APIError("Размер PPTX должен быть не больше 25 МБ.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            names = [member.filename for member in members]
            if len(members) > 5000 or len(names) != len(set(names)):
                raise APIError("В шаблоне слишком много файлов или есть повторяющиеся записи.")
            if not {"[Content_Types].xml", "ppt/presentation.xml"}.issubset(names):
                raise APIError("Этот файл не является презентацией PPTX.")
            expanded_size = 0
            for member in members:
                path = PurePosixPath(member.filename)
                if path.is_absolute() or ".." in path.parts or member.flag_bits & 1:
                    raise APIError("Шаблон содержит неподдерживаемые записи.")
                expanded_size += member.file_size
                if member.file_size > 50 * 1024 * 1024:
                    raise APIError("В шаблоне есть слишком большой встроенный файл.")
            if expanded_size > MAX_UNCOMPRESSED_BYTES:
                raise APIError("Распакованный шаблон слишком большой. Уменьшите размер изображений.")
            if archive.testzip() is not None:
                raise APIError("PPTX повреждён. Сохраните его заново в PowerPoint.")
        presentation = Presentation(io.BytesIO(data))
        if not 1 <= len(presentation.slides) <= MAX_TEMPLATE_SLIDES:
            raise APIError("Шаблон должен содержать от 1 до 250 слайдов.")
        slides = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts = [
                shape.text[:5000]
                for shape in slide.shapes
                if shape.has_text_frame and shape.text.strip()
            ][:30]
            title_shape = slide.shapes.title
            title = title_shape.text if title_shape is not None else (texts[0] if texts else "")
            slides.append(
                {
                    "index": index,
                    "title": title[:180] or f"Слайд {index}",
                    "texts": texts,
                    "placeholder_count": sum(
                        shape.is_placeholder and shape.has_text_frame for shape in slide.shapes
                    ),
                }
            )
        if require_placeholders and not any(slide["placeholder_count"] for slide in slides):
            raise APIError(
                "В шаблоне нет текстовых заполнителей. "
                "Добавьте заголовок или область текста через макет PowerPoint."
            )
        return {
            "name": name,
            "slide_count": len(slides),
            "width": presentation.slide_width / Inches(1),
            "height": presentation.slide_height / Inches(1),
            "slides": slides,
        }
    except APIError:
        raise
    except Exception as error:
        raise APIError("Не удалось открыть PPTX. Проверьте формат и целостность файла.") from error


def make_example() -> bytes:
    """Создать редактируемый пример с настоящими текстовыми заполнителями."""
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    contents = (
        ("Квартальный обзор", "Команда продукта\nИтоги работы и планы на следующий квартал"),
        ("Что изменилось", "Понятный первый запуск\nЕдиная база знаний\nОбратная связь клиентов"),
        ("Ключевые результаты", "Обновили знакомство с продуктом\nПодготовили справочный раздел\nПровели интервью с клиентами"),
        ("Следующий квартал", "Проверить вводный сценарий\nДополнить базу знаний\nСобрать обратную связь в продукте"),
        ("Главное", "Понятный путь к первому результату\nИзменения на основе обратной связи клиентов"),
    )
    for index, (title, text) in enumerate(contents, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F5F3EC")
        heading = slide.shapes.title
        heading.left, heading.top = Inches(0.8), Inches(1.05)
        heading.width, heading.height = Inches(11.6), Inches(1.2)
        heading.text = title
        body = slide.placeholders[1]
        body.left, body.top = Inches(0.85), Inches(2.8)
        body.width, body.height = Inches(10.9), Inches(3.25)
        body.text = text
        for shape, size in ((heading, 42), (body, 25)):
            for paragraph in shape.text_frame.paragraphs:
                paragraph.font.name = "Arial"
                paragraph.font.size = Pt(size)
                paragraph.font.color.rgb = RGBColor.from_string("252724")
                paragraph.font.bold = shape is heading
                paragraph.space_after = Pt(18)
        accent = slide.shapes.add_shape(1, Inches(0.85), Inches(0.65), Inches(0.48), Inches(0.055))
        accent.fill.solid()
        accent.fill.fore_color.rgb = RGBColor.from_string("D56739")
        accent.line.fill.background()
        footer = slide.shapes.add_textbox(Inches(11.55), Inches(6.8), Inches(0.8), Inches(0.3))
        footer.text_frame.paragraphs[0].text = f"{index:02d} / 05"
        footer.text_frame.paragraphs[0].font.size = Pt(10)
        footer.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string("74766E")
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


@dataclass
class Job:
    id: str
    directory: Path
    template_id: str
    max_slides: int | None
    status: str = "running"
    stage: str = "parser"
    error: str | None = None
    error_code: str | None = None
    slide_count: int | None = None
    progress: dict[str, Any] | None = None
    process: subprocess.Popen[str] | None = field(default=None, repr=False)
    presentation: dict[str, Any] | None = field(default=None, repr=False)

    def public(self) -> dict[str, Any]:
        result: dict[str, Any] = {"id": self.id, "status": self.status, "stage": self.stage}
        if self.error is not None:
            result["error"] = self.error
            if self.error_code in CONTENT_ERRORS:
                result["error_code"] = self.error_code
        if self.slide_count is not None:
            result["slide_count"] = self.slide_count
        if self.progress is not None:
            result["progress"] = dict(self.progress)
        return result

    def observe_output(self, line: str) -> None:
        """Извлечь только известные маркеры; тексты и ошибки LLM остаются приватными."""
        line = line.rstrip("\r\n")
        marker = line.split(" ", 1)[0]
        if marker in STAGE_MARKERS:
            stage = STAGE_MARKERS[marker]
            if stage != self.stage:
                self.stage = stage
                self.progress = None
                self.error_code = None
            return
        if self.stage != "content":
            return
        if line.startswith(CONTENT_ERROR_MARKER):
            code = line[len(CONTENT_ERROR_MARKER):]
            if code in CONTENT_ERRORS:
                self.error_code = code
            return
        log = CONTENT_LOG.fullmatch(line)
        if log is None:
            return
        logger, level, message = log.groups()
        if logger == "app.graph.fast":
            if level == "INFO" and message in FAST_CONTENT_PHASES:
                self.progress = {"phase": FAST_CONTENT_PHASES[message]}
            elif (
                level == "WARNING" and self.progress is not None
                and message == "Повтор пакетного запроса после проверки"
            ):
                self.progress = {**self.progress, "retrying": True}
            return
        if logger == "app.graph.nodes" and level == "INFO":
            if message in CONTENT_PHASES:
                self.progress = {"phase": CONTENT_PHASES[message]}
                return
            slide = re.fullmatch(
                r"--- Генерация контента для слайда ([1-9][0-9]{0,2})/([1-9][0-9]{0,2}) ---",
                message,
            )
            if slide is not None:
                current, total = map(int, slide.groups())
                if current <= total <= MAX_TEMPLATE_SLIDES:
                    self.progress = {"phase": "slides", "current": current, "total": total}
                return
        if self.progress is None:
            return
        phase = self.progress["phase"]
        retrying = False
        if logger == "app.graph.nodes":
            if level == "WARNING" and phase in PHASE_RETRIES:
                retry = re.fullmatch(PHASE_RETRIES[phase], message)
                retrying = retry is not None and (
                    phase != "slides" or int(retry[1]) <= MAX_TEMPLATE_SLIDES
                )
            elif level == "INFO" and phase == "validation":
                retrying = re.fullmatch(
                    r"Будет повторная генерация \(retry [1-5]\)", message,
                ) is not None
        elif level == "WARNING":
            model = {"analysis": "ScriptAnalysis", "planning": "SlidePlan"}.get(phase)
            prefixes = [
                r"Временная ошибка LLM HTTP (?:429|500|502|503|504), ",
                r"Временная ошибка OpenRouter HTTP (?:429|500|502|503|504), ",
                r"Временная ошибка GigaChat HTTP (?:500|502|503|504), ",
            ]
            if model is not None:
                prefixes.append(f"Некорректный ответ LLM для {model}, ")
            if phase == "slides":
                prefixes.append("Некорректный JSON-ответ LLM, ")
            for prefix in prefixes:
                suffix = r" через [1248] с" if prefix.startswith("Временная") else r": [^\r\n]*"
                retry = re.fullmatch(prefix + r"повтор ([1-5])/([1-5])" + suffix, message)
                if retry is not None and int(retry[1]) <= int(retry[2]):
                    retrying = True
                    break
        if retrying:
            self.progress = {**self.progress, "retrying": True}


@dataclass(frozen=True)
class PreviewTask:
    """Источник изображений: загруженный шаблон или готовый результат задания."""

    namespace: str
    source: Path
    metadata: dict[str, Any]


class Workspace:
    """Файлы и задания одной сессии локального сервера."""

    def __init__(
        self, *, renderer: PreviewRenderer | None = None,
        storage: FileServiceStorage | None = None, data_dir: Path | None = None,
    ) -> None:
        if (storage is None) != (data_dir is None):
            raise ValueError("Хранилище и постоянный каталог должны задаваться вместе")
        self.storage = storage
        self.catalog = None
        self.temporary = None
        self.renderer = None
        try:
            self._initialize(data_dir, renderer)
        except Exception:
            if self.renderer is not None:
                self.renderer.close()
            if self.temporary is not None:
                self.temporary.cleanup()
            if self.catalog is not None:
                self.catalog.close()
            raise

    def _initialize(self, data_dir: Path | None, renderer: PreviewRenderer | None) -> None:
        self.catalog = WebCatalog(data_dir) if data_dir is not None else None
        self.temporary = tempfile.TemporaryDirectory(prefix="exposlides-web-")
        self.root = Path(self.temporary.name)
        self.token = secrets.token_urlsafe(32)
        self.lock = threading.RLock()
        self.templates: dict[str, dict[str, Any]] = {}
        self.jobs: dict[str, Job] = {}
        self.workers: list[threading.Thread] = []
        self.example_id: str | None = None
        self.closed = False
        self.renderer = renderer if renderer is not None else PreviewRenderer()
        self.preview_files: dict[tuple[str, str], tuple[Path, ...]] = {}
        self.preview_queue: queue.Queue[PreviewTask | None] = queue.Queue()
        self.preview_worker: threading.Thread | None = None
        self._restored_previews: set[tuple[str, str]] = set()
        self._session_templates = 0
        self._session_jobs = 0
        if self.catalog is not None:
            self._restore_library()
        if self.renderer.available:
            self.preview_worker = threading.Thread(target=self._render_previews, daemon=True)
            self.preview_worker.start()

    def _restore_library(self) -> None:
        """Восстановить каталог; PPTX скачиваются только при использовании."""
        for identifier, record in self.catalog.library.templates.items():
            metadata = record.metadata.model_dump()
            metadata["preview"] = {"status": "unavailable", "message": PREVIEW_UNAVAILABLE}
            self.templates[identifier] = metadata
            self._restored_previews.add(("templates", identifier))
        for identifier, record in self.catalog.library.jobs.items():
            metadata = record.metadata.model_dump()
            metadata["preview"] = {"status": "unavailable", "message": RESULT_PREVIEW_UNAVAILABLE}
            directory = self.root / identifier
            directory.mkdir()
            self.jobs[identifier] = Job(
                identifier, directory, record.template_id, record.max_slides,
                status="completed", stage="complete", slide_count=metadata["slide_count"],
                presentation=metadata,
            )
            self._restored_previews.add(("jobs", identifier))

    def library(self) -> dict[str, Any]:
        with self.lock:
            return {
                "persistent": self.catalog is not None,
                "templates": [
                    {**metadata, "preview": dict(metadata["preview"])}
                    for metadata in reversed(list(self.templates.values()))
                ],
                "jobs": [
                    {**job.public(), "name": self.templates[job.template_id]["name"]}
                    for job in reversed(list(self.jobs.values())) if job.status == "completed"
                ],
            }

    def _cached_file(self, namespace: str, identifier: str) -> Path:
        """Локальная копия нужна CLI и предпросмотру; источник — версия в хранилище."""
        path = (
            self.root / f"{identifier}.pptx" if namespace == "templates"
            else self.jobs[identifier].directory / "result.pptx"
        )
        if not path.is_file() and self.catalog is not None:
            records = (
                self.catalog.library.templates if namespace == "templates"
                else self.catalog.library.jobs
            )
            record = records[identifier]
            data = self.storage.download(record.file.stored_file())
            inspect_template(data, record.metadata.name)
            self._write_cache(path, data)
        return path

    @staticmethod
    def _write_cache(path: Path, data: bytes) -> None:
        """Неполная локальная копия никогда не становится доступным PPTX."""
        temporary = None
        try:
            descriptor, name = tempfile.mkstemp(prefix=".download-", dir=path.parent)
            temporary = Path(name)
            with os.fdopen(descriptor, "wb") as stream:
                stream.write(data)
            os.replace(temporary, path)
        except OSError as error:
            raise StorageError("Не удалось записать локальную копию файла.") from error
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)

    def _restore_preview(self, namespace: str, identifier: str, metadata: dict) -> None:
        key = (namespace, identifier)
        if key in self._restored_previews and self.renderer.available:
            path = self._cached_file(namespace, identifier)
            self._queue_preview(namespace, path, metadata)
            self._restored_previews.remove(key)

    def download_result(self, job_id: str) -> bytes:
        with self.lock:
            job = self.get_job(job_id)
            if job.status != "completed":
                raise APIError("Презентация ещё не готова.", 409)
            return self._cached_file("jobs", job_id).read_bytes()

    def add_template(self, name: str, data: bytes) -> dict[str, Any]:
        metadata = inspect_template(data, name)
        template_id = uuid4().hex
        with self.lock:
            if self.closed:
                raise APIError("Сессия завершена. Перезапустите интерфейс.", 503)
            if self._session_templates >= 30:
                raise APIError("В этой сессии уже 30 шаблонов. Перезапустите интерфейс.", 409)
            metadata["id"] = template_id
            path = self.root / f"{template_id}.pptx"
            self._write_cache(path, data)
            try:
                if self.catalog is not None:
                    file = self.storage.upload(name, data)
                    self.catalog.save_template(template_id, metadata, file)
            except Exception:
                path.unlink(missing_ok=True)
                raise
            self.templates[template_id] = metadata
            self._session_templates += 1
            self._queue_preview("templates", path, metadata)
            return {**metadata, "preview": dict(metadata["preview"])}

    def _queue_preview(self, namespace: str, source: Path, metadata: dict[str, Any]) -> None:
        """Вызывается под блокировкой после сохранения и проверки исходного PPTX."""
        unavailable = (
            PREVIEW_UNAVAILABLE if namespace == "templates" else RESULT_PREVIEW_UNAVAILABLE
        )
        metadata["preview"] = (
            {"status": "pending"} if self.renderer.available
            else {"status": "unavailable", "message": unavailable}
        )
        if self.renderer.available:
            self.preview_queue.put(PreviewTask(namespace, source, metadata))

    def example(self) -> dict[str, Any]:
        with self.lock:
            if self.example_id is None:
                metadata = self.add_template("Квартальный обзор.pptx", make_example())
                self.example_id = metadata["id"]
            return {**self.templates[self.example_id], "script": EXAMPLE_SCRIPT}

    def get_preview(self, template_id: str) -> dict[str, Any]:
        with self.lock:
            if template_id not in self.templates:
                raise APIError("Шаблон не найден в этой сессии.", 404)
            self._restore_preview("templates", template_id, self.templates[template_id])
            return dict(self.templates[template_id]["preview"])

    def get_preview_image(self, template_id: str, index: int) -> bytes:
        with self.lock:
            preview = self.get_preview(template_id)
            if not 1 <= index <= self.templates[template_id]["slide_count"]:
                raise APIError("Слайд не найден.", 404)
            if preview["status"] != "ready":
                raise APIError("Предпросмотр слайда недоступен.", 409)
            return self.preview_files[("templates", template_id)][index - 1].read_bytes()

    def get_presentation(self, job_id: str) -> dict[str, Any]:
        with self.lock:
            job = self.get_job(job_id)
            if job.status != "completed" or job.presentation is None:
                raise APIError("Презентация ещё не готова.", 409)
            self._restore_preview("jobs", job_id, job.presentation)
            return {**job.presentation, "preview": dict(job.presentation["preview"])}

    def get_job_preview(self, job_id: str) -> dict[str, Any]:
        return self.get_presentation(job_id)["preview"]

    def get_job_preview_image(self, job_id: str, index: int) -> bytes:
        with self.lock:
            presentation = self.get_presentation(job_id)
            if not 1 <= index <= presentation["slide_count"]:
                raise APIError("Слайд не найден.", 404)
            if presentation["preview"]["status"] != "ready":
                raise APIError("Предпросмотр слайда недоступен.", 409)
            return self.preview_files[("jobs", job_id)][index - 1].read_bytes()

    def _render_previews(self) -> None:
        """Одна очередь не запускает несколько офисных конвертеров одновременно."""
        while True:
            task = self.preview_queue.get()
            try:
                if task is None:
                    return
                with self.lock:
                    if self.closed:
                        continue
                    identifier = task.metadata["id"]
                    count = task.metadata["slide_count"]
                directory = self.root / "previews" / task.namespace / identifier
                try:
                    directory.mkdir(parents=True)
                    images = self.renderer.render(task.source, directory, count)
                    resolved = tuple(image.resolve() for image in images)
                    if len(resolved) != count or len(set(resolved)) != count or any(
                        not image.is_relative_to(directory.resolve())
                        or not image.is_file() or image.suffix != ".png"
                        for image in resolved
                    ):
                        raise ValueError("Incomplete preview")
                    with self.lock:
                        if self.closed:
                            continue
                        self.preview_files[(task.namespace, identifier)] = resolved
                        task.metadata["preview"] = {
                            "status": "ready",
                            "slides": [
                                f"/api/{task.namespace}/{identifier}/slides/{index}.png"
                                for index in range(1, count + 1)
                            ],
                        }
                except Exception:
                    shutil.rmtree(directory, ignore_errors=True)
                    with self.lock:
                        task.metadata["preview"] = {
                            "status": "failed",
                            "message": (
                                PREVIEW_FAILED if task.namespace == "templates"
                                else RESULT_PREVIEW_FAILED
                            ),
                        }
            finally:
                self.preview_queue.task_done()

    def create_job(self, payload: dict[str, Any]) -> dict[str, Any]:
        template_id = payload.get("template_id")
        script = payload.get("script")
        max_slides = payload.get("max_slides")
        if not isinstance(template_id, str):
            raise APIError("Выберите PPTX-шаблон.")
        if not isinstance(script, str) or not script.strip():
            raise APIError("Добавьте текст для презентации.")
        if len(script) > MAX_SCRIPT_LENGTH:
            raise APIError("Текст слишком длинный. Допускается не больше 100 000 символов.")
        if max_slides is not None and (
            type(max_slides) is not int or not 1 <= max_slides <= MAX_TEMPLATE_SLIDES
        ):
            raise APIError("Количество слайдов должно быть целым числом от 1 до 250.")
        with self.lock:
            if self.closed:
                raise APIError("Сессия завершена. Перезапустите интерфейс.", 503)
            if template_id not in self.templates:
                raise APIError("Шаблон не найден. Загрузите его ещё раз.", 404)
            template_slide_count = self.templates[template_id]["slide_count"]
            if max_slides is not None and max_slides > template_slide_count:
                raise APIError(
                    f"В шаблоне только {template_slide_count} слайдов. "
                    "Уменьшите количество или выберите другой шаблон."
                )
            if self._session_jobs >= 50:
                raise APIError("В этой сессии уже 50 запусков. Перезапустите интерфейс.", 409)
            self._cached_file("templates", template_id)
            job_id = uuid4().hex
            directory = self.root / job_id
            directory.mkdir()
            (directory / "script.txt").write_text(script, encoding="utf-8")
            job = Job(job_id, directory, template_id, max_slides)
            self.jobs[job_id] = job
            self._session_jobs += 1
            worker = threading.Thread(target=self._run_job, args=(job,), daemon=True)
            self.workers.append(worker)
            worker.start()
            return job.public()

    def get_job(self, job_id: str) -> Job:
        with self.lock:
            if job_id not in self.jobs:
                raise APIError("Презентация не найдена в этой сессии.", 404)
            return self.jobs[job_id]

    def _complete_job(self, job: Job) -> None:
        """Открыть готовый файл и опубликовать результат без ожидания изображений."""
        result = job.directory / "result.pptx"
        metadata = inspect_template(result.read_bytes(), "Презентация.pptx")
        with self.lock:
            if self.closed:
                return
            metadata["id"] = job.id
            if self.catalog is not None:
                file = self.storage.upload("Презентация.pptx", result.read_bytes(), task_id=job.id)
                self.catalog.save_result(
                    job.id, metadata, file, template_id=job.template_id, max_slides=job.max_slides,
                )
            job.presentation = metadata
            job.slide_count = metadata["slide_count"]
            job.stage = "complete"
            job.progress = None
            job.status = "completed"
            self._queue_preview("jobs", result, metadata)

    def _run_job(self, job: Job) -> None:
        command = [
            sys.executable, "-m", "exposlides",
            "--template", str(self.root / f"{job.template_id}.pptx"),
            "--script", str(job.directory / "script.txt"),
            "--output", str(job.directory / "result.pptx"),
            "--generation-mode", "fast",
        ]
        if job.max_slides is not None:
            command.extend(["--max-slides", str(job.max_slides)])
        try:
            with self.lock:
                if self.closed:
                    return
                process = subprocess.Popen(
                    command,
                    cwd=REPOSITORY_ROOT,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    start_new_session=True,
                )
                job.process = process
            if process.stdout is not None:
                with (job.directory / "job.log").open("w", encoding="utf-8") as job_log:
                    for line in process.stdout:
                        job_log.write(line)
                        job_log.flush()
                        with self.lock:
                            job.observe_output(line)
                process.stdout.close()
            if process.wait() != 0:
                raise RuntimeError("pipeline failed")
            self._complete_job(job)
        except StorageError:
            with self.lock:
                job.error = STORAGE_SAVE_FAILED
                job.error_code = None
                job.status = "failed"
        except Exception:
            with self.lock:
                if job.stage == "content":
                    job.error_code = job.error_code or "unknown"
                    job.error = CONTENT_ERRORS[job.error_code]
                else:
                    job.error = STAGE_ERRORS.get(job.stage, STAGE_ERRORS["builder"])
                job.status = "failed"
        finally:
            with self.lock:
                job.process = None

    def close(self) -> None:
        with self.lock:
            self.closed = True
            processes = [job.process for job in self.jobs.values() if job.process is not None]
        self.preview_queue.put(None)
        self.renderer.close()
        if self.preview_worker is not None:
            self.preview_worker.join(timeout=5)
        for process in processes:
            try:
                # У завершившегося CLI ещё могут работать дочерние сервисы.
                os.killpg(process.pid, signal.SIGTERM)
            except ProcessLookupError:
                pass
        for worker in self.workers:
            worker.join(timeout=5)
        for process in processes:
            try:
                # Удалить оставшуюся группу, если сервис проигнорировал SIGTERM.
                os.killpg(process.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass
        for worker in self.workers:
            if worker.is_alive():
                worker.join(timeout=5)
        if not any(worker.is_alive() for worker in self.workers) and (
            self.preview_worker is None or not self.preview_worker.is_alive()
        ):
            self.temporary.cleanup()
            if self.catalog is not None:
                self.catalog.close()
            if self.storage is not None:
                self.storage.close()


class WebServer(ThreadingHTTPServer):
    daemon_threads = True

    def __init__(self, address: tuple[str, int], workspace: Workspace | None = None) -> None:
        self.workspace = workspace or Workspace()
        super().__init__(address, WebHandler)

    def server_close(self) -> None:
        super().server_close()
        self.workspace.close()


class WebHandler(BaseHTTPRequestHandler):
    server: WebServer

    def log_message(self, format: str, *args: Any) -> None:
        # Тексты запросов, локальные пути и ответы сервисов не попадают в HTTP-лог.
        return

    def _guard(self, *, write: bool = False) -> None:
        host = self.headers.get("Host", "")
        allowed = {f"127.0.0.1:{self.server.server_port}", f"localhost:{self.server.server_port}"}
        if host not in allowed:
            raise APIError("Этот интерфейс доступен только на локальном компьютере.", 403)
        origin = self.headers.get("Origin")
        if origin is not None and origin != f"http://{host}":
            raise APIError("Запрос с другого сайта отклонён.", 403)
        if write and not hmac.compare_digest(
            self.headers.get("X-ExpoSlides-Token", ""), self.server.workspace.token
        ):
            raise APIError("Сессия устарела. Обновите страницу.", 403)

    def _send(self, data: bytes, content_type: str, status: int = 200, **headers: str) -> None:
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "no-referrer")
        self.send_header(
            "Content-Security-Policy",
            "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; "
            "img-src 'self' data: blob:; connect-src 'self'; frame-ancestors 'none'; base-uri 'none'",
        )
        for key, value in headers.items():
            self.send_header(key.replace("_", "-"), value)
        self.end_headers()
        self.wfile.write(data)

    def _json(self, data: Any, status: int = 200) -> None:
        self._send(json.dumps(data, ensure_ascii=False).encode("utf-8"),
                   "application/json; charset=utf-8", status)

    def _body(self) -> dict[str, Any]:
        if self.headers.get_content_type() != "application/json":
            raise APIError("Ожидались данные в формате JSON.", 415)
        if self.headers.get("Transfer-Encoding") is not None:
            raise APIError("Неподдерживаемый формат передачи данных.")
        try:
            size = int(self.headers.get("Content-Length", "0"))
        except ValueError as error:
            raise APIError("Не удалось определить размер запроса.") from error
        if not 0 < size <= MAX_REQUEST_BYTES:
            raise APIError("Запрос слишком большой или пустой.", 413)
        try:
            payload = json.loads(self.rfile.read(size))
        except (UnicodeError, json.JSONDecodeError) as error:
            raise APIError("Не удалось прочитать данные запроса.") from error
        if not isinstance(payload, dict):
            raise APIError("Ожидался объект с данными запроса.")
        return payload

    def do_GET(self) -> None:
        try:
            self._guard()
            path = urlsplit(self.path).path
            if path in STATIC_ROUTES:
                filename, content_type = STATIC_ROUTES[path]
                self._send((STATIC_ROOT / filename).read_bytes(), content_type)
            elif path == "/api/session":
                self._json({"token": self.server.workspace.token})
            elif path == "/api/library":
                self._json(self.server.workspace.library())
            elif path == "/api/example":
                self._json(self.server.workspace.example())
            elif path.startswith("/api/templates/"):
                match = re.fullmatch(
                    r"/api/templates/([a-f0-9]{32})/(preview|slides/([1-9][0-9]{0,2})\.png)",
                    path,
                )
                if match is None:
                    raise APIError("Страница не найдена.", 404)
                template_id, route, index = match.groups()
                if route == "preview":
                    self._json(self.server.workspace.get_preview(template_id))
                else:
                    self._send(
                        self.server.workspace.get_preview_image(template_id, int(index)),
                        "image/png",
                    )
            elif path.startswith("/api/jobs/"):
                match = re.fullmatch(
                    r"/api/jobs/([a-f0-9]{32})"
                    r"(?:/(download|presentation|preview|slides/([1-9][0-9]{0,2})\.png))?",
                    path,
                )
                if match is None:
                    raise APIError("Страница не найдена.", 404)
                job_id, route, index = match.groups()
                job = self.server.workspace.get_job(job_id)
                with self.server.workspace.lock:
                    if route is None:
                        self._json(job.public())
                    elif route == "presentation":
                        self._json(self.server.workspace.get_presentation(job_id))
                    elif route == "preview":
                        self._json(self.server.workspace.get_job_preview(job_id))
                    elif index is not None:
                        self._send(
                            self.server.workspace.get_job_preview_image(job_id, int(index)),
                            "image/png",
                        )
                    elif job.status != "completed":
                        raise APIError("Презентация ещё не готова.", 409)
                    else:
                        self._send(
                            self.server.workspace.download_result(job_id),
                            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                            Content_Disposition='attachment; filename="ExpoSlides.pptx"',
                        )
            else:
                raise APIError("Страница не найдена.", 404)
        except StorageError:
            self._json({"error": STORAGE_UNAVAILABLE}, 503)
        except APIError as error:
            self._json({"error": str(error)}, error.status)
        except (BrokenPipeError, ConnectionResetError):
            pass
        except Exception:
            self._json({"error": "Не удалось открыть данные. Обновите страницу."}, 500)

    def do_POST(self) -> None:
        try:
            self._guard()
            path = urlsplit(self.path).path
            if path not in {"/api/templates", "/api/jobs"}:
                raise APIError("Страница не найдена.", 404)
            payload = self._body()
            # При старой сессии сначала дочитываем ограниченное тело загрузки:
            # ранний ответ закрывает соединение, пока браузер ещё отправляет файл.
            self._guard(write=True)
            if path == "/api/jobs":
                self._json(self.server.workspace.create_job(payload), 202)
                return
            name, encoded = payload.get("name"), payload.get("data")
            if not isinstance(name, str) or not name.lower().endswith(".pptx"):
                raise APIError("Выберите файл с расширением .pptx.")
            basename = name.replace("\\", "/").rsplit("/", 1)[-1]
            name = basename[:-5][:175] + ".pptx"
            if not isinstance(encoded, str):
                raise APIError("Не удалось прочитать загруженный файл.")
            try:
                data = base64.b64decode(encoded, validate=True)
            except (binascii.Error, ValueError) as error:
                raise APIError("Не удалось прочитать загруженный файл.") from error
            self._json(self.server.workspace.add_template(name, data), 201)
        except StorageError:
            self._json({"error": STORAGE_UNAVAILABLE}, 503)
        except APIError as error:
            self._json({"error": str(error)}, error.status)
        except (BrokenPipeError, ConnectionResetError):
            pass
        except Exception:
            self._json({"error": "Не удалось сохранить данные. Попробуйте ещё раз."}, 500)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Локальный веб-интерфейс ExpoSlides")
    parser.add_argument("--port", type=int, default=8765, help="Локальный порт (по умолчанию 8765)")
    parser.add_argument(
        "--file-service", default=os.environ.get("EXPOSLIDES_FILE_SERVICE"),
        help="Адрес file-service, например 127.0.0.1:50051",
    )
    parser.add_argument(
        "--data-dir", type=Path, default=os.environ.get("EXPOSLIDES_DATA_DIR"),
        help="Постоянный каталог библиотеки (вместе с --file-service)",
    )
    args = parser.parse_args(argv)
    if not 1 <= args.port <= 65535:
        parser.error("порт должен быть числом от 1 до 65535")
    if bool(args.file_service) != bool(args.data_dir):
        parser.error("--file-service и --data-dir должны задаваться вместе")
    storage = None
    workspace = None
    try:
        if args.file_service:
            storage = FileServiceStorage(args.file_service)
            workspace = Workspace(storage=storage, data_dir=args.data_dir.expanduser().resolve())
        server = WebServer(("127.0.0.1", args.port), workspace=workspace)
    except (OSError, StorageError, ValueError):
        if workspace is not None:
            workspace.close()
        elif storage is not None:
            storage.close()
        print(
            "Не удалось запустить интерфейс. Проверьте порт, адрес хранилища "
            "и постоянный каталог.", file=sys.stderr,
        )
        return 1
    print(f"ExpoSlides: http://127.0.0.1:{args.port}", flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
