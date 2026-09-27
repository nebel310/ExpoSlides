"""Редактирование подписей нативного SmartArt по совместимому образцу шаблона.

Новые topology/layout не синтезируются: число узлов и их связи сохраняются.
Неизвестное соответствие dataModel и кешированного drawing отклоняется явно.
"""

from __future__ import annotations

from copy import deepcopy
from typing import Any

from lxml import etree
from pptx.opc.package import XmlPart
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn

from exposlides.design_models import PlacedBlock
from exposlides.design_pptx_parts import NativeBuildError

DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
DSP = "http://schemas.microsoft.com/office/drawing/2008/diagram"


def capture_smartart(slide: Any, blocks: list[PlacedBlock]) -> dict[str, tuple[Any, int]]:
    requested = [block for block in blocks if block.kind == "smartart"]
    if not requested:
        return {}
    if len(requested) > 1:
        raise NativeBuildError("Пока поддерживается один SmartArt на выходном слайде")
    candidates = []

    def visit(shapes) -> None:
        for shape in shapes:
            if shape._element.find(f".//{{{DGM}}}relIds") is not None:
                candidates.append(shape)
            if hasattr(shape, "shapes"):
                visit(shape.shapes)

    visit(slide.shapes)
    block = requested[0]
    if block.source_shape_id is not None:
        candidates = [shape for shape in candidates if shape.shape_id == block.source_shape_id]
    if len(candidates) != 1:
        raise NativeBuildError(
            "Для SmartArt нужен один однозначно выбранный нативный образец в исходном слайде"
        )
    source = candidates[0]
    return {block.id: (deepcopy(source._element), source.shape_id)}


def _replace_text(element: Any, text: str) -> None:
    texts = list(element.iter(qn("a:t")))
    if not texts:
        raise NativeBuildError("В узле SmartArt отсутствует редактируемый текст")
    texts[0].text = text
    for part in texts[1:]:
        part.text = ""


def _store_xml(part: Any, xml: Any) -> None:
    data = etree.tostring(xml, xml_declaration=True, encoding="UTF-8", standalone=True)
    if isinstance(part, XmlPart):
        part._element = parse_xml(data)
    else:
        part.blob = data


def add_smartart(slide: Any, block: PlacedBlock, frame: Any) -> None:
    rel_ids = frame.find(f".//{{{DGM}}}relIds")
    rid = rel_ids.get(qn("r:dm")) if rel_ids is not None else None
    if rid is None:
        raise NativeBuildError("SmartArt не содержит ссылки на dataModel")
    data_part = slide.part.related_part(rid)
    data_xml = etree.fromstring(data_part.blob)
    points = list(data_xml.iter(f"{{{DGM}}}pt"))
    nodes = [
        point
        for point in points
        if point.get("type", "node") == "node"
        and any((text.text or "").strip() for text in point.iter(qn("a:t")))
    ]
    if len(nodes) != len(block.items):
        raise NativeBuildError(
            f"SmartArt сохраняет число узлов образца: нужно {len(nodes)}, получено {len(block.items)}"
        )
    labels = {node.get("modelId"): label for node, label in zip(nodes, block.items, strict=True)}
    if None in labels or len(labels) != len(nodes):
        raise NativeBuildError("SmartArt содержит неоднозначные идентификаторы узлов")
    aliases = {identifier: identifier for identifier in labels}
    for connection in data_xml.iter(f"{{{DGM}}}cxn"):
        if connection.get("type") == "presOf" and connection.get("srcId") in labels:
            aliases[connection.get("destId")] = connection.get("srcId")
    for point in points:
        properties = point.find(f"{{{DGM}}}prSet")
        if properties is not None and properties.get("presAssocID") in labels:
            aliases[point.get("modelId")] = properties.get("presAssocID")
    drawings = [
        relationship.target_part
        for relationship in data_part.rels.values()
        if not relationship.is_external and relationship.reltype.endswith("/diagramDrawing")
    ]
    if not drawings:
        raise NativeBuildError("SmartArt без кешированного drawing пока не поддерживается")
    updates = []
    matched = set()
    for drawing in drawings:
        xml = etree.fromstring(drawing.blob)
        for shape in xml.iter(f"{{{DSP}}}sp"):
            texts = list(shape.iter(qn("a:t")))
            if not any((text.text or "").strip() for text in texts):
                continue
            node_id = aliases.get(shape.get("modelId"))
            if node_id not in labels:
                raise NativeBuildError("Не удалось сопоставить подпись drawing с узлом SmartArt")
            _replace_text(shape, labels[node_id])
            matched.add(node_id)
        updates.append((drawing, xml))
    if matched != set(labels):
        raise NativeBuildError("Кешированное изображение SmartArt содержит не все узлы dataModel")
    for point in points:
        node_id = aliases.get(point.get("modelId"))
        if node_id in labels and list(point.iter(qn("a:t"))):
            _replace_text(point, labels[node_id])
    _store_xml(data_part, data_xml)
    for drawing, xml in updates:
        _store_xml(drawing, xml)
    properties = frame.find(".//" + qn("p:cNvPr"))
    if properties is None:
        raise NativeBuildError("Некорректная рамка SmartArt")
    properties.set("id", str(slide.shapes._next_shape_id))
    properties.set("name", "exposlides:" + block.id)
    slide.shapes._spTree.insert_element_before(frame, "p:extLst")
    shape = slide.shapes._shape_factory(frame)
    shape.left, shape.top = block.box.left, block.box.top
    shape.width, shape.height = block.box.width, block.box.height
