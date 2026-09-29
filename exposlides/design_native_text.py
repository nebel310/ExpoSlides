"""Заполнение существующих текстовых фигур с сохранением нативного оформления."""

from __future__ import annotations

import re
from collections import Counter
from copy import deepcopy
from typing import Any

from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

from exposlides.design_models import PlacedBlock
from exposlides.design_pptx_parts import NativeBuildError

TEXT_KINDS = {"title", "text", "page_number"}
RUN_TAGS = {qn("a:r"), qn("a:fld")}
BODY_TYPES = {"BODY", "SUBTITLE", "OBJECT", "CHART", "TABLE", "BITMAP", "ORG_CHART",
              "MEDIA_CLIP", "PICTURE"}
TITLE_TYPES = {"TITLE", "CENTER_TITLE"}


def native_shapes(shapes):
    """ID фигуры уникален внутри слайда, в том числе для потомков групп."""
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from native_shapes(shape.shapes)


def bound_text_shapes(slide, blocks: list[PlacedBlock]) -> dict[str, Any]:
    sources = {shape.shape_id: shape for shape in native_shapes(slide.shapes)}
    result, used = {}, set()
    for block in blocks:
        if block.kind not in TEXT_KINDS or block.source_shape_id is None:
            continue
        if block.source_shape_id in used:
            raise NativeBuildError("Одна исходная текстовая фигура назначена нескольким блокам")
        shape = sources.get(block.source_shape_id)
        if shape is None or not shape.has_text_frame:
            raise NativeBuildError(f"Не найдена исходная текстовая фигура {block.source_shape_id}")
        result[block.id] = shape
        used.add(block.source_shape_id)
    return result


def retained_shape_ids(shapes) -> set[int]:
    """Удаление группы недопустимо, если в ней заполняется исходный текстовый слот."""
    result = set()
    for shape in shapes:
        result.add(shape.shape_id)
        parent = shape._element.getparent()
        while parent is not None and parent.tag == qn("p:grpSp"):
            properties = parent.find(qn("p:nvGrpSpPr"))
            identity = properties.find(qn("p:cNvPr")) if properties is not None else None
            if identity is not None:
                result.add(int(identity.get("id")))
            parent = parent.getparent()
    return result


def _list_properties(style, level):
    if style is not None:
        for name in (f"a:lvl{level+1}pPr", "a:defPPr"):
            properties = style.find(qn(name))
            if properties is not None:
                yield properties


def _paragraph_level(paragraph):
    # Paragraph.level создаёт pPr при чтении; layout/master должны оставаться неизменными.
    properties = paragraph._p.find(qn("a:pPr"))
    return int(properties.get("lvl", "0")) if properties is not None else 0


def _shape_properties(shape, level):
    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            if _paragraph_level(paragraph) == level:
                properties = paragraph._p.find(qn("a:pPr"))
                if properties is not None:
                    yield properties
                break
        yield from _list_properties(shape.text_frame._txBody.find(qn("a:lstStyle")), level)


def _inherited_paragraph_properties(shape, paragraph):
    level = _paragraph_level(paragraph)
    properties = paragraph._p.find(qn("a:pPr"))
    if properties is not None:
        yield properties
    yield from _list_properties(shape.text_frame._txBody.find(qn("a:lstStyle")), level)
    slide = shape.part.slide
    master = slide.slide_layout.slide_master
    kind = shape.placeholder_format.type.name if shape.is_placeholder else None
    if shape.is_placeholder:
        layout_shape = slide.slide_layout.placeholders.get(idx=shape.placeholder_format.idx)
        if layout_shape is not None:
            yield from _shape_properties(layout_shape, level)
            kind = layout_shape.placeholder_format.type.name
        master_kind = "TITLE" if kind in TITLE_TYPES else "BODY" if kind in BODY_TYPES else kind
        for master_shape in master.placeholders:
            if master_shape.placeholder_format.type.name == master_kind:
                yield from _shape_properties(master_shape, level)
                break
    master_styles = master._element.find(qn("p:txStyles"))
    style_name = "p:titleStyle" if kind in TITLE_TYPES else "p:bodyStyle" if kind in BODY_TYPES else "p:otherStyle"
    if master_styles is not None:
        yield from _list_properties(master_styles.find(qn(style_name)), level)
    presentation = slide.part.package.presentation_part.presentation
    yield from _list_properties(presentation._element.find(qn("p:defaultTextStyle")), level)


def _inherited_defaults(shape, paragraph):
    for properties in _inherited_paragraph_properties(shape, paragraph):
        defaults = properties.find(qn("a:defRPr"))
        if defaults is not None:
            yield defaults


def native_plain_text_margin(shape, paragraph) -> int | None:
    """Убирает остаточный выступ без маркера, сохраняя начало первой строки.

    Настоящие списки, положительный абзацный отступ и нелевое выравнивание
    остаются как в шаблоне. Чтение наследуемых свойств не меняет layout/master.
    """
    attributes, bullet = {}, None
    bullet_tags = {qn(name) for name in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buBlip")}
    for properties in _inherited_paragraph_properties(shape, paragraph):
        for name in ("marL", "indent", "algn"):
            if name not in attributes and properties.get(name) is not None:
                attributes[name] = properties.get(name)
        if bullet is None:
            bullet = next((child.tag for child in properties if child.tag in bullet_tags), None)
    indent, margin = int(attributes.get("indent", "0")), int(attributes.get("marL", "0"))
    if (bullet == qn("a:buNone") and attributes.get("algn", "l") == "l"
            and indent < 0 and margin + indent >= 0):
        return margin + indent
    return None


def native_run_size(shape, paragraph, run=None) -> float | None:
    """Размер по цепочке run → paragraph/list → layout/master → presentation."""
    element = getattr(run, "_r", run)
    if element is not None:
        properties = element.find(qn("a:rPr"))
        if properties is not None and properties.get("sz") is not None:
            return int(properties.get("sz"))/100
    else:
        end = paragraph._p.find(qn("a:endParaRPr"))
        if end is not None and end.get("sz") is not None:
            return int(end.get("sz"))/100
    for properties in _inherited_defaults(shape, paragraph):
        if properties.get("sz") is not None:
            return int(properties.get("sz"))/100
    return None


def native_reference_size(shape, paragraph_count: int | None = None) -> float | None:
    """Размер исходных прототипов, которые действительно получат новый текст.

    Без количества сохраняется прежний доминирующий размер всей фигуры.
    Для заполнения используется максимум доминирующих размеров выбранных
    абзацев, совпадающий с базовым размером в плане; дополнительные абзацы
    повторяют последний исходный прототип и максимум не меняют.
    """
    if paragraph_count is not None and paragraph_count < 1:
        raise ValueError("Количество заполняемых абзацев должно быть положительным")
    paragraphs = list(shape.text_frame.paragraphs)
    if paragraph_count is not None:
        paragraphs = paragraphs[:paragraph_count]
    weights, paragraph_sizes = Counter(), []
    for paragraph in paragraphs:
        local_weights = Counter()
        runs = [child for child in paragraph._p if child.tag in RUN_TAGS]
        for run in runs or [None]:
            size = native_run_size(shape, paragraph, run)
            if size is not None:
                text = "".join(node.text or "" for node in run.iter(qn("a:t"))) if run is not None else ""
                weight = max(1, len(text))
                weights[size] += weight
                local_weights[size] += weight
        if local_weights:
            paragraph_sizes.append(local_weights.most_common(1)[0][0])
    if paragraph_count is not None:
        return max(paragraph_sizes) if paragraph_sizes else None
    return weights.most_common(1)[0][0] if weights else None


def _segments(text: str, originals: list) -> list[str]:
    weights = [len("".join(node.text or "" for node in run.iter(qn("a:t")))) for run in originals]
    if not any(weights):
        weights = [1, *[0]*(len(originals)-1)]
    words = sorted({0, len(text), *[match.end() for match in re.finditer(r"\s+", text)]})
    boundaries, cumulative = [0], 0
    for weight in weights[:-1]:
        cumulative += weight
        target = len(text)*cumulative/sum(weights)
        available = [position for position in words if position >= boundaries[-1]]
        boundaries.append(min(available, key=lambda position: (abs(position-target), position)))
    boundaries.append(len(text))
    return [text[start:end] for start, end in zip(boundaries, boundaries[1:])]


def _scaled_sizes(element, ratio):
    if element is not None and ratio < 1:
        for properties in element.iter():
            if properties.get("sz") is not None:
                properties.set("sz", str(max(100, round(int(properties.get("sz"))*ratio))))


def fill_native_text(shape, block: PlacedBlock) -> None:
    """Заполняет txBody, убирает выступ без маркера и при необходимости уменьшает кегль."""
    frame = shape.text_frame
    prototypes = list(frame.paragraphs)
    lines = [line for item in (block.items or [block.text]) for line in item.splitlines()] or [""]
    reference = native_reference_size(shape, paragraph_count=len(lines))
    ratio = min(1, block.style.size/reference) if reference else 1
    replacements = []
    for index, text in enumerate(lines):
        source = prototypes[min(index, len(prototypes)-1)]
        paragraph = deepcopy(source._p)
        margin = native_plain_text_margin(shape, source) if block.kind == "text" else None
        if margin is not None:
            properties = paragraph.get_or_add_pPr()
            properties.set("marL", str(margin))
            properties.set("indent", "0")
        originals = [child for child in source._p if child.tag in RUN_TAGS]
        sizes = [native_run_size(shape, source, run) for run in originals]
        if not originals:
            originals = [OxmlElement("a:r")]
            sizes = [native_run_size(shape, source)]
        for child in list(paragraph):
            if child.tag in RUN_TAGS or child.tag == qn("a:br"):
                paragraph.remove(child)
        _scaled_sizes(paragraph.find(qn("a:pPr")), ratio)
        _scaled_sizes(paragraph.find(qn("a:endParaRPr")), ratio)
        for original, segment, size in zip(originals, _segments(text, originals), sizes, strict=True):
            run = deepcopy(original)
            texts = list(run.iter(qn("a:t")))
            if texts:
                texts[0].text = segment
                for part in texts[1:]:
                    part.text = ""
            else:
                content = OxmlElement("a:t")
                content.text = segment
                run.append(content)
            if ratio < 1 and size is not None:
                properties = run.find(qn("a:rPr"))
                if properties is None:
                    properties = OxmlElement("a:rPr")
                    run.insert(0, properties)
                properties.set("sz", str(max(100, round(size*100*ratio))))
            paragraph.insert_element_before(run, "a:endParaRPr", "a:extLst")
        replacements.append(paragraph)
    for paragraph in list(frame._txBody):
        if paragraph.tag == qn("a:p"):
            frame._txBody.remove(paragraph)
    for paragraph in replacements:
        frame._txBody.append(paragraph)
