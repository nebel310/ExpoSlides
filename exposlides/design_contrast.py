"""Проверка цвета заполненного текста относительно фактического фона PPTX."""

from __future__ import annotations

from collections.abc import Callable
from typing import Any

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

from exposlides.design_native_text import _inherited_defaults

_FILL_TAGS = {qn("a:" + name) for name in
              ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill")}


def _theme(slide: Any) -> tuple[Any, dict[str, str]]:
    master = slide.slide_layout.slide_master
    try:
        theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
    except KeyError:
        theme = None
    mapping = master._element.find(qn("p:clrMap"))
    colors = dict(mapping.attrib) if mapping is not None else {}
    for owner in (slide.slide_layout, slide):
        override = owner._element.find("./" + qn("p:clrMapOvr") + "/" + qn("a:overrideClrMapping"))
        if override is not None:
            colors.update(override.attrib)
    return theme, colors


def _color(node: Any, theme: Any, mapping: dict[str, str]) -> str | None:
    if node is None:
        return None
    if node.tag == qn("a:schemeClr"):
        token = mapping.get(node.get("val"), node.get("val"))
        source = theme.find(".//" + qn("a:clrScheme") + "/" + qn("a:" + token)) if theme is not None else None
        value = _color(source[0], theme, {}) if source is not None and len(source) else None
    elif node.tag == qn("a:srgbClr"):
        value = node.get("val")
    elif node.tag == qn("a:sysClr"):
        value = node.get("lastClr")
    else:
        return None
    if not value or len(value) != 6:
        return None
    try:
        channels = [int(value[i:i+2], 16) / 255 for i in (0, 2, 4)]
    except ValueError:
        return None
    for transform in node:
        amount = int(transform.get("val", "0")) / 100000
        if transform.tag == qn("a:tint"):
            channels = [v + (1-v)*amount for v in channels]
        elif transform.tag == qn("a:shade"):
            channels = [v*amount for v in channels]
        elif transform.tag == qn("a:alpha") and amount == 1:
            continue
        else:
            # Не подменяем прозрачный/сложный цвет приблизительным RGB.
            return None
    return "".join(f"{round(max(0, min(1, v))*255):02X}" for v in channels)


def _fill(properties: Any, theme: Any, mapping: dict[str, str]) -> tuple[bool, str | None]:
    if properties is not None:
        for child in properties:
            if child.tag == qn("a:noFill"):
                return False, None
            if child.tag in _FILL_TAGS:
                color = _color(child[0], theme, mapping) if child.tag == qn("a:solidFill") and len(child) else None
                return True, color
    return False, None


def _shape_fill(shape: Any, theme: Any, mapping: dict[str, str]) -> tuple[bool, str | None]:
    properties = shape._element.find(qn("p:spPr"))
    if properties is not None and any(child.tag in _FILL_TAGS for child in properties):
        return _fill(properties, theme, mapping)
    if shape.shape_type == 13 or shape.has_chart or shape.has_table or hasattr(shape, "shapes"):
        return True, None
    reference = shape._element.find("./" + qn("p:style") + "/" + qn("a:fillRef"))
    if reference is not None and reference.get("idx", "0") != "0":
        # Стиль может задавать градиент или эффекты: используем безопасную подложку.
        return True, None
    if shape.is_placeholder:
        layout = shape.part.slide.slide_layout.placeholders.get(idx=shape.placeholder_format.idx)
        if layout is not None and layout._element is not shape._element:
            return _fill(layout._element.find(qn("p:spPr")), theme, mapping)
    return False, None


def _background(slide: Any, theme: Any, mapping: dict[str, str]) -> str | None:
    for owner in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        background = owner._element.find("./" + qn("p:cSld") + "/" + qn("p:bg"))
        if background is None:
            continue
        properties = background.find(qn("p:bgPr"))
        if properties is not None:
            return _fill(properties, theme, mapping)[1]
        reference = background.find(qn("p:bgRef"))
        if reference is not None and theme is not None:
            index = int(reference.get("idx", "0")) - 1001
            styles = theme.find(".//" + qn("a:bgFillStyleLst"))
            if styles is not None and 0 <= index < len(styles):
                fill = styles[index]
                if fill.tag == qn("a:solidFill") and len(fill):
                    color = fill[0]
                    if color.tag == qn("a:schemeClr") and color.get("val") == "phClr":
                        if len(color) == 0 and len(reference):
                            return _color(reference[0], theme, mapping)
                    return _color(color, theme, mapping)
        return None
    return "FFFFFF"


def _behind(slide: Any, target: Any, theme: Any, mapping: dict[str, str]) -> str | None:
    present, color = _shape_fill(target, theme, mapping)
    if present:
        return color
    # Координаты внутри трансформированной группы нельзя сравнивать с координатами слайда.
    if target._element.getparent().tag == qn("p:grpSp"):
        return None
    color = _background(slide, theme, mapping)
    for owner in (slide.slide_layout.slide_master, slide.slide_layout, slide):
        if owner is not slide and slide._element.get("showMasterSp") == "0":
            continue
        for shape in owner.shapes:
            if shape._element is target._element:
                return color
            if owner is not slide and shape.is_placeholder:
                continue
            if (shape.left >= target.left + target.width or target.left >= shape.left + shape.width
                    or shape.top >= target.top + target.height or target.top >= shape.top + shape.height):
                continue
            present, fill = _shape_fill(shape, theme, mapping)
            if not present:
                continue
            contains = (shape.left <= target.left and shape.top <= target.top
                        and shape.left + shape.width >= target.left + target.width
                        and shape.top + shape.height >= target.top + target.height)
            geometry = shape._element.find("./" + qn("p:spPr") + "/" + qn("a:prstGeom"))
            rectangular = geometry is None or geometry.get("prst") == "rect"
            color = fill if contains and rectangular and not shape.rotation else None
    return color


def repair_text_contrast(slide: Any, shape: Any, fallback: str,
                         choose: Callable[[str, str], str]) -> None:
    """Исправляет только нечитаемые runs; сложный фон закрывает белой подложкой."""
    theme, mapping = _theme(slide)
    background = _behind(slide, shape, theme, mapping)
    if background is None:
        background = "FFFFFF"
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(background)
    for paragraph in shape.text_frame.paragraphs:
        # Включая поля номера страницы, не представленные в paragraph.runs.
        for run in paragraph._p:
            if run.tag not in {qn("a:r"), qn("a:fld")}:
                continue
            properties = run.find(qn("a:rPr"))
            color = None
            for defaults in [properties, *_inherited_defaults(shape, paragraph)]:
                present, resolved = _fill(defaults, theme, mapping)
                if present:
                    color = resolved
                    break
            original = color or fallback
            corrected = choose(background, original)
            if color is not None and corrected == color:
                continue
            if properties is None:
                properties = OxmlElement("a:rPr")
                run.insert(0, properties)
            for child in list(properties):
                if child.tag in _FILL_TAGS:
                    properties.remove(child)
            fill = OxmlElement("a:solidFill")
            rgb = OxmlElement("a:srgbClr")
            rgb.set("val", corrected)
            fill.append(rgb)
            properties.insert_element_before(fill, "a:effectLst", "a:effectDag", "a:highlight",
                                             "a:uLnTx", "a:uLn", "a:uFillTx", "a:uFill",
                                             "a:latin", "a:ea", "a:cs", "a:sym", "a:hlinkClick",
                                             "a:hlinkMouseOver", "a:rtl", "a:extLst")
