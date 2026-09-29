"""Изолированные операции OPC для независимых копий слайдов внутри шаблона."""

from __future__ import annotations

from pathlib import PurePosixPath
from typing import Any

from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.oxml import serialize_part_xml
from pptx.opc.package import XmlPart
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement


class NativeBuildError(ValueError):
    """Шаблон или план нельзя безопасно собрать в редактируемую презентацию."""


def _prune_source_navigation(presentation: Any) -> None:
    """Убрать навигацию исходной колоды, не подходящую новым копиям слайдов.

    Показы, диапазоны показа и разделы описывают исходные слайды. Новая колода
    может повторять образцы и менять порядок, поэтому начинается с полного показа
    без исходного разбиения. Сам загруженный PPTX при этом не меняется.
    """
    xml = presentation.part._element
    shows = xml.find(qn("p:custShowLst"))
    if shows is not None:
        xml.remove(shows)
    section_tag = "{http://schemas.microsoft.com/office/powerpoint/2010/main}sectionLst"
    for sections in list(xml.iter(section_tag)):
        parent = sections.getparent()
        parent.remove(sections)
        # Другие расширения, в том числе соседи списка разделов, сохраняются.
        if parent.tag == qn("p:ext") and len(parent) == 0:
            extensions = parent.getparent()
            extensions.remove(parent)
            if extensions.tag == qn("p:extLst") and len(extensions) == 0:
                extensions.getparent().remove(extensions)

    try:
        properties = presentation.part.part_related_by(RT.PRES_PROPS)
    except KeyError:
        return
    properties_xml = parse_xml(properties.blob)
    show = properties_xml.find(qn("p:showPr"))
    if show is None:
        return
    selectors = [child for child in show if child.tag in {qn("p:custShow"), qn("p:sldRg")}]
    if not selectors:
        return
    position = min(show.index(child) for child in selectors)
    for selector in selectors:
        show.remove(selector)
    if show.find(qn("p:sldAll")) is None:
        show.insert(position, OxmlElement("p:sldAll"))
    if isinstance(properties, XmlPart):
        properties._element = properties_xml
    else:
        properties.blob = serialize_part_xml(properties_xml)


def _prune_custom_show_actions(presentation: Any) -> None:
    """Ссылки на удалённые показы убираются и в копиях, и в общих макетах."""
    for part in presentation.part.package.iter_parts():
        if isinstance(part, XmlPart):
            for element in list(part._element.iter(qn("a:hlinkClick"), qn("a:hlinkHover"))):
                action = element.get("action", "").partition("?")[0].lower()
                if action == "ppaction://customshow":
                    element.getparent().remove(element)


class _PartCloner:
    def __init__(self, presentation: Any, selected_parts: list[Any]) -> None:
        self.package = presentation.part.package
        self.used_names = {str(part.partname) for part in self.package.iter_parts()}
        self.primary: dict[Any, Any] = {}
        self.clones = []
        for source in selected_parts:
            clone = self._empty_clone(source)
            self.clones.append(clone)
            self.primary.setdefault(source, clone)

    def _empty_clone(self, source: Any) -> Any:
        path = PurePosixPath(str(source.partname))
        counter = 1
        while True:
            candidate = str(path.with_name(f"{path.stem}-exposlides-{counter}{path.suffix}"))
            if candidate not in self.used_names:
                break
            counter += 1
        self.used_names.add(candidate)
        return type(source).load(PackURI(candidate), source.content_type, self.package, source.blob)

    def _copy_relationships(self, source: Any, clone: Any, memo: dict[Any, Any]) -> None:
        targets = {}
        for relationship in source.rels.values():
            if relationship.is_external:
                continue
            target = relationship.target_part
            if relationship.reltype in {
                RT.SLIDE_LAYOUT,
                RT.SLIDE_MASTER,
                RT.NOTES_MASTER,
                RT.THEME,
                RT.IMAGE,
            }:
                copied = target
            elif target in memo:
                copied = memo[target]
            elif relationship.reltype == RT.SLIDE:
                copied = self.primary.get(target)
                if copied is None:
                    raise NativeBuildError(
                        "Шаблон содержит внутреннюю ссылку на слайд, отсутствующий в плане"
                    )
            else:
                # У каждого экземпляра собственные chart/workbook/SmartArt/notes parts.
                copied = self._empty_clone(target)
                memo[target] = copied
                self._copy_relationships(target, copied, memo)
            targets[target.partname] = copied
        clone.load_rels_from_xml(parse_xml(source.rels.xml), targets)
        if len(clone.rels) != len(source.rels):
            raise NativeBuildError("При клонировании потеряна связь между частями PPTX")

    def populate(self, sources: list[Any]) -> list[Any]:
        for source, clone in zip(sources, self.clones, strict=True):
            self._copy_relationships(source, clone, {source: clone})
        return self.clones


def clone_selected_slides(presentation: Any, source_indices: list[int]) -> list[Any]:
    """Заменяет список слайдов независимыми копиями в заданном порядке."""
    sources = list(presentation.slides)
    if any(index < 1 or index > len(sources) for index in source_indices):
        raise NativeBuildError("План ссылается на отсутствующий слайд шаблона")
    selected = [sources[index - 1].part for index in source_indices]
    clones = _PartCloner(presentation, selected).populate(selected)
    _prune_source_navigation(presentation)
    slide_ids = presentation.slides._sldIdLst
    old_ids = list(slide_ids)
    for element in old_ids:
        # drop_rel допускает одну оставшуюся ссылку: проверяем до удаления записи,
        # чтобы неизвестное расширение не получило переиспользованный rId копии.
        references = sum(
            node.get(qn("r:id")) == element.rId
            for node in presentation.part._element.iter()
        )
        if references > 1:
            raise NativeBuildError("PPTX содержит дополнительные ссылки на исходный слайд")
        slide_ids.remove(element)
        presentation.part.drop_rel(element.rId)
        if element.rId in presentation.part.rels:
            raise NativeBuildError("PPTX содержит дополнительные ссылки на исходный слайд")
    for clone in clones:
        rid = presentation.part.relate_to(clone, RT.SLIDE)
        slide_ids.add_sldId(rid)
        # Тайминги могут ссылаться на удаляемые фигуры. Новая колода статическая.
        timing = clone._element.find(qn("p:timing"))
        if timing is not None:
            clone._element.remove(timing)
    _prune_custom_show_actions(presentation)
    return [part.slide for part in clones]


def remove_shapes(slide: Any, shape_ids: list[int]) -> None:
    """Удаляет только явно выбранные фигуры, сохраняя соседей во вложенных группах."""
    requested = set(shape_ids)
    found: set[int] = set()

    def visit(shapes) -> None:
        for shape in list(shapes):
            if shape.shape_id in requested:
                found.add(shape.shape_id)
                # Также учитываем явно перечисленных потомков удаляемой группы.
                if hasattr(shape, "shapes"):
                    collect(shape.shapes)
                shape._element.getparent().remove(shape._element)
            elif hasattr(shape, "shapes"):
                visit(shape.shapes)

    def collect(shapes) -> None:
        for shape in shapes:
            if shape.shape_id in requested:
                found.add(shape.shape_id)
            if hasattr(shape, "shapes"):
                collect(shape.shapes)

    visit(slide.shapes)
    missing = requested - found
    if missing:
        raise NativeBuildError(f"Не найдены заменяемые фигуры: {sorted(missing)}")


def isolate_slide_layout(presentation: Any, slide: Any) -> Any:
    """Собственный layout/master для замены фото без изменения соседних слайдов."""
    from pptx.oxml.xmlchemy import OxmlElement

    layout = slide.slide_layout.part
    master = slide.slide_layout.slide_master.part
    cloner = _PartCloner(presentation, [layout, master])
    new_layout, new_master = cloner.populate([layout, master])

    def relink(source, target, replacements):
        targets = {rel.target_part.partname: replacements.get(rel.target_part, rel.target_part)
                   for rel in source.rels.values() if not rel.is_external}
        target.load_rels_from_xml(parse_xml(source.rels.xml), targets)

    relink(layout, new_layout, {master: new_master})
    relink(master, new_master, {layout: new_layout})
    # В новом master регистрируем только собственный layout; тема и декор прежние.
    listing = new_master._element.find(qn("p:sldLayoutIdLst"))
    if listing is None:
        raise NativeBuildError("У мастера отсутствует реестр макетов")
    for entry in list(listing):
        rid = entry.get(qn("r:id"))
        if new_master.related_part(rid) is not new_layout:
            listing.remove(entry)
            new_master.drop_rel(rid)
    if len(listing) != 1:
        raise NativeBuildError("Исходный макет не зарегистрирован в мастере")
    relink(slide.part, slide.part, {layout: new_layout})
    rid = presentation.part.relate_to(new_master, RT.SLIDE_MASTER)
    masters = presentation.part._element.get_or_add_sldMasterIdLst()
    entry = OxmlElement("p:sldMasterId")
    entry.set("id", str(max([int(e.get("id", 2147483647)) for e in masters]
                            + [2147483647]) + 1))
    entry.set(qn("r:id"), rid)
    masters.append(entry)
    return new_layout.slide_layout
