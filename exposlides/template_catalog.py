"""Рабочая копия с образцами из неиспользованных нативных layout шаблона."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _empty_panel(layout, width: int, height: int) -> bool:
    """Большой пустой прямоугольник вне текста — незаполненная область иллюстрации."""
    text = [shape for shape in layout.placeholders
            if shape.placeholder_format.type.name in {"TITLE", "CENTER_TITLE", "BODY", "SUBTITLE", "OBJECT"}]
    for shape in layout.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or shape.is_placeholder:
            continue
        if shape.has_text_frame and shape.text.strip():
            continue
        geometry = shape._element.find(".//" + qn("a:prstGeom"))
        if geometry is None or geometry.get("prst") != "rect":
            continue
        if shape.width * shape.height < width * height * .25:
            continue
        if not any(max(shape.left, slot.left) < min(shape.left + shape.width, slot.left + slot.width)
                   and max(shape.top, slot.top) < min(shape.top + shape.height, slot.top + slot.height)
                   for slot in text):
            return True
    return False


def prepare_template_catalog(source: Path, directory: Path) -> Path:
    """Добавляет только реальные текстовые макеты; загруженный оригинал неизменен."""
    presentation = Presentation(source)
    used = {slide.slide_layout.part.partname for slide in presentation.slides}
    added = []
    # Все masters, а не только slide_layouts первого master.
    for master in presentation.slide_masters:
        for layout in master.slide_layouts:
            if layout.part.partname in used:
                continue
            if _empty_panel(layout, presentation.slide_width, presentation.slide_height):
                continue
            kinds = [shape.placeholder_format.type.name for shape in layout.placeholders]
            if sum(kind in {"TITLE", "CENTER_TITLE"} for kind in kinds) != 1:
                continue
            if not any(kind in {"BODY", "SUBTITLE", "OBJECT"} for kind in kinds):
                continue
            # Пустая область фото/диаграммы не становится заполненным образцом.
            if any(kind not in {"TITLE", "CENTER_TITLE", "BODY", "SUBTITLE", "OBJECT",
                                "FOOTER", "HEADER", "DATE", "SLIDE_NUMBER"} for kind in kinds):
                continue
            presentation.slides.add_slide(layout)
            added.append({"source_slide_index": len(presentation.slides),
                          "layout_part": str(layout.part.partname), "name": layout.name})
            used.add(layout.part.partname)
    if not added:
        return source
    directory.mkdir(parents=True, exist_ok=True)
    output = directory / "layout-template.pptx"
    if output.resolve() == source.resolve():
        raise ValueError("Каталог макетов не должен перезаписывать исходник")
    presentation.save(output)
    assert len(Presentation(output).slides) == len(presentation.slides)
    (directory / "layout-template.json").write_text(json.dumps({
        "source_sha256": _sha(source), "catalog_sha256": _sha(output), "added": added,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    return output


def resolve_template_catalog(source: Path, directory: Path, profile_sha256: str) -> Path:
    """Старые профили используют исходник; новые — проверенную рабочую копию."""
    metadata = directory / "layout-template.json"
    if not metadata.is_file() or profile_sha256 == _sha(source):
        return source
    manifest = json.loads(metadata.read_text(encoding="utf-8"))
    output = directory / "layout-template.pptx"
    if (manifest["source_sha256"] != _sha(source)
            or manifest["catalog_sha256"] != profile_sha256
            or not output.is_file() or _sha(output) != profile_sha256):
        raise ValueError("Каталог макетов не соответствует исходному шаблону и профилю")
    return output
